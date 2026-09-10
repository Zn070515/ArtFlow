from __future__ import annotations

import hashlib
import json
import re
import secrets
from contextlib import contextmanager
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from typing import Iterable, Mapping

from accounts.services import require_current_admin, require_current_staff
from common.authority import (
    CONTEST_ROUND_STATE,
    JUDGE_PANEL_STATE,
    JUDGE_SCORE_SUBMISSION,
    JUDGE_SESSION_STATE,
    ROUND_PERFORMANCE_STATE,
    SCORE_SUMMARY_RECALCULATE,
    STAGE_RESULT_CONFIRM,
    TEST_DATA_CLEANUP,
    VOTE_SESSION_STATE,
    authority_write,
)
from common.business_rules import (
    ensure_activity_unlocked,
    ensure_lifecycle_consistent,
    ensure_round_unlocked,
)
from common.lifecycle import runtime_approved_singers, runtime_is_test, scope_runtime
from common.models import AuditLog
from common.test_data import lock_activity_for_runtime_data
from core.models import Activity
from core.policies import ActivityAction, ensure_activity_action_allowed
from core.services import lock_activity_for_action
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.db.models import F, Max, OuterRef, Q, QuerySet, Subquery
from django.utils import timezone
from ruleset.compiler import ExecutionPlan, compile_version
from ruleset.resolver import (
    OutcomeCode,
    ResolveInput,
    ResolveResult,
    ResolverState,
    checkpoint_inputs_fingerprint,
    inputs_fingerprint,
    resolve,
    resolve_to_checkpoint,
    stage_consumed_duel_keys,
    stage_consumed_scopes,
)
from voting.models import VoteSession

from .models import (
    Award,
    CompositeResult,
    ContestRound,
    DuelDecision,
    Judge,
    JudgeScoreReceipt,
    JudgeSeat,
    JudgeSeatGrant,
    JudgeSession,
    ManualDecision,
    Performance,
    PerformanceGroup,
    PerformanceRunState,
    RoundEntry,
    RoundJudge,
    RoundPanelSnapshot,
    RoundPanelSnapshotMember,
    RubricCriterion,
    ScoreRecord,
    ScoreSummary,
    ScoreWriteReceipt,
    ScoringRubric,
    SingerRegistration,
    StageAwardDecision,
    StageDecision,
    StageResult,
    _authorize_award_materialization,
    _authorize_duel_write,
    _authorize_manual_write,
    _award_materialization_authorized,
    _duel_write_authorized,
    _manual_write_authorized,
)


class ResultClosureCode(StrEnum):
    """Stable machine-readable reasons why an activity is not closed."""

    NO_CURRENT_FROZEN_RULESET = "no_current_frozen_ruleset"
    RULESET_BINDING_INVALID = "ruleset_binding_invalid"
    RAW_FACTS_INCOMPLETE = "raw_facts_incomplete"
    RAW_FACTS_UNLOCKED = "raw_facts_unlocked"
    RULE_REVIEW_REQUIRED = "rule_review_required"
    UPSTREAM_CONFIRMATION_PENDING = "upstream_confirmation_pending"
    STALE_CANDIDATE = "stale_candidate"
    STAGE_CONFIRMATION_PENDING = "stage_confirmation_pending"
    ACTIVITY_OPERATIONALLY_LOCKED = "activity_operationally_locked"
    SCHOOL_EXTERNAL_EVIDENCE_PENDING = "school_external_evidence_pending"


@dataclass(frozen=True)
class StageClosure:
    stage_key: str
    current_result_id: int | None
    result_version: int | None
    status: str | None
    confirmable: bool
    reasons: tuple[str, ...]
    input_fingerprint: str | None
    required_raw_facts: dict[str, tuple[int, ...]]
    raw_facts_locked: bool
    upstream_confirmed: bool
    official_awards: int
    round_entries: int
    blocking_reasons: tuple[ResultClosureCode, ...]


@dataclass(frozen=True)
class ResultClosure:
    activity_id: int
    ruleset_version_id: int | None
    ruleset_authority_hash: str | None
    stages: tuple[StageClosure, ...]
    closeable: bool
    blocking_reasons: tuple[ResultClosureCode, ...]


def result_closure_as_dict(closure: ResultClosure) -> dict[str, object]:
    """Serialize closure status without exposing model objects or private payloads."""
    return {
        "activity_id": closure.activity_id,
        "ruleset_version_id": closure.ruleset_version_id,
        "ruleset_authority_hash": closure.ruleset_authority_hash,
        "stages": [
            {
                "stage_key": stage.stage_key,
                "current_result_id": stage.current_result_id,
                "result_version": stage.result_version,
                "status": stage.status,
                "confirmable": stage.confirmable,
                "reasons": list(stage.reasons),
                "input_fingerprint": stage.input_fingerprint,
                "required_raw_facts": {
                    key: list(values) for key, values in stage.required_raw_facts.items()
                },
                "raw_facts_locked": stage.raw_facts_locked,
                "upstream_confirmed": stage.upstream_confirmed,
                "official_awards": stage.official_awards,
                "round_entries": stage.round_entries,
                "blocking_reasons": [code.value for code in stage.blocking_reasons],
            }
            for stage in closure.stages
        ],
        "closeable": closure.closeable,
        "blocking_reasons": [code.value for code in closure.blocking_reasons],
    }


def _eligible_singers(contest_round: ContestRound) -> QuerySet[SingerRegistration]:
    return SingerRegistration.objects.filter(round_entries__round=contest_round).order_by(
        "round_entries__running_order", "pk"
    )


def _active_judges(contest_round: ContestRound) -> QuerySet[Judge]:
    return Judge.objects.filter(round_assignments__round=contest_round).order_by("pk")


def authoritative_panel_judges(contest_round: ContestRound) -> QuerySet[Judge]:
    """Return the one judge set allowed to participate in this round's scoring matrix.

    A formal panel snapshot is authoritative even when a judge's mutable activity
    profile is later deactivated. Legacy rounds without a panel snapshot retain the
    prepared ``RoundJudge`` fallback for compatibility.
    """
    snapshot = (
        RoundPanelSnapshot.objects.filter(round=contest_round)
        .exclude(state=RoundPanelSnapshot.State.SUPERSEDED)
        .order_by("-version")
        .first()
    )
    if snapshot is None:
        return _active_judges(contest_round)
    return (
        Judge.objects.filter(
            panel_snapshot_members__panel_snapshot=snapshot,
            panel_snapshot_members__is_active=True,
        )
        .distinct()
        .order_by("pk")
    )


def _round_source_singers(contest_round: ContestRound) -> list[SingerRegistration]:
    source = contest_round.effective_roster_source()
    if source == ContestRound.RosterSource.APPROVED:
        return list(runtime_approved_singers(contest_round.activity).order_by("pk"))
    if source == ContestRound.RosterSource.LEGACY:
        previous_round = (
            ContestRound.objects.filter(activity=contest_round.activity)
            .filter(
                Q(sequence__lt=contest_round.sequence)
                | Q(sequence=contest_round.sequence, pk__lt=contest_round.pk)
            )
            .order_by("-sequence", "-pk")
            .first()
        )
        if previous_round is None:
            raise ValidationError("后续轮次必须先有上一轮比赛。")
        ensure_round_final_for_advancement(previous_round)
        return list(
            scope_runtime(
                SingerRegistration.objects.filter(
                    activity=contest_round.activity,
                    summaries__round=previous_round,
                    summaries__is_advanced=True,
                    summaries__is_test_data=runtime_is_test(contest_round.activity),
                ),
                contest_round.activity,
            ).order_by("pk")
        )
    if not contest_round.roster_source_stage:
        raise ValidationError("STAGE 来源轮次必须配置上游赛段。")
    upstream = (
        StageResult.objects.filter(
            activity=contest_round.activity,
            stage_key=contest_round.roster_source_stage,
            is_test_data=runtime_is_test(contest_round.activity),
        )
        .order_by("-pk")
        .first()
    )
    if upstream is None or upstream.status != StageResult.Status.CONFIRMED:
        raise ValidationError("上游赛段未核定，禁止准备该轮次。")
    singers = list(
        SingerRegistration.objects.filter(
            round_entries__round=contest_round, activity=contest_round.activity
        ).order_by("pk")
    )
    if not singers:
        raise ValidationError("尚未生成该轮晋级名单，请先重算并确认对应赛段结果。")
    return singers


def _order_singers_for_round(
    contest_round: ContestRound,
    singers: list[SingerRegistration],
    *,
    allow_unordered_manual=False,
):
    """Apply the declared order policy before the entry snapshot is frozen."""
    policy = contest_round.order_policy
    if policy == ContestRound.OrderPolicy.REGISTRATION_ORDER:
        return sorted(singers, key=lambda singer: singer.pk)
    if policy in (ContestRound.OrderPolicy.DRAW, ContestRound.OrderPolicy.MANUAL):
        existing = dict(
            RoundEntry.objects.filter(round=contest_round, singer__in=singers).values_list(
                "singer_id", "running_order"
            )
        )
        if set(existing) == {singer.pk for singer in singers} and all(
            order is not None for order in existing.values()
        ):
            return sorted(singers, key=lambda singer: (existing[singer.pk], singer.pk))
        if policy == ContestRound.OrderPolicy.MANUAL and not allow_unordered_manual:
            raise ValidationError("人工顺序轮次必须先通过正式服务设置完整出场顺序。")
        ordered = list(singers)
        if policy == ContestRound.OrderPolicy.DRAW:
            secrets.SystemRandom().shuffle(ordered)
        return ordered
    previous = (
        ContestRound.objects.filter(activity=contest_round.activity)
        .filter(
            Q(sequence__lt=contest_round.sequence)
            | Q(sequence=contest_round.sequence, pk__lt=contest_round.pk)
        )
        .order_by("-sequence", "-pk")
        .first()
    )
    if previous is None:
        raise ValidationError("按上一轮排名排序时必须存在上一轮比赛。")
    scores = dict(
        ScoreSummary.objects.filter(round=previous, singer__in=singers).values_list(
            "singer_id", "average_score"
        )
    )
    if any(singer.pk not in scores for singer in singers):
        raise ValidationError("按上一轮成绩排序时，所有选手都必须有已确定的上一轮成绩。")
    direction = 1 if policy == ContestRound.OrderPolicy.PREVIOUS_RANK_ASC else -1
    grouped: dict[Decimal, list[SingerRegistration]] = {}
    for singer in singers:
        grouped.setdefault(scores[singer.pk], []).append(singer)
    if any(len(group) > 1 for group in grouped.values()):
        if contest_round.tie_order_policy == ContestRound.TieOrderPolicy.REVIEW:
            raise ValidationError("上一轮成绩存在同分，必须先声明同分出场顺序。")
        for group in grouped.values():
            if len(group) < 2:
                continue
            if contest_round.tie_order_policy == ContestRound.TieOrderPolicy.DRAW:
                secrets.SystemRandom().shuffle(group)
            else:
                group.sort(key=lambda singer: singer.pk)
    return [singer for score in sorted(grouped, reverse=direction < 0) for singer in grouped[score]]


@transaction.atomic
def prepare_round(contest_round: ContestRound, operator) -> ContestRound:
    current_operator = require_current_staff(operator)
    locked_activity = lock_activity_for_action(contest_round.activity, ActivityAction.SCORE)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if locked_round.activity_id != locked_activity.pk:
        raise PermissionDenied("轮次不属于当前活动。")
    if locked_round.status != ContestRound.Status.DRAFT:
        raise ValidationError("比赛轮次只能从草稿状态准备。")

    source = locked_round.effective_roster_source()
    singers = _round_source_singers(locked_round)

    judges = list(
        Judge.objects.filter(activity=locked_round.activity, is_active=True).order_by("pk")
    )
    singers = _order_singers_for_round(locked_round, singers)
    if not singers or not judges:
        raise ValidationError("准备比赛轮次需要至少一名选手和一名活跃评委。")

    if locked_round.scoring_mode == ContestRound.ScoringMode.DROP_HIGH_LOW and len(judges) < 3:
        raise ValidationError("当前评分方式至少需要 3 名评委。")

    existing_entries = list(
        RoundEntry.objects.filter(round=locked_round).values_list("singer_id", "running_order")
    )
    manual_snapshot = locked_round.order_policy == ContestRound.OrderPolicy.MANUAL and (
        {singer_id for singer_id, _ in existing_entries} == {singer.pk for singer in singers}
        and all(order is not None for _, order in existing_entries)
    )
    if source != ContestRound.RosterSource.STAGE and not manual_snapshot:
        if existing_entries:
            RoundEntry.objects.filter(round=locked_round).delete()
        RoundEntry.objects.bulk_create(
            [
                RoundEntry(round=locked_round, singer=singer, running_order=index)
                for index, singer in enumerate(singers, start=1)
            ]
        )
    else:
        entries = RoundEntry.objects.filter(round=locked_round)
        entries.update(running_order=None)
        for index, singer in enumerate(singers, start=1):
            entries.filter(singer=singer).update(running_order=index)
    RoundJudge.objects.bulk_create(
        [RoundJudge(round=locked_round, judge=judge) for judge in judges]
    )
    locked_round.status = ContestRound.Status.PREPARED
    locked_round.is_locked = False
    with authority_write(CONTEST_ROUND_STATE):
        locked_round.save(update_fields=["status", "is_locked"])
    AuditLog.objects.create(
        operator=current_operator,
        action_type=AuditLog.ActionType.UPDATE_STATUS,
        target=f"ContestRound:{locked_round.pk}",
        new_value=json.dumps(
            {
                "round": locked_round.pk,
                "status": ContestRound.Status.PREPARED,
                "entries": len(singers),
                "judges": len(judges),
            },
            ensure_ascii=False,
        ),
    )
    return locked_round


@transaction.atomic
def set_round_running_order(contest_round, singer_ids, operator) -> ContestRound:
    """Persist a complete manual running-order snapshot on a draft round."""
    current_operator = require_current_staff(operator)
    locked_activity = lock_activity_for_action(contest_round.activity, ActivityAction.SCORE)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if locked_round.activity_id != locked_activity.pk:
        raise PermissionDenied("轮次不属于当前活动。")
    if locked_round.status != ContestRound.Status.DRAFT:
        raise ValidationError("只能为草稿轮次设置人工出场顺序。")
    if locked_round.order_policy != ContestRound.OrderPolicy.MANUAL:
        raise ValidationError("只有人工顺序轮次可以设置人工出场顺序。")
    singers = _round_source_singers(locked_round)
    ordered_ids = [str(singer_id) for singer_id in singer_ids]
    expected_ids = [str(singer.pk) for singer in singers]
    if len(ordered_ids) != len(set(ordered_ids)) or set(ordered_ids) != set(expected_ids):
        raise ValidationError("人工出场顺序必须完整覆盖该轮次的全部选手，且不能重复。")
    singer_by_id = {str(singer.pk): singer for singer in singers}
    ordered_pk_ids = [int(singer_id) for singer_id in ordered_ids]
    entries = RoundEntry.objects.filter(round=locked_round)
    entries.exclude(singer_id__in=ordered_pk_ids).delete()
    entries.update(running_order=None)
    existing_ids = set(entries.values_list("singer_id", flat=True))
    RoundEntry.objects.bulk_create(
        [
            RoundEntry(round=locked_round, singer=singer_by_id[singer_id], running_order=index)
            for index, singer_id in enumerate(ordered_ids, start=1)
            if singer_by_id[singer_id].pk not in existing_ids
        ]
    )
    for index, singer_id in enumerate(ordered_ids, start=1):
        entries.filter(singer_id=singer_by_id[singer_id].pk).update(running_order=index)
    AuditLog.objects.create(
        operator=current_operator,
        action_type=AuditLog.ActionType.UPDATE_STATUS,
        target=f"ContestRound:{locked_round.pk}",
        new_value=json.dumps(
            {"order_policy": locked_round.order_policy, "running_order": ordered_ids},
            ensure_ascii=False,
        ),
        note="set_round_running_order",
    )
    return locked_round


@transaction.atomic
def set_round_groups(contest_round, group_specs, operator) -> list[PerformanceGroup]:
    """Persist a complete, auditable group assignment for a draft round.

    Group membership is an operator-supplied fact, not a resolver convenience. The
    command therefore requires a complete partition of the round roster and writes
    groups and performances together while the Activity is locked.
    """
    current_operator = require_current_staff(operator)
    locked_activity = lock_activity_for_action(contest_round.activity, ActivityAction.SCORE)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if locked_round.activity_id != locked_activity.pk:
        raise PermissionDenied("轮次不属于当前活动。")
    if locked_round.status != ContestRound.Status.DRAFT:
        raise ValidationError("只能为草稿轮次设置分组。")
    if not isinstance(group_specs, list) or not group_specs:
        raise ValidationError("至少需要一个分组。")

    singers = _round_source_singers(locked_round)
    singer_by_id = {str(singer.pk): singer for singer in singers}
    seen: set[str] = set()
    normalized: list[tuple[str, list[SingerRegistration]]] = []
    names: set[str] = set()
    for index, spec in enumerate(group_specs, start=1):
        if not isinstance(spec, dict):
            raise ValidationError(f"第 {index} 个分组必须是对象。")
        name = str(spec.get("name") or "").strip()
        if not name or name in names:
            raise ValidationError("分组名称必须非空且不能重复。")
        raw_ids = spec.get("singer_ids")
        if not isinstance(raw_ids, list) or not raw_ids:
            raise ValidationError(f"分组 {name} 必须包含选手。")
        group_singers: list[SingerRegistration] = []
        for raw_id in raw_ids:
            singer_id = str(raw_id)
            if singer_id not in singer_by_id:
                raise ValidationError(f"选手 {singer_id} 不属于当前轮次。")
            if singer_id in seen:
                raise ValidationError(f"选手 {singer_id} 被重复分组。")
            seen.add(singer_id)
            group_singers.append(singer_by_id[singer_id])
        names.add(name)
        normalized.append((name, group_singers))
    if seen != set(singer_by_id):
        missing = sorted(set(singer_by_id) - seen)
        raise ValidationError(f"分组必须覆盖全部选手，缺少：{missing}。")

    Performance.objects.filter(round=locked_round).delete()
    PerformanceGroup.objects.filter(round=locked_round).delete()
    created: list[PerformanceGroup] = []
    for sequence, (name, group_singers) in enumerate(normalized, start=1):
        group = PerformanceGroup.objects.create(
            activity=locked_activity,
            round=locked_round,
            name=name,
            sequence=sequence,
            is_test_data=runtime_is_test(locked_activity),
        )
        Performance.objects.bulk_create(
            [
                Performance(
                    activity=locked_activity,
                    round=locked_round,
                    singer=singer,
                    group=group,
                    sequence=offset,
                    is_test_data=runtime_is_test(locked_activity),
                )
                for offset, singer in enumerate(group_singers, start=1)
            ]
        )
        created.append(group)
    AuditLog.objects.create(
        operator=current_operator,
        action_type=AuditLog.ActionType.OTHER,
        target=f"ContestRound:{locked_round.pk}",
        new_value=json.dumps(
            {
                "groups": [
                    {
                        "name": group.name,
                        "singer_ids": [p.singer_id for p in group.performances.all()],
                    }
                    for group in created
                ]
            },
            ensure_ascii=False,
        ),
        note="set_round_groups",
    )
    return created


@transaction.atomic
def create_scoring_rubric(activity, *, name, description, criteria, operator) -> ScoringRubric:
    """Provision a scoring rubric and all criteria as one audited operator command."""
    current_operator = require_current_staff(operator)
    locked_activity = lock_activity_for_action(activity, ActivityAction.SCORE)
    name = str(name or "").strip()
    if not name:
        raise ValidationError("评分标准名称不能为空。")
    if not isinstance(criteria, list) or not criteria:
        raise ValidationError("评分标准至少需要一个评分项。")
    normalized: list[dict] = []
    seen: set[str] = set()
    for index, item in enumerate(criteria, start=1):
        if not isinstance(item, dict):
            raise ValidationError(f"第 {index} 个评分项必须是对象。")
        criterion_name = str(item.get("name") or "").strip()
        if not criterion_name or criterion_name in seen:
            raise ValidationError("评分项名称必须非空且不能重复。")
        try:
            max_score = Decimal(str(item.get("max_score")))
        except (InvalidOperation, TypeError, ValueError):
            raise ValidationError(f"评分项 {criterion_name} 的最高分无效。")
        if max_score <= 0:
            raise ValidationError(f"评分项 {criterion_name} 的最高分必须大于 0。")
        seen.add(criterion_name)
        normalized.append(
            {
                "name": criterion_name,
                "max_score": max_score,
                "description": str(item.get("description") or "").strip(),
                "sequence": index,
            }
        )
    rubric = ScoringRubric.objects.create(
        activity=locked_activity,
        name=name,
        description=str(description or "").strip(),
        is_test_data=runtime_is_test(locked_activity),
    )
    RubricCriterion.objects.bulk_create(
        [RubricCriterion(rubric=rubric, **item) for item in normalized]
    )
    AuditLog.objects.create(
        operator=current_operator,
        action_type=AuditLog.ActionType.OTHER,
        target=f"ScoringRubric:{rubric.pk}",
        new_value=json.dumps(
            {"name": rubric.name, "criteria": normalized},
            ensure_ascii=False,
            default=str,
        ),
        note="create_scoring_rubric",
    )
    return rubric


@transaction.atomic
def reset_test_round_snapshots(
    contest_round: ContestRound,
    operator,
    *,
    test_only: bool = False,
    owned_singer_ids: set[int],
    owned_judge_ids: set[int],
) -> ContestRound:
    lock_activity_for_runtime_data(contest_round.activity)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if not test_only:
        raise ValidationError("Test round reset requires explicit test-only opt-in.")
    if not locked_round.activity.is_test_mode:
        raise ValidationError("Only test-mode activities can reset prepared round snapshots.")
    if (
        RoundEntry.objects.filter(round=locked_round)
        .exclude(singer_id__in=owned_singer_ids)
        .exists()
        or RoundEntry.objects.filter(round=locked_round, singer__is_test_data=False).exists()
        or RoundJudge.objects.filter(round=locked_round)
        .exclude(judge_id__in=owned_judge_ids)
        .exists()
    ):
        raise ValidationError("Test reset requires test-owned snapshot parents.")

    old_value = json.dumps(
        {
            "status": locked_round.status,
            "is_locked": locked_round.is_locked,
            "entries": locked_round.entries.count(),
            "judges": locked_round.round_judges.count(),
        },
        ensure_ascii=False,
    )
    locked_round.status = ContestRound.Status.DRAFT
    locked_round.is_locked = False
    with authority_write(CONTEST_ROUND_STATE):
        locked_round.save(update_fields=["status", "is_locked"])
    RoundEntry.objects.filter(round=locked_round).delete()
    RoundJudge.objects.filter(round=locked_round).delete()
    AuditLog.objects.create(
        operator=operator,
        action_type=AuditLog.ActionType.OTHER,
        target=f"ContestRound:{locked_round.pk}",
        old_value=old_value,
        new_value=json.dumps(
            {"status": ContestRound.Status.DRAFT, "is_locked": False, "entries": 0, "judges": 0},
            ensure_ascii=False,
        ),
        note="Demo test round reset",
    )
    return locked_round


@transaction.atomic
def reset_round_snapshots(
    contest_round: ContestRound, operator, *, reason: str = ""
) -> ContestRound:
    """Reset a round to DRAFT, discarding its score matrix and snapshots.

    For a test-mode activity, every parent carried into the reset must itself be
    test-marked; otherwise a test rehearsal snapshot is silently wiping formal
    rows. A formal activity is intentionally allowed to unwind via this path.
    """
    lock_activity_for_runtime_data(contest_round.activity)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    activity = locked_round.activity
    if activity.is_test_mode and (
        RoundEntry.objects.filter(round=locked_round, singer__is_test_data=False).exists()
        or ScoreRecord.objects.filter(round=locked_round, is_test_data=False).exists()
        or ScoreSummary.objects.filter(round=locked_round, is_test_data=False).exists()
    ):
        raise ValidationError("Test round reset requires test-owned scores and singers.")

    old_value = json.dumps(
        {
            "status": locked_round.status,
            "is_locked": locked_round.is_locked,
            "entries": locked_round.entries.count(),
            "judges": locked_round.round_judges.count(),
        },
        ensure_ascii=False,
    )
    locked_round.status = ContestRound.Status.DRAFT
    locked_round.is_locked = False
    with authority_write(CONTEST_ROUND_STATE):
        locked_round.save(update_fields=["status", "is_locked"])
    test_panel_snapshots = RoundPanelSnapshot.objects.filter(
        round=locked_round,
        is_test_data=True,
    )
    if activity.is_test_mode and test_panel_snapshots.exists():
        from entry_access.models import AccessGrant, EphemeralSession

        # Judge panel snapshots retain their source RoundJudge rows with PROTECT.  Test
        # cleanup is the one explicit lifecycle that may unwind the complete disposable
        # graph; formal snapshots remain immutable and therefore keep the old protection.
        with authority_write(JUDGE_SCORE_SUBMISSION):
            JudgeScoreReceipt.objects.filter(panel_snapshot__in=test_panel_snapshots).delete()
        with authority_write(JUDGE_SESSION_STATE):
            JudgeSession.objects.filter(panel_snapshot__in=test_panel_snapshots).delete()
            JudgeSeatGrant.objects.filter(panel_snapshot__in=test_panel_snapshots).delete()
        with authority_write(JUDGE_PANEL_STATE):
            JudgeSeat.objects.filter(panel_member__panel_snapshot__in=test_panel_snapshots).delete()
            RoundPanelSnapshotMember.objects.filter(
                panel_snapshot__in=test_panel_snapshots
            ).delete()
            RoundPanelSnapshot.objects.filter(pk__in=test_panel_snapshots).delete()
        with authority_write(TEST_DATA_CLEANUP):
            EphemeralSession.objects.filter(round=locked_round).delete()
        with authority_write(TEST_DATA_CLEANUP):
            AccessGrant.objects.filter(round=locked_round).delete()
        with authority_write(ROUND_PERFORMANCE_STATE):
            PerformanceRunState.objects.filter(round=locked_round).delete()
        Performance.objects.filter(round=locked_round).delete()
        PerformanceGroup.objects.filter(round=locked_round).delete()
    ScoreRecord.objects.filter(round=locked_round).delete()
    with authority_write(SCORE_SUMMARY_RECALCULATE):
        ScoreSummary.objects.filter(round=locked_round).delete()
    RoundEntry.objects.filter(round=locked_round).delete()
    RoundJudge.objects.filter(round=locked_round).delete()
    AuditLog.objects.create(
        operator=operator,
        action_type=AuditLog.ActionType.OTHER,
        target=f"ContestRound:{locked_round.pk}",
        old_value=old_value,
        new_value=json.dumps(
            {"status": ContestRound.Status.DRAFT, "is_locked": False, "entries": 0, "judges": 0},
            ensure_ascii=False,
        ),
        note=f"reset_round_snapshots={locked_round.pk} {reason}".strip(),
    )
    return locked_round


@transaction.atomic
def reset_round_to_draft(contest_round: ContestRound, actor, *, reason: str = "") -> ContestRound:
    """Admin-unwind a prepared/scoring/locked round back to DRAFT.

    Rejects when a later round already consumes this round's advancement, since
    resetting upstream would orphan the downstream snapshot. Requires a reason.
    """
    current_actor = require_current_admin(actor)
    if not reason.strip():
        raise ValidationError("重置轮次必须填写原因。")
    lock_activity_for_runtime_data(contest_round.activity)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if locked_round.activity.is_locked:
        raise PermissionDenied("活动结果已锁定，无法重置轮次。")
    if locked_round.status == ContestRound.Status.DRAFT:
        return locked_round
    ensure_round_not_consumed_by_confirmed_stage(locked_round)
    for other in (
        ContestRound.objects.select_for_update()
        .filter(activity=locked_round.activity)
        .exclude(pk=locked_round.pk)
        .order_by("pk")
    ):
        if other.status != ContestRound.Status.DRAFT and (
            (other.sequence, other.pk) > (locked_round.sequence, locked_round.pk)
        ):
            raise ValidationError("后续轮次仍在使用本轮结果，重置前必须先清空后续轮次。")
    return reset_round_snapshots(locked_round, current_actor, reason=reason)


def downstream_rounds(contest_round: ContestRound) -> QuerySet[ContestRound]:
    """Rounds that consume this round's advancement, if any.

    ``(sequence, pk)`` is the ordering authority: any round positioned later is a
    consumer. Use ``.first()`` on the result (both fields are defined) to pick the
    immediate downstream round.
    """
    return (
        ContestRound.objects.filter(activity=contest_round.activity)
        .filter(
            Q(sequence__gt=contest_round.sequence)
            | Q(sequence=contest_round.sequence, pk__gt=contest_round.pk)
        )
        .order_by("sequence", "pk")
    )


def ensure_round_final_for_advancement(contest_round: ContestRound) -> None:
    """Require a complete, locked upstream round before building downstream.

    A later-round snapshot must never consume a provisional result, or the two
    rounds would hold two different truths for the same advancement decision.
    """
    if contest_round.status != ContestRound.Status.LOCKED or not contest_round.is_locked:
        raise ValidationError("上游轮次尚未锁定，无法生成后续轮次。")
    if not expected_score_cells(contest_round):
        raise ValidationError("上游轮次没有可用的评分矩阵。")
    if missing_score_cells(contest_round):
        raise ValidationError("上游轮次仍有未录完的分数。")
    entry_count = RoundEntry.objects.filter(round=contest_round).count()
    summary_count = ScoreSummary.objects.filter(round=contest_round).count()
    if entry_count == 0 or summary_count != entry_count:
        raise ValidationError("上游轮次晋级名单不完整。")
    if ScoreSummary.objects.filter(round=contest_round, rank__lte=0).exists():
        raise ValidationError("上游轮次存在无效排名。")
    if contest_round.advancement_status == ContestRound.AdvancementStatus.NEEDS_REVIEW:
        raise ValidationError("上游轮次晋级名单存在同分，请先人工核定晋级人选。")


def validate_score(value: object) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValidationError("分数必须是 0 到 100 之间的数字。")
    try:
        score = Decimal(str(value).strip())
    except (AttributeError, InvalidOperation, ValueError):
        raise ValidationError("分数必须是 0 到 100 之间的数字。") from None
    if not score.is_finite() or score < 0 or score > 100:
        raise ValidationError("分数必须是 0 到 100 之间的数字。")
    if max(0, -int(score.as_tuple().exponent)) > 2:
        raise ValidationError("分数最多保留两位小数。")
    return score


def expected_score_cells(contest_round: ContestRound) -> list[tuple[int, int]]:
    singers = list(_eligible_singers(contest_round).values_list("pk", flat=True))
    judges = list(authoritative_panel_judges(contest_round).values_list("pk", flat=True))
    return [(singer_id, judge_id) for singer_id in singers for judge_id in judges]


def missing_score_cells(contest_round: ContestRound) -> list[dict[str, int | str]]:
    singers = list(_eligible_singers(contest_round))
    judges = list(authoritative_panel_judges(contest_round))
    existing = set(
        ScoreRecord.objects.filter(round=contest_round).values_list("singer_id", "judge_id")
    )
    return [
        {
            "singer_id": singer.pk,
            "singer_name": singer.name,
            "judge_id": judge.pk,
            "judge_name": judge.name,
        }
        for singer in singers
        for judge in judges
        if (singer.pk, judge.pk) not in existing
    ]


def ensure_score_records_belong_to_round(
    contest_round: ContestRound, pairs: Iterable[tuple[int, int]]
) -> None:
    singer_ids = set(_eligible_singers(contest_round).values_list("pk", flat=True))
    judge_ids = set(authoritative_panel_judges(contest_round).values_list("pk", flat=True))
    for singer_id, judge_id in pairs:
        if singer_id not in singer_ids or judge_id not in judge_ids:
            raise ValidationError("选手和评委必须属于当前比赛活动。")


def recalculate_round(contest_round: ContestRound) -> None:
    with authority_write(SCORE_SUMMARY_RECALCULATE):
        _recalculate_round(contest_round)


def _recalculate_round(contest_round: ContestRound) -> None:
    missing = missing_score_cells(contest_round)
    if missing:
        ScoreSummary.objects.filter(round=contest_round).delete()
        if contest_round.advancement_status != ContestRound.AdvancementStatus.AUTO:
            contest_round.advancement_status = ContestRound.AdvancementStatus.AUTO
            with authority_write(CONTEST_ROUND_STATE):
                contest_round.save(update_fields=["advancement_status"])
        return

    singers = _eligible_singers(contest_round)
    judge_ids = authoritative_panel_judges(contest_round).values_list("pk", flat=True)
    for singer in singers:
        singer_scores = list(
            ScoreRecord.objects.filter(
                round=contest_round, singer=singer, judge_id__in=judge_ids
            ).values_list("score", flat=True)
        )
        if (
            contest_round.scoring_mode == ContestRound.ScoringMode.DROP_HIGH_LOW
            and len(singer_scores) >= 3
        ):
            singer_scores.remove(max(singer_scores))
            singer_scores.remove(min(singer_scores))
        average = sum(singer_scores, Decimal("0")) / len(singer_scores)
        ScoreSummary.objects.update_or_create(
            round=contest_round,
            singer=singer,
            defaults={
                "average_score": average,
                "is_test_data": contest_round.activity.is_test_mode,
            },
        )

    ScoreSummary.objects.filter(round=contest_round).exclude(singer__in=singers).delete()
    summaries = list(
        ScoreSummary.objects.filter(round=contest_round).order_by("-average_score", "singer_id")
    )
    advance_count = contest_round.advance_count
    tie_at_boundary = False
    if advance_count and advance_count < len(summaries):
        boundary_score = summaries[advance_count - 1].average_score
        following_score = summaries[advance_count].average_score
        tie_at_boundary = boundary_score == following_score
    for rank, summary in enumerate(summaries, start=1):
        summary.rank = rank
        summary.is_advanced = bool(advance_count and rank <= advance_count)
        summary.save(update_fields=["rank", "is_advanced"])

    new_status = (
        ContestRound.AdvancementStatus.NEEDS_REVIEW
        if tie_at_boundary
        else ContestRound.AdvancementStatus.AUTO
    )
    if contest_round.advancement_status != new_status:
        contest_round.advancement_status = new_status
        with authority_write(CONTEST_ROUND_STATE):
            contest_round.save(update_fields=["advancement_status"])


@transaction.atomic
def finalize_advancement(
    contest_round: ContestRound,
    selected_singer_ids: Iterable[int],
    actor,
    *,
    note: str = "",
) -> ContestRound:
    """Record a staff-confirmed advancement selection after a boundary tie.

    The score matrix computes a rank-based proposal, but when two singers tie at
    the advancement boundary the system must not silently break the tie by
    database order. Staff explicitly choose who advances; that choice is frozen
    for downstream round consumption and audited.
    """
    current_actor = require_current_admin(actor)
    lock_activity_for_action(contest_round.activity, ActivityAction.SCORE)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if locked_round.status == ContestRound.Status.DRAFT:
        raise ValidationError("请先准备并完成比赛轮次后再核定晋级名单。")
    ensure_round_unlocked(locked_round)
    if missing_score_cells(locked_round):
        raise ValidationError("评分尚未完成，无法核定晋级名单。")
    entrants = set(_eligible_singers(locked_round).values_list("pk", flat=True))
    selected = {int(singer_id) for singer_id in selected_singer_ids}
    if not selected:
        raise ValidationError("请选择至少一名晋级选手。")
    if not selected <= entrants:
        raise ValidationError("晋级选手必须属于当前轮次。")
    advance_count = locked_round.advance_count
    if advance_count and len(selected) != advance_count:
        raise ValidationError(f"晋级核定人数必须等于该轮晋级名额（{advance_count} 人）。")

    old_advanced = set(
        ScoreSummary.objects.filter(round=locked_round, is_advanced=True).values_list(
            "singer_id", flat=True
        )
    )
    if (
        selected == old_advanced
        and locked_round.advancement_status == ContestRound.AdvancementStatus.FINALIZED
    ):
        return locked_round

    with authority_write(SCORE_SUMMARY_RECALCULATE):
        ScoreSummary.objects.filter(round=locked_round).update(is_advanced=False)
        ScoreSummary.objects.filter(round=locked_round, singer_id__in=selected).update(
            is_advanced=True
        )
    locked_round.advancement_status = ContestRound.AdvancementStatus.FINALIZED
    with authority_write(CONTEST_ROUND_STATE):
        locked_round.save(update_fields=["advancement_status"])
    AuditLog.objects.create(
        operator=current_actor,
        action_type=AuditLog.ActionType.FINALIZE_ADVANCEMENT,
        target=f"ContestRound:{locked_round.pk}",
        old_value=json.dumps({"advanced": sorted(old_advanced)}, ensure_ascii=False),
        new_value=json.dumps({"advanced": sorted(selected)}, ensure_ascii=False),
        note=note,
    )
    return locked_round


@transaction.atomic
def apply_scores(
    contest_round: ContestRound,
    score_values: Mapping[tuple[int, int], object],
    operator,
    *,
    note: str = "",
) -> list[dict[str, int | str | None]]:
    current_operator = require_current_staff(operator)
    lock_activity_for_action(contest_round.activity, ActivityAction.SCORE)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if locked_round.status == ContestRound.Status.DRAFT:
        raise ValidationError("请先准备比赛轮次后再评分。")
    if locked_round.status == ContestRound.Status.LOCKED:
        raise ValidationError("该比赛轮次已锁定。")
    if locked_round.is_locked or locked_round.activity.is_locked:
        raise ValidationError("该比赛轮次或活动已锁定。")

    pairs = list(score_values)
    ensure_score_records_belong_to_round(locked_round, pairs)
    normalized = {pair: validate_score(value) for pair, value in score_values.items()}
    changes: list[dict[str, int | str | None]] = []
    for (singer_id, judge_id), score in normalized.items():
        record = ScoreRecord.objects.filter(
            round=locked_round, singer_id=singer_id, judge_id=judge_id
        ).first()
        old_score = record.score if record else None
        if (
            record is not None
            and record.score == score
            and record.is_test_data == locked_round.activity.is_test_mode
        ):
            continue
        ScoreRecord.objects.update_or_create(
            round=locked_round,
            singer_id=singer_id,
            judge_id=judge_id,
            defaults={
                "score": score,
                "is_test_data": locked_round.activity.is_test_mode,
            },
        )
        changes.append(
            {
                "singer": singer_id,
                "judge": judge_id,
                "old": str(old_score) if old_score is not None else None,
                "new": str(score),
            }
        )

    recalculate_round(locked_round)
    if changes:
        if locked_round.status == ContestRound.Status.PREPARED:
            locked_round.status = ContestRound.Status.SCORING
        locked_round.score_version += 1
        with authority_write(CONTEST_ROUND_STATE):
            locked_round.save(update_fields=["status", "score_version"])
        AuditLog.objects.create(
            operator=current_operator,
            action_type=AuditLog.ActionType.ENTER_SCORE,
            target=f"ContestRound:{locked_round.pk}",
            new_value=json.dumps(
                {"round": locked_round.pk, "changes": changes}, ensure_ascii=False
            ),
            note=note,
        )
    return changes


class StaleScoreVersionError(Exception):
    """Raised when a rapid-entry save's ``base_version`` no longer matches the round.

    Signals a stale editor so the caller can surface a 409 and a fresh grid instead of
    silently overwriting another staff member's column.
    """

    def __init__(self, current_version: int):
        self.current_version = current_version
        super().__init__(f"评分已过期，当前版本为 {current_version}。")


class IdempotencyConflictError(Exception):
    """A rapid-score command ID was already used for a different command payload."""

    reason_code = "IDEMPOTENCY_CONFLICT"

    def __init__(self):
        super().__init__(self.reason_code)


RAPID_SCORE_OPERATION = "rapid_score_apply"


def _validated_score_command_id(command_id: object) -> str:
    if not isinstance(command_id, str):
        raise ValidationError("缺少 command_id。")
    command_id = command_id.strip()
    if not command_id or len(command_id) > 64:
        raise ValidationError("command_id 无效。")
    return command_id


def _score_command_payload_hash(
    contest_round: ContestRound,
    base_version: int,
    score_values: Mapping[tuple[int, int], object],
) -> str:
    """Hash a canonical, non-sensitive description of the score command.

    The receipt never persists the cells themselves.  IDs, the optimistic version, and
    the operation context make an accidental reuse across activity/round/version a
    conflict instead of an authority bypass.
    """
    payload = {
        "activity_id": contest_round.activity_id,
        "base_version": base_version,
        "cells": [
            {"judge_id": judge_id, "score": str(score).strip(), "singer_id": singer_id}
            for (singer_id, judge_id), score in sorted(score_values.items())
        ],
        "operation": RAPID_SCORE_OPERATION,
        "round_id": contest_round.pk,
    }
    canonical = json.dumps(payload, ensure_ascii=True, separators=(",", ":"), sort_keys=True)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


@transaction.atomic
def apply_scores_if_version(
    round_id: int,
    base_version: int | None,
    score_values: Mapping[tuple[int, int], object],
    operator,
    *,
    command_id: str,
    note: str = "",
) -> dict:
    """Apply a sparse cell save under the M0 Activity→Round lock order (M1-H).

    The version check, score application, and version bump all happen inside one
    transaction that takes the Activity lock first, then the ContestRound lock — never
    ``Round → Activity``. A ``base_version`` mismatch raises
    :class:`StaleScoreVersionError` (409) rather than being written over. Returns the
    new ``score_version`` and whether the matrix is now complete.
    """
    if base_version is None:
        raise ValidationError("缺少 base_version。")
    command_id = _validated_score_command_id(command_id)
    current_operator = require_current_staff(operator)
    try:
        parsed_base_version = int(base_version)
    except (TypeError, ValueError):
        raise ValidationError("base_version 无效。") from None
    contest_round = ContestRound.objects.only("activity_id").get(pk=round_id)
    locked_activity = Activity.objects.select_for_update().get(pk=contest_round.activity_id)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=round_id)
    )
    if locked_round.activity_id != locked_activity.pk:
        raise ValidationError("评分轮次活动上下文不一致。")
    payload_hash = _score_command_payload_hash(locked_round, parsed_base_version, score_values)
    try:
        with transaction.atomic():
            receipt = ScoreWriteReceipt.objects.create(
                command_id=command_id,
                operator=current_operator,
                operation=RAPID_SCORE_OPERATION,
                payload_hash=payload_hash,
                result_version=0,
                result_payload={},
                status=ScoreWriteReceipt.Status.PENDING,
            )
    except IntegrityError:
        receipt = ScoreWriteReceipt.objects.select_for_update().get(command_id=command_id)
        if (
            receipt.operator_id != current_operator.pk
            or receipt.payload_hash != payload_hash
            or receipt.operation != RAPID_SCORE_OPERATION
        ):
            raise IdempotencyConflictError() from None
        if receipt.status != ScoreWriteReceipt.Status.SUCCEEDED:
            raise ValidationError("评分写入回执仍在处理中。")
        return receipt.result_payload
    ensure_activity_unlocked(locked_activity)
    ensure_activity_action_allowed(locked_activity, ActivityAction.SCORE)
    if locked_round.status == ContestRound.Status.DRAFT:
        raise ValidationError("请先准备比赛轮次后再录入评分。")
    if parsed_base_version != locked_round.score_version:
        raise StaleScoreVersionError(locked_round.score_version)
    apply_scores(locked_round, score_values, current_operator, note=note)
    locked_round.refresh_from_db()
    result = {
        "status": ScoreWriteReceipt.Status.SUCCEEDED,
        "reason_code": "SCORES_APPLIED",
        "version": locked_round.score_version,
        "matrix_complete": not missing_score_cells(locked_round),
    }
    receipt.result_version = locked_round.score_version
    receipt.result_payload = result
    receipt.status = ScoreWriteReceipt.Status.SUCCEEDED
    receipt.full_clean(validate_unique=False)
    receipt.save(update_fields=["result_version", "result_payload", "status"])
    return result


@transaction.atomic
def lock_round(contest_round: ContestRound, operator) -> ContestRound:
    """Freeze a prepared/scoring round once its score matrix is complete."""
    current_operator = require_current_staff(operator)
    lock_activity_for_action(contest_round.activity, ActivityAction.SCORE)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if (
        locked_round.status not in {ContestRound.Status.PREPARED, ContestRound.Status.SCORING}
        or locked_round.is_locked
    ):
        raise PermissionDenied("该比赛轮次已锁定。")
    if not expected_score_cells(locked_round):
        raise PermissionDenied("当前轮次没有可锁定的完整评分矩阵。")
    if missing_score_cells(locked_round):
        raise PermissionDenied("仍有未完成的评委评分，不能锁定结果。")
    if locked_round.advancement_status == ContestRound.AdvancementStatus.NEEDS_REVIEW:
        raise PermissionDenied("晋级线存在同分，请先人工核定晋级名单。")
    locked_round.is_locked = True
    locked_round.status = ContestRound.Status.LOCKED
    with authority_write(CONTEST_ROUND_STATE):
        locked_round.save(update_fields=["is_locked", "status"])
    AuditLog.objects.create(
        operator=current_operator,
        action_type=AuditLog.ActionType.RELOCK_RESULT,
        target=f"ContestRound:{locked_round.pk}",
        old_value="unlocked",
        new_value="locked",
    )
    return locked_round


@transaction.atomic
def unlock_round(contest_round: ContestRound, actor, *, note: str = "") -> ContestRound:
    """Unlock an upstream locked round only if no downstream round consumes it."""
    current_actor = require_current_admin(actor)
    lock_activity_for_action(contest_round.activity)
    locked_round = (
        ContestRound.objects.select_for_update().select_related("activity").get(pk=contest_round.pk)
    )
    if locked_round.status == ContestRound.Status.SCORING and not locked_round.is_locked:
        # Stage-result unlock may already release this input round. Treat a repeated
        # operator click as an idempotent command rather than reporting a false failure.
        return locked_round
    if locked_round.status != ContestRound.Status.LOCKED or not locked_round.is_locked:
        raise PermissionDenied("该比赛轮次未锁定。")
    ensure_round_not_consumed_by_confirmed_stage(locked_round)
    for other in (
        ContestRound.objects.select_for_update()
        .filter(activity=locked_round.activity, pk__in=downstream_rounds(locked_round))
        .order_by("pk")
    ):
        if other.status != ContestRound.Status.DRAFT:
            raise PermissionDenied("后续轮次仍在使用本轮结果，解锁前必须先清空后续轮次。")
    locked_round.is_locked = False
    locked_round.status = ContestRound.Status.SCORING
    with authority_write(CONTEST_ROUND_STATE):
        locked_round.save(update_fields=["is_locked", "status"])
    AuditLog.objects.create(
        operator=current_actor,
        action_type=AuditLog.ActionType.UNLOCK_RESULT,
        target=f"ContestRound:{locked_round.pk}",
        old_value="locked",
        new_value="unlocked",
        note=note,
    )
    return locked_round


_META_SHEET_NAME = "ArtFlowMeta"
_SCORE_WORKBOOK_SCHEMA_VERSION = "1"
# §35: when the ArtFlowMeta sheet is present, every field is mandatory. A missing
# field means the metadata was damaged or hand-edited — the importer must reject
# rather than fail-open into the legacy name parser and import the wrong rows.
_META_REQUIRED_KEYS = (
    "schema_version",
    "activity_id",
    "round_id",
    "entry_ids",
    "judge_ids",
    "snapshot_fingerprint",
)


def _read_artflow_meta(workbook) -> dict[str, str] | None:
    """Return the meta sheet as a dict, or ``None`` when the sheet is entirely absent.

    ``None`` (no ArtFlowMeta sheet) is the only signal that lets a workbook take the
    explicit legacy name-based import path. A present-but-empty sheet yields ``{}``
    so the caller treats it as metadata present but incomplete → reject.
    """
    if _META_SHEET_NAME not in workbook.sheetnames:
        return None
    meta_worksheet = workbook[_META_SHEET_NAME]
    data: dict[str, str] = {}
    for row in meta_worksheet.iter_rows(values_only=True):
        if not row or len(row) < 2 or row[0] is None:
            continue
        key = str(row[0]).strip()
        value = row[1]
        if key:
            data[key] = str(value).strip() if value is not None else ""
    return data


def _split_ids(value: str | None) -> list[str]:
    if not value:
        return []
    return [token for token in (str(value).strip().split(",")) if token]


def snapshot_fingerprint(
    *,
    schema_version: str,
    activity_id,
    round_id,
    ruleset_version: str,
    entry_ids: list[str],
    judge_ids: list[str],
) -> str:
    """Deterministic fingerprint of the snapshot a score workbook was built from.

    Any drift in the rosters or the bound ruleset (or tampering with the meta
    sheet) changes the fingerprint, so the importer can reject a stale or edited
    workbook outright with zero partial mutation.
    """
    payload = json.dumps(
        {
            "schema_version": str(schema_version),
            "activity_id": str(activity_id),
            "round_id": str(round_id),
            "ruleset_version": str(ruleset_version or ""),
            "entry_ids": [str(x) for x in entry_ids],
            "judge_ids": [str(x) for x in judge_ids],
        },
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _version_binding(version) -> dict:
    """The frozen version's binding snapshot, falling back to the ruleset for legacy data.

    A frozen version is self-authoritative: once bound at freeze, the runtime reads the
    binding from the version and never from the still-mutable :class:`ContestRuleset`
    fields. Pre-R1 frozen versions carried no snapshot, so we fall back to the ruleset's
    fields — the data migration backfills those, and this fallback only survives for
    unmigrated rows.
    """
    binding = version.binding or {}
    if not binding.get("stage_key") and not binding.get("round_keys"):
        ruleset = version.ruleset
        if ruleset.round_keys or ruleset.stage_key:
            binding = {
                "stage_key": ruleset.stage_key or "",
                "round_keys": dict(ruleset.round_keys or {}),
                "vote_keys": dict(ruleset.vote_keys or {}),
                "group_keys": dict(ruleset.group_keys or {}),
                "audience_keys": dict(ruleset.audience_keys or {}),
                "announcement_blocks": list(ruleset.announcement_blocks or []),
                "announcement_blocks_by_checkpoint": dict(
                    ruleset.announcement_blocks_by_checkpoint or {}
                ),
            }
    return binding


def bound_ruleset_version_label(contest_round: ContestRound) -> str:
    """Label the frozen ruleset version bound to a round via its binding snapshot, if any."""
    from ruleset.models import ContestRuleset, RulesetVersion

    for ruleset in ContestRuleset.objects.filter(activity_id=contest_round.activity_id):
        version = (
            ruleset.versions.filter(is_current=True, status=RulesetVersion.Status.FROZEN)
            .order_by("-version")
            .first()
        )
        binding = _version_binding(version) if version is not None else {}
        if not binding.get("round_keys"):
            binding = {"round_keys": dict(ruleset.round_keys or {})}
        for round_pk in binding.get("round_keys", {}).values():
            if str(round_pk) == str(contest_round.pk):
                return f"{ruleset.pk}.{version.version}" if version else f"{ruleset.pk}.draft"
    return ""


def _resolve_judge_header(
    header: str, judge_by_id: dict[int, Judge], judge_by_name: dict[str, Judge]
) -> Judge | None:
    header = header.strip()
    if not header:
        return None
    match = re.match(r"^J(\d+)\b", header)
    if match:
        return judge_by_id.get(int(match.group(1)))
    return judge_by_name.get(header)


def _parse_id_authority_workbook(rows, headers, meta, contest_round: ContestRound):
    """Parse a workbook generated with its ArtFlowMeta sheet.

    Singer and judge IDs are the authority; the name column is only a hint and
    is never used for lookup. A mismatched activity/round is rejected outright.
    """
    errors: list[str] = []
    if meta.get("activity_id") and str(contest_round.activity_id) != meta["activity_id"]:
        errors.append("评分表属于其他活动。")
    if meta.get("round_id") and str(contest_round.pk) != meta["round_id"]:
        errors.append("评分表属于其他轮次。")

    if meta.get("schema_version") and meta["schema_version"] != _SCORE_WORKBOOK_SCHEMA_VERSION:
        errors.append(f"评分表 schema_version 不受支持（{meta['schema_version']}）。")

    active_judges = list(authoritative_panel_judges(contest_round))
    judge_by_id = {judge.pk: judge for judge in active_judges}
    judge_by_name = {judge.name: judge for judge in active_judges}
    singers = {singer.pk: singer for singer in _eligible_singers(contest_round)}

    meta_entries = _split_ids(meta.get("entry_ids"))
    current_entries = sorted(str(pk) for pk in singers)
    if meta_entries and set(current_entries) != set(meta_entries):
        errors.append("评分表选手名单过期或与当前轮次不一致。")

    meta_judges = _split_ids(meta.get("judge_ids"))
    current_judges = sorted(str(pk) for pk in judge_by_id)
    if meta_judges and set(current_judges) != set(meta_judges):
        errors.append("评分表评委名单过期或与当前轮次不一致。")

    if meta.get("ruleset_version"):
        bound = bound_ruleset_version_label(contest_round)
        if bound and meta["ruleset_version"] != bound:
            errors.append("评分表所属赛制版本已过期。")

    if meta.get("snapshot_fingerprint"):
        expected = snapshot_fingerprint(
            schema_version=meta.get("schema_version", _SCORE_WORKBOOK_SCHEMA_VERSION),
            activity_id=meta.get("activity_id", str(contest_round.activity_id)),
            round_id=meta.get("round_id", str(contest_round.pk)),
            ruleset_version=meta.get("ruleset_version", ""),
            entry_ids=meta_entries,
            judge_ids=meta_judges,
        )
        if expected != meta["snapshot_fingerprint"]:
            errors.append("评分表指纹不匹配（快照已过期或被修改）。")

    valid_judges: dict[int, Judge] = {}
    for column in range(2, len(headers)):
        header = headers[column]
        if not header:
            continue
        judge = _resolve_judge_header(header, judge_by_id, judge_by_name)
        if judge is None:
            errors.append(f"第 {column + 1} 列评委不存在: {header}")
        else:
            valid_judges[column] = judge

    scores = {}
    for row_number, row in enumerate(rows, start=2):
        if not row:
            continue
        singer_value = row[0] if len(row) > 0 else None
        try:
            singer_id = int(str(singer_value).strip()) if singer_value is not None else None
        except (TypeError, ValueError):
            singer_id = None
        if singer_id is None or singer_id not in singers:
            errors.append(f"第 {row_number} 行选手不属于当前轮次: {singer_value}")
            continue
        for column, score_value in enumerate(row[2:], start=2):
            if score_value is None or str(score_value).strip() == "":
                continue
            judge = valid_judges.get(column)
            if judge is None:
                errors.append(f"第 {row_number} 行存在没有评委标题的分数。")
                continue
            pair = (singer_id, judge.pk)
            if pair in scores:
                errors.append(f"分数重复: 选手{singer_id} {judge.name}")
                continue
            try:
                scores[pair] = validate_score(score_value)
            except ValidationError as error:
                errors.extend(
                    f"第 {row_number} 行分数错误: 选手{singer_id} {judge.name} ({message})"
                    for message in error.messages
                )
    return scores, errors


def _parse_name_based_workbook(rows, headers, contest_round: ContestRound):
    """Legacy fallback that matches singers and judges by exact name only."""
    judge_names = headers[1:]
    duplicate_judge_names = {name for name in judge_names if name and judge_names.count(name) > 1}
    errors = [f"评委姓名重复: {name}" for name in sorted(duplicate_judge_names)]
    judges_by_name = {
        judge.name: judge
        for judge in authoritative_panel_judges(contest_round)
        if judge.name not in duplicate_judge_names
    }
    singers = list(_eligible_singers(contest_round))
    singer_names = {singer.name for singer in singers}
    duplicate_singer_names = {
        name for name in singer_names if sum(singer.name == name for singer in singers) > 1
    }
    errors.extend(f"选手姓名重复: {name}" for name in sorted(duplicate_singer_names))
    singer_by_name = {
        singer.name: singer for singer in singers if singer.name not in duplicate_singer_names
    }
    scores = {}
    for row_number, row in enumerate(rows, start=2):
        if not row:
            continue
        singer_value = row[0] if len(row) > 0 else None
        singer_name = str(singer_value).strip() if singer_value is not None else ""
        if not singer_name:
            continue
        singer = singer_by_name.get(singer_name)
        if singer is None:
            errors.append(f"第 {row_number} 行选手不存在或姓名不唯一: {singer_name}")
            continue
        for column, score_value in enumerate(row[1:], start=1):
            if score_value is None or str(score_value).strip() == "":
                continue
            if column >= len(headers) or not headers[column]:
                errors.append(f"第 {row_number} 行存在没有评委标题的分数。")
                continue
            judge_name = headers[column]
            judge = judges_by_name.get(judge_name)
            if judge is None:
                errors.append(f"第 {row_number} 列评委不存在或姓名不唯一: {judge_name}")
                continue
            pair = (singer.pk, judge.pk)
            if pair in scores:
                errors.append(f"分数重复: {singer_name} {judge_name}")
                continue
            try:
                scores[pair] = validate_score(score_value)
            except ValidationError as error:
                errors.extend(
                    f"第 {row_number} 行分数错误: {singer_name} {judge_name} ({message})"
                    for message in error.messages
                )
    return scores, errors


def parse_score_workbook(uploaded_file, contest_round: ContestRound):
    from openpyxl import load_workbook

    workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
    worksheet = workbook.active
    if worksheet is None:
        return {}, ["评分表没有工作表。"]
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return {}, ["评分表不能为空。"]

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    if not headers or not headers[0]:
        return {}, ["评分表第一列必须是选手姓名。"]

    meta = _read_artflow_meta(workbook)
    if meta is not None:
        # ArtFlowMeta present → the id-authority path is mandatory, not a choice.
        # Any missing field (incl. round_id) previously downgraded to the name parser;
        # §35 demands a hard reject instead so a tampered sheet can never import wrong rows.
        missing = [key for key in _META_REQUIRED_KEYS if not meta.get(key)]
        if missing:
            return {}, [
                "评分表 ArtFlowMeta 缺失必需字段: "
                + ", ".join(missing)
                + "（拒绝降级为姓名导入）。"
            ]
        return _parse_id_authority_workbook(rows, headers, meta, contest_round)
    # No ArtFlowMeta sheet → explicit legacy name-based import (unchanged).
    return _parse_name_based_workbook(rows, headers, contest_round)


# --- M1-F: bind a frozen ruleset to the database, resolve, and persist. ---------


def _consumed_entry_round(version, activity, bound_rounds, checkpoint) -> ContestRound | None:
    """The highest-sequence consumed bound round for a checkpoint, or None.

    A checkpoint's dependency closure names the round keys it reads; the round whose
    snapshot declares the advancing roster is the last (largest ``(sequence, pk)``) of
    those. Non-checkpoint (full) stages are not roster-scoped, so return None and let the
    binder use the whole approved roster instead.
    """
    if checkpoint is None:
        return None
    consumed_round_keys, _, _, _ = stage_consumed_scopes(version.definition, checkpoint)
    candidates = [bound_rounds[k] for k in consumed_round_keys if k in bound_rounds]
    if not candidates:
        return None
    candidates.sort(key=lambda round_: (round_.sequence, round_.pk))
    return candidates[-1]


class StageRosterNotMaterialized(ValidationError):
    """A checkpoint's consumed STAGE/LEGACY round has no materialised entry roster yet.

    A formal resolve/confirm surfaces this as a hard error so staff confirm the upstream
    stage first; the progressive auto-resolve trigger treats it as "not ready yet" and
    skips the stage rather than crashing the whole probe loop.
    """


def _stage_entry_roster(version, activity, bound_rounds, *, checkpoint=None) -> tuple[str, ...]:
    """The entry roster for a stage bind.

    A checkpoint stage screens a round snapshot — the consumed bound round with the
    highest ``(sequence, pk)`` — whose ``RoundEntry`` is exactly the advancing set (the
    reviewer's "final ruleset entry must contain only the finalists"). When the stage is
    not checkpoint-scoped, or the consumed round has no materialized entry, fall back to
    the activity's approved singers so non-checkpoint and golden paths are unchanged.
    """
    entry_round = _consumed_entry_round(version, activity, bound_rounds, checkpoint)
    if entry_round is not None:
        entry_ids = list(
            RoundEntry.objects.filter(round=entry_round)
            .order_by("pk")
            .values_list("singer_id", flat=True)
        )
        if entry_ids:
            return tuple(str(pk) for pk in entry_ids)
        # The consumed round owns a subset roster (STAGE / LEGACY) that has not been
        # materialized yet. Refusing to widen to the whole approved list keeps the
        # "only the finalists" guarantee. An APPROVED round legitimately rosters every
        # approved singer regardless, so only that case falls back.
        if entry_round.effective_roster_source() != ContestRound.RosterSource.APPROVED:
            raise StageRosterNotMaterialized("尚未生成该轮晋级名单，请先重算并确认对应赛段结果。")
    return tuple(
        str(pk)
        for pk in scope_runtime(
            SingerRegistration.objects.filter(
                activity=activity,
                pre_status=SingerRegistration.PreStatus.APPROVED,
            ).order_by("pk"),
            activity,
        ).values_list("pk", flat=True)
    )


def bind_resolve_input(
    version,
    activity,
    *,
    round_keys: Mapping[str, ContestRound],
    vote_scores=None,
    group_of=None,
    manual=None,
    duel_decisions=None,
    checkpoint=None,
) -> ResolveInput:
    """Load a :class:`ResolveInput` from the DB for a frozen ruleset.

    A checkpoint stage's roster is the advancing snapshot of its consumed round (see
    :func:`_stage_entry_roster`), so the entry reflects *who advanced*, not the whole
    approved list. Non-checkpoint stages keep the approved singers. Round scores are
    read raw from :class:`ScoreRecord` grouped round -> singer -> judge scores, so the
    engine owns the scoring mode. Vote scores arrive already normalized (§16.9): the
    caller supplies the single Decimal per contestant — the binder never re-derives a
    raw-vote score.
    """
    roster = _stage_entry_roster(version, activity, round_keys, checkpoint=checkpoint)
    round_scores = {}
    for rkey, contest_round in round_keys.items():
        grouped: dict[str, list[Decimal]] = {}
        records = (
            ScoreRecord.objects.filter(round=contest_round)
            .select_related("singer", "judge")
            .order_by("singer_id", "judge__pk")
        )
        for rec in records:
            grouped.setdefault(str(rec.singer_id), []).append(rec.score)
        round_scores[rkey] = {k: tuple(v) for k, v in grouped.items()}
    return ResolveInput(
        roster=roster,
        round_scores=round_scores,
        vote_scores=vote_scores or {},
        group_of=group_of or {},
        manual=manual or {},
        duel_decisions=duel_decisions or {},
    )


def _source_group_of(activity, binding) -> dict[str, dict[str, str]]:
    """Source the ``group_of`` ResolveInput from the binding's ``group_keys``.

    Each ``by`` key resolves to a ContestRound; a singer's group is the
    ``PerformanceGroup`` of their ``Performance`` in that round.
    """
    group_keys = binding.get("group_keys") or {}
    if not group_keys:
        return {}
    out: dict[str, dict[str, str]] = {}
    for by, round_pk in group_keys.items():
        perfs = Performance.objects.filter(round_id=round_pk, group__isnull=False).select_related(
            "group"
        )
        entry: dict[str, str] = {}
        for p in perfs:
            assert p.group is not None
            entry[str(p.singer_id)] = p.group.name
        out[by] = entry
    return out


def _source_vote_scores(activity, binding) -> dict[str, dict[str, Decimal]]:
    """Source the ``vote_scores`` ResolveInput from the binding's ``vote_keys``.

    Each ``vote_source`` key resolves to a VoteSession; per-singer RAW vote count is
    returned (an authoritative fact). The loader never fabricates a normalized 0-10/100
    score from votes — converting votes to points belongs to an explicit VoteScoringRule,
    not the data loader. An empty session is left unbound so the resolver holds.
    """
    from voting.models import VoteOption, VoteRecord

    vote_keys = binding.get("vote_keys") or {}
    if not vote_keys:
        return {}
    test_flag = runtime_is_test(activity)
    out: dict[str, dict[str, Decimal]] = {}
    for source, vs_pk in vote_keys.items():
        # M1-R9-Final: initialise every candidate in the session to 0 first, so a zero-vote
        # candidate is a legal 0 rather than a "missing" result (the resolver would otherwise
        # HOLD on an absent entry). An orphan entry cast against a deleted option still counts.
        counts: dict[str, int] = {
            str(o.singer_id): 0 for o in VoteOption.objects.filter(vote_session_id=vs_pk)
        }
        records = VoteRecord.objects.filter(
            vote_session_id=vs_pk, is_test_data=test_flag
        ).select_related("vote_option")
        for rec in records:
            sid = str(rec.vote_option.singer_id)
            counts[sid] = counts.get(sid, 0) + 1
        if not counts:
            continue
        out[source] = {sid: Decimal(n) for sid, n in counts.items()}
    return out


def _source_audience_scores(activity, binding) -> dict[str, dict[str, Decimal]]:
    """Source the audience-scores ResolveInput from the binding's ``audience_keys``.

    Each ``audience_key`` maps to an ``AudienceScore.stage_key`` grouping. Per-singer
    scores are read as authoritative 0-100 values (scale ``hundred``) and returned under
    the definition's vote_source key, so the resolver's aggregate can combine them with
    judge scores by weight. An empty set is left unbound so the resolver holds.
    """
    from .models import AudienceScore

    audience_keys = binding.get("audience_keys") or {}
    if not audience_keys:
        return {}
    test_flag = runtime_is_test(activity)
    out: dict[str, dict[str, Decimal]] = {}
    for source_key, set_name in audience_keys.items():
        rows = AudienceScore.objects.filter(
            activity=activity, stage_key=set_name, is_test_data=test_flag
        ).select_related("singer")
        mapping: dict[str, Decimal] = {str(r.singer_id): r.score for r in rows}
        if not mapping:
            continue
        out[source_key] = mapping
    return out


def _combined_vote_scores(activity, binding) -> dict[str, dict[str, Decimal]]:
    """Merge raw vote counts and audience-scores into the resolver's vote channel."""
    vote_scores = _source_vote_scores(activity, binding)
    vote_scores.update(_source_audience_scores(activity, binding))
    return vote_scores


def _source_manual(version, activity) -> dict[str, dict[str, tuple[str, ...]]]:
    """Source the ``manual`` ResolveInput from the version's ManualDecision rows.

    Returns ``{manual_key: {group_or_empty: tuple_of_chosen_pks}}`` so the resolver's
    group (and flat ``""``) MANUAL_SELECT paths both read it directly.
    """
    decisions = ManualDecision.objects.filter(ruleset_version=version, activity=activity).order_by(
        "manual_key", "group", "pk"
    )
    out: dict[str, dict[str, tuple[str, ...]]] = {}
    for d in decisions:
        out.setdefault(d.manual_key, {})[d.group] = tuple(str(c) for c in (d.chosen or []))
    return out


def _source_duel_decisions(version, activity) -> dict[str, dict[str, str]]:
    """Source persisted pairwise winners for the frozen ruleset version."""
    decisions = DuelDecision.objects.filter(ruleset_version=version, activity=activity).order_by(
        "duel_key", "pair_key", "pk"
    )
    out: dict[str, dict[str, str]] = {}
    for decision in decisions:
        out.setdefault(decision.duel_key, {})[decision.pair_key] = decision.winner
    return out


def _plan_from_version(version) -> ExecutionPlan:
    if version.execution_plan:
        return ExecutionPlan.from_dict(json.loads(version.execution_plan))
    report, plan = compile_version(version)
    if not report.passes():
        raise ValidationError("无法解析无效赛制版本。")
    assert plan is not None
    return plan


# The resolver is a pure "machine computed" verdict. A resolved result is only
# "ready to confirm" (§36-37); it becomes final/locked when a staff member 核定.
_RESOLVER_STATUS = {
    ResolverState.HOLD: StageResult.Status.HOLD,
    ResolverState.REVIEW: StageResult.Status.REVIEW,
    ResolverState.READY: StageResult.Status.READY_TO_CONFIRM,
}


@transaction.atomic
def persist_stage_result(version, activity, result, *, stage_key, computed_by):
    """Write a :class:`ResolveResult` into StageResult + decisions + composites.

    Idempotency + versioning: the identity is ``(ruleset_version, input_fingerprint)``
    (the immutable authority, so two versions with the same definition but different
    binding never collide). Recomputing the same frozen ruleset over the same raw facts
    reuses the existing StageResult (a CONFIRMED one is returned as-is since it is
    immutable); any input change produces a new input_fingerprint and therefore a new
    versioned StageResult
    whose ``result_version`` is one past the stage's latest. This is what makes repeated
    recompute safe (no IntegrityError, no silent overwrite of a published result).
    """
    if version.ruleset.activity_id != activity.pk:
        raise ValidationError("Ruleset version must belong to the result activity.")
    is_test = runtime_is_test(activity)
    singer_by_key = {str(s.pk): s for s in SingerRegistration.objects.filter(activity=activity)}
    missing = [d.contestant for d in result.decisions if d.contestant not in singer_by_key]
    if missing:
        raise ValidationError(f"无法将结果写回选手：{missing}")

    # Identity is the immutable RulesetVersion (not the definition-only content_hash):
    # two versions with the same definition but different binding no longer collide
    # (P0-3). The authority_hash is retained for audit / external alignment.
    authority_hash = getattr(version, "authority_hash", "") or result.content_hash
    fingerprint = result.input_fingerprint
    status = _RESOLVER_STATUS.get(result.status)
    if status is None:
        raise ValidationError(f"无法映射解析器状态：{result.status}")
    # M1-R9-Final (current-version semantics): reuse only the *latest* StageResult, and only
    # when it is grounded on the same ruleset version and the same input facts. An earlier
    # same-fingerprint row (A → v1, then a corrected B → v2, then back to A) must NOT be
    # re-adopted — the correct A is a new v3, so confirm's latest-version check still holds.
    latest = (
        StageResult.objects.filter(activity=activity, stage_key=stage_key)
        .order_by("-result_version", "-pk")
        .first()
    )
    if (
        latest is not None
        and latest.ruleset_version_id == version.pk
        and latest.input_fingerprint == fingerprint
    ):
        if latest.status == StageResult.Status.CONFIRMED:
            return latest
        latest.status = status
        latest.reasons = list(result.reasons)
        latest.created_by = computed_by
        latest.save(update_fields=["status", "reasons", "created_by"])
        StageDecision.objects.filter(stage_result=latest).delete()
        CompositeResult.objects.filter(stage_result=latest).delete()
        StageAwardDecision.objects.filter(stage_result=latest).delete()
        _create_children(latest, result, singer_by_key, is_test)
        materialize_round_entry_from_stage(latest, operator=computed_by)
        return latest

    last_version = (
        StageResult.objects.filter(activity=activity, stage_key=stage_key).aggregate(
            m=Max("result_version")
        )["m"]
        or 0
    )
    stage = StageResult.objects.create(
        activity=activity,
        ruleset_version=version,
        created_by=computed_by,
        stage_key=stage_key,
        status=status,
        reasons=list(result.reasons),
        ruleset_hash=authority_hash,
        input_fingerprint=fingerprint,
        schema_version=result.schema_version,
        plan_version=result.plan_version,
        result_version=last_version + 1,
        is_test_data=is_test,
    )
    _create_children(stage, result, singer_by_key, is_test)
    materialize_round_entry_from_stage(stage, operator=computed_by)
    return stage


def _create_children(stage, result, singer_by_key, is_test):
    """Bulk-create a result's decisions + composites onto a stage result."""
    StageDecision.objects.bulk_create(
        [
            StageDecision(
                stage_result=stage,
                singer=singer_by_key[d.contestant],
                outcome_code=d.outcome_code.value,
                source_node=d.source_node,
                rank=d.rank,
                score=d.score,
                reason=d.reason,
                is_test_data=is_test,
            )
            for d in result.decisions
        ]
    )
    CompositeResult.objects.bulk_create(
        [
            CompositeResult(
                stage_result=stage,
                singer=singer_by_key[c.contestant],
                node_key=c.node_key,
                value=c.value,
                components=[
                    {
                        "source": source,
                        "weight": str(weight),
                        "value": str(value),
                        "contribution": str(contribution),
                    }
                    for (source, weight, value, contribution) in c.components
                ],
                is_test_data=is_test,
            )
            for c in result.composites
            if c.contestant in singer_by_key
        ]
    )
    StageAwardDecision.objects.bulk_create(
        [
            StageAwardDecision(
                activity=stage.activity,
                stage_result=stage,
                singer=singer_by_key[award.contestant],
                name=award.award,
                is_test_data=is_test,
                source_node=award.source_node,
            )
            for award in result.awards
            if award.contestant in singer_by_key
        ]
    )


@transaction.atomic
def materialize_stage_awards(stage: StageResult, *, operator) -> list[Award]:
    """Publish computed award candidates only as part of stage confirmation."""
    locked_stage = StageResult.objects.select_for_update().get(pk=stage.pk)
    if locked_stage.status != StageResult.Status.CONFIRMED:
        raise ValidationError("只有已核定赛段才能发布正式奖项。")
    candidates = list(
        StageAwardDecision.objects.filter(stage_result=locked_stage).select_related("singer")
    )
    existing = list(Award.objects.filter(source_stage_result=locked_stage))
    if existing:
        return existing
    with _authorized_award_materialization():
        return Award.objects.bulk_create(
            [
                Award(
                    activity=locked_stage.activity,
                    singer=candidate.singer,
                    name=candidate.name,
                    is_test_data=candidate.is_test_data,
                    source_stage_result=locked_stage,
                    source_award_decision=candidate,
                    source_node=candidate.source_node,
                )
                for candidate in candidates
            ]
        )


# Outcome codes that carry a contestant forward into the round bound to that stage.
# ``pending`` (tie awaiting adjudication) and ``eliminated`` are NOT advancing.
_ADVANCING_OUTCOME_CODES = frozenset(
    {
        OutcomeCode.DIRECT.value,
        OutcomeCode.ADVANCED.value,
        OutcomeCode.WILDCARD.value,
        OutcomeCode.FINALIST.value,
        OutcomeCode.REPECHAGE.value,
    }
)


@transaction.atomic
def materialize_round_entry_from_stage(stage: StageResult, *, operator=None) -> ContestRound | None:
    """Auto-generate a DRAFT round's entry roster from a resolved stage.

    The StageDecision table is the advancement authority (M1-INTEGRATION-1): every
    contestant with an advancing outcome code moves into the ``STAGE`` round bound to
    this stage (``roster_source_stage == stage.stage_key``). Only resolve on a stage
    that has actually decided something (READY_TO_CONFIRM / CONFIRMED), so a HOLD or
    REVIEW result never writes a partial roster. The round entry is reconciled to the
    *current* advancing set: a re-resolved stage (new ``result_version``) prunes
    superseded advancers and adds new ones, so a round still in DRAFT never silently
    retains a contestant who no longer advances. Returns the target round, or None
    when no DRAFT STAGE round consumes this stage.
    """
    if stage.status not in {
        StageResult.Status.READY_TO_CONFIRM,
        StageResult.Status.CONFIRMED,
    }:
        return None
    advancers = list(
        stage.decisions.filter(outcome_code__in=_ADVANCING_OUTCOME_CODES)
        .select_related("singer")
        .order_by("rank", "pk")
    )
    if not advancers:
        return None
    target = (
        ContestRound.objects.select_for_update()
        .filter(
            activity=stage.activity,
            roster_source=ContestRound.RosterSource.STAGE,
            roster_source_stage=stage.stage_key,
            status=ContestRound.Status.DRAFT,
        )
        .order_by("sequence", "pk")
        .first()
    )
    if target is None:
        return None
    target_singers = [decision.singer for decision in advancers]
    existing_orders = dict(
        RoundEntry.objects.filter(round=target, singer__in=target_singers).values_list(
            "singer_id", "running_order"
        )
    )
    manual_unordered = target.order_policy == ContestRound.OrderPolicy.MANUAL and not (
        set(existing_orders) == {singer.pk for singer in target_singers}
        and all(order is not None for order in existing_orders.values())
    )
    ordered_singers = _order_singers_for_round(
        target, target_singers, allow_unordered_manual=manual_unordered
    )
    running_order_by_singer = (
        {}
        if manual_unordered
        else {singer.pk: index for index, singer in enumerate(ordered_singers, start=1)}
    )
    entry_qs = RoundEntry.objects.filter(round=target)
    existing_ids = set(entry_qs.values_list("singer_id", flat=True))
    target_ids = {decision.singer_id for decision in advancers}
    removed = existing_ids - target_ids
    to_create = [
        RoundEntry(
            round=target,
            singer=decision.singer,
            running_order=running_order_by_singer.get(decision.singer_id),
        )
        for decision in advancers
        if decision.singer_id not in existing_ids
    ]
    if removed:
        entry_qs.filter(singer_id__in=removed).delete()
    if running_order_by_singer:
        entry_qs.update(running_order=None)
    if to_create:
        RoundEntry.objects.bulk_create(to_create)
    for singer_id, index in running_order_by_singer.items():
        entry_qs.filter(singer_id=singer_id).update(running_order=index)
    if removed or to_create:
        AuditLog.objects.create(
            operator=operator,
            action_type=AuditLog.ActionType.OTHER,
            target=f"ContestRound:{target.pk}",
            new_value=json.dumps(
                {
                    "stage_key": stage.stage_key,
                    "added_entries": len(to_create),
                    "removed_entries": len(removed),
                    "advancers": len(advancers),
                },
                ensure_ascii=False,
            ),
            note="materialize_round_entry_from_stage",
        )
    return target


def _definition_checkpoint_keys(definition) -> set[str]:
    obj = definition if isinstance(definition, dict) else json.loads(definition)
    return {c["key"] for c in (obj.get("checkpoints") or ())}


def _stage_is_checkpoint(version, stage_key: str) -> bool:
    """True when ``stage_key`` names a declared checkpoint in the frozen definition."""
    return stage_key in _definition_checkpoint_keys(version.definition)


def _bound_inputs(version, activity, *, checkpoint=None) -> tuple[ResolveInput, dict]:
    """Rebuild the current :class:`ResolveInput` from the frozen version's binding.

    This is the authority reference for staleness re-checking: it re-reads the raw
    facts (round scores, vote counts, group map, manual picks) exactly as the runtime
    would today. The version's frozen binding drives which rounds/votes/groups are read,
    so a binding-only change already alters the recomputed fingerprint. ``checkpoint``
    is threaded so a checkpoint stage's roster matches the one bound at persist time.
    """
    binding = _version_binding(version)
    raw_keys = binding.get("round_keys") or {}
    round_ids = list(raw_keys.values())
    rounds = {r.pk: r for r in ContestRound.objects.filter(pk__in=round_ids, activity=activity)}
    bound = {key: rounds[rid] for key, rid in raw_keys.items()}
    inputs = bind_resolve_input(
        version,
        activity,
        round_keys=bound,
        vote_scores=_combined_vote_scores(activity, binding),
        group_of=_source_group_of(activity, binding),
        manual=_source_manual(version, activity),
        duel_decisions=_source_duel_decisions(version, activity),
        checkpoint=checkpoint,
    )
    return inputs, binding


def _current_input_fingerprint(version, activity, stage_key: str) -> str:
    """The input fingerprint the current raw facts would produce for this stage.

    Checkpoint stages scope the projection to the closure's consumed keys (so future
    stages never shift an earlier stage's identity); full (non-checkpoint) stages use
    the whole-input fingerprint. A mismatch against a stored ``input_fingerprint`` means
    the result is stale.
    """
    checkpoint = stage_key if _stage_is_checkpoint(version, stage_key) else None
    inputs, _ = _bound_inputs(version, activity, checkpoint=checkpoint)
    if checkpoint:
        return checkpoint_inputs_fingerprint(version.definition, inputs, stage_key)
    return inputs_fingerprint(inputs)


def _ensure_stage_dependencies_final(version, activity, stage_key: str) -> None:
    """Require the raw facts a stage consumed to be at rest before it can be frozen.

    A CONFIRMED stage is the point of no return, so the scores it read must already be
    locked. Round sources require the bound :class:`ContestRound` to be locked; vote
    sources require the bound :class:`VoteSession` to be locked. Group maps are a view
    over a round, so their round lock is covered above.
    """
    from voting.models import VoteSession

    binding = _version_binding(version)
    checkpoint = stage_key if _stage_is_checkpoint(version, stage_key) else None
    round_keys, vote_keys, _, _ = stage_consumed_scopes(version.definition, checkpoint)
    bound_round_ids = [
        rid for k, rid in (binding.get("round_keys") or {}).items() if k in round_keys
    ]
    if bound_round_ids:
        open_rounds = ContestRound.objects.filter(
            pk__in=bound_round_ids, activity=activity
        ).exclude(is_locked=True)
        if open_rounds.exists():
            raise ValidationError("该赛段依赖的淘汰轮次尚未锁定，请先锁定轮次后再核定。")
    bound_vote_ids = [vid for k, vid in (binding.get("vote_keys") or {}).items() if k in vote_keys]
    if bound_vote_ids:
        open_votes = VoteSession.objects.filter(pk__in=bound_vote_ids, activity=activity).exclude(
            is_locked=True
        )
        if open_votes.exists():
            raise ValidationError("该赛段依赖的投票时段尚未锁定，请先锁定后再核定。")


_CONSUMED_BY_CONFIRMED_MSG = "该原始数据已被已核定赛段结果使用，请先由管理员解锁赛段结果。"


def _stage_consumed_facts(stage) -> dict[str, set]:
    """The raw pks a confirmed stage consumed, read back from its binding snapshot.

    This is the inverse of :func:`_ensure_stage_dependencies_final`: where that freeze
    gate requires the facts a stage *will* read to be at rest, this gate refuses to
    *un-wind* a round/vote/manual that a CONFIRMED stage already read. Both map the
    stage's closure keys (``round_keys``/``vote_keys``/``group_keys``/``manual_keys``)
    onto physical pks via the frozen version's binding.
    """
    version = stage.ruleset_version
    binding = _version_binding(version)
    checkpoint = stage.stage_key if _stage_is_checkpoint(version, stage.stage_key) else None
    round_keys, vote_keys, group_keys, manual_keys = stage_consumed_scopes(
        version.definition, checkpoint
    )
    duel_keys = stage_consumed_duel_keys(version.definition, checkpoint)
    round_map = binding.get("round_keys") or {}
    group_map = binding.get("group_keys") or {}
    vote_map = binding.get("vote_keys") or {}
    audience_map = binding.get("audience_keys") or {}
    round_pks = {round_map.get(k) for k in round_keys} | {group_map.get(k) for k in group_keys}
    round_pks.discard(None)
    # Audience sources ride the vote channel but are keyed by audience_keys, not a
    # VoteSession; split them out so the vote lock/finality check does not mis-read them.
    audience_source_keys = {k for k in vote_keys if k in audience_map}
    real_vote_keys = set(vote_keys) - audience_source_keys
    vote_pks = {vote_map.get(k) for k in real_vote_keys}
    vote_pks.discard(None)
    audience_names = {audience_map[k] for k in audience_source_keys}
    return {
        "rounds": round_pks,
        "votes": vote_pks,
        "audience": audience_names,
        "manual": set(manual_keys),
        "duels": set(duel_keys),
    }


def ensure_round_not_consumed_by_confirmed_stage(contest_round) -> None:
    """Reject un-winding a round that a CONFIRMED stage of its activity already read."""
    consumed = StageResult.objects.filter(
        activity_id=contest_round.activity_id, status=StageResult.Status.CONFIRMED
    ).values("pk")
    for stage_pk in consumed:
        stage = StageResult.objects.get(pk=stage_pk["pk"])
        if contest_round.pk in _stage_consumed_facts(stage)["rounds"]:
            raise ValidationError(_CONSUMED_BY_CONFIRMED_MSG)


def ensure_vote_not_consumed_by_confirmed_stage(vote_session) -> None:
    """Reject un-winding a VoteSession that a CONFIRMED stage of its activity already read."""
    consumed = StageResult.objects.filter(
        activity_id=vote_session.activity_id, status=StageResult.Status.CONFIRMED
    ).values("pk")
    for stage_pk in consumed:
        stage = StageResult.objects.get(pk=stage_pk["pk"])
        if vote_session.pk in _stage_consumed_facts(stage)["votes"]:
            raise ValidationError(_CONSUMED_BY_CONFIRMED_MSG)


def ensure_audience_not_consumed_by_confirmed_stage(activity, stage_key: str) -> None:
    """Reject re-entering an audience set that a CONFIRMED stage of its activity already read."""
    activity_id = activity.pk if hasattr(activity, "pk") else int(activity)
    consumed = StageResult.objects.filter(
        activity_id=activity_id, status=StageResult.Status.CONFIRMED
    ).values("pk")
    for stage_pk in consumed:
        stage = StageResult.objects.get(pk=stage_pk["pk"])
        if stage_key in _stage_consumed_facts(stage)["audience"]:
            raise ValidationError(_CONSUMED_BY_CONFIRMED_MSG)


def ensure_manual_not_consumed_by_confirmed_stage(version, manual_key) -> None:
    """Reject mutating a manual pick that a CONFIRMED stage of the same version already read."""
    consumed = StageResult.objects.filter(
        ruleset_version_id=version.pk, status=StageResult.Status.CONFIRMED
    ).values("pk")
    for stage_pk in consumed:
        stage = StageResult.objects.get(pk=stage_pk["pk"])
        if manual_key in _stage_consumed_facts(stage)["manual"]:
            raise ValidationError(_CONSUMED_BY_CONFIRMED_MSG)


def _manual_node_keys(version) -> frozenset[str]:
    """All MANUAL_SELECT node keys declared in the frozen definition."""
    return frozenset(stage_consumed_scopes(version.definition, None)[3])


def _duel_node_keys(version) -> frozenset[str]:
    """All DUEL node keys declared in the frozen definition."""
    return frozenset(stage_consumed_duel_keys(version.definition, None))


@contextmanager
def _authorized_manual_write():
    """Bracket a ManualDecision write with the production write-authorization guard.

    The guard is thread-local (models.py), so an untrusted direct ``.save()`` on a
    :class:`ManualDecision` is refused; only a formal service that holds the Activity lock
    may toggle it. Prior state is restored so a nested authorized write never leaks a
    spurious grant out of scope.
    """
    prior = _manual_write_authorized()
    _authorize_manual_write(True)
    try:
        yield
    finally:
        _authorize_manual_write(prior)


@contextmanager
def _authorized_duel_write():
    """Bracket a DuelDecision write with its formal service authority guard."""
    prior = _duel_write_authorized()
    _authorize_duel_write(True)
    try:
        yield
    finally:
        _authorize_duel_write(prior)


@contextmanager
def _authorized_award_materialization():
    """Allow only the stage-confirm service to create official stage awards."""
    prior = _award_materialization_authorized()
    _authorize_award_materialization(True)
    try:
        yield
    finally:
        _authorize_award_materialization(prior)


def ensure_duel_not_consumed_by_confirmed_stage(version, duel_key, pair_key=None) -> None:
    """Reject changing a DUEL input already read by a confirmed stage."""
    consumed = StageResult.objects.filter(
        ruleset_version_id=version.pk, status=StageResult.Status.CONFIRMED
    ).values("pk")
    for stage_pk in consumed:
        stage = StageResult.objects.get(pk=stage_pk["pk"])
        if duel_key in _stage_consumed_facts(stage)["duels"]:
            raise ValidationError(_CONSUMED_BY_CONFIRMED_MSG)


@transaction.atomic
def set_duel_decision(
    version, *, duel_key, pair_key, winner, created_by, activity=None
) -> DuelDecision:
    """Upsert one explicit DUEL winner for a frozen ruleset version."""
    from ruleset.models import RulesetVersion

    current_actor = require_current_admin(created_by)
    activity = activity or version.ruleset.activity
    locked_activity = lock_activity_for_action(activity)
    locked = RulesetVersion.objects.select_for_update().get(pk=version.pk)
    if locked.ruleset.activity_id != locked_activity.pk:
        raise ValidationError("赛制版本不属于当前活动。")
    if locked.status != RulesetVersion.Status.FROZEN:
        raise ValidationError("仅可更新已冻结赛制版本的对决决定。")
    if duel_key not in _duel_node_keys(locked):
        raise ValidationError(f"赛制未声明 DUEL 节点：{duel_key}。")
    pair_key = str(pair_key or "")
    winner = str(winner or "")
    existing = DuelDecision.objects.filter(
        ruleset_version=locked, duel_key=duel_key, pair_key=pair_key
    ).first()
    with _authorized_duel_write():
        decision, _created = DuelDecision.objects.update_or_create(
            ruleset_version=locked,
            duel_key=duel_key,
            pair_key=pair_key,
            defaults={
                "activity": locked_activity,
                "winner": winner,
                "created_by": current_actor,
                "is_test_data": runtime_is_test(locked_activity),
            },
        )
    AuditLog.objects.create(
        operator=current_actor,
        action_type=AuditLog.ActionType.DUEL_DECISION,
        target=f"RulesetVersion:{locked.pk}",
        old_value=json.dumps(
            {
                "duel_key": duel_key,
                "pair_key": pair_key,
                "winner": existing.winner if existing else None,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        new_value=json.dumps(
            {"duel_key": duel_key, "pair_key": pair_key, "winner": winner},
            ensure_ascii=False,
            sort_keys=True,
        ),
    )
    return decision


@transaction.atomic
def delete_duel_decision(
    version, *, duel_key, pair_key, activity=None, deleted_by=None
) -> DuelDecision | None:
    """Remove one explicit DUEL winner through the formal authority service."""
    from ruleset.models import RulesetVersion

    current_actor = require_current_admin(deleted_by)
    activity = activity or version.ruleset.activity
    locked_activity = lock_activity_for_action(activity)
    locked = RulesetVersion.objects.select_for_update().get(pk=version.pk)
    if locked.ruleset.activity_id != locked_activity.pk:
        raise ValidationError("赛制版本不属于当前活动。")
    if locked.status != RulesetVersion.Status.FROZEN:
        raise ValidationError("仅可删除已冻结赛制版本的对决决定。")
    if duel_key not in _duel_node_keys(locked):
        raise ValidationError(f"赛制未声明 DUEL 节点：{duel_key}。")
    decision = DuelDecision.objects.filter(
        ruleset_version=locked, duel_key=duel_key, pair_key=str(pair_key or "")
    ).first()
    if decision is None:
        return None
    old_value = {
        "duel_key": decision.duel_key,
        "pair_key": decision.pair_key,
        "winner": decision.winner,
    }
    with _authorized_duel_write():
        decision.delete()
    AuditLog.objects.create(
        operator=current_actor,
        action_type=AuditLog.ActionType.DUEL_DECISION,
        target=f"RulesetVersion:{locked.pk}",
        old_value=json.dumps(old_value, ensure_ascii=False, sort_keys=True),
        new_value="",
    )
    return decision


@transaction.atomic
def set_manual_decision(
    version, *, manual_key, group, chosen, created_by, activity=None
) -> ManualDecision:
    """Upsert a staff member's manual pick for a MANUAL_SELECT node (§31 closure).

    M0 lock order (Activity → RulesetVersion/ManualDecision): lock the Activity FOR
    UPDATE (the authoritative serialization point, same as confirm/unlock/recompute),
    then re-read the RulesetVersion fresh inside the lock. Re-validates that the version
    is frozen, belongs to the activity, and that ``manual_key`` is a declared MANUAL_SELECT
    node; the model's ``save`` re-checks the consumed-by-confirmed guard inside this
    window, so a pick a CONFIRMED stage already read can never be edited even under a
    concurrent recompute.
    """
    from ruleset.models import RulesetVersion

    current_actor = require_current_admin(created_by)
    activity = activity or version.ruleset.activity
    locked_activity = lock_activity_for_action(activity)
    locked = RulesetVersion.objects.select_for_update().get(pk=version.pk)
    if locked.ruleset.activity_id != locked_activity.pk:
        raise ValidationError("赛制版本不属于当前活动。")
    if locked.status != RulesetVersion.Status.FROZEN:
        raise ValidationError("仅可更新已冻结赛制版本的人工选择。")
    if manual_key not in _manual_node_keys(locked):
        raise ValidationError(f"赛制未声明 MANUAL_SELECT 节点：{manual_key}。")
    chosen = [str(c) for c in (chosen or [])]
    with _authorized_manual_write():
        decision, _created = ManualDecision.objects.update_or_create(
            ruleset_version=locked,
            manual_key=manual_key,
            group=group or "",
            defaults={
                "activity": locked_activity,
                "chosen": chosen,
                "created_by": current_actor,
                "is_test_data": runtime_is_test(locked_activity),
            },
        )
    AuditLog.objects.create(
        operator=current_actor,
        action_type=AuditLog.ActionType.MANUAL_DECISION,
        target=f"RulesetVersion:{locked.pk}",
        old_value="",
        new_value=json.dumps(
            {
                "manual_key": manual_key,
                "group": group or "",
                "chosen": chosen,
                "created_by": current_actor.pk,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
    )
    return decision


@transaction.atomic
def delete_manual_decision(
    version, *, manual_key, group="", activity=None, deleted_by=None
) -> ManualDecision | None:
    """Remove a staff member's manual pick for a MANUAL_SELECT node (§31 closure).

    Same lock order and frozen/version/manual-key validation as :func:`set_manual_decision`;
    the model's ``delete`` re-checks the consumed-by-confirmed guard inside the lock.
    Returns the deleted row (or ``None`` if no such pick exists).
    """
    from ruleset.models import RulesetVersion

    current_actor = require_current_admin(deleted_by)
    activity = activity or version.ruleset.activity
    locked_activity = lock_activity_for_action(activity)
    locked = RulesetVersion.objects.select_for_update().get(pk=version.pk)
    if locked.ruleset.activity_id != locked_activity.pk:
        raise ValidationError("赛制版本不属于当前活动。")
    if locked.status != RulesetVersion.Status.FROZEN:
        raise ValidationError("仅可删除已冻结赛制版本的人工选择。")
    if manual_key not in _manual_node_keys(locked):
        raise ValidationError(f"赛制未声明 MANUAL_SELECT 节点：{manual_key}。")
    decision = ManualDecision.objects.filter(
        ruleset_version=locked, manual_key=manual_key, group=group or ""
    ).first()
    if decision is None:
        return None
    with _authorized_manual_write():
        decision.delete()
    AuditLog.objects.create(
        operator=current_actor,
        action_type=AuditLog.ActionType.MANUAL_DECISION,
        target=f"RulesetVersion:{locked.pk}",
        old_value=json.dumps(
            {
                "manual_key": manual_key,
                "group": group or "",
                "chosen": [str(c) for c in (decision.chosen or [])],
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
        new_value="",
    )
    return decision


@transaction.atomic
def create_manual_award(activity, *, singer, name: str, operator) -> Award:
    """Create an explicitly manual official award under admin authority."""
    current_operator = require_current_admin(operator)
    locked_activity = lock_activity_for_action(activity, ActivityAction.MANAGE_AWARD)
    award_name = str(name or "").strip()
    if not award_name:
        raise ValidationError("奖项名称不能为空。")
    locked_singer = SingerRegistration.objects.select_for_update().get(pk=singer.pk)
    if locked_singer.activity_id != locked_activity.pk:
        raise ValidationError("奖项选手必须属于当前活动。")
    ensure_lifecycle_consistent(locked_activity, locked_singer, label="Award 选手")
    award = Award.objects.create(
        activity=locked_activity,
        singer=locked_singer,
        name=award_name,
        source_node="MANUAL",
        is_test_data=locked_activity.is_test_mode,
    )
    AuditLog.objects.create(
        operator=current_operator,
        action_type=AuditLog.ActionType.OTHER,
        target=f"Award:{award.pk}",
        new_value=json.dumps(
            {"name": award.name, "source": "MANUAL", "activity": locked_activity.pk},
            ensure_ascii=False,
        ),
    )
    return award


@transaction.atomic
def confirm_stage_result(stage: StageResult, *, confirmed_by):
    """核定并锁定 a resolved stage result into its final handcard state (§36-37).

    A resolver ``READY_TO_CONFIRM`` result is machine-computed but still mutable —
    another score edit would recompute it. The staff 核定 action is the point of no
    return: once CONFIRMED the result (and its decisions/composites) is immutable so
    a host can safely copy the handcard.

    M0 lock order and finality (M1-R6): lock the Activity first (the authoritative
    serialization point), then re-read the StageResult (never trust a pre-lock stale
    object), confirm it is the stage's *current* candidate, recompute the current input
    fingerprint and require it to match (no stale results frozen over changed facts), and
    require the consumed raw facts to be at rest (rounds/votes locked). Idempotent: a
    re-confirm of an already-locked row returns it as-is.
    """
    current_actor = require_current_admin(confirmed_by)
    locked_activity = lock_activity_for_action(stage.activity, ActivityAction.PUBLISH_RESULT)
    locked = StageResult.objects.select_for_update().get(pk=stage.pk)
    if locked.activity_id != locked_activity.pk:
        raise ValidationError("赛段结果不属于当前活动。")
    if locked.status not in {
        StageResult.Status.READY_TO_CONFIRM,
        StageResult.Status.CONFIRMED,
    }:
        raise ValidationError("仅可核定已解析到“待核定”状态的赛段结果。")
    # M1-R9 (§五): 核定 only finalizes a result grounded in the *current* FROZEN
    # authority. A superseded (demoted) or draft version is not the running authority —
    # finalizing against it would let a host copy a handcard that no longer matches what
    # the activity publishes. Re-read with FOR UPDATE so the currency decision is taken
    # under the Activity serialization lock held above.
    from ruleset.models import RulesetVersion

    version = RulesetVersion.objects.select_for_update().get(pk=locked.ruleset_version_id)
    if version.status != RulesetVersion.Status.FROZEN or not version.is_current:
        raise ValidationError("只能核定基于当前冻结赛制版本生成的结果。")
    latest = StageResult.objects.filter(
        activity=locked.activity, stage_key=locked.stage_key
    ).aggregate(m=Max("result_version"))["m"]
    if locked.result_version != latest:
        raise ValidationError("该赛段存在更新的结果版本，请核定最新结果。")
    try:
        current_fingerprint = _current_input_fingerprint(version, locked.activity, locked.stage_key)
    except ValueError as error:
        raise ValidationError(str(error))
    if current_fingerprint != locked.input_fingerprint:
        raise ValidationError("该结果已过期（原始分数/票数已变化），请重新计算后核定。")
    _ensure_stage_dependencies_final(version, locked.activity, locked.stage_key)
    if locked.status == StageResult.Status.CONFIRMED:
        return locked
    locked.status = StageResult.Status.CONFIRMED
    locked.confirmed_by = current_actor
    locked.confirmed_at = timezone.now()
    with authority_write(STAGE_RESULT_CONFIRM):
        locked.save(update_fields=["status", "confirmed_by", "confirmed_at"])
    AuditLog.objects.create(
        operator=current_actor,
        action_type=AuditLog.ActionType.CONFIRM_STAGE_RESULT,
        target=f"StageResult:{locked.pk}",
        old_value=StageResult.Status.READY_TO_CONFIRM,
        new_value=json.dumps(
            {
                "status": StageResult.Status.CONFIRMED,
                "stage_key": locked.stage_key,
                "result_version": locked.result_version,
                "ruleset_version": version.pk,
                "authority_hash": version.authority_hash,
                "input_fingerprint": locked.input_fingerprint,
                "confirmed_by": current_actor.pk,
                "plan_version": locked.plan_version,
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
    )
    materialize_stage_awards(locked, operator=current_actor)
    materialize_round_entry_from_stage(locked, operator=current_actor)
    return locked


def _release_stage_consumed_facts(stage: StageResult, *, operator) -> None:
    """Release raw facts only after their consuming stage leaves CONFIRMED.

    This keeps the correction path explicit: a score/vote cannot be edited merely because
    a caller bypassed the normal service, but an audited stage unlock re-opens its inputs.
    Facts shared by another confirmed stage remain protected.
    """
    from voting.models import VoteSession

    consumed = _stage_consumed_facts(stage)
    confirmed = StageResult.objects.filter(
        activity=stage.activity, status=StageResult.Status.CONFIRMED
    ).exclude(pk=stage.pk)
    for other in confirmed:
        other_facts = _stage_consumed_facts(other)
        if consumed["rounds"] & other_facts["rounds"]:
            raise PermissionDenied("该轮次仍被其他已核定赛段使用，不能解锁原始评分。")
        if consumed["votes"] & other_facts["votes"]:
            raise PermissionDenied("该投票仍被其他已核定赛段使用，不能解锁原始记录。")
        if consumed["audience"] & other_facts["audience"]:
            raise PermissionDenied("该观众分仍被其他已核定赛段使用，不能解锁原始记录。")

    for contest_round in ContestRound.objects.select_for_update().filter(
        activity=stage.activity, pk__in=consumed["rounds"], is_locked=True
    ):
        contest_round.is_locked = False
        contest_round.status = ContestRound.Status.SCORING
        with authority_write(CONTEST_ROUND_STATE):
            contest_round.save(update_fields=["is_locked", "status"])
        AuditLog.objects.create(
            operator=operator,
            action_type=AuditLog.ActionType.UNLOCK_RESULT,
            target=f"ContestRound:{contest_round.pk}",
            old_value="locked",
            new_value="unlocked",
            note=f"stage_unlock:{stage.pk}",
        )
    for vote_session in VoteSession.objects.select_for_update().filter(
        activity=stage.activity, pk__in=consumed["votes"], is_locked=True
    ):
        vote_session.is_locked = False
        with authority_write(VOTE_SESSION_STATE):
            vote_session.save(update_fields=["is_locked"])
        AuditLog.objects.create(
            operator=operator,
            action_type=AuditLog.ActionType.UNLOCK_RESULT,
            target=f"VoteSession:{vote_session.pk}",
            old_value="locked",
            new_value="unlocked",
            note=f"stage_unlock:{stage.pk}",
        )


@transaction.atomic
def unlock_stage_result(stage: StageResult, *, operator, note: str = "") -> StageResult:
    """Admin unlock a CONFIRMED stage result so its raw facts can be corrected (§38).

    Reverting to ``READY_TO_CONFIRM`` releases the finality that freezing the result
    imposed: the round / vote / manual edits that the result consumed are editable again,
    and the stage can be re-resolved and re-confirmed. Audits the change.
    """
    current_operator = require_current_admin(operator)
    lock_activity_for_action(stage.activity)
    locked = StageResult.objects.select_for_update().get(pk=stage.pk)
    if locked.status != StageResult.Status.CONFIRMED:
        raise PermissionDenied("仅可解锁已核定并锁定的赛段结果。")
    downstream_started = ContestRound.objects.filter(
        activity=locked.activity,
        roster_source=ContestRound.RosterSource.STAGE,
        roster_source_stage=locked.stage_key,
        status__in=[
            ContestRound.Status.PREPARED,
            ContestRound.Status.SCORING,
            ContestRound.Status.LOCKED,
        ],
    ).exists()
    if downstream_started:
        raise PermissionDenied("已存在开始进行的下游轮次，禁止解锁上游赛段。")
    locked.confirmed_by = None
    locked.confirmed_at = None
    locked.status = StageResult.Status.READY_TO_CONFIRM
    with authority_write(STAGE_RESULT_CONFIRM):
        locked.save(update_fields=["confirmed_by", "confirmed_at", "status"])
    _release_stage_consumed_facts(locked, operator=current_operator)
    AuditLog.objects.create(
        operator=current_operator,
        action_type=AuditLog.ActionType.UNLOCK_STAGE_RESULT,
        target=f"StageResult:{locked.pk}",
        old_value=StageResult.Status.CONFIRMED,
        new_value=StageResult.Status.READY_TO_CONFIRM,
        note=note,
    )
    return locked


@transaction.atomic
def run_ruleset(
    version,
    activity,
    *,
    stage_key,
    computed_by,
    round_keys: Mapping[str, ContestRound],
    vote_scores=None,
    group_of=None,
    manual=None,
    duel_decisions=None,
    checkpoint=None,
    preview: bool = False,
) -> "StageResult | ResolveResult":
    """Single generic entry: bind -> resolve -> persist. Returns the StageResult.

    For a non-preview (formal) call this is the persisted :class:`StageResult`; for a
    ``preview`` it is the in-memory :class:`ResolveResult` and nothing is persisted.

    ``checkpoint`` (§11-15) is an optional named boundary in the definition. When given,
    only the checkpoint's dependency closure is resolved (future-stage inputs no longer
    force a HOLD) and ``stage_key`` should be the checkpoint's key, so each stage persists
    as its own progressive :class:`StageResult`.

    ``preview`` (M1-R9 §七) marks the call as a staff comparison, not a FORMAL
    publication: the formal path (default, used by :func:`recompute_activity_result`) only
    executes the *current* FROZEN authority and persists a :class:`StageResult`; a preview
    may run a historical (superseded) frozen version side-by-side for staff to compare, but
    it is ISOLATED — it returns the in-memory :class:`ResolveResult` and never writes a
    StageResult, so a comparison can never pollute the result board.
    """
    from ruleset.models import RulesetVersion

    if version.ruleset.activity_id != activity.pk:
        raise ValidationError("Ruleset version must belong to the result activity.")
    if version.status != RulesetVersion.Status.FROZEN:
        raise ValidationError("Only a frozen ruleset version may be executed.")
    if not version.is_current and not preview:
        raise ValidationError("Only the current frozen ruleset version produces a formal result.")
    # A stage_key that names a declared checkpoint must resolve *as* that checkpoint:
    # confirm_stage_result re-derives the projection purely from stage_key, so the
    # persisted input_fingerprint has to be computed on the same (scoped) roster.
    if checkpoint is None and _stage_is_checkpoint(version, stage_key):
        checkpoint = stage_key
    if duel_decisions is None:
        duel_decisions = _source_duel_decisions(version, activity)
    inputs = bind_resolve_input(
        version,
        activity,
        round_keys=round_keys,
        vote_scores=vote_scores,
        group_of=group_of,
        manual=manual,
        duel_decisions=duel_decisions,
        checkpoint=checkpoint,
    )
    plan = _plan_from_version(version)
    if checkpoint:
        result = resolve_to_checkpoint(version.definition, inputs, checkpoint, plan=plan)
    else:
        result = resolve(version.definition, inputs, plan=plan)
    if preview:
        return result
    return persist_stage_result(
        version, activity, result, stage_key=stage_key, computed_by=computed_by
    )


def _current_frozen_version(activity):
    """The activity's single current frozen ruleset version, or ``None``.

    A benign "not set up yet" state (no ruleset, or no current frozen version) returns
    ``None`` so the auto-resolve trigger can skip silently; real ambiguity (more than one
    auto ruleset) still fails loudly, mirroring :func:`recompute_activity_result`.
    """
    from ruleset.models import ContestRuleset, RulesetVersion

    candidates = ContestRuleset.objects.filter(activity=activity, stage_key__gt="").order_by("pk")
    if candidates.count() > 1:
        raise ValidationError("该活动存在多个自动重算赛制，请先清理重复。")
    ruleset = candidates.first()
    if ruleset is None:
        return None
    return (
        ruleset.versions.filter(is_current=True, status=RulesetVersion.Status.FROZEN)
        .order_by("-version")
        .first()
    )


def _definition_checkpoints_ordered(version) -> tuple[str, ...]:
    """Declared checkpoint keys in declaration order, or ``()`` for a non-progressive ruleset."""
    if not version.definition:
        return ()
    obj = (
        version.definition
        if isinstance(version.definition, dict)
        else json.loads(version.definition)
    )
    return tuple(c["key"] for c in (obj.get("checkpoints") or ()) if c.get("key"))


def official_stage_award_queryset(activity: Activity | None = None) -> QuerySet[Award]:
    """Return manual Awards plus only current confirmed stage-sourced Awards."""
    from ruleset.models import RulesetVersion

    latest_version = (
        StageResult.objects.filter(
            activity_id=OuterRef("source_stage_result__activity_id"),
            stage_key=OuterRef("source_stage_result__stage_key"),
        )
        .order_by("-result_version", "-pk")
        .values("result_version")[:1]
    )
    stage_awards = Award.objects.filter(
        source_stage_result__isnull=False,
        source_stage_result__activity_id=F("activity_id"),
        source_stage_result__status=StageResult.Status.CONFIRMED,
        source_stage_result__result_version=Subquery(latest_version),
        source_stage_result__ruleset_version__is_current=True,
        source_stage_result__ruleset_version__status=RulesetVersion.Status.FROZEN,
        source_award_decision__activity_id=F("activity_id"),
        source_award_decision__stage_result_id=F("source_stage_result_id"),
    )
    base = Award.objects.all() if activity is None else Award.objects.filter(activity=activity)
    return base.filter(
        Q(source_stage_result__isnull=True) | Q(pk__in=stage_awards)
    ).select_related("singer", "source_stage_result", "source_award_decision")


def _coerce_bound_ids(values: Iterable[object]) -> set[int]:
    ids: set[int] = set()
    for value in values:
        if value is None:
            continue
        text = str(value)
        if not text.isdigit():
            raise ValueError("规则绑定包含无法识别的原始数据 ID。")
        ids.add(int(text))
    return ids


def _closure_required_raw_facts(version, activity, stage_key: str) -> dict[str, tuple[int, ...]]:
    """Resolve the physical round/vote IDs consumed by one stage."""
    checkpoint = stage_key if _stage_is_checkpoint(version, stage_key) else None
    round_keys, vote_keys, group_keys, _manual_keys = stage_consumed_scopes(
        version.definition, checkpoint
    )
    binding = _version_binding(version)
    round_map = binding.get("round_keys") or {}
    group_map = binding.get("group_keys") or {}
    vote_map = binding.get("vote_keys") or {}
    audience_map = binding.get("audience_keys") or {}
    round_ids = _coerce_bound_ids(
        [round_map.get(key) for key in round_keys]
        + [group_map.get(key) for key in group_keys]
    )
    audience_keys = {key for key in vote_keys if key in audience_map}
    vote_ids = _coerce_bound_ids([vote_map.get(key) for key in vote_keys - audience_keys])
    round_qs = ContestRound.objects.filter(activity=activity, pk__in=round_ids)
    vote_qs = VoteSession.objects.filter(activity=activity, pk__in=vote_ids)
    if round_qs.count() != len(round_ids) or vote_qs.count() != len(vote_ids):
        raise ValueError("规则绑定的原始轮次或投票不存在。")
    return {
        "rounds": tuple(sorted(round_ids)),
        "votes": tuple(sorted(vote_ids)),
    }


def _closure_upstream_stage_keys(
    version, activity, required_round_ids: tuple[int, ...]
) -> tuple[str, ...]:
    return tuple(
        ContestRound.objects.filter(
            activity=activity,
            pk__in=required_round_ids,
            roster_source=ContestRound.RosterSource.STAGE,
        )
        .exclude(roster_source_stage="")
        .values_list("roster_source_stage", flat=True)
        .distinct()
    )


def _closure_stage(
    activity,
    version,
    stage_key: str,
    latest: StageResult | None,
) -> StageClosure:
    blockers: list[ResultClosureCode] = []
    if activity.is_locked:
        blockers.append(ResultClosureCode.ACTIVITY_OPERATIONALLY_LOCKED)
    reasons: tuple[str, ...] = tuple(latest.reasons or ()) if latest else ()
    required: dict[str, tuple[int, ...]] = {}
    raw_facts_locked = True
    upstream_confirmed = True
    try:
        required = _closure_required_raw_facts(version, activity, stage_key)
    except (KeyError, TypeError, ValueError, ValidationError):
        blockers.append(ResultClosureCode.RULESET_BINDING_INVALID)

    if required:
        round_ids = required.get("rounds", ())
        vote_ids = required.get("votes", ())
        rounds = ContestRound.objects.filter(activity=activity, pk__in=round_ids)
        votes = VoteSession.objects.filter(activity=activity, pk__in=vote_ids)
        if rounds.count() != len(round_ids) or votes.count() != len(vote_ids):
            blockers.append(ResultClosureCode.RAW_FACTS_INCOMPLETE)
        raw_facts_locked = not rounds.exclude(is_locked=True).exists() and not votes.exclude(
            is_locked=True
        ).exists()
        if not raw_facts_locked:
            blockers.append(ResultClosureCode.RAW_FACTS_UNLOCKED)
        upstream_keys = _closure_upstream_stage_keys(version, activity, round_ids)
        if upstream_keys:
            upstream_confirmed = all(
                StageResult.objects.filter(
                    activity=activity,
                    stage_key=upstream_key,
                    status=StageResult.Status.CONFIRMED,
                )
                .order_by("-result_version", "-pk")
                .first()
                is not None
                for upstream_key in upstream_keys
            )
            if not upstream_confirmed:
                blockers.append(ResultClosureCode.UPSTREAM_CONFIRMATION_PENDING)

    if latest is None:
        blockers.append(ResultClosureCode.RAW_FACTS_INCOMPLETE)
    else:
        if latest.ruleset_version_id != version.pk or latest.ruleset_hash != version.authority_hash:
            blockers.append(ResultClosureCode.STALE_CANDIDATE)
        else:
            try:
                if latest.input_fingerprint != _current_input_fingerprint(
                    version, activity, stage_key
                ):
                    blockers.append(ResultClosureCode.STALE_CANDIDATE)
            except (KeyError, TypeError, ValueError, ValidationError):
                blockers.append(ResultClosureCode.RULESET_BINDING_INVALID)
        if latest.status == StageResult.Status.HOLD:
            blockers.append(ResultClosureCode.RAW_FACTS_INCOMPLETE)
        elif latest.status == StageResult.Status.REVIEW:
            blockers.append(ResultClosureCode.RULE_REVIEW_REQUIRED)
        elif latest.status not in {
            StageResult.Status.READY_TO_CONFIRM,
            StageResult.Status.CONFIRMED,
        }:
            blockers.append(ResultClosureCode.RULESET_BINDING_INVALID)
        if latest.status == StageResult.Status.READY_TO_CONFIRM:
            blockers.append(ResultClosureCode.STAGE_CONFIRMATION_PENDING)

    ordered_blockers = tuple(dict.fromkeys(blockers))
    confirmable = (
        latest is not None
        and latest.status == StageResult.Status.READY_TO_CONFIRM
        and not set(ordered_blockers) - {ResultClosureCode.STAGE_CONFIRMATION_PENDING}
    )
    official_awards = 0
    round_entries = 0
    if latest is not None and latest.status == StageResult.Status.CONFIRMED:
        official_awards = official_stage_award_queryset(activity).filter(
            source_stage_result=latest
        ).count()
        round_entries = RoundEntry.objects.filter(
            round__activity=activity,
            round__roster_source=ContestRound.RosterSource.STAGE,
            round__roster_source_stage=stage_key,
        ).count()
    return StageClosure(
        stage_key=stage_key,
        current_result_id=latest.pk if latest else None,
        result_version=latest.result_version if latest else None,
        status=latest.status if latest else None,
        confirmable=confirmable,
        reasons=reasons,
        input_fingerprint=latest.input_fingerprint if latest else None,
        required_raw_facts=required,
        raw_facts_locked=raw_facts_locked,
        upstream_confirmed=upstream_confirmed,
        official_awards=official_awards,
        round_entries=round_entries,
        blocking_reasons=ordered_blockers,
    )


def build_result_closure(activity, *, stage_key: str | None = None) -> ResultClosure:
    """Build a read-only, activity-scoped result closure report."""
    version = _current_frozen_version(activity)
    if version is None:
        return ResultClosure(
            activity_id=activity.pk,
            ruleset_version_id=None,
            ruleset_authority_hash=None,
            stages=(),
            closeable=False,
            blocking_reasons=(ResultClosureCode.NO_CURRENT_FROZEN_RULESET,),
        )
    blockers: list[ResultClosureCode] = []
    if activity.is_locked:
        blockers.append(ResultClosureCode.ACTIVITY_OPERATIONALLY_LOCKED)
    try:
        stage_keys = _definition_checkpoints_ordered(version)
        if not stage_keys:
            stage_keys = ((_version_binding(version).get("stage_key") or "").strip(),)
        stage_keys = tuple(key for key in stage_keys if key)
        if not stage_keys:
            return ResultClosure(
                activity_id=activity.pk,
                ruleset_version_id=version.pk,
                ruleset_authority_hash=version.authority_hash or None,
                stages=(),
                closeable=False,
                blocking_reasons=(ResultClosureCode.RULESET_BINDING_INVALID,),
            )
        if stage_key is not None:
            if stage_key not in stage_keys:
                return ResultClosure(
                    activity_id=activity.pk,
                    ruleset_version_id=version.pk,
                    ruleset_authority_hash=version.authority_hash or None,
                    stages=(),
                    closeable=False,
                    blocking_reasons=(ResultClosureCode.RULESET_BINDING_INVALID,),
                )
            stage_keys = (stage_key,)
    except (KeyError, TypeError, ValueError, ValidationError):
        return ResultClosure(
            activity_id=activity.pk,
            ruleset_version_id=version.pk,
            ruleset_authority_hash=version.authority_hash or None,
            stages=(),
            closeable=False,
            blocking_reasons=(ResultClosureCode.RULESET_BINDING_INVALID,),
        )
    stages: list[StageClosure] = []
    for key in stage_keys:
        latest = (
            StageResult.objects.filter(activity=activity, stage_key=key)
            .order_by("-result_version", "-pk")
            .first()
        )
        stages.append(_closure_stage(activity, version, key, latest))
    for stage in stages:
        blockers.extend(stage.blocking_reasons)
    ordered_blockers = tuple(dict.fromkeys(blockers))
    return ResultClosure(
        activity_id=activity.pk,
        ruleset_version_id=version.pk,
        ruleset_authority_hash=version.authority_hash or None,
        stages=tuple(stages),
        closeable=bool(stages) and not ordered_blockers,
        blocking_reasons=ordered_blockers,
    )


def _run_ruleset_args(version, activity) -> dict:
    """The bound rounds + sources a formal/preview run of the frozen version consumes."""
    binding = _version_binding(version)
    raw_keys = binding.get("round_keys") or {}
    round_ids = list(raw_keys.values())
    rounds = {r.pk: r for r in ContestRound.objects.filter(pk__in=round_ids, activity=activity)}
    missing = [rid for rid in round_ids if rid not in rounds]
    if missing:
        raise ValidationError(f"赛制绑定的比赛轮次不存在：{missing}")
    bound = {key: rounds[rid] for key, rid in raw_keys.items()}
    return {
        "round_keys": bound,
        "vote_scores": _combined_vote_scores(activity, binding),
        "group_of": _source_group_of(activity, binding),
        "manual": _source_manual(version, activity),
        "duel_decisions": _source_duel_decisions(version, activity),
    }


def _current_resolve_status(activity) -> str | None:
    """The status of the activity's most-recent target stage result, or ``None``.

    The frontend shows this label after an input batch, so it must reflect the true current
    state — not merely whether the just-run batch resolved something new. A no-op re-save of
    an already-READY stage therefore keeps ``ready_to_confirm`` instead of reading as
    "not computed".
    """
    try:
        version = _current_frozen_version(activity)
    except ValidationError:
        # A read-only status label must not surface an authority ambiguity as a 500.
        return None
    if version is None:
        return None
    checkpoints = _definition_checkpoints_ordered(version)
    if checkpoints:
        target_keys = list(checkpoints)
    else:
        full_stage = (_version_binding(version).get("stage_key") or "").strip()
        target_keys = [full_stage] if full_stage else []
    target_keys = [k for k in target_keys if k]
    if not target_keys:
        return None
    latest = (
        StageResult.objects.filter(activity=activity, stage_key__in=target_keys)
        .order_by("-result_version", "-pk")
        .first()
    )
    return latest.status if latest else None


@transaction.atomic
def maybe_resolve_checkpoints(activity, operator) -> list[str]:
    """Publish a READY_TO_CONFIRM StageResult for every satisfiable stage.

    Called after each onsite input batch (score entry, audience entry, vote/round lock).
    It probes — without persisting — whether a stage's dependency closure is now complete.
    Targets are the declared checkpoints in definition order; a ruleset with no declared
    checkpoints falls back to its single binding ``stage_key`` (a full resolve). A probe
    that resolves READY and whose current facts differ from the latest published result is
    published via :func:`recompute_activity_result`; a HOLD/REVIEW (still-missing input) or
    an already-current result is skipped so an incomplete stage never pollutes the board and
    a re-run never duplicates.
    """
    current_operator = require_current_staff(operator)
    lock_activity_for_action(activity)
    version = _current_frozen_version(activity)
    if version is None:
        return []
    checkpoints = _definition_checkpoints_ordered(version)
    targets: list[tuple[str, str | None]] = []
    if checkpoints:
        targets = [(cp, cp) for cp in checkpoints]
    else:
        full_stage = (_version_binding(version).get("stage_key") or "").strip()
        targets = [(full_stage, None)] if full_stage else []
    resolved: list[str] = []
    for stage_key, checkpoint in targets:
        args = _run_ruleset_args(version, activity)
        try:
            probe = run_ruleset(
                version,
                activity,
                stage_key=stage_key,
                computed_by=current_operator,
                checkpoint=checkpoint,
                preview=True,
                **args,
            )
        except StageRosterNotMaterialized:
            # The stage's consumed STAGE/LEGACY roster round has not been materialized
            # because its upstream stage is still unresolved. That is "not ready yet",
            # not a config error — skip and let a later input batch publish it.
            continue
        if probe.status != ResolverState.READY:
            continue
        latest = (
            StageResult.objects.filter(activity=activity, stage_key=stage_key)
            .order_by("-result_version", "-pk")
            .first()
        )
        if latest is not None and latest.input_fingerprint == probe.input_fingerprint:
            continue
        recompute_activity_result(activity, current_operator, checkpoint=checkpoint)
        resolved.append(stage_key)
    return resolved


@transaction.atomic
def recompute_activity_result(
    activity,
    computed_by,
    *,
    ruleset=None,
    checkpoint=None,
) -> StageResult:
    """Re-resolve an activity's current ruleset against its bound rounds (M1-H).

    The ruleset carries the production binding: ``stage_key`` and ``round_keys``
    (round key -> ContestRound pk). The frozen version is authoritative: the round
    mapping is always read from the version's binding snapshot, never overridden.
    ``checkpoint`` (§11-15) targets a single declared stage and publishes it as its
    own :class:`StageResult` (stage_key = the checkpoint key). Returns the persisted
    :class:`StageResult`.

    M1-R9 (§二): this is the AUTHORITATIVE publication path. It runs inside one
    ``transaction.atomic`` block holding ``Activity FOR UPDATE`` (via
    :func:`lock_activity_for_action`), so every raw fact (round scores, vote counts,
    group map, manual picks) is read in a single serialization window and the StageResult
    is persisted within it. Two concurrent recomputes of one activity are therefore
    linearized — ``result_version`` cannot be minted twice.
    """
    from ruleset.models import ContestRuleset, RulesetVersion

    # Activity FOR UPDATE is the serialization point for every formal recompute. It also
    # re-validates that the activity is not operationally locked, exactly as the other
    # formal result mutations do (confirm/unlock).
    lock_activity_for_action(activity)

    if ruleset is None:
        candidates = ContestRuleset.objects.filter(activity=activity, stage_key__gt="").order_by(
            "pk"
        )
        # §16/§32: one activity owns exactly one contest ruleset; versions hold the history.
        # A stray duplicate is a data error, so fail loudly instead of silently picking the
        # newest and revealing an ambiguous authority.
        if candidates.count() > 1:
            raise ValidationError("该活动存在多个自动重算赛制，请先清理重复。")
        ruleset = candidates.first()
    if ruleset is None:
        raise ValidationError("该活动尚未绑定可自动重算的赛制。")
    version = (
        ruleset.versions.filter(is_current=True, status=RulesetVersion.Status.FROZEN)
        .order_by("-version")
        .first()
    )
    if version is None:
        raise ValidationError("该赛制没有可用的当前冻结版本。")
    # The frozen version is authoritative: read its binding snapshot (never the still
    # mutable ContestRuleset rounding fields, never an injected override).
    binding = _version_binding(version)
    raw_keys = binding.get("round_keys") or {}
    round_ids = list(raw_keys.values())
    rounds = {r.pk: r for r in ContestRound.objects.filter(pk__in=round_ids, activity=activity)}
    missing = [rid for rid in round_ids if rid not in rounds]
    if missing:
        raise ValidationError(f"赛制绑定的比赛轮次不存在：{missing}")
    bound = {key: rounds[rid] for key, rid in raw_keys.items()}
    stage_key = checkpoint or binding.get("stage_key") or ruleset.stage_key or ""
    result = run_ruleset(
        version,
        activity,
        stage_key=stage_key,
        computed_by=computed_by,
        round_keys=bound,
        vote_scores=_combined_vote_scores(activity, binding),
        group_of=_source_group_of(activity, binding),
        manual=_source_manual(version, activity),
        duel_decisions=_source_duel_decisions(version, activity),
        checkpoint=checkpoint,
    )
    assert isinstance(result, StageResult)
    return result


_OUTCOME_LABELS = {
    "direct": "直接晋级",
    "advanced": "晋级",
    "repechage": "复活赛",
    "pending": "待定",
    "eliminated": "淘汰",
    "wildcard": "外卡",
    "finalist": "最终晋级",
}


def stage_decisions_by_blocks(stage_result: StageResult) -> list[dict]:
    """Group a stage result's decisions into handcard announcement blocks.

    Uses the ruleset's ``announcement_blocks`` config (list of ``{label,
    outcome_codes}``) when present; otherwise falls back to grouping by
    ``outcome_code``. Decisions are ordered by ``rank`` then pk within a block
    so a staff member can read the result card top-to-bottom.
    """
    decisions = list(stage_result.decisions.select_related("singer").order_by("rank", "pk"))
    version = stage_result.ruleset_version
    binding = _version_binding(version) if version else {}
    # A checkpoint stage prefers its own blocks; a global set is the legacy fallback.
    blocks_config = (binding.get("announcement_blocks") or []) if binding else []
    if binding and _stage_is_checkpoint(version, stage_result.stage_key):
        blocks_config = (binding.get("announcement_blocks_by_checkpoint") or {}).get(
            stage_result.stage_key
        ) or blocks_config
    by_code: dict[str, list[StageDecision]] = {}
    for decision in decisions:
        by_code.setdefault(decision.outcome_code, []).append(decision)

    if blocks_config:
        blocks = []
        consumed: set[int] = set()
        for block in blocks_config:
            label = block.get("label") or "公告"
            members = []
            for code in block.get("outcome_codes") or []:
                for decision in by_code.get(code, []):
                    if decision.pk not in consumed:
                        members.append(decision)
                        consumed.add(decision.pk)
            if members:
                blocks.append({"label": label, "decisions": members})
        leftover = [d for d in decisions if d.pk not in consumed]
        if leftover:
            blocks.append({"label": "其他", "decisions": leftover})
        return blocks

    ordered_codes: list[str] = []
    for decision in decisions:
        if decision.outcome_code not in ordered_codes:
            ordered_codes.append(decision.outcome_code)
    return [
        {"label": _OUTCOME_LABELS.get(code, code), "decisions": by_code[code]}
        for code in ordered_codes
    ]
