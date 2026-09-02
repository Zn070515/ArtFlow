import json
import threading
import time
from contextlib import contextmanager
from decimal import Decimal
from io import BytesIO
from unittest import skipUnless

from accounts.models import User
from common.authority import RULESET_FREEZE, STAGE_RESULT_CONFIRM, authority_write
from common.models import AuditLog
from common.test_data import clear_activity_test_data
from core.models import Activity
from core.policies import ActivityAction
from core.services import lock_activity_for_action
from django.contrib import admin
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import (
    IntegrityError,
    OperationalError,
    close_old_connections,
    connection,
    transaction,
)
from django.db.models.deletion import ProtectedError
from django.test import RequestFactory, TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from exports.services import build_score_template_workbook
from files.models import MaterialCheck, MaterialRequirement, SubmissionFile
from files.services import review_material_check, store_submission_file
from openpyxl import Workbook
from ruleset.models import ContestRuleset, RulesetVersion
from staff_panel.views import activity_material_requirements
from voting.models import VoteSession

from .admin import ContestRoundAdmin, RoundEntryAdmin, RoundJudgeAdmin
from .models import (
    AudienceScore,
    CompositeResult,
    ContestRound,
    Judge,
    RoundEntry,
    RoundJudge,
    ScoreRecord,
    ScoreSummary,
    SingerRegistration,
    StageDecision,
    StageResult,
)
from .services import (
    apply_scores,
    downstream_rounds,
    ensure_round_final_for_advancement,
    expected_score_cells,
    finalize_advancement,
    lock_round,
    missing_score_cells,
    parse_score_workbook,
    prepare_round,
    recalculate_round,
    reset_round_snapshots,
    reset_round_to_draft,
    reset_test_round_snapshots,
    unlock_round,
    validate_score,
)
from .views import my_registration_detail


@contextmanager
def _manual_write_ctx():
    """Enable the ManualDecision production write-guard while a test seeds a row directly.

    R9-3 makes a bare ``ManualDecision.save()`` a production error; test setup (and the
    model-level clean/unique_together tests) need to create rows directly, so this toggles
    the thread-local guard for the duration and restores the prior state on exit.
    """
    from .models import _authorize_manual_write, _manual_write_authorized

    prior = _manual_write_authorized()
    _authorize_manual_write(True)
    try:
        yield
    finally:
        _authorize_manual_write(prior)


def _promote_version_to_current(version):
    """Demote any prior current RulesetVersion, then mark ``version`` the single current.

    核定 (R9-5 §五) finalizes only a result grounded in the *current* FROZEN authority, so a
    test that confirms must first make its frozen version the one current runner. Uses
    ``_base_manager`` (same as the production freeze demotion) to bypass the frozen
    immutability guard when demoting a prior current version.
    """
    RulesetVersion._base_manager.filter(ruleset_id=version.ruleset_id).exclude(
        pk=version.pk
    ).update(is_current=False)
    RulesetVersion._base_manager.filter(pk=version.pk).update(is_current=True)
    # The queryset update above leaves the in-memory object stale; run_ruleset reads
    # ``version.is_current`` directly, so refresh it or the is_current gate misfires.
    version.refresh_from_db()


class ScoringServiceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="singer", password="pass")
        self.activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=self.user,
            name="Singer",
            student_id="20260001",
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        self.judge = Judge.objects.create(activity=self.activity, name="Judge")

    def make_singer(self, *, student_id, activity=None, name=None, user=None):
        user = user or User.objects.create_user(username=f"singer-{student_id}", password="pass")
        return SingerRegistration.objects.create(
            activity=activity or self.activity,
            user=user,
            name=name or f"Singer {student_id}",
            student_id=student_id,
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )

    def test_round_entry_rejects_singer_from_another_activity(self):
        other_activity = Activity.objects.create(
            title="Other", activity_type=Activity.Type.SINGER_CONTEST
        )
        foreign = self.make_singer(activity=other_activity, student_id="foreign")

        with self.assertRaises(ValidationError):
            RoundEntry.objects.create(round=self.round, singer=foreign)

    def test_round_judge_rejects_judge_from_another_activity(self):
        other_activity = Activity.objects.create(
            title="Other", activity_type=Activity.Type.SINGER_CONTEST
        )
        foreign = Judge.objects.create(activity=other_activity, name="Foreign Judge")

        with self.assertRaises(ValidationError):
            RoundJudge.objects.create(round=self.round, judge=foreign)

    def test_round_snapshot_pairs_are_unique(self):
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)

        with self.assertRaises(IntegrityError), transaction.atomic():
            RoundEntry.objects.create(round=self.round, singer=self.singer)
        with self.assertRaises(IntegrityError), transaction.atomic():
            RoundJudge.objects.create(round=self.round, judge=self.judge)

    def test_round_admin_lists_status_lock_and_snapshot_counts(self):
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])

        round_admin = ContestRoundAdmin(ContestRound, admin.site)

        self.assertEqual(
            round_admin.list_display,
            [
                "name",
                "activity",
                "round_type",
                "scoring_mode",
                "status",
                "is_locked",
                "entry_count",
                "judge_count",
            ],
        )
        self.assertEqual(round_admin.entry_count(self.round), 1)
        self.assertEqual(round_admin.judge_count(self.round), 1)
        request = RequestFactory().get("/admin/")
        request.user = self.user
        form = round_admin.get_form(request, self.round)
        self.assertNotIn("status", form.base_fields)
        self.assertNotIn("is_locked", form.base_fields)

    def test_test_round_reset_requires_explicit_opt_in(self):
        self.singer.is_test_data = True
        self.singer.save(update_fields=["is_test_data"])
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save(update_fields=["status"])

        with self.assertRaisesMessage(
            ValidationError, "Test round reset requires explicit test-only opt-in."
        ):
            reset_test_round_snapshots(
                self.round,
                self.user,
                owned_singer_ids={self.singer.pk},
                owned_judge_ids={self.judge.pk},
            )

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.PREPARED)
        self.assertTrue(RoundEntry.objects.filter(round=self.round, singer=self.singer).exists())
        self.assertTrue(RoundJudge.objects.filter(round=self.round, judge=self.judge).exists())

    def test_prepared_round_rejects_snapshot_creates(self):
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()
        singer = self.make_singer(activity=self.activity, student_id="prepared-create")
        judge = Judge.objects.create(activity=self.activity, name="Prepared Create Judge")

        with self.assertRaises(ValidationError):
            RoundEntry.objects.create(round=self.round, singer=singer)
        with self.assertRaises(ValidationError):
            RoundJudge.objects.create(round=self.round, judge=judge)

    def test_prepared_round_rejects_snapshot_updates(self):
        entry = RoundEntry.objects.create(round=self.round, singer=self.singer)
        round_judge = RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()
        entry.singer = self.make_singer(activity=self.activity, student_id="prepared-update")
        round_judge.judge = Judge.objects.create(
            activity=self.activity, name="Prepared Update Judge"
        )

        with self.assertRaises(ValidationError):
            entry.save()
        with self.assertRaises(ValidationError):
            round_judge.save()

    def test_prepared_round_rejects_snapshot_deletes(self):
        entry = RoundEntry.objects.create(round=self.round, singer=self.singer)
        round_judge = RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()

        with self.assertRaises(ValidationError):
            entry.delete()
        with self.assertRaises(ValidationError):
            round_judge.delete()

        self.assertTrue(RoundEntry.objects.filter(pk=entry.pk).exists())
        self.assertTrue(RoundJudge.objects.filter(pk=round_judge.pk).exists())

    def test_admin_bulk_delete_rejects_prepared_snapshots(self):
        entry = RoundEntry.objects.create(round=self.round, singer=self.singer)
        round_judge = RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()
        request = RequestFactory().post("/admin/")

        with self.assertRaises(ValidationError):
            RoundEntryAdmin(RoundEntry, admin.site).delete_queryset(
                request, RoundEntry.objects.filter(pk=entry.pk)
            )
        with self.assertRaises(ValidationError):
            RoundJudgeAdmin(RoundJudge, admin.site).delete_queryset(
                request, RoundJudge.objects.filter(pk=round_judge.pk)
            )

        self.assertTrue(RoundEntry.objects.filter(pk=entry.pk).exists())
        self.assertTrue(RoundJudge.objects.filter(pk=round_judge.pk).exists())

    def test_prepared_round_rejects_snapshot_queryset_deletes(self):
        entry = RoundEntry.objects.create(round=self.round, singer=self.singer)
        round_judge = RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()

        with self.assertRaises(ValidationError):
            RoundEntry.objects.filter(pk=entry.pk).delete()
        with self.assertRaises(ValidationError):
            RoundJudge.objects.filter(pk=round_judge.pk).delete()

    def test_prepared_round_rejects_snapshot_queryset_updates(self):
        entry = RoundEntry.objects.create(round=self.round, singer=self.singer)
        round_judge = RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()
        singer = self.make_singer(activity=self.activity, student_id="prepared-qset-update")
        judge = Judge.objects.create(activity=self.activity, name="Prepared Queryset Update Judge")

        with self.assertRaises(ValidationError):
            RoundEntry.objects.filter(pk=entry.pk).update(singer_id=singer.pk)
        with self.assertRaises(ValidationError):
            RoundJudge.objects.filter(pk=round_judge.pk).update(judge_id=judge.pk)

    def test_prepared_round_rejects_snapshot_bulk_creates(self):
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()
        singer = self.make_singer(activity=self.activity, student_id="prepared-bulk-create")
        judge = Judge.objects.create(activity=self.activity, name="Prepared Bulk Create Judge")

        with self.assertRaises(ValidationError):
            RoundEntry.objects.bulk_create([RoundEntry(round=self.round, singer=singer)])
        with self.assertRaises(ValidationError):
            RoundJudge.objects.bulk_create([RoundJudge(round=self.round, judge=judge)])

    def test_prepared_snapshots_protect_parents_from_deletion(self):
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()

        with self.assertRaises(ProtectedError):
            self.round.delete()
        with self.assertRaises(ProtectedError):
            self.singer.delete()
        with self.assertRaises(ProtectedError):
            self.judge.delete()

    def test_draft_round_allows_snapshot_bulk_operations(self):
        entry = RoundEntry.objects.bulk_create([RoundEntry(round=self.round, singer=self.singer)])[
            0
        ]
        round_judge = RoundJudge.objects.bulk_create(
            [RoundJudge(round=self.round, judge=self.judge)]
        )[0]
        singer = self.make_singer(activity=self.activity, student_id="draft-qset-update")
        judge = Judge.objects.create(activity=self.activity, name="Draft Queryset Update Judge")

        self.assertEqual(
            RoundEntry.objects.filter(pk=entry.pk).update(singer_id=singer.pk),
            1,
        )
        self.assertEqual(
            RoundJudge.objects.filter(pk=round_judge.pk).update(judge_id=judge.pk),
            1,
        )
        self.assertEqual(RoundEntry.objects.filter(pk=entry.pk).delete()[0], 1)
        self.assertEqual(RoundJudge.objects.filter(pk=round_judge.pk).delete()[0], 1)

    def test_base_manager_rejects_prepared_snapshot_bulk_mutations(self):
        entry = RoundEntry.objects.create(round=self.round, singer=self.singer)
        round_judge = RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save()
        singer = self.make_singer(activity=self.activity, student_id="prepared-base-mgr")
        judge = Judge.objects.create(activity=self.activity, name="Prepared Base Manager Judge")

        with self.assertRaises(ValidationError):
            RoundEntry._base_manager.filter(pk=entry.pk).update(singer_id=singer.pk)
        with self.assertRaises(ValidationError):
            RoundJudge._base_manager.filter(pk=round_judge.pk).delete()
        with self.assertRaises(ValidationError):
            RoundEntry._base_manager.bulk_create([RoundEntry(round=self.round, singer=singer)])
        with self.assertRaises(ValidationError):
            RoundJudge._base_manager.bulk_create([RoundJudge(round=self.round, judge=judge)])

    def test_draft_snapshots_allow_parent_deletion(self):
        entry = RoundEntry.objects.create(round=self.round, singer=self.singer)
        round_judge = RoundJudge.objects.create(round=self.round, judge=self.judge)

        self.round.delete()

        self.assertFalse(RoundEntry.objects.filter(pk=entry.pk).exists())
        self.assertFalse(RoundJudge.objects.filter(pk=round_judge.pk).exists())

        singer_round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
        )
        singer = self.make_singer(activity=self.activity, student_id="draft-parent-singer")
        singer_entry = RoundEntry.objects.create(round=singer_round, singer=singer)

        singer.delete()

        self.assertFalse(RoundEntry.objects.filter(pk=singer_entry.pk).exists())

        other_activity = Activity.objects.create(
            title="Other", activity_type=Activity.Type.SINGER_CONTEST
        )
        judge_round = ContestRound.objects.create(
            activity=other_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        judge = Judge.objects.create(activity=other_activity, name="Draft Parent Judge")
        round_judge = RoundJudge.objects.create(round=judge_round, judge=judge)

        judge.delete()

        self.assertFalse(RoundJudge.objects.filter(pk=round_judge.pk).exists())

    def test_prepare_preliminary_round_snapshots_approved_non_test_singers(self):
        test_singer = self.make_singer(student_id="prepared-test")
        test_singer.is_test_data = True
        test_singer.save(update_fields=["is_test_data"])
        unapproved_singer = self.make_singer(student_id="prepared-unapproved")
        unapproved_singer.pre_status = SingerRegistration.PreStatus.SUBMITTED
        unapproved_singer.save(update_fields=["pre_status"])

        prepared_round = prepare_round(self.round, self.user)

        self.assertEqual(
            list(RoundEntry.objects.filter(round=self.round).values_list("singer_id", flat=True)),
            [self.singer.pk],
        )
        self.assertEqual(prepared_round.status, ContestRound.Status.PREPARED)
        self.assertFalse(prepared_round.is_locked)
        self.assertTrue(
            AuditLog.objects.filter(
                operator=self.user,
                action_type=AuditLog.ActionType.UPDATE_STATUS,
                target=f"ContestRound:{self.round.pk}",
            ).exists()
        )

    def _lock_scored_round(self, round, singer_count, advance_count):
        for index in range(singer_count):
            self.make_singer(student_id=f"2027{index:04d}")
        if advance_count:
            round.advance_count = advance_count
            round.save(update_fields=["advance_count"])
        prepare_round(round, self.user)
        judge_ids = list(RoundJudge.objects.filter(round=round).values_list("judge_id", flat=True))
        singer_ids = list(
            RoundEntry.objects.filter(round=round).values_list("singer_id", flat=True)
        )
        apply_scores(
            round,
            {
                (singer_id, judge_id): 95 - index
                for index, singer_id in enumerate(singer_ids)
                for judge_id in judge_ids
            },
            self.user,
        )
        round.status = ContestRound.Status.LOCKED
        round.is_locked = True
        round.save(update_fields=["status", "is_locked"])

    def _prepare_boundary_tie_round(self, advance_count=2):
        """Score a round where the last passing score ties the first failing one."""
        self.round.advance_count = advance_count
        self.round.save(update_fields=["advance_count"])
        extras = [self.make_singer(student_id=f"tie{idx}", name=f"Tie {idx}") for idx in range(3)]
        prepare_round(self.round, self.user)
        judge_id = list(
            RoundJudge.objects.filter(round=self.round).values_list("judge_id", flat=True)
        )[0]
        apply_scores(
            self.round,
            {
                (self.singer.pk, judge_id): 95,
                (extras[0].pk, judge_id): 90,
                (extras[1].pk, judge_id): 90,
                (extras[2].pk, judge_id): 80,
            },
            self.user,
        )
        return self.singer, extras[0], extras[1], extras[2]

    def test_advancement_boundary_tie_marks_round_needs_review_and_blocks_downstream(self):
        self._prepare_boundary_tie_round()
        self.round.refresh_from_db()
        self.assertEqual(self.round.advancement_status, ContestRound.AdvancementStatus.NEEDS_REVIEW)

        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])
        semifinal = ContestRound.objects.create(
            activity=self.activity, round_type=ContestRound.RoundType.SEMI_FINAL
        )
        with self.assertRaisesMessage(ValidationError, "请先人工核定晋级人选"):
            prepare_round(semifinal, self.user)

    def test_finalize_advancement_records_manual_selection_and_audits(self):
        singer, extra0, extra1, extra2 = self._prepare_boundary_tie_round()
        chosen = [singer.pk, extra0.pk]
        result = finalize_advancement(self.round, chosen, self.user)
        self.assertEqual(result.advancement_status, ContestRound.AdvancementStatus.FINALIZED)
        advanced = set(
            ScoreSummary.objects.filter(round=self.round, is_advanced=True).values_list(
                "singer_id", flat=True
            )
        )
        self.assertEqual(advanced, set(chosen))
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.FINALIZE_ADVANCEMENT,
                operator=self.user,
                target=f"ContestRound:{self.round.pk}",
            ).exists()
        )

    def test_finalize_advancement_rejects_singer_outside_round(self):
        self._prepare_boundary_tie_round()
        outsider = self.make_singer(student_id="outsider")
        with self.assertRaisesMessage(ValidationError, "晋级选手必须属于当前轮次"):
            finalize_advancement(self.round, [outsider.pk], self.user)

    def test_finalize_advancement_rejects_locked_round(self):
        self._prepare_boundary_tie_round()
        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])
        with self.assertRaises(PermissionDenied):
            finalize_advancement(self.round, [self.singer.pk], self.user)

    def test_finalize_advancement_rejects_locked_activity(self):
        self._prepare_boundary_tie_round()
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        with self.assertRaises(PermissionDenied):
            finalize_advancement(self.round, [self.singer.pk], self.user)

    def test_finalize_advancement_requires_exact_advance_count(self):
        self._prepare_boundary_tie_round()
        with self.assertRaisesMessage(
            ValidationError, "晋级核定人数必须等于该轮晋级名额（2 人）。"
        ):
            finalize_advancement(self.round, [self.singer.pk], self.user)

    def test_advancement_without_boundary_tie_stays_auto(self):
        self._lock_scored_round(self.round, singer_count=4, advance_count=2)
        self.round.refresh_from_db()
        self.assertEqual(self.round.advancement_status, ContestRound.AdvancementStatus.AUTO)

    def test_prepare_semifinal_round_uses_only_previous_advancers(self):
        self._lock_scored_round(self.round, singer_count=20, advance_count=10)

        semifinal = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
        )
        prepare_round(semifinal, self.user)

        self.assertEqual(RoundEntry.objects.filter(round=semifinal).count(), 10)

    def test_prepare_semifinal_rejects_when_upstream_not_locked(self):
        self._lock_scored_round(self.round, singer_count=5, advance_count=0)
        self.round.status = ContestRound.Status.SCORING
        self.round.is_locked = False
        self.round.save(update_fields=["status", "is_locked"])

        semifinal = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
        )
        with self.assertRaisesMessage(ValidationError, "上游轮次尚未锁定，无法生成后续轮次。"):
            prepare_round(semifinal, self.user)

    def test_ensure_round_final_rejects_unlocked_upstream(self):
        with self.assertRaisesMessage(ValidationError, "上游轮次尚未锁定，无法生成后续轮次。"):
            ensure_round_final_for_advancement(self.round)

    def test_ensure_round_final_rejects_empty_score_matrix(self):
        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])

        with self.assertRaisesMessage(ValidationError, "上游轮次没有可用的评分矩阵。"):
            ensure_round_final_for_advancement(self.round)

    def test_downstream_rounds_empty_for_semifinal(self):
        semifinal = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
        )
        self.assertEqual(downstream_rounds(semifinal).count(), 0)

    def test_downstream_rounds_finds_semifinal_for_preliminary(self):
        semifinal = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
        )
        self.assertEqual(list(downstream_rounds(self.round)), [semifinal])

    def test_prepare_round_snapshots_active_judges(self):
        second_judge = Judge.objects.create(activity=self.activity, name="Second Judge")
        Judge.objects.create(activity=self.activity, name="Inactive Judge", is_active=False)

        prepare_round(self.round, self.user)

        self.assertEqual(
            list(RoundJudge.objects.filter(round=self.round).values_list("judge_id", flat=True)),
            [self.judge.pk, second_judge.pk],
        )

    def test_prepare_round_rejects_drop_high_low_with_fewer_than_three_judges(self):
        self.round.scoring_mode = ContestRound.ScoringMode.DROP_HIGH_LOW
        self.round.save()
        Judge.objects.create(activity=self.activity, name="Second Judge")

        with self.assertRaises(ValidationError):
            prepare_round(self.round, self.user)

        self.assertEqual(self.round.status, ContestRound.Status.DRAFT)

    def test_prepare_round_allows_drop_high_low_with_three_judges(self):
        self.round.scoring_mode = ContestRound.ScoringMode.DROP_HIGH_LOW
        self.round.save()
        Judge.objects.create(activity=self.activity, name="Second Judge")
        Judge.objects.create(activity=self.activity, name="Third Judge")

        prepared = prepare_round(self.round, self.user)

        self.assertEqual(prepared.status, ContestRound.Status.PREPARED)

    def test_prepare_round_is_rejected_after_preparation(self):
        prepare_round(self.round, self.user)

        with self.assertRaises(ValidationError):
            prepare_round(self.round, self.user)

    def test_expected_cells_include_approved_singer_and_active_judge(self):
        prepare_round(self.round, self.user)
        self.assertEqual(expected_score_cells(self.round), [(self.singer.pk, self.judge.pk)])

    def test_missing_cells_report_absent_score_record(self):
        prepare_round(self.round, self.user)

        self.assertEqual(
            missing_score_cells(self.round),
            [
                {
                    "singer_id": self.singer.pk,
                    "singer_name": "Singer",
                    "judge_id": self.judge.pk,
                    "judge_name": "Judge",
                }
            ],
        )

        ScoreRecord.objects.create(round=self.round, singer=self.singer, judge=self.judge, score=90)
        self.assertEqual(missing_score_cells(self.round), [])

    def test_new_global_judge_does_not_change_prepared_matrix(self):
        prepare_round(self.round, self.user)
        new_judge = Judge.objects.create(activity=self.activity, name="Late Judge")

        self.assertEqual(expected_score_cells(self.round), [(self.singer.pk, self.judge.pk)])
        self.assertEqual(
            missing_score_cells(self.round),
            [
                {
                    "singer_id": self.singer.pk,
                    "singer_name": self.singer.name,
                    "judge_id": self.judge.pk,
                    "judge_name": self.judge.name,
                }
            ],
        )
        self.assertNotIn((self.singer.pk, new_judge.pk), expected_score_cells(self.round))

    def test_registration_status_change_does_not_change_prepared_roster(self):
        prepare_round(self.round, self.user)
        self.singer.pre_status = SingerRegistration.PreStatus.SUBMITTED
        self.singer.save(update_fields=["pre_status"])

        self.assertEqual(expected_score_cells(self.round), [(self.singer.pk, self.judge.pk)])
        self.assertEqual(
            missing_score_cells(self.round),
            [
                {
                    "singer_id": self.singer.pk,
                    "singer_name": self.singer.name,
                    "judge_id": self.judge.pk,
                    "judge_name": self.judge.name,
                }
            ],
        )

    def test_recalculate_ignores_score_from_judge_outside_snapshot(self):
        prepare_round(self.round, self.user)
        inactive_judge = Judge.objects.create(
            activity=self.activity, name="Inactive Judge", is_active=False
        )
        ScoreRecord.objects.create(round=self.round, singer=self.singer, judge=self.judge, score=90)
        ScoreRecord.objects.create(
            round=self.round, singer=self.singer, judge=inactive_judge, score=0
        )

        recalculate_round(self.round)

        self.assertEqual(
            ScoreSummary.objects.get(round=self.round, singer=self.singer).average_score,
            Decimal("90"),
        )

    def test_apply_scores_rejects_draft_round(self):
        with self.assertRaisesMessage(ValidationError, "请先准备比赛轮次"):
            apply_scores(
                self.round,
                {(self.singer.pk, self.judge.pk): "90"},
                self.user,
            )

    def test_apply_scores_transitions_prepared_round_to_scoring(self):
        prepare_round(self.round, self.user)

        apply_scores(
            self.round,
            {(self.singer.pk, self.judge.pk): "90"},
            self.user,
        )

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.SCORING)

    def test_apply_scores_keeps_prepared_round_prepared_when_submission_is_empty(self):
        prepare_round(self.round, self.user)

        apply_scores(self.round, {}, self.user)

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.PREPARED)

    def test_apply_scores_keeps_prepared_round_prepared_when_scores_are_unchanged(self):
        prepare_round(self.round, self.user)
        ScoreRecord.objects.create(
            round=self.round,
            singer=self.singer,
            judge=self.judge,
            score=90,
            is_test_data=False,
        )

        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "90"}, self.user)

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.PREPARED)

    def test_apply_scores_rejects_locked_round_status(self):
        prepare_round(self.round, self.user)
        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])

        with self.assertRaisesMessage(ValidationError, "该比赛轮次已锁定"):
            apply_scores(
                self.round,
                {(self.singer.pk, self.judge.pk): "90"},
                self.user,
            )

    def test_validate_score_rejects_invalid_values(self):
        for value in ("", "abc", "-0.01", "100.01", "NaN", "Infinity", "1.234"):
            with self.subTest(value=value):
                with self.assertRaises(ValidationError):
                    validate_score(value)

    def test_validate_score_returns_decimal_for_valid_values(self):
        self.assertEqual(validate_score("99.50"), Decimal("99.50"))

    def test_apply_scores_is_atomic_when_one_cell_is_invalid(self):
        prepare_round(self.round, self.user)

        with self.assertRaises(ValidationError):
            apply_scores(
                self.round,
                {(self.singer.pk, self.judge.pk): "91", (999999, self.judge.pk): "92"},
                self.user,
            )

        self.assertFalse(ScoreRecord.objects.exists())
        self.assertFalse(
            AuditLog.objects.filter(action_type=AuditLog.ActionType.ENTER_SCORE).exists()
        )

    def test_apply_scores_records_edit_details_and_recalculates(self):
        prepare_round(self.round, self.user)

        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "91"}, self.user)
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "92.50"}, self.user)

        self.assertEqual(
            ScoreRecord.objects.get(round=self.round, singer=self.singer, judge=self.judge).score,
            Decimal("92.50"),
        )
        audit = (
            AuditLog.objects.filter(action_type=AuditLog.ActionType.ENTER_SCORE)
            .order_by("-pk")
            .first()
        )
        assert audit is not None
        self.assertIn('"old": "91.00"', audit.new_value)
        self.assertIn('"new": "92.50"', audit.new_value)

    def test_parse_score_workbook_reports_late_invalid_cell_without_writing(self):
        second_user = User.objects.create_user(username="singer-two", password="pass")
        second_singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=second_user,
            name="Second Singer",
            student_id="20260002",
            college="College",
            class_name="Class",
            phone="13800000001",
            song_name="Song 2",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        prepare_round(self.round, self.user)
        workbook = Workbook()
        worksheet = workbook.active
        assert worksheet is not None
        worksheet.append(["选手\\评委", self.judge.name])
        worksheet.append([self.singer.name, 90])
        worksheet.append([second_singer.name, "bad"])
        upload = BytesIO()
        workbook.save(upload)
        upload.seek(0)

        scores, errors = parse_score_workbook(upload, self.round)

        self.assertEqual(scores, {(self.singer.pk, self.judge.pk): Decimal("90")})
        self.assertEqual(len(errors), 1)
        self.assertFalse(ScoreRecord.objects.exists())

    def _parse_workbook(self, workbook):
        upload = BytesIO()
        workbook.save(upload)
        upload.seek(0)
        return parse_score_workbook(upload, self.round)

    def _set_meta(self, workbook, key, value):
        for row in workbook["ArtFlowMeta"].iter_rows():
            if row and row[0].value == key:
                row[1].value = str(value)
                return workbook
        raise AssertionError(f"{key} not found in ArtFlowMeta")

    def _meta_dict(self, workbook):
        data = {}
        for row in workbook["ArtFlowMeta"].iter_rows():
            if row and len(row) >= 2 and row[0].value is not None:
                data[str(row[0].value)] = str(row[1].value) if row[1].value is not None else ""
        return data

    def test_score_template_emits_ids_and_hidden_artflow_meta(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        ws = wb.active
        assert ws is not None
        self.assertEqual(ws.cell(row=1, column=1).value, "选手ID")
        self.assertEqual(ws.cell(row=1, column=2).value, "姓名")
        self.assertEqual(ws.cell(row=1, column=3).value, f"J{self.judge.pk} {self.judge.name}")
        self.assertEqual(ws.cell(row=2, column=1).value, self.singer.pk)
        self.assertIn("ArtFlowMeta", wb.sheetnames)
        self.assertEqual(wb["ArtFlowMeta"].sheet_state, "hidden")
        self.assertEqual(
            [cell.value for cell in wb["ArtFlowMeta"][1]],
            ["schema_version", "1"],
        )

    def test_parse_score_workbook_uses_id_as_authority(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        ws = wb.active
        assert ws is not None
        ws.cell(row=2, column=3).value = 91
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(errors, [])
        self.assertEqual(scores, {(self.singer.pk, self.judge.pk): Decimal("91")})

    def test_parse_score_workbook_rejects_wrong_activity(self):
        prepare_round(self.round, self.user)
        other = Activity.objects.create(
            title="Other",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=False,
        )
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "activity_id", other.pk)
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("其他活动" in error for error in errors))

    def test_parse_score_workbook_rejects_wrong_round(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "round_id", self.round.pk + 1)
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("其他轮次" in error for error in errors))

    def test_parse_score_workbook_rejects_singer_outside_round(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        ws = wb.active
        assert ws is not None
        ws.append([999999, "Ghost", ""])
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("不属于当前轮次" in error for error in errors))

    def test_score_template_meta_embeds_snapshot_fingerprint_and_ruleset_version(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        meta = self._meta_dict(wb)
        self.assertEqual(meta["schema_version"], "1")
        self.assertEqual(meta["activity_id"], str(self.activity.pk))
        self.assertEqual(meta["round_id"], str(self.round.pk))
        self.assertEqual(meta["entry_ids"], str(self.singer.pk))
        self.assertEqual(meta["judge_ids"], str(self.judge.pk))
        self.assertIn("ruleset_version", meta)
        self.assertIn("snapshot_fingerprint", meta)
        self.assertTrue(meta["snapshot_fingerprint"])

    def test_parse_score_workbook_accepts_matching_snapshot(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        ws = wb.active
        assert ws is not None
        ws.cell(row=2, column=3).value = 88
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(errors, [])
        self.assertEqual(scores, {(self.singer.pk, self.judge.pk): Decimal("88")})

    def test_parse_score_workbook_rejects_stale_entry_snapshot(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "entry_ids", "1,999999")
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("选手名单" in error for error in errors))

    def test_parse_score_workbook_rejects_stale_judge_snapshot(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "judge_ids", "1,999999")
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("评委名单" in error for error in errors))

    def test_parse_score_workbook_rejects_tampered_fingerprint(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "snapshot_fingerprint", "deadbeef")
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("指纹不匹配" in error for error in errors))

    def test_parse_score_workbook_rejects_unsupported_schema_version(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "schema_version", "9")
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("schema_version" in error for error in errors))

    def test_parse_score_workbook_rejects_meta_missing_round_id(self):
        # §35: an ArtFlowMeta sheet without round_id used to downgrade to the legacy
        # name parser (fail-open). It must now reject instead of importing wrong rows.
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "round_id", "")
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(any("缺失必需字段" in error and "round_id" in error for error in errors))

    def test_parse_score_workbook_rejects_meta_missing_fingerprint(self):
        prepare_round(self.round, self.user)
        wb = build_score_template_workbook(self.round)
        self._set_meta(wb, "snapshot_fingerprint", "")
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {})
        self.assertTrue(
            any("缺失必需字段" in error and "snapshot_fingerprint" in error for error in errors)
        )

    def test_parse_score_workbook_legacy_path_without_meta_sheet(self):
        # A workbook with no ArtFlowMeta sheet still takes the explicit name-based path.
        prepare_round(self.round, self.user)
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["选手\\评委", self.judge.name])
        ws.append([self.singer.name, 95])
        scores, errors = self._parse_workbook(wb)
        self.assertEqual(scores, {(self.singer.pk, self.judge.pk): Decimal("95")})
        self.assertEqual(errors, [])


class SingerUploadViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="applicant", password="pass")
        self.activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )

    def test_apply_rejects_disallowed_file_without_creating_registration(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("singer_contest:apply"),
            {
                "activity_id": self.activity.pk,
                "name": "Singer",
                "student_id": "20260001",
                "college": "College",
                "class_name": "Class",
                "phone": "13800000000",
                "song_name": "Song",
                "accompaniment": SimpleUploadedFile("song.exe", b"bad"),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "文件类型")
        self.assertFalse(SingerRegistration.objects.exists())

    def test_apply_rejects_duplicate_registration_for_same_user(self):
        self.client.force_login(self.user)
        payload = {
            "activity_id": self.activity.pk,
            "name": "Singer",
            "student_id": "20260001",
            "college": "College",
            "class_name": "Class",
            "phone": "13800000000",
            "song_name": "Song",
        }

        first = self.client.post(reverse("singer_contest:apply"), payload)
        self.assertEqual(first.status_code, 302)

        second = self.client.post(reverse("singer_contest:apply"), payload)
        self.assertEqual(second.status_code, 200)
        self.assertContains(second, "请勿重复提交")
        self.assertEqual(
            SingerRegistration.objects.filter(activity=self.activity, user=self.user).count(), 1
        )

    def test_duplicate_registration_by_student_id_is_rejected_at_db_level(self):
        SingerRegistration.objects.create(
            activity=self.activity,
            user=self.user,
            name="First Singer",
            student_id="20260001",
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
        )
        other_user = User.objects.create_user(username="applicant-two", password="pass")

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                SingerRegistration.objects.create(
                    activity=self.activity,
                    user=other_user,
                    name="Second Singer",
                    student_id="20260001",
                    college="College",
                    class_name="Class",
                    phone="13800000001",
                    song_name="Song 2",
                )


class RoundResetServiceTests(TestCase):
    def setUp(self):
        self.actor = User.objects.create_user(username="reset-actor", password="pass")
        self.activity = Activity.objects.create(
            title="Test contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=True,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username="reset-singer", password="pass"),
            name="Singer",
            student_id="20260001",
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )
        self.judge = Judge.objects.create(activity=self.activity, name="Judge")

    def _prepare_round_with_scores(self):
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)
        ScoreRecord.objects.create(
            round=self.round,
            singer=self.singer,
            judge=self.judge,
            score=Decimal("90.00"),
            is_test_data=True,
        )
        ScoreSummary.objects.create(
            round=self.round,
            singer=self.singer,
            average_score=Decimal("90.000"),
            rank=1,
            is_advanced=True,
            is_test_data=True,
        )
        self.round.status = ContestRound.Status.PREPARED
        self.round.save(update_fields=["status", "is_locked"])

    def test_reset_round_snapshots_returns_test_round_to_draft(self):
        self._prepare_round_with_scores()

        reset_round_snapshots(self.round, self.actor, reason="reset test round")

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.DRAFT)
        self.assertFalse(self.round.is_locked)
        self.assertFalse(RoundEntry.objects.filter(round=self.round).exists())
        self.assertFalse(RoundJudge.objects.filter(round=self.round).exists())
        self.assertFalse(ScoreRecord.objects.filter(round=self.round).exists())
        self.assertFalse(ScoreSummary.objects.filter(round=self.round).exists())
        self.assertTrue(AuditLog.objects.filter(target=f"ContestRound:{self.round.pk}").exists())

    def test_reset_round_snapshots_rejects_formal_parents_in_test_activity(self):
        self.singer.is_test_data = False
        self.singer.save(update_fields=["is_test_data"])
        self._prepare_round_with_scores()

        with self.assertRaises(ValidationError):
            reset_round_snapshots(self.round, self.actor, reason="reset formal parent")

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.PREPARED)

    def test_reset_round_to_draft_requires_reason(self):
        self._prepare_round_with_scores()
        with self.assertRaises(ValidationError):
            reset_round_to_draft(self.round, self.actor, reason="  ")

    def test_reset_round_to_draft_rejects_downstream_active(self):
        semifinal = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
            status=ContestRound.Status.PREPARED,
        )
        self._prepare_round_with_scores()

        with self.assertRaises(ValidationError):
            reset_round_to_draft(self.round, self.actor, reason="reset preliminary")

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.PREPARED)
        self.assertEqual(semifinal.status, ContestRound.Status.PREPARED)

    def test_reset_round_to_draft_resets_prepared_round(self):
        self._prepare_round_with_scores()

        reset_round_to_draft(self.round, self.actor, reason="admin unwind")

        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.DRAFT)
        self.assertFalse(RoundEntry.objects.filter(round=self.round).exists())

    def test_reset_round_to_draft_is_no_op_on_draft_round(self):
        reset_round_to_draft(self.round, self.actor, reason="already draft")
        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.DRAFT)


class ParticipantRegistrationFlowTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="participant", password="pass")
        self.other = User.objects.create_user(username="other", password="pass")
        self.activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )

    def make_registration(self, *, user=None, student_id="20260001"):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=user or self.user,
            name=f"Singer {student_id}",
            student_id=student_id,
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
        )

    def test_my_registrations_lists_all_activities(self):
        other_activity = Activity.objects.create(
            title="Other Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
        )
        self.make_registration(student_id="20260001")
        SingerRegistration.objects.create(
            activity=other_activity,
            user=self.user,
            name="Another",
            student_id="20260002",
            college="College",
            class_name="Class",
            phone="13800000001",
            song_name="Song 2",
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("singer_contest:my_registrations"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.activity.title)
        self.assertContains(response, other_activity.title)

    def test_my_registration_detail_is_user_scoped(self):
        registration = self.make_registration()
        self.client.force_login(self.other)
        response = self.client.get(
            reverse("singer_contest:my_registration_detail", args=[registration.pk])
        )
        self.assertEqual(response.status_code, 404)

    def test_my_registration_detail_get_does_not_mutate_material_checks(self):
        registration = self.make_registration()
        MaterialCheck.objects.create(
            singer_registration=registration,
            item_name="基本信息",
            status=MaterialCheck.Status.UPLOADED,
            sort_order=0,
        )
        MaterialCheck.objects.create(
            singer_registration=registration,
            item_name="伴奏文件",
            status=MaterialCheck.Status.APPROVED,
            review_note="ok",
            sort_order=2,
        )
        self.client.force_login(self.user)
        before = list(
            registration.material_checks.order_by("pk").values_list(
                "item_name", "status", "sort_order", "review_note"
            )
        )
        response = self.client.get(
            reverse("singer_contest:my_registration_detail", args=[registration.pk])
        )
        self.assertEqual(response.status_code, 200)
        after = list(
            registration.material_checks.order_by("pk").values_list(
                "item_name", "status", "sort_order", "review_note"
            )
        )
        self.assertEqual(before, after)

    def test_archived_activity_get_produces_no_material_check_mutation(self):
        registration = self.make_registration()
        self.activity.phase = Activity.Phase.ARCHIVED
        self.activity.save(update_fields=["phase"])
        self.client.force_login(self.user)
        before = registration.material_checks.count()
        response = self.client.get(
            reverse("singer_contest:my_registration_detail", args=[registration.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(registration.material_checks.count(), before)

    def test_participant_can_edit_own_registration_during_registration_open(self):
        registration = self.make_registration()
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("singer_contest:my_registration_detail", args=[registration.pk]),
            {"phone": "13900000000", "song_name": "New Song"},
        )
        self.assertEqual(response.status_code, 302)
        registration.refresh_from_db()
        self.assertEqual(registration.phone, "13900000000")
        self.assertEqual(registration.song_name, "New Song")
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UPDATE_REGISTRATION,
                target=f"SingerRegistration:{registration.pk}",
            ).exists()
        )

    def test_participant_cannot_edit_when_phase_is_live(self):
        self.activity.phase = Activity.Phase.LIVE
        self.activity.save()
        registration = self.make_registration()
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("singer_contest:my_registration_detail", args=[registration.pk]),
            {"phone": "13900000000"},
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "不允许修改报名")
        registration.refresh_from_db()
        self.assertEqual(registration.phone, "13800000000")

    def test_participant_cannot_edit_a_locked_activity(self):
        self.activity.is_locked = True
        self.activity.save()
        registration = self.make_registration()
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("singer_contest:my_registration_detail", args=[registration.pk]),
            {"phone": "13900000000"},
        )
        self.assertEqual(response.status_code, 200)
        registration.refresh_from_db()
        self.assertEqual(registration.phone, "13800000000")

    def test_apply_rejects_locked_activity(self):
        self.activity.is_locked = True
        self.activity.save()
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("singer_contest:apply"),
            {
                "activity_id": self.activity.pk,
                "name": "Singer",
                "student_id": "20260001",
                "college": "College",
                "class_name": "Class",
                "phone": "13800000000",
                "song_name": "Song",
            },
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(SingerRegistration.objects.exists())

    def test_apply_rejects_activity_not_in_registration_open(self):
        self.activity.phase = Activity.Phase.REGISTRATION_CLOSED
        self.activity.save(update_fields=["phase"])
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("singer_contest:apply"),
            {
                "activity_id": self.activity.pk,
                "name": "Singer",
                "student_id": "20260001",
                "college": "College",
                "class_name": "Class",
                "phone": "13800000000",
                "song_name": "Song",
            },
        )
        self.assertEqual(response.status_code, 404)
        self.assertFalse(SingerRegistration.objects.exists())


class LockActivityForActionTests(TestCase):
    """The canonical activity-first lock helper re-validates state under the lock."""

    def setUp(self):
        self.activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )

    def test_returns_authoritative_activity(self):
        with transaction.atomic():
            locked = lock_activity_for_action(self.activity)
        self.assertEqual(locked.pk, self.activity.pk)
        self.assertFalse(locked.is_locked)

    def test_rejects_locked_activity(self):
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        with transaction.atomic():
            with self.assertRaisesMessage(PermissionDenied, "Activity results are locked."):
                lock_activity_for_action(self.activity)

    def test_rejects_action_not_allowed_in_phase(self):
        self.activity.phase = Activity.Phase.REGISTRATION_CLOSED
        self.activity.save(update_fields=["phase"])
        with transaction.atomic():
            with self.assertRaises(PermissionDenied):
                lock_activity_for_action(self.activity, ActivityAction.SUBMIT_REGISTRATION)


class RoundLockTOCTOUTests(TestCase):
    """M0-P guards: activity-first lock and re-validation of round mutations."""

    def setUp(self):
        self.user = User.objects.create_user(username="round-actor", password="pass")
        self.activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username="round-singer", password="pass"),
            name="Singer",
            student_id="20260001",
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        self.judge = Judge.objects.create(activity=self.activity, name="Judge")

    def prepare_matrix(self):
        prepare_round(self.round, self.user)
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): 90}, self.user)
        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])

    def test_prepare_round_rejects_locked_activity(self):
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        with self.assertRaisesMessage(PermissionDenied, "Activity results are locked."):
            prepare_round(self.round, self.user)
        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.DRAFT)

    def test_prepare_round_rejects_phase_without_score_authority(self):
        # M0-W: prepare_round must be its own authority — it re-checks the SCORE
        # action policy after acquiring the Activity lock, not before.
        self.activity.phase = Activity.Phase.DRAFT
        self.activity.save(update_fields=["phase"])
        with self.assertRaises(PermissionDenied):
            prepare_round(self.round, self.user)
        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.DRAFT)
        self.assertEqual(self.round.entries.count(), 0)
        self.assertEqual(self.round.round_judges.count(), 0)
        self.assertFalse(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UPDATE_STATUS,
                target=f"ContestRound:{self.round.pk}",
            ).exists()
        )

    def test_reset_round_to_draft_rejects_locked_activity(self):
        self.round.status = ContestRound.Status.PREPARED
        self.round.save(update_fields=["status"])
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        with self.assertRaisesMessage(PermissionDenied, "活动结果已锁定，无法重置轮次。"):
            reset_round_to_draft(self.round, self.user, reason="admin unwind")
        self.round.refresh_from_db()
        self.assertEqual(self.round.status, ContestRound.Status.PREPARED)

    def test_lock_round_rejects_already_locked_round(self):
        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])
        with self.assertRaisesMessage(PermissionDenied, "该比赛轮次已锁定。"):
            lock_round(self.round, self.user)

    def test_lock_round_freezes_complete_round_and_audits(self):
        prepare_round(self.round, self.user)
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): 90}, self.user)
        locked = lock_round(self.round, self.user)
        self.assertTrue(locked.is_locked)
        self.assertEqual(locked.status, ContestRound.Status.LOCKED)
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.RELOCK_RESULT,
                target=f"ContestRound:{self.round.pk}",
            ).exists()
        )

    def test_unlock_round_rejects_active_downstream(self):
        self.prepare_matrix()
        ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
            status=ContestRound.Status.PREPARED,
        )
        with self.assertRaisesMessage(
            PermissionDenied, "后续轮次仍在使用本轮结果，解锁前必须先清空后续轮次。"
        ):
            unlock_round(self.round, self.user)
        self.round.refresh_from_db()
        self.assertTrue(self.round.is_locked)

    def test_unlock_round_restores_scoring_state_and_audits(self):
        self.prepare_matrix()
        unlocked = unlock_round(self.round, self.user, note="admin unwind")
        self.assertFalse(unlocked.is_locked)
        self.assertEqual(unlocked.status, ContestRound.Status.SCORING)
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UNLOCK_RESULT,
                target=f"ContestRound:{self.round.pk}",
            ).exists()
        )


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class ActivityFirstLockConcurrencyTests(TransactionTestCase):
    """Concurrent M0-P races must never produce a state the lock order forbids."""

    def setUp(self):
        self.user = User.objects.create_user(username="concurrency-actor", password="pass")
        self.activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username="concurrency-singer", password="pass"),
            name="Singer",
            student_id="20260001",
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        self.judge = Judge.objects.create(activity=self.activity, name="Judge")

    def _lock_round(self):
        prepare_round(self.round, self.user)
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): 90}, self.user)
        self.round.status = ContestRound.Status.LOCKED
        self.round.is_locked = True
        self.round.save(update_fields=["status", "is_locked"])

    def test_concurrent_unlock_and_downstream_prepare_never_leave_orphan(self):
        self._lock_round()
        semifinal = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
        )
        results: dict[str, str] = {}

        def do_unlock():
            close_old_connections()
            try:
                unlock_round(self.round, self.user)
                results["unlock"] = "ok"
            except (PermissionDenied, ValidationError):
                results["unlock"] = "rejected"
            except Exception:
                results["unlock"] = "error"
            finally:
                close_old_connections()

        def do_prepare():
            close_old_connections()
            try:
                prepare_round(semifinal, self.user)
                results["prepare"] = "ok"
            except (PermissionDenied, ValidationError):
                results["prepare"] = "rejected"
            except Exception:
                results["prepare"] = "error"
            finally:
                close_old_connections()

        threads = [threading.Thread(target=do_unlock), threading.Thread(target=do_prepare)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)

        self.round.refresh_from_db()
        semifinal.refresh_from_db()
        downstream_active = semifinal.status != ContestRound.Status.DRAFT
        orphan = downstream_active and not self.round.is_locked
        self.assertFalse(
            orphan,
            "a PREPARED downstream round cannot coexist with an unlocked upstream",
        )


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class PrepareRoundScoreAuthorityConcurrencyTests(TransactionTestCase):
    """M0-W: a no-SCORE phase commit must prevent prepare_round from landing."""

    def setUp(self):
        self.user = User.objects.create_user(username="prepare-actor", password="pass")
        self.activity = Activity.objects.create(
            title="Prepared Score Authority",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username="prepare-singer", password="pass"),
            name="Singer",
            student_id="20260001",
            college="College",
            class_name="Class",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        Judge.objects.create(activity=self.activity, name="Judge", is_active=True)

    def test_prepare_never_lands_after_no_score_phase_commits(self):
        phased_held = threading.Event()
        release_lock = threading.Event()
        holder_error: dict[str, object] = {}

        def transition_to_no_score_phase():
            try:
                with transaction.atomic():
                    activity = Activity.objects.select_for_update().get(pk=self.activity.pk)
                    activity.phase = Activity.Phase.REVIEWING
                    activity.save(update_fields=["phase"])
                    phased_held.set()
                    release_lock.wait(timeout=10)
            except Exception as error:  # pragma: no cover - diagnostic only
                holder_error["error"] = error
            finally:
                close_old_connections()

        holder = threading.Thread(target=transition_to_no_score_phase)
        holder.start()
        self.assertTrue(phased_held.wait(timeout=10))

        prepare_result: dict[str, object] = {}

        def try_prepare():
            close_old_connections()
            try:
                prepare_round(self.round, self.user)
                prepare_result["done"] = True
            except PermissionDenied as error:
                prepare_result["rejected"] = str(error)
            except Exception as error:  # pragma: no cover - diagnostic only
                prepare_result["error"] = repr(error)
            finally:
                close_old_connections()

        worker = threading.Thread(target=try_prepare)
        worker.start()
        time.sleep(1)  # let the worker block on the Activity row lock
        release_lock.set()
        holder.join(timeout=10)
        worker.join(timeout=10)

        self.assertFalse(holder_error, holder_error)
        self.activity.refresh_from_db()
        self.round.refresh_from_db()
        self.assertEqual(self.activity.phase, Activity.Phase.REVIEWING)
        self.assertNotIn("done", prepare_result, prepare_result)
        self.assertIn("rejected", prepare_result, prepare_result)
        self.assertEqual(self.round.status, ContestRound.Status.DRAFT)
        self.assertEqual(self.round.entries.count(), 0)
        self.assertEqual(self.round.round_judges.count(), 0)


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class ActivityOwnedMutationBoundaryTests(TransactionTestCase):
    """M0-T: an Activity lock commit must block every activity-owned mutation.

    O1 participant edit, O2 singer review, O3 material-requirement change. Each
    races a live Activity row lock against the mutation and asserts the mutation
    never lands after the lock commits.
    """

    def setUp(self):
        self.user = User.objects.create_user(username="boundary-participant", password="pass")
        self.staff = User.objects.create_user(
            username="boundary-staff", password="pass", role=User.Role.STAFF
        )
        self.activity = Activity.objects.create(
            title="Boundary Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.reg = SingerRegistration.objects.create(
            activity=self.activity,
            user=self.user,
            name="Boundary Singer",
            student_id="20260050",
            college="College",
            class_name="Class",
            phone="13800000050",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )

    def test_o1_participant_edit_never_lands_after_activity_lock(self):
        lock_held = threading.Event()
        release_lock = threading.Event()
        holder_error: dict[str, object] = {}

        def hold_activity_lock():
            try:
                with transaction.atomic():
                    activity = Activity.objects.select_for_update().get(pk=self.activity.pk)
                    activity.is_locked = True
                    activity.save(update_fields=["is_locked"])
                    lock_held.set()
                    release_lock.wait(timeout=10)
            except Exception as error:  # pragma: no cover - diagnostic only
                holder_error["error"] = error
            finally:
                close_old_connections()

        holder = threading.Thread(target=hold_activity_lock)
        holder.start()
        self.assertTrue(lock_held.wait(timeout=10))

        edit_result: dict[str, object] = {}

        def try_participant_edit():
            close_old_connections()
            try:
                request = RequestFactory().post("/x", {"phone": "13800000051"})
                request.user = self.user
                my_registration_detail(request, self.reg.pk)
                edit_result["done"] = True
            except Exception as error:  # pragma: no cover - diagnostic only
                edit_result["error"] = repr(error)
            finally:
                close_old_connections()

        editor = threading.Thread(target=try_participant_edit)
        editor.start()
        time.sleep(1)  # let the editor block on the Activity row lock
        release_lock.set()
        holder.join(timeout=10)
        editor.join(timeout=10)

        self.assertFalse(holder_error, holder_error)
        self.activity.refresh_from_db()
        self.reg.refresh_from_db()
        self.assertTrue(self.activity.is_locked)
        self.assertEqual(self.reg.phone, "13800000050")
        self.assertEqual(edit_result.get("done"), True, edit_result)

    def test_o2_singer_review_never_lands_after_activity_lock(self):
        check = MaterialCheck.objects.create(
            singer_registration=self.reg,
            item_name="Accompaniment",
            status=MaterialCheck.Status.UPLOADED,
        )
        lock_held = threading.Event()
        release_lock = threading.Event()
        holder_error: dict[str, object] = {}

        def hold_activity_lock():
            try:
                with transaction.atomic():
                    activity = Activity.objects.select_for_update().get(pk=self.activity.pk)
                    activity.is_locked = True
                    activity.save(update_fields=["is_locked"])
                    lock_held.set()
                    release_lock.wait(timeout=10)
            except Exception as error:  # pragma: no cover - diagnostic only
                holder_error["error"] = error
            finally:
                close_old_connections()

        holder = threading.Thread(target=hold_activity_lock)
        holder.start()
        self.assertTrue(lock_held.wait(timeout=10))

        review_result: dict[str, object] = {}

        def try_review():
            close_old_connections()
            try:
                review_material_check(
                    check,
                    status=MaterialCheck.Status.APPROVED,
                    note="ok",
                    actor=self.staff,
                )
                review_result["done"] = True
            except PermissionDenied:
                review_result["rejected"] = True
            except Exception as error:  # pragma: no cover - diagnostic only
                review_result["error"] = repr(error)
            finally:
                close_old_connections()

        reviewer = threading.Thread(target=try_review)
        reviewer.start()
        time.sleep(1)
        release_lock.set()
        holder.join(timeout=10)
        reviewer.join(timeout=10)

        self.assertFalse(holder_error, holder_error)
        self.activity.refresh_from_db()
        check.refresh_from_db()
        self.assertTrue(self.activity.is_locked)
        self.assertEqual(check.status, MaterialCheck.Status.UPLOADED)
        self.assertEqual(review_result.get("rejected"), True, review_result)

    def test_o3_material_requirement_never_lands_after_activity_lock(self):
        lock_held = threading.Event()
        release_lock = threading.Event()
        holder_error: dict[str, object] = {}

        def hold_activity_lock():
            try:
                with transaction.atomic():
                    activity = Activity.objects.select_for_update().get(pk=self.activity.pk)
                    activity.is_locked = True
                    activity.save(update_fields=["is_locked"])
                    lock_held.set()
                    release_lock.wait(timeout=10)
            except Exception as error:  # pragma: no cover - diagnostic only
                holder_error["error"] = error
            finally:
                close_old_connections()

        holder = threading.Thread(target=hold_activity_lock)
        holder.start()
        self.assertTrue(lock_held.wait(timeout=10))

        mutate_result: dict[str, object] = {}

        def try_mutate_requirements():
            close_old_connections()
            try:
                request = RequestFactory().post(
                    "/x",
                    {
                        "applies_to": MaterialRequirement.AppliesTo.SINGER,
                        "item_name": "New Requirement",
                    },
                )
                request.user = self.staff
                activity_material_requirements(request, self.activity.pk)
                mutate_result["done"] = True
            except PermissionDenied:
                mutate_result["rejected"] = True
            except Exception as error:  # pragma: no cover - diagnostic only
                mutate_result["error"] = repr(error)
            finally:
                close_old_connections()

        mutator = threading.Thread(target=try_mutate_requirements)
        mutator.start()
        time.sleep(1)
        release_lock.set()
        holder.join(timeout=10)
        mutator.join(timeout=10)

        self.assertFalse(holder_error, holder_error)
        self.activity.refresh_from_db()
        self.assertTrue(self.activity.is_locked)
        self.assertFalse(
            MaterialRequirement.objects.filter(
                activity=self.activity, item_name="New Requirement"
            ).exists()
        )
        self.assertEqual(mutate_result.get("rejected"), True, mutate_result)

    def test_o5_test_cleanup_vs_participant_upload_no_deadlock(self):
        test_activity = Activity.objects.create(
            title="Test Cleanup",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        test_user = User.objects.create_user(username="cleanup-participant", password="pass")
        test_reg = SingerRegistration.objects.create(
            activity=test_activity,
            user=test_user,
            name="Cleanup Singer",
            student_id="20260060",
            college="College",
            class_name="Class",
            phone="13800000060",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        results: dict[str, str] = {}

        def run_cleanup():
            close_old_connections()
            try:
                clear_activity_test_data(test_activity, operator=self.staff)
                results["clear"] = "ok"
            except OperationalError as error:
                results["clear"] = "deadlock" if "deadlock" in str(error).lower() else "error"
            except Exception:
                results["clear"] = "error"
            finally:
                close_old_connections()

        def run_upload():
            close_old_connections()
            try:
                store_submission_file(
                    owner=test_reg,
                    uploaded_file=SimpleUploadedFile("a.mp3", b"x"),
                    purpose=SubmissionFile.Purpose.ACCOMPANIMENT,
                    uploaded_by=test_user,
                )
                results["upload"] = "ok"
            except OperationalError as error:
                results["upload"] = "deadlock" if "deadlock" in str(error).lower() else "error"
            except Exception:
                results["upload"] = "error"
            finally:
                close_old_connections()

        threads = [threading.Thread(target=run_cleanup), threading.Thread(target=run_upload)]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=20)

        self.assertNotEqual(results.get("clear"), "deadlock", results)
        self.assertNotEqual(results.get("upload"), "deadlock", results)


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class RulesetActivityLockConcurrencyTests(TransactionTestCase):
    """R0 §44: ruleset edits/clones serialize on the Activity lock (M0 authority).

    A ruleset mutation must never land after an Activity lock (or ARCHIVED) commits
    — the mutation re-validates the lock inside ``lock_activity_for_action``.
    """

    def setUp(self):
        import json as _json

        from ruleset.models import ContestRuleset, RulesetTemplate, RulesetVersion

        self.admin = User.objects.create_user(
            username="r0-rule-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="R0 Ruleset Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="R0规则",
            is_test_data=False,
            created_by=self.admin,
        )
        definition = _json.dumps(
            {
                "schema_version": 1,
                "nodes": [
                    {"key": "assess", "type": "ASSESS", "source": "entry", "round": "r1"},
                    {"key": "ranked", "type": "RANK", "source": "assess", "descending": True},
                ],
            }
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=definition,
                version=1,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                created_by=self.admin,
            )
        self.template = RulesetTemplate.objects.create(
            name="院十佳", definition=definition, created_by=self.admin
        )

    def _hold_lock(self, *, phase=None, locked=False):
        lock_held = threading.Event()
        release_lock = threading.Event()
        holder_error: dict[str, object] = {}

        def hold():
            try:
                with transaction.atomic():
                    act = Activity.objects.select_for_update().get(pk=self.activity.pk)
                    updated = []
                    if phase is not None:
                        act.phase = phase
                        updated.append("phase")
                    if locked:
                        act.is_locked = True
                        updated.append("is_locked")
                    act.save(update_fields=updated)
                    lock_held.set()
                    release_lock.wait(timeout=10)
            except Exception as error:  # pragma: no cover - diagnostic only
                holder_error["error"] = repr(error)
            finally:
                close_old_connections()

        holder = threading.Thread(target=hold)
        holder.start()
        self.assertTrue(lock_held.wait(timeout=10))
        return holder, release_lock, holder_error

    def test_ruleset_edit_never_lands_after_activity_lock(self):
        from staff_panel.views import ruleset_edit

        holder, release_lock, holder_error = self._hold_lock(locked=True)
        edit_result: dict[str, object] = {}

        def try_edit():
            close_old_connections()
            try:
                request = RequestFactory().post("/x", {"action": "add", "new_type": "ASSESS"})
                request.user = self.admin
                ruleset_edit(request, self.version.pk)
                edit_result["done"] = True
            except PermissionDenied as error:
                edit_result["rejected"] = str(error)
            except Exception as error:  # pragma: no cover - diagnostic only
                edit_result["error"] = repr(error)
            finally:
                close_old_connections()

        worker = threading.Thread(target=try_edit)
        worker.start()
        time.sleep(1)
        release_lock.set()
        holder.join(timeout=10)
        worker.join(timeout=10)

        self.assertFalse(holder_error, holder_error)
        self.activity.refresh_from_db()
        self.version.refresh_from_db()
        self.assertTrue(self.activity.is_locked)
        self.assertNotIn("done", edit_result, edit_result)
        self.assertIn("rejected", edit_result, edit_result)
        nodes = json.loads(self.version.definition)["nodes"]
        self.assertEqual(len(nodes), 2)

    def test_ruleset_clone_never_lands_after_activity_lock(self):
        from staff_panel.views import ruleset_clone_from_template

        holder, release_lock, holder_error = self._hold_lock(locked=True)
        clone_result: dict[str, object] = {}

        def try_clone():
            close_old_connections()
            try:
                request = RequestFactory().post("/x", {"activity": self.activity.pk, "name": "X"})
                request.user = self.admin
                ruleset_clone_from_template(request, self.template.pk)
                clone_result["done"] = True
            except PermissionDenied as error:
                clone_result["rejected"] = str(error)
            except Exception as error:  # pragma: no cover - diagnostic only
                clone_result["error"] = repr(error)
            finally:
                close_old_connections()

        worker = threading.Thread(target=try_clone)
        worker.start()
        time.sleep(1)
        release_lock.set()
        holder.join(timeout=10)
        worker.join(timeout=10)

        self.assertFalse(holder_error, holder_error)
        self.activity.refresh_from_db()
        self.assertTrue(self.activity.is_locked)
        self.assertNotIn("done", clone_result, clone_result)
        self.assertIn("rejected", clone_result, clone_result)


class ParticipantApplyVisibilityTests(TestCase):
    def setUp(self):
        self.participant = User.objects.create_user(username="apply-participant", password="pass")
        self.staff = User.objects.create_user(
            username="apply-staff", password="pass", role=User.Role.STAFF
        )
        self.formal = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.testing = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )

    def test_participant_apply_hides_test_activity(self):
        self.client.force_login(self.participant)
        response = self.client.get(reverse("singer_contest:apply"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.formal.title)
        self.assertNotContains(response, self.testing.title)

    def test_participant_cannot_submit_to_test_activity(self):
        self.client.force_login(self.participant)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("singer_contest:apply"),
            {
                "activity_id": str(self.testing.pk),
                "name": "Intruder",
                "student_id": "20269999",
                "college": "College",
                "class_name": "Class",
                "song_name": "Song",
            },
        )
        self.assertEqual(response.status_code, 404)
        self.assertFalse(SingerRegistration.objects.filter(activity=self.testing).exists())

    def test_staff_apply_can_preview_test_activity(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("singer_contest:apply"))
        self.assertContains(response, self.testing.title)

    def test_participant_apply_hides_video_field_for_formal_activity(self):
        self.client.force_login(self.participant)
        response = self.client.get(reverse("singer_contest:apply"))
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "performance_video")

    def test_staff_apply_shows_video_field_for_test_activity(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("singer_contest:apply"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "performance_video")

    def test_participant_apply_rejects_oversize_video_for_formal_activity(self):
        self.client.force_login(self.participant)
        response = self.client.post(
            reverse("singer_contest:apply"),
            {
                "activity_id": str(self.formal.pk),
                "name": "Singer",
                "student_id": "20260001",
                "college": "College",
                "class_name": "Class",
                "song_name": "Song",
                "performance_video": SimpleUploadedFile(
                    "clip.mp4", b"x" * (200 * 1024 * 1024), content_type="video/mp4"
                ),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "正式活动不支持大视频直传。")
        self.assertFalse(SingerRegistration.objects.filter(activity=self.formal).exists())


class StageResultModelTests(TestCase):
    """Characterization of the M1-F result models (StageResult + children)."""

    def setUp(self):
        self.user = User.objects.create_user(username="stage-staff", password="pass")
        self.activity = Activity.objects.create(
            title="院十佳",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        definition = json.dumps(
            {
                "schema_version": 1,
                "nodes": [
                    {"key": "assess", "type": "ASSESS", "source": "entry", "round": "r1"},
                    {"key": "ranked", "type": "RANK", "source": "assess", "descending": True},
                ],
            }
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="院十佳规则", is_test_data=True
        )
        self.version = RulesetVersion.objects.create(
            ruleset=self.ruleset, definition=definition, is_current=False
        )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=self.user,
            name="选手甲",
            student_id="20260001",
            college="学院",
            class_name="班级",
            song_name="歌曲",
        )

    def _result(self, **overrides):
        payload = {
            "activity": self.activity,
            "ruleset_version": self.version,
            "created_by": self.user,
            "stage_key": "院十佳",
            "ruleset_hash": self.version.content_hash,
            "is_test_data": True,
        }
        payload.update(overrides)
        return StageResult.objects.create(**payload)

    def test_stage_result_roundtrips_with_fks(self):
        result = self._result()
        self.assertEqual(result.activity, self.activity)
        self.assertEqual(result.ruleset_version, self.version)
        self.assertEqual(result.status, StageResult.Status.HOLD)
        self.assertTrue(result.ruleset_hash)
        self.assertEqual(result.schema_version, 1)

    def test_children_roundtrip_with_parent(self):
        result = self._result()
        decision = StageDecision.objects.create(
            stage_result=result,
            singer=self.singer,
            outcome_code="direct",
            rank=1,
            score=Decimal("92.46"),
            is_test_data=True,
        )
        composite = CompositeResult.objects.create(
            stage_result=result,
            singer=self.singer,
            node_key="stage2",
            value=Decimal("92.46"),
            components=[
                {"source": "stage1", "weight": "0.6", "value": "88.10", "contribution": "52.86"}
            ],
            is_test_data=True,
        )
        self.assertEqual(result.decisions.get(pk=decision.pk).stage_result, result)
        self.assertEqual(result.composites.get(pk=composite.pk).singer, self.singer)
        self.assertEqual(StageResult.Status.CONFIRMED.value, StageResult.Status.CONFIRMED.value)

    def test_test_marker_mismatch_rejected(self):
        with self.assertRaises(ValidationError):
            self._result(is_test_data=False)

    def test_child_test_marker_mismatch_rejected(self):
        result = self._result()
        with self.assertRaises(ValidationError):
            StageDecision(
                stage_result=result,
                singer=self.singer,
                outcome_code="direct",
                is_test_data=False,
            ).save()

    def test_reads_immutable_after_ready(self):
        result = self._result()
        # A HOLD result can still be updated to CONFIRMED (with its confirming trail).
        with authority_write(STAGE_RESULT_CONFIRM):
            StageResult.objects.filter(pk=result.pk).update(  # type: ignore[misc]
                status=StageResult.Status.CONFIRMED,
                confirmed_at=timezone.now(),
                confirmed_by=self.user,
            )
        # Once CONFIRMED, further update/delete is blocked.
        with self.assertRaises(ValidationError):
            StageResult.objects.filter(pk=result.pk).update(status=StageResult.Status.HOLD)
        with self.assertRaises(ValidationError):
            StageResult.objects.filter(pk=result.pk).delete()

    def test_child_immutable_when_parent_ready(self):
        result = self._result()
        decision = StageDecision.objects.create(
            stage_result=result,
            singer=self.singer,
            outcome_code="direct",
            rank=1,
            score=Decimal("92.46"),
            is_test_data=True,
        )
        with authority_write(STAGE_RESULT_CONFIRM):
            StageResult.objects.filter(pk=result.pk).update(  # type: ignore[misc]
                status=StageResult.Status.CONFIRMED,
                confirmed_at=timezone.now(),
                confirmed_by=self.user,
            )
        with self.assertRaises(ValidationError):
            StageDecision.objects.filter(pk=decision.pk).update(rank=2)

    def test_unique_activity_stage_key_hash(self):
        self._result()
        with self.assertRaises(IntegrityError):
            self._result()


class StageResolverBindingTests(TestCase):
    """M1-F binder + run_ruleset/persist over a small DB-backed fixture."""

    def setUp(self):
        self.user = User.objects.create_user(
            username="binder-staff", password="pass", role=User.Role.STAFF
        )
        self.activity = Activity.objects.create(
            title="院十佳",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        definition = json.dumps(
            {
                "schema_version": 1,
                "nodes": [
                    {"key": "assess", "type": "ASSESS", "source": "entry", "round": "r1"},
                    {"key": "ranked", "type": "RANK", "source": "assess", "descending": True},
                ],
            }
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="院十佳规则", is_test_data=True
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=definition,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="初赛",
        )
        self.judge = Judge.objects.create(activity=self.activity, name="评委甲")
        self.singers = [self._make_singer(i) for i in range(1, 4)]

    def _make_singer(self, index):
        singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"binder-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"2026100{index}",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )
        ScoreRecord.objects.create(
            round=self.round,
            singer=singer,
            judge=self.judge,
            score=Decimal(80 + index),
            is_test_data=True,
        )
        return singer

    def test_bind_resolve_input_builds_roster_and_round_scores(self):
        from .services import bind_resolve_input

        inputs = bind_resolve_input(self.version, self.activity, round_keys={"r1": self.round})
        self.assertEqual(inputs.roster, tuple(str(s.pk) for s in self.singers))
        expected = {str(s.pk): (Decimal(80 + i),) for i, s in enumerate(self.singers, 1)}
        self.assertEqual(inputs.round_scores["r1"], expected)
        self.assertEqual(inputs.vote_scores, {})

    def test_run_ruleset_persists_stage_result_and_decisions(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.stage_key, "选拔")  # type: ignore[union-attr]
        self.assertEqual(stage.ruleset_version, self.version)  # type: ignore[union-attr]
        self.assertTrue(stage.ruleset_hash)  # type: ignore[union-attr]
        self.assertTrue(stage.input_fingerprint)
        self.assertEqual(stage.plan_version, 1)
        self.assertEqual(stage.result_version, 1)
        self.assertEqual(stage.decisions.count(), 3)  # type: ignore[call-arg]
        self.assertEqual(stage.composites.count(), 0)  # type: ignore[call-arg]
        for decision in stage.decisions.all():  # type: ignore[union-attr]
            self.assertEqual(decision.outcome_code, "eliminated")
            self.assertIsNotNone(decision.score)

    def test_run_ruleset_rejects_version_from_other_activity(self):
        from .services import run_ruleset

        other = Activity.objects.create(
            title="其他",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        with self.assertRaises(ValidationError):
            run_ruleset(
                self.version,
                other,
                stage_key="选拔",
                computed_by=self.user,
                round_keys={"r1": self.round},
            )

    def test_run_ruleset_rejects_draft_version(self):
        """A DRAFT ruleset version is not an execution authority (R0 boundary)."""
        from .services import run_ruleset

        draft = RulesetVersion.objects.create(
            ruleset=self.ruleset,
            definition=self.version.definition,
            version=2,
            is_current=False,
            status=RulesetVersion.Status.DRAFT,
        )
        with self.assertRaises(ValidationError):
            run_ruleset(
                draft,
                self.activity,
                stage_key="选拔",
                computed_by=self.user,
                round_keys={"r1": self.round},
            )

    def test_run_ruleset_accepts_frozen_same_activity_version(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        self.assertEqual(stage.ruleset_version, self.version)  # type: ignore[union-attr]
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)

    def test_run_ruleset_rejects_superseded_frozen_unless_preview(self):
        """M1-R8 gate #2: a frozen version that is no longer current cannot produce a
        formal StageResult; only an explicit staff comparison (preview) may run it."""
        from ruleset.models import RulesetVersion

        from .services import run_ruleset

        # self.version is the current frozen v1; simulate a version that was frozen then
        # superseded by a newer authority (non-current, still FROZEN).
        with authority_write(RULESET_FREEZE):
            stopped = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔", "round_keys": {"r1": self.round.pk}},
            )
        with self.assertRaises(ValidationError):
            run_ruleset(
                stopped,
                self.activity,
                stage_key="选拔",
                computed_by=self.user,
                round_keys={"r1": self.round},
            )
        # A preview may still run it for a staff side-by-side comparison, but it is
        # ISOLATED (M1-R9 §七): it returns the in-memory ResolveResult and never writes
        # a formal StageResult, so a historical version can't pollute the result board.
        from ruleset.resolver import ResolverState

        stage = run_ruleset(
            stopped,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=True,
        )
        self.assertEqual(stage.status, ResolverState.READY)
        self.assertEqual(
            StageResult.objects.filter(activity=self.activity, ruleset_version=stopped).count(),
            0,
        )

    def test_persist_idempotent_same_input_reuses_stage(self):
        """§18: identical ruleset + input → idempotent reuse, no duplicate/version bump."""
        from .services import run_ruleset

        first = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        second = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        self.assertEqual(first.pk, second.pk)  # type: ignore[union-attr]
        self.assertEqual(second.result_version, 1)
        stage_count = StageResult.objects.filter(activity=self.activity, stage_key="选拔").count()
        self.assertEqual(stage_count, 1)
        self.assertEqual(second.decisions.count(), 3)  # type: ignore[call-arg]

    def test_persist_changed_input_bumps_result_version(self):
        """§18: an input change yields a new input_fingerprint and a new versioned result."""
        from .services import run_ruleset

        first = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        self.assertEqual(first.result_version, 1)
        # Correct a single score (the §16 scenario): the raw facts change.
        rec = ScoreRecord.objects.filter(round=self.round).first()
        rec.score += Decimal("0.50")  # type: ignore[union-attr]
        rec.save()  # type: ignore[union-attr]
        second = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        self.assertNotEqual(first.pk, second.pk)  # type: ignore[union-attr]
        self.assertEqual(second.result_version, 2)
        self.assertEqual(
            StageResult.objects.filter(activity=self.activity, stage_key="选拔").count(), 2
        )
        # The first result keeps its resolved-but-unconfirmed state; the new one is
        # a separate version, so the old one is never silently overwritten.
        first_still_resolved = StageResult.objects.filter(
            pk=first.pk,  # type: ignore[union-attr]
            status=StageResult.Status.READY_TO_CONFIRM,
        ).exists()
        self.assertTrue(first_still_resolved)

    def test_persist_identity_uses_ruleset_version_not_hash(self):
        """R5-4: two frozen versions with an identical definition (same content_hash)
        but a different binding no longer collide in the StageResult identity.

        The old identity was ``(activity, stage_key, ruleset_hash, input_fingerprint)``;
        since ``ruleset_hash`` is definition-only, a binding-only change between versions
        would silently reuse the prior result. The identity now keys on the immutable
        ``ruleset_version``, so both rows coexist.
        """
        from ruleset.resolver import resolve

        from .services import bind_resolve_input, persist_stage_result

        with authority_write(RULESET_FREEZE):
            v2 = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
            )
        self.assertEqual(v2.content_hash, self.version.content_hash)

        inputs = bind_resolve_input(self.version, self.activity, round_keys={"r1": self.round})
        result = resolve(self.version.definition, inputs)
        stage1 = persist_stage_result(
            self.version, self.activity, result, stage_key="选拔", computed_by=self.user
        )
        stage2 = persist_stage_result(
            v2, self.activity, result, stage_key="选拔", computed_by=self.user
        )
        self.assertNotEqual(stage1.pk, stage2.pk)
        self.assertEqual(stage1.ruleset_version, self.version)
        self.assertEqual(stage2.ruleset_version, v2)
        # Same definition hash + same input_fingerprint → the old identity collides; the
        # new one (keyed on ruleset_version) keeps both rows.
        self.assertEqual(stage1.ruleset_hash, stage2.ruleset_hash)
        self.assertEqual(
            StageResult.objects.filter(activity=self.activity, stage_key="选拔").count(), 2
        )

    def test_persist_ready_same_input_returned_unchanged(self):
        """§18: recomputing an already-resolved result over identical facts reuses it."""
        from .services import run_ruleset

        first = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        refreshed = run_ruleset(
            self.version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
        )
        self.assertEqual(refreshed.pk, first.pk)  # type: ignore[union-attr]
        self.assertEqual(refreshed.status, StageResult.Status.READY_TO_CONFIRM)
        stage_count = StageResult.objects.filter(activity=self.activity, stage_key="选拔").count()
        self.assertEqual(stage_count, 1)
        self.assertEqual(first.decisions.count(), 3)  # type: ignore[call-arg]

    def test_persist_maps_resolver_ready_to_ready_to_confirm(self):
        """§36-37: the resolver's "ready" persists as READY_TO_CONFIRM, never CONFIRMED."""
        from ruleset.resolver import resolve

        from .services import bind_resolve_input, persist_stage_result

        inputs = bind_resolve_input(self.version, self.activity, round_keys={"r1": self.round})
        result = resolve(self.version.definition, inputs)
        self.assertEqual(result.status.value, "ready")
        stage = persist_stage_result(
            self.version, self.activity, result, stage_key="选拔", computed_by=self.user
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertNotEqual(stage.status, StageResult.Status.CONFIRMED)

    def test_confirm_stage_result_locks_and_is_idempotent(self):
        """§36-37/§38: 核定 locks a resolved result; re-confirm reuses the locked row.

        M1-R6 finality: confirm re-locks the Activity (the authoritative serialization
        point), refuses a stale result over changed facts, and requires the consumed
        round to already be locked. The frozen version records its own binding so the
        service can recompute the current input fingerprint as the authority reference.
        """
        from .services import confirm_stage_result, run_ruleset

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        self.round.is_locked = True
        self.round.status = ContestRound.Status.LOCKED
        self.round.save(update_fields=["is_locked", "status"])
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔", "round_keys": {"r1": self.round.pk}},
            )
        # 核定 (R9-5) finalizes only a result grounded in the *current* FROZEN authority, so
        # make this frozen version the single current runner before confirming.
        _promote_version_to_current(version)
        stage = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.ruleset_version, version)  # type: ignore[union-attr]
        confirmed = confirm_stage_result(stage, confirmed_by=self.user)  # type: ignore[arg-type]
        self.assertEqual(confirmed.status, StageResult.Status.CONFIRMED)
        self.assertEqual(confirmed.confirmed_by, self.user)
        self.assertIsNotNone(confirmed.confirmed_at)
        again = confirm_stage_result(StageResult.objects.get(pk=stage.pk), confirmed_by=self.user)  # type: ignore[union-attr]
        self.assertEqual(again.pk, stage.pk)  # type: ignore[union-attr]
        self.assertEqual(again.status, StageResult.Status.CONFIRMED)
        self.assertEqual(again.confirmed_at, confirmed.confirmed_at)
        # A CONFIRMED result is immutable: recomputing identical facts returns it as-is.
        refreshed = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        self.assertEqual(refreshed.pk, stage.pk)  # type: ignore[union-attr]
        self.assertEqual(refreshed.status, StageResult.Status.CONFIRMED)

    def test_confirm_stage_result_rejects_unresolved(self):
        """§36-37: only a READY_TO_CONFIRM result may be 核定."""
        from .services import confirm_stage_result

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        stage = StageResult.objects.create(
            activity=self.activity,
            ruleset_version=self.version,
            created_by=self.user,
            stage_key="选拔",
            status=StageResult.Status.HOLD,
            reasons=["缺分数"],
            is_test_data=True,
        )
        with self.assertRaises(ValidationError):
            confirm_stage_result(stage, confirmed_by=self.user)
        stage.refresh_from_db()
        self.assertEqual(stage.status, StageResult.Status.HOLD)

    def test_confirm_rejects_stale_result_over_changed_scores(self):
        """§38: a result whose raw facts changed after computation is no longer current."""
        from .services import confirm_stage_result, run_ruleset

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        self.round.is_locked = True
        self.round.status = ContestRound.Status.LOCKED
        self.round.save(update_fields=["is_locked", "status"])
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔", "round_keys": {"r1": self.round.pk}},
            )
        # 核定 (R9-5) finalizes only a result grounded in the *current* FROZEN authority, so
        # make this frozen version the single current runner before confirming.
        _promote_version_to_current(version)
        stage = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        # A correct score edit changes the current input fingerprint.
        rec = ScoreRecord.objects.filter(round=self.round).first()
        rec.score += Decimal("0.25")  # type: ignore[union-attr]
        rec.save(update_fields=["score"])  # type: ignore[union-attr]
        with self.assertRaises(ValidationError):
            confirm_stage_result(stage, confirmed_by=self.user)  # type: ignore[arg-type]
        stage.refresh_from_db()  # type: ignore[union-attr]
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)

    def test_confirm_requires_consumed_round_locked(self):
        """§38: a stage whose source round is unlocked cannot be frozen."""
        from .services import confirm_stage_result, run_ruleset

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔", "round_keys": {"r1": self.round.pk}},
            )
        # 核定 (R9-5) finalizes only a result grounded in the *current* FROZEN authority, so
        # make this frozen version the single current runner before confirming.
        _promote_version_to_current(version)
        stage = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        with self.assertRaises(ValidationError):
            confirm_stage_result(stage, confirmed_by=self.user)  # type: ignore[arg-type]
        stage.refresh_from_db()  # type: ignore[union-attr]
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)

    def test_unlock_stage_result_reverts_confirmed(self):
        """§38: an admin unlock reverts a CONFIRMED result to READY_TO_CONFIRM."""
        from .services import confirm_stage_result, run_ruleset, unlock_stage_result

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        self.round.is_locked = True
        self.round.status = ContestRound.Status.LOCKED
        self.round.save(update_fields=["is_locked", "status"])
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔", "round_keys": {"r1": self.round.pk}},
            )
        # 核定 (R9-5) finalizes only a result grounded in the *current* FROZEN authority, so
        # make this frozen version the single current runner before confirming.
        _promote_version_to_current(version)
        stage = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        confirm_stage_result(stage, confirmed_by=self.user)  # type: ignore[arg-type]
        unlocked = unlock_stage_result(stage, operator=self.user, note="核对录错了")  # type: ignore[arg-type]
        self.assertEqual(unlocked.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertIsNone(unlocked.confirmed_by)
        self.assertIsNone(unlocked.confirmed_at)
        # Unlocked results are editable again: a score edit recomputes a new version.
        rec = ScoreRecord.objects.filter(round=self.round).first()
        rec.score += Decimal("0.50")  # type: ignore[union-attr]
        rec.save(update_fields=["score"])  # type: ignore[union-attr]
        second = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        self.assertNotEqual(second.pk, stage.pk)  # type: ignore[union-attr]
        self.assertEqual(second.result_version, 2)

    def test_confirm_rejects_when_newer_version_exists(self):
        """§38: a result that is no longer the stage's latest candidate cannot be frozen."""
        from .services import confirm_stage_result, run_ruleset

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        self.round.is_locked = True
        self.round.status = ContestRound.Status.LOCKED
        self.round.save(update_fields=["is_locked", "status"])
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔", "round_keys": {"r1": self.round.pk}},
            )
        # 核定 (R9-5) finalizes only a result on the current FROZEN authority, so make
        # this frozen version the single current runner before persisting (M1-R9 §七).
        _promote_version_to_current(version)
        first = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        rec = ScoreRecord.objects.filter(round=self.round).first()
        rec.score += Decimal("0.25")  # type: ignore[union-attr]
        rec.save(update_fields=["score"])  # type: ignore[union-attr]
        second = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        self.assertEqual(second.result_version, 2)
        with self.assertRaises(ValidationError):
            confirm_stage_result(first, confirmed_by=self.user)  # type: ignore[arg-type]
        first.refresh_from_db()  # type: ignore[union-attr]
        self.assertEqual(first.status, StageResult.Status.READY_TO_CONFIRM)

    def test_confirm_rejects_result_on_superseded_version(self):
        """M1-R9 (§五): 核定 only finalizes a result grounded in the *current* FROZEN
        authority. A result on a demoted (non-current) frozen version is rejected so a host
        cannot copy a handcard that no longer matches what the activity publishes."""
        from .services import confirm_stage_result

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        self.round.is_locked = True
        self.round.status = ContestRound.Status.LOCKED
        self.round.save(update_fields=["is_locked", "status"])
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self.version.definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔", "round_keys": {"r1": self.round.pk}},
            )
        # M1-R9 §七 isolates a preview (in-memory, no persist), so a demoted version's
        # only persisted result is one written BEFORE it was demoted. Simulate exactly
        # that historical row: it is READY_TO_CONFIRM but grounded on a non-current
        # authority, which 核定 must reject.
        stage = StageResult.objects.create(
            activity=self.activity,
            ruleset_version=version,
            created_by=self.user,
            stage_key="选拔",
            status=StageResult.Status.READY_TO_CONFIRM,
            reasons=[],
            is_test_data=True,
        )
        with self.assertRaises(ValidationError):
            confirm_stage_result(stage, confirmed_by=self.user)
        stage.refresh_from_db()
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)

    def test_stage_decisions_blocks_prefer_checkpoint_override(self):
        """§十七/P1: a checkpoint stage's handcard uses its own blocks, not the global set."""
        from .services import stage_decisions_by_blocks

        definition = json.dumps(
            {
                "schema_version": 1,
                "checkpoints": [{"key": "stage2", "output": "ranked"}],
                "nodes": [
                    {"key": "assess", "type": "ASSESS", "source": "entry", "round": "r1"},
                    {"key": "ranked", "type": "RANK", "source": "assess", "descending": True},
                ],
            }
        )
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=definition,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={
                    "stage_key": "院十佳",
                    "round_keys": {"r1": self.round.pk},
                    "announcement_blocks": [{"label": "全局", "outcome_codes": ["direct"]}],
                    "announcement_blocks_by_checkpoint": {
                        "stage2": [{"label": "赛段专属", "outcome_codes": ["advanced"]}]
                    },
                },
            )
        stage = StageResult.objects.create(
            activity=self.activity,
            ruleset_version=version,
            created_by=self.user,
            stage_key="stage2",
            status=StageResult.Status.READY_TO_CONFIRM,
            reasons=[],
            is_test_data=True,
        )
        StageDecision.objects.create(
            stage_result=stage,
            singer=self.singers[0],
            outcome_code="advanced",
            rank=1,
            is_test_data=True,
        )
        StageDecision.objects.create(
            stage_result=stage,
            singer=self.singers[1],
            outcome_code="direct",
            rank=2,
            is_test_data=True,
        )
        blocks = stage_decisions_by_blocks(stage)
        labels = [b["label"] for b in blocks]
        self.assertIn("赛段专属", labels)
        self.assertNotIn("全局", labels)
        self.assertIn("其他", labels)
        self.assertEqual([d.outcome_code for d in blocks[0]["decisions"]], ["advanced"])


def _schidui_definition():
    """§11.3 院十佳 weights (30/60/10, 60/40, 30/50/20) as one forward-only graph."""
    from ruleset.templates import GOLDEN_SCHIDUI

    return GOLDEN_SCHIDUI


def _xiaofeng_definition():
    """§12.5 校十佳屏峰 chain as one forward-only graph (no 2025 special-casing)."""
    from ruleset.templates import GOLDEN_XIAOFENG

    return GOLDEN_XIAOFENG


def _roster_bridge_definition():
    """Two-stage forward chain: R1 picks a top-3, the final round picks a top-1."""
    return json.dumps(
        {
            "schema_version": 1,
            "nodes": [
                {"key": "assess_r1", "type": "ASSESS", "source": "entry", "round": "r1"},
                {"key": "rank1", "type": "RANK", "source": "assess_r1", "descending": True},
                {"key": "top3", "type": "SELECT", "source": "rank1", "count": 3},
                {"key": "assess_f", "type": "ASSESS", "source": "top3", "round": "final"},
                {"key": "rank2", "type": "RANK", "source": "assess_f", "descending": True},
                {"key": "top1", "type": "SELECT", "source": "rank2", "count": 1},
            ],
            "checkpoints": [
                {"key": "stage1", "output": "top3"},
                {"key": "stage2", "output": "top1"},
            ],
        },
        ensure_ascii=False,
    )


class RoundEntryBridgeTests(TestCase):
    """M1-INTEGRATION-1: StageDecision drives the next round's entry roster.

    Covers the reviewer's acceptance: 复赛 TopN -> StageDecision -> automatically
    materialise the ``RoundEntry`` of the round bound to that stage, and the final
    ruleset's bind rosters exactly those advancers — never the whole approved list.
    """

    def setUp(self):
        self.admin = User.objects.create_user(
            username="roster-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="Roster Bridge",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="Roster规则", is_test_data=True
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=_roster_bridge_definition(),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        self.judge = Judge.objects.create(activity=self.activity, name="评委A")
        self.singers = [self._singer(i) for i in range(1, 7)]
        self.prelim = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="初赛",
            sequence=1,
        )
        self.final = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
            name="决赛",
            sequence=2,
            roster_source=ContestRound.RosterSource.STAGE,
            roster_source_stage="stage1",
        )
        self._seed_r1()

    def _singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"roster-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"9000{index:02d}",
            college="学院",
            class_name="班级",
            song_name="歌",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def _score(self, round_, singers, scorer):
        for idx, singer in enumerate(singers, 1):
            ScoreRecord.objects.create(
                round=round_,
                singer=singer,
                judge=self.judge,
                score=Decimal(scorer(idx)),
                is_test_data=True,
            )

    def _seed_r1(self):
        # Descending scores: singer 1 (idx0) highest -> singer 3 (idx2) third.
        self._score(self.prelim, self.singers, lambda i: 100 - i)

    def _entry_ids(self, round_):
        return list(
            RoundEntry.objects.filter(round=round_)
            .order_by("pk")
            .values_list("singer_id", flat=True)
        )

    def test_stage1_bridges_next_round_entry_to_top3(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.prelim, "final": self.final},
            checkpoint="stage1",
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)  # type: ignore[union-attr]
        self.assertEqual(stage.stage_key, "stage1")  # type: ignore[union-attr]
        self.assertEqual(self._entry_ids(self.final), [s.pk for s in self.singers[:3]])

    def test_final_ruleset_entry_is_top3_not_all_approved(self):
        from .services import bind_resolve_input, run_ruleset

        run_ruleset(
            self.version,
            self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.prelim, "final": self.final},
            checkpoint="stage1",
        )
        inputs = bind_resolve_input(
            self.version,
            self.activity,
            round_keys={"r1": self.prelim, "final": self.final},
            checkpoint="stage2",
        )
        self.assertEqual(inputs.roster, tuple(str(s.pk) for s in self.singers[:3]))

    def test_stage_round_seeds_from_roster_not_is_advanced(self):
        from .services import prepare_round

        RoundEntry.objects.bulk_create(
            [RoundEntry(round=self.final, singer=s) for s in self.singers[:3]]
        )
        # A conflicting upstream is_advanced marker must NOT leak into the STAGE round.
        ScoreSummary.objects.create(
            round=self.prelim,
            singer=self.singers[3],
            average_score=Decimal("99.0"),
            rank=1,
            is_advanced=True,
            is_test_data=True,
        )
        prepared = prepare_round(self.final, self.admin)
        self.assertEqual(self._entry_ids(self.final), [s.pk for s in self.singers[:3]])
        self.assertEqual(prepared.status, ContestRound.Status.PREPARED)

    def test_stage_round_without_roster_raises(self):
        from .services import prepare_round

        with self.assertRaisesMessage(ValidationError, "尚未生成该轮晋级名单"):
            prepare_round(self.final, self.admin)

    def test_revising_stage_reconciles_round_entry(self):
        """A re-resolved stage must prune a superseded advancer from the round entry."""
        from .services import run_ruleset

        run_ruleset(
            self.version,
            self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.prelim, "final": self.final},
            checkpoint="stage1",
        )
        self.assertEqual(self._entry_ids(self.final), [s.pk for s in self.singers[:3]])

        # Flip the ranking so a different trio advances (descending scores 99..94).
        ScoreRecord.objects.filter(round=self.prelim).delete()
        self._score(self.prelim, list(reversed(self.singers)), lambda i: 100 - i)
        run_ruleset(
            self.version,
            self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.prelim, "final": self.final},
            checkpoint="stage1",
        )
        ids = set(self._entry_ids(self.final))
        self.assertEqual(ids, {self.singers[3].pk, self.singers[4].pk, self.singers[5].pk})
        # The original trio (s1..s3) all advanced before; none may be retained now.
        self.assertTrue({s.pk for s in self.singers[:3]}.isdisjoint(ids))

    def test_checkpoint_bind_without_materialized_subset_round_raises(self):
        """A checkpoint over a not-yet-materialized STAGE round must not widen to all."""
        from .services import bind_resolve_input

        with self.assertRaisesMessage(ValidationError, "尚未生成该轮晋级名单"):
            bind_resolve_input(
                self.version,
                self.activity,
                round_keys={"r1": self.prelim, "final": self.final},
                checkpoint="stage2",
            )

    def test_checkpoint_named_stage_key_persists_scoped_fingerprint(self):
        """A checkpoint-named stage_key resolves as that checkpoint even without it."""
        from .services import run_ruleset

        no_checkpoint = run_ruleset(
            self.version,
            self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.prelim, "final": self.final},
        )
        explicit = run_ruleset(
            self.version,
            self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.prelim, "final": self.final},
            checkpoint="stage1",
        )
        self.assertEqual(no_checkpoint.input_fingerprint, explicit.input_fingerprint)


class GoldenSchiduiDbTests(TestCase):
    """M1-F golden 院十佳 end-to-end: 15 singers through the DB, run_ruleset, verify.

    Scores are deterministic and synthetic (five identical judge scores per round so
    the engine's mean reproduces the target exactly); audience scores are pre-normalized
    Decimals passed as ready values — the binder never re-derives a raw-vote score
    (§16.9). No 2025-value special-casing lives in Python: the weights and the
    15->10->5->3 cutoffs all live in the frozen definition.
    """

    JUDGE_COUNT = 5

    def setUp(self):
        self.admin = User.objects.create_user(
            username="golden-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="院十佳",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="院十佳规则", is_test_data=True
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=_schidui_definition(),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        self.judges = [
            Judge.objects.create(activity=self.activity, name=f"评委{chr(0x41 + i)}")
            for i in range(self.JUDGE_COUNT)
        ]
        self.rounds = {}
        for rkey in ("r1", "r2", "r3", "r4"):
            self.rounds[rkey] = ContestRound.objects.create(
                activity=self.activity,
                round_type=ContestRound.RoundType.PRELIMINARY,
                name=rkey,
            )
        self.singers = [self._make_singer(i) for i in range(1, 16)]
        self._seed_scores()

    def _make_singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"golden-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"2026001{index:02d}",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def _score(self, round_key, singers, scorer):
        for idx, singer in enumerate(singers, 1):
            for judge in self.judges:
                ScoreRecord.objects.create(
                    round=self.rounds[round_key],
                    singer=singer,
                    judge=judge,
                    score=Decimal(scorer(idx)),
                    is_test_data=True,
                )

    def _seed_scores(self):
        # R1/R2: the whole entry roster; R3: top-10 advancees; R4: top-5 advancees.
        self._score("r1", self.singers, lambda i: 100 - i)
        self._score("r2", self.singers, lambda i: 90 - i)
        self._score("r3", self.singers[:10], lambda i: 100 - i)
        self._score("r4", self.singers[:5], lambda i: 100 - i)

    def _vote_scores(self):
        audience1 = {str(s.pk): Decimal("50") for s in self.singers}
        audience4 = {str(self.singers[i].pk): Decimal(90 - (i + 1)) for i in range(5)}
        return {"audience1": audience1, "audience4": audience4}

    def _round_keys(self):
        return dict(self.rounds)

    def test_golden_schidui_end_to_end_ready(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="院十佳",
            computed_by=self.admin,
            round_keys=self._round_keys(),
            vote_scores=self._vote_scores(),
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.stage_key, "院十佳")  # type: ignore[union-attr]
        self.assertEqual(stage.ruleset_version, self.version)  # type: ignore[union-attr]
        self.assertEqual(stage.plan_version, 1)
        self.assertEqual(stage.decisions.count(), 15)  # type: ignore[call-arg]

        # Composite scoping is roster-scoped: stage1 all 15, stage2 top-10, final top-5.
        self.assertEqual(stage.composites.filter(node_key="stage1").count(), 15)  # type: ignore[union-attr]
        self.assertEqual(stage.composites.filter(node_key="stage2").count(), 10)  # type: ignore[union-attr]
        self.assertEqual(stage.composites.filter(node_key="final").count(), 5)  # type: ignore[union-attr]

        # Top-3 advance directly; the rest who only reached a lower roster are direct
        # too (origin tag), and those never selected are eliminated.
        by_singer = {d.singer_id: d for d in stage.decisions.all()}  # type: ignore[union-attr]
        for idx in range(3):
            decision = by_singer[self.singers[idx].pk]
            self.assertEqual(decision.outcome_code, "direct")
            self.assertEqual(decision.rank, idx + 1)
        for idx in range(3, 10):
            self.assertEqual(by_singer[self.singers[idx].pk].outcome_code, "direct")
        for idx in range(10, 15):
            self.assertEqual(by_singer[self.singers[idx].pk].outcome_code, "eliminated")

    def test_golden_schidui_checkpoint_publishes_stage1_with_only_first_facts(self):
        """§11-15: the stage1 checkpoint publishes READY when only R1/R2/audience1 exist."""
        from .services import run_ruleset

        audience1 = {str(s.pk): Decimal("50") for s in self.singers}
        # Bind only the first-stage rounds; R3/R4/audience4 rows exist in DB but are unbound.
        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.rounds["r1"], "r2": self.rounds["r2"]},
            vote_scores={"audience1": audience1},
            checkpoint="stage1",
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.stage_key, "stage1")  # type: ignore[union-attr]
        self.assertEqual(stage.ruleset_version, self.version)  # type: ignore[union-attr]
        self.assertEqual(stage.decisions.count(), 15)  # type: ignore[call-arg]
        # Only the stage1 composite is emitted; stage2/final are not in the closure.
        self.assertEqual(stage.composites.filter(node_key="stage1").count(), 15)  # type: ignore[union-attr]
        self.assertEqual(stage.composites.filter(node_key="stage2").count(), 0)  # type: ignore[union-attr]
        self.assertEqual(stage.composites.filter(node_key="final").count(), 0)  # type: ignore[union-attr]
        by_singer = {d.singer_id: d for d in stage.decisions.all()}  # type: ignore[union-attr]
        for idx in range(10):
            self.assertEqual(by_singer[self.singers[idx].pk].outcome_code, "direct")
        for idx in range(10, 15):
            self.assertEqual(by_singer[self.singers[idx].pk].outcome_code, "eliminated")

    def test_golden_schidui_checkpoint_reuse_until_full_resolve(self):
        """§11-15: re-resolving the same stage1 facts yields the same READY row (idempotent)."""
        from .services import run_ruleset

        audience1 = {str(s.pk): Decimal("50") for s in self.singers}
        kwargs = dict(
            version=self.version,
            activity=self.activity,
            stage_key="stage1",
            computed_by=self.admin,
            round_keys={"r1": self.rounds["r1"], "r2": self.rounds["r2"]},
            vote_scores={"audience1": audience1},
            checkpoint="stage1",
        )
        first = run_ruleset(**kwargs)
        second = run_ruleset(**kwargs)
        self.assertEqual(first.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(first.pk, second.pk)  # type: ignore[union-attr]
        self.assertEqual(first.result_version, second.result_version)

    def _composite_map(self, stage, node_key):
        return {c.singer_id: c.value for c in stage.composites.filter(node_key=node_key)}

    def test_golden_schidui_layered_decimals(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="院十佳",
            computed_by=self.admin,
            round_keys=self._round_keys(),
            vote_scores=self._vote_scores(),
        )

        # Hand-computed §11.3 values keyed by creation order (pk order).
        stage1 = self._composite_map(stage, "stage1")
        self.assertEqual(stage1[self.singers[0].pk], Decimal("88.1"))
        self.assertEqual(stage1[self.singers[4].pk], Decimal("84.5"))
        self.assertEqual(stage1[self.singers[14].pk], Decimal("75.5"))

        stage2 = self._composite_map(stage, "stage2")
        self.assertEqual(stage2[self.singers[0].pk], Decimal("92.46"))
        self.assertEqual(stage2[self.singers[9].pk], Decimal("84.00"))

        final = self._composite_map(stage, "final")
        self.assertEqual(final[self.singers[0].pk], Decimal("97.0"))
        self.assertEqual(final[self.singers[4].pk], Decimal("93.0"))

    def test_golden_schidui_persisted_result_roundtrip(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="院十佳",
            computed_by=self.admin,
            round_keys=self._round_keys(),
            vote_scores=self._vote_scores(),
        )
        self.assertTrue(stage.ruleset_hash)  # type: ignore[union-attr]
        self.assertTrue(stage.input_fingerprint)
        self.assertEqual(stage.schema_version, 1)
        self.assertEqual(stage.result_version, 1)
        decision = stage.decisions.first()  # type: ignore[union-attr]
        self.assertIsNotNone(decision.singer_id)  # type: ignore[union-attr]
        self.assertIsNotNone(decision.score)  # type: ignore[union-attr]
        composite = stage.composites.filter(node_key="stage1").first()  # type: ignore[union-attr]
        self.assertEqual(len(composite.components), 3)  # type: ignore[union-attr]
        self.assertIn("source", composite.components[0])  # type: ignore[union-attr]


class GoldenSchiduiXiaofengDbTests(TestCase):
    """M1-G golden 校十佳屏峰 end-to-end (DB-bound).

    20 singers -> 5 initial groups -> top1/group direct (5) -> leftover R1 top12
    -> R2 top7 -> merge 12 -> final groups -> manual 0~2/group -> fill to 6. All
    control flow (counts, groups, quota) lives in the frozen definition; binder and
    run_ruleset are generic. Finalists are read from ``node_values['filled']`` because
    every merged contestant already carries DIRECT/REPECHAGE via first-writer-wins.
    """

    JUDGE_COUNT = 5

    def setUp(self):
        self.admin = User.objects.create_user(
            username="xf-golden-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="校十佳屏峰",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="校十佳屏峰规则", is_test_data=True
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=_xiaofeng_definition(),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        self.judges = [
            Judge.objects.create(activity=self.activity, name=f"评委{chr(0x41 + i)}")
            for i in range(self.JUDGE_COUNT)
        ]
        self.rounds = {}
        for rkey in ("r1", "r2"):
            self.rounds[rkey] = ContestRound.objects.create(
                activity=self.activity,
                round_type=ContestRound.RoundType.PRELIMINARY,
                name=rkey,
            )
        self.singers = [self._make_singer(i) for i in range(1, 21)]
        self._seed_scores()

    def _make_singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"xf-golden-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"2026010{index:02d}",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def _score(self, round_key, singers, scorer):
        for idx, singer in enumerate(singers, 1):
            for judge in self.judges:
                ScoreRecord.objects.create(
                    round=self.rounds[round_key],
                    singer=singer,
                    judge=judge,
                    score=Decimal(scorer(idx)),
                    is_test_data=True,
                )

    def _seed_scores(self):
        self._score("r1", self.singers, lambda i: 100 - i)
        self._score("r2", self.singers, lambda i: 200 - i)

    def _group_of(self):
        initial = {
            str(singer.pk): group
            for group, start in (("G1", 0), ("G2", 4), ("G3", 8), ("G4", 12), ("G5", 16))
            for singer in self.singers[start : start + 4]
        }
        final = {
            str(singer.pk): group
            for group, positions in (
                ("F1", (0, 1, 2, 3)),
                ("F2", (4, 5, 6, 7)),
                ("F3", (8, 9, 12, 16)),
            )
            for position in positions
            for singer in (self.singers[position],)
        }
        return {"initial_group": initial, "final_group": final}

    def _manual(self):
        # F1/F2 full (2 each), F3 short by one -> FILL_TO_QUOTA must pull c13 in.
        return {
            "manual": {
                "F1": (str(self.singers[0].pk), str(self.singers[1].pk)),
                "F2": (str(self.singers[4].pk), str(self.singers[5].pk)),
                "F3": (str(self.singers[8].pk),),
            }
        }

    def test_golden_xiaofeng_end_to_end_ready(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="校十佳屏峰",
            computed_by=self.admin,
            round_keys=dict(self.rounds),
            group_of=self._group_of(),
            manual=self._manual(),
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.stage_key, "校十佳屏峰")  # type: ignore[union-attr]
        self.assertEqual(stage.ruleset_version, self.version)  # type: ignore[union-attr]
        self.assertEqual(stage.plan_version, 1)
        self.assertEqual(stage.decisions.count(), 20)  # type: ignore[call-arg]

        by_singer = {d.singer_id: d for d in stage.decisions.all()}  # type: ignore[union-attr]
        # 5 direct (top1/group), 12 repechage, 3 never-selected eliminated.
        codes = [d.outcome_code for d in stage.decisions.all()]  # type: ignore[union-attr]
        self.assertEqual(codes.count("direct"), 5)
        self.assertEqual(codes.count("repechage"), 12)
        self.assertEqual(codes.count("eliminated"), 3)
        # Direct winners are creation order 1/5/9/13/17, each seated by their global R1 rank.
        for idx, expected_rank in ((0, 1), (4, 5), (8, 9), (12, 13), (16, 17)):
            decision = by_singer[self.singers[idx].pk]
            self.assertEqual(decision.outcome_code, "direct")
            self.assertEqual(decision.rank, expected_rank)
        # Eliminated are the three lowest leftovers: creation order 18/19/20.
        for idx in (17, 18, 19):
            self.assertEqual(by_singer[self.singers[idx].pk].outcome_code, "eliminated")

    def test_golden_xiaofeng_fill_to_quota_finalists(self):
        from .services import run_ruleset

        stage = run_ruleset(
            self.version,
            self.activity,
            stage_key="校十佳屏峰",
            computed_by=self.admin,
            round_keys=dict(self.rounds),
            group_of=self._group_of(),
            manual=self._manual(),
        )
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)

        # Re-resolve to inspect the fill GROUP_MAP: F3 was short one and must be topped
        # up from the merged pool (c13). Finalists are read from node_values['filled'].
        from ruleset.resolver import resolve

        result = resolve(self.version.definition, self._resolve_input())
        self.assertEqual(result.status.value, "ready")
        self.assertEqual(
            result.node_values["filled"],
            {
                "F1": [str(self.singers[0].pk), str(self.singers[1].pk)],
                "F2": [str(self.singers[4].pk), str(self.singers[5].pk)],
                "F3": [str(self.singers[8].pk), str(self.singers[12].pk)],
            },
        )
        finalists = [c for group in result.node_values["filled"].values() for c in group]  # type: ignore[attr-defined]
        self.assertEqual(len(finalists), 6)
        self.assertEqual(len(set(finalists)), 6)

    def test_recompute_auto_sources_group_and_manual_ready(self):
        """§32: recompute auto-sources group_of + manual from DB (no hand-injected kwargs)."""
        from .models import Performance, PerformanceGroup
        from .services import recompute_activity_result

        # Reuse the setUp ruleset (UNIQUE(activity)); demote its v1 current version so
        # this v2 authority becomes the single current FROZEN runner for recompute.
        ruleset = self.ruleset
        ruleset.round_keys = {"r1": self.rounds["r1"].pk, "r2": self.rounds["r2"].pk}
        ruleset.save(update_fields=["round_keys"])
        RulesetVersion._base_manager.filter(pk=self.version.pk).update(is_current=False)
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=ruleset,
                version=2,
                definition=_xiaofeng_definition(),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                binding={
                    "stage_key": "校十佳屏峰",
                    "round_keys": {"r1": self.rounds["r1"].pk, "r2": self.rounds["r2"].pk},
                    "vote_keys": {},
                    "group_keys": {
                        "initial_group": self.rounds["r1"].pk,
                        "final_group": self.rounds["r2"].pk,
                    },
                    "announcement_blocks": [],
                },
            )
        # Persist the group_of facts as Performance rows keyed by the binding's group_keys.
        group_of = self._group_of()
        for key, round_ in (
            ("initial_group", self.rounds["r1"]),
            ("final_group", self.rounds["r2"]),
        ):
            for singer in self.singers:
                group_name = group_of[key].get(str(singer.pk))
                if group_name is None:
                    continue
                pg, _ = PerformanceGroup.objects.get_or_create(
                    activity=self.activity, round=round_, name=group_name, is_test_data=True
                )
                Performance.objects.create(
                    activity=self.activity,
                    round=round_,
                    singer=singer,
                    group=pg,
                    is_test_data=True,
                )
        # Persist the manual facts as ManualDecision rows via the formal R9-3 service.
        from .services import set_manual_decision

        for group_name, chosen in self._manual()["manual"].items():
            set_manual_decision(
                version,
                manual_key="manual",
                group=group_name,
                chosen=[str(c) for c in chosen],
                created_by=self.admin,
            )

        stage = recompute_activity_result(self.activity, self.admin, ruleset=ruleset)
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.stage_key, "校十佳屏峰")
        self.assertEqual(stage.ruleset_version, version)
        self.assertEqual(stage.decisions.count(), 20)
        codes = [d.outcome_code for d in stage.decisions.all()]
        self.assertEqual(codes.count("direct"), 5)
        self.assertEqual(codes.count("repechage"), 12)
        self.assertEqual(codes.count("eliminated"), 3)

    def _resolve_input(self):
        from .services import bind_resolve_input

        return bind_resolve_input(
            self.version,
            self.activity,
            round_keys=dict(self.rounds),
            group_of=self._group_of(),
            manual=self._manual(),
        )


class RapidEntryServiceTests(TestCase):
    """M1-H rapid-entry backstage: stale-edit version bump + auto re-resolve service."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="rapid-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="录分活动",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="r1",
        )
        self.singer = self._make_singer(1)
        self.judge = Judge.objects.create(activity=self.activity, name="评委A")
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save(update_fields=["status"])

    def _make_singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"rapid-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"2026m1h{index:02d}",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def test_apply_scores_bumps_score_version(self):
        from .services import apply_scores

        self.assertEqual(self.round.score_version, 0)
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "95"}, self.admin)
        self.round.refresh_from_db()
        self.assertEqual(self.round.score_version, 1)
        # Re-saving the same value is a no-op and must not bump again.
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "95"}, self.admin)
        self.round.refresh_from_db()
        self.assertEqual(self.round.score_version, 1)

    def test_recompute_activity_result_persists_ready(self):
        import json

        from ruleset.models import ContestRuleset, RulesetVersion

        from .services import recompute_activity_result

        ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="录分规则",
            is_test_data=True,
            stage_key="院十佳",
            round_keys={"r1": self.round.pk},
        )
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=ruleset,
                definition=json.dumps(
                    {
                        "schema_version": 1,
                        "nodes": [
                            {
                                "key": "assess_r1",
                                "type": "ASSESS",
                                "source": "entry",
                                "round": "r1",
                            },
                            {
                                "key": "rank1",
                                "type": "RANK",
                                "source": "assess_r1",
                                "descending": True,
                            },
                            {"key": "top1", "type": "SELECT", "source": "rank1", "count": 1},
                        ],
                    }
                ),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        # Seed the single cell so the composite resolves (no HOLD).
        from .services import apply_scores

        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "90"}, self.admin)

        stage = recompute_activity_result(self.activity, self.admin, ruleset=ruleset)
        self.assertEqual(stage.stage_key, "院十佳")
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.ruleset_version, version)
        self.assertEqual(stage.decisions.count(), 1)
        decision = stage.decisions.get()
        self.assertEqual(decision.outcome_code, "direct")

    def test_frozen_binding_authoritative_over_ruleset_mutation(self):
        """M1-R1: after freeze, editing the ContestRuleset binding does not change the result.

        A frozen version carries its binding snapshot; the runtime reads it from the
        version, never the still-mutable ContestRuleset round_keys / stage_key.
        """
        import json

        from ruleset.models import ContestRuleset, RulesetVersion

        from .services import apply_scores, recompute_activity_result

        ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="权威规则",
            is_test_data=True,
            stage_key="院十佳",
            round_keys={"r1": self.round.pk},
        )
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=ruleset,
                definition=json.dumps(
                    {
                        "schema_version": 1,
                        "nodes": [
                            {
                                "key": "assess_r1",
                                "type": "ASSESS",
                                "source": "entry",
                                "round": "r1",
                            },
                            {
                                "key": "rank1",
                                "type": "RANK",
                                "source": "assess_r1",
                                "descending": True,
                            },
                            {"key": "top1", "type": "SELECT", "source": "rank1", "count": 1},
                        ],
                    }
                ),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                binding={
                    "stage_key": "院十佳",
                    "round_keys": {"r1": self.round.pk},
                    "announcement_blocks": [],
                },
            )
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "90"}, self.admin)

        # The frozen version is authoritative: rewriting the ruleset binding must be inert.
        ruleset.stage_key = "篡改赛段"
        ruleset.round_keys = {}
        ruleset.save()

        stage = recompute_activity_result(self.activity, self.admin, ruleset=ruleset)
        self.assertEqual(stage.stage_key, "院十佳")
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.ruleset_version, version)
        self.assertEqual(stage.decisions.count(), 1)

    def test_recompute_with_checkpoint_publishes_that_stage_key(self):
        """§11-15: a checkpoint recompute persists its own READY stage, not the whole graph."""
        import json

        from ruleset.models import ContestRuleset, RulesetVersion

        from .services import apply_scores, recompute_activity_result

        ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="检查点规则",
            is_test_data=True,
            stage_key="院十佳",
            round_keys={"r1": self.round.pk},
        )
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=ruleset,
                definition=json.dumps(
                    {
                        "schema_version": 1,
                        "checkpoints": [{"key": "stage1", "output": "top1"}],
                        "nodes": [
                            {
                                "key": "assess_r1",
                                "type": "ASSESS",
                                "source": "entry",
                                "round": "r1",
                            },
                            {
                                "key": "rank1",
                                "type": "RANK",
                                "source": "assess_r1",
                                "descending": True,
                            },
                            {"key": "top1", "type": "SELECT", "source": "rank1", "count": 1},
                        ],
                    }
                ),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "90"}, self.admin)

        stage = recompute_activity_result(
            self.activity, self.admin, ruleset=ruleset, checkpoint="stage1"
        )
        self.assertEqual(stage.stage_key, "stage1")
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(stage.ruleset_version, version)
        self.assertEqual(stage.decisions.count(), 1)
        self.assertEqual(stage.decisions.get().outcome_code, "direct")


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class RecomputeActivityResultConcurrencyTests(TransactionTestCase):
    """M1-R9 §二: concurrent recomputes of one activity are linearized by the Activity
    FOR UPDATE lock, so ``result_version`` is never minted twice (the DB unique is the
    final guard). Identical facts converge to a single row; no thread raises IntegrityError."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="recompute-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="并发重算活动",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PENDING,
            is_test_mode=True,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="r1",
        )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username="recompute-s1", password="pass"),
            name="选手一",
            student_id="2026rc001",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )
        self.judge = Judge.objects.create(activity=self.activity, name="评委A")
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save(update_fields=["status"])
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="并发规则",
            is_test_data=True,
            stage_key="院十佳",
            round_keys={"r1": self.round.pk},
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=json.dumps(
                    {
                        "schema_version": 1,
                        "nodes": [
                            {
                                "key": "assess_r1",
                                "type": "ASSESS",
                                "source": "entry",
                                "round": "r1",
                            },
                            {
                                "key": "rank1",
                                "type": "RANK",
                                "source": "assess_r1",
                                "descending": True,
                            },
                            {"key": "top1", "type": "SELECT", "source": "rank1", "count": 1},
                        ],
                    }
                ),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "院十佳", "round_keys": {"r1": self.round.pk}},
            )
        from .services import apply_scores

        apply_scores(self.round, {(self.singer.pk, self.judge.pk): "90"}, self.admin)

    def test_concurrent_identical_recompute_is_linearized(self):
        from .services import recompute_activity_result

        outcomes = []
        guard = threading.Lock()

        def do_recompute():
            close_old_connections()
            try:
                recompute_activity_result(self.activity, self.admin, ruleset=self.ruleset)
                with guard:
                    outcomes.append("ok")
            except Exception as exc:  # noqa: BLE001 — record any failure for the assert below
                with guard:
                    outcomes.append(f"error:{type(exc).__name__}")
            finally:
                close_old_connections()

        threads = [threading.Thread(target=do_recompute) for _ in range(4)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(timeout=30)

        self.assertEqual(len(outcomes), 4, outcomes)
        self.assertFalse(any(o.startswith("error") for o in outcomes), outcomes)
        versions = list(
            StageResult.objects.filter(activity=self.activity, stage_key="院十佳").values_list(
                "activity_id", "stage_key", "result_version"
            )
        )
        self.assertEqual(len(versions), len(set(versions)), "result_version must be unique")
        self.assertEqual(len(versions), 1, "identical facts converge to a single row")


class BindingSourceHelperTests(TestCase):
    """§32 binding-sourcing helpers: group_of / vote_scores / manual are read from DB."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="src-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="源活动",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="源规则", is_test_data=True
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="r1",
        )
        self.singers = [
            SingerRegistration.objects.create(
                activity=self.activity,
                user=User.objects.create_user(username=f"src-s{i}", password="pass"),
                name=f"选手{i}",
                student_id=f"src{i:02d}",
                college="学院",
                class_name="班级",
                song_name="歌曲",
                pre_status=SingerRegistration.PreStatus.APPROVED,
                is_test_data=True,
            )
            for i in range(1, 4)
        ]

    def test_source_group_of_reads_performance_group(self):
        from .models import Performance, PerformanceGroup
        from .services import _source_group_of

        pg = PerformanceGroup.objects.create(
            activity=self.activity, round=self.round, name="G1", is_test_data=True
        )
        Performance.objects.create(
            activity=self.activity,
            round=self.round,
            singer=self.singers[0],
            group=pg,
            is_test_data=True,
        )
        out = _source_group_of(self.activity, {"group_keys": {"by": self.round.pk}})
        self.assertEqual(out, {"by": {str(self.singers[0].pk): "G1"}})

    def test_source_group_of_missing_key_returns_empty(self):
        from .services import _source_group_of

        self.assertEqual(_source_group_of(self.activity, {"group_keys": {}}), {})

    def _vote_session(self):
        from django.utils import timezone
        from voting.models import VoteSession

        return VoteSession.objects.create(
            activity=self.activity,
            name="大众投票",
            passcode="0000",
            start_time=timezone.now(),
            end_time=timezone.now() + timezone.timedelta(hours=1),  # type: ignore[attr-defined]
            is_test_data=True,
        )

    def test_source_vote_scores_returns_raw_counts(self):
        from voting.models import VoteOption, VoteRecord

        from .services import _source_vote_scores

        vs = self._vote_session()
        opts = [
            VoteOption.objects.create(vote_session=vs, singer=s, is_test_data=True)
            for s in self.singers
        ]
        # 2 + 1 + 1 votes across the three singers: raw counts, never a normalized score.
        for key, opt in (("a", opts[0]), ("b", opts[0]), ("c", opts[1]), ("d", opts[2])):
            VoteRecord.objects.create(
                vote_session=vs,
                vote_option=opt,
                browser_session_key=key,
                ip_address="127.0.0.1",
                is_test_data=True,
            )
        out = _source_vote_scores(self.activity, {"vote_keys": {"audience": vs.pk}})
        self.assertEqual(out["audience"][str(self.singers[0].pk)], Decimal("2"))
        self.assertEqual(out["audience"][str(self.singers[1].pk)], Decimal("1"))
        self.assertEqual(out["audience"][str(self.singers[2].pk)], Decimal("1"))

    def test_source_vote_scores_empty_session_skips(self):

        from .services import _source_vote_scores

        vs = self._vote_session()
        out = _source_vote_scores(self.activity, {"vote_keys": {"audience": vs.pk}})
        self.assertEqual(out, {})

    def test_source_manual_returns_group_and_flat_keys(self):
        from .models import ManualDecision
        from .services import _source_manual

        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition='{"schema_version": 1, "nodes": [{"key": "roster", "type": "ROSTER"}]}',
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        with _manual_write_ctx():
            ManualDecision.objects.create(
                activity=self.activity,
                ruleset_version=version,
                manual_key="manual",
                group="G1",
                chosen=[str(self.singers[0].pk)],
                is_test_data=True,
            )
            ManualDecision.objects.create(
                activity=self.activity,
                ruleset_version=version,
                manual_key="manual",
                group="",
                chosen=[str(self.singers[1].pk)],
                is_test_data=True,
            )
        out = _source_manual(version, self.activity)
        self.assertEqual(
            out,
            {
                "manual": {
                    "G1": (str(self.singers[0].pk),),
                    "": (str(self.singers[1].pk),),
                }
            },
        )


class ManualDecisionModelTests(TestCase):
    """ManualDecision clean + unique_together (§31 closure for §32 MANUAL_SELECT)."""

    def setUp(self):
        self.activity = Activity.objects.create(
            title="手动",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.other = Activity.objects.create(
            title="他活动",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="规则", is_test_data=True
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition='{"schema_version": 1, "nodes": [{"key": "roster", "type": "ROSTER"}]}',
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username="md-s", password="pass"),
            name="选手",
            student_id="md01",
            college="学院",
            class_name="班级",
            song_name="歌",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def test_chosen_must_belong_to_activity(self):
        from .models import ManualDecision

        foreign = SingerRegistration.objects.create(
            activity=self.other,
            user=User.objects.create_user(username="md-f", password="pass"),
            name="他",
            student_id="md02",
            college="学院",
            class_name="班级",
            song_name="歌",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )
        with self.assertRaises(ValidationError):
            with _manual_write_ctx():
                ManualDecision.objects.create(
                    activity=self.activity,
                    ruleset_version=self.version,
                    manual_key="manual",
                    group="",
                    chosen=[str(foreign.pk)],
                    is_test_data=True,
                )

    def test_unique_together_enforced(self):
        from .models import ManualDecision

        with _manual_write_ctx():
            ManualDecision.objects.create(
                activity=self.activity,
                ruleset_version=self.version,
                manual_key="manual",
                group="G1",
                chosen=[str(self.singer.pk)],
                is_test_data=True,
            )
        with self.assertRaises(IntegrityError), transaction.atomic():
            with _manual_write_ctx():
                ManualDecision.objects.create(
                    activity=self.activity,
                    ruleset_version=self.version,
                    manual_key="manual",
                    group="G1",
                    chosen=[],
                    is_test_data=True,
                )


class ManualDecisionServiceTests(TestCase):
    """M1-R9 §三: ManualDecision mutation must run through the formal service (authority).

    A bare ``ManualDecision.save()`` is refused unless the thread-local production guard
    is on; only ``set_manual_decision`` / ``delete_manual_decision`` hold the Activity lock
    and re-check the consumed-by-confirmed guard inside it. Covers upsert, delete, the
    freeze/activity/manual-key validations, and the audit record.
    """

    def setUp(self):
        self.admin = User.objects.create_user(
            username="md-svc-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="手动服务",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.other = Activity.objects.create(
            title="其他活动",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="手动规则", is_test_data=True
        )
        self.version = self._frozen(self.ruleset, self._manual_definition())
        self.singers = [self._make_singer(i) for i in range(3)]

    @staticmethod
    def _manual_definition():
        return {
            "schema_version": 1,
            "nodes": [
                {"key": "roster", "type": "ROSTER"},
                {
                    "key": "manual",
                    "type": "MANUAL_SELECT",
                    "source": "roster",
                    "groups": 2,
                    "quota": 2,
                },
            ],
        }

    def _frozen(self, ruleset, definition):
        with authority_write(RULESET_FREEZE):
            return RulesetVersion.objects.create(
                ruleset=ruleset,
                definition=json.dumps(definition, ensure_ascii=False),
                version=RulesetVersion.objects.filter(ruleset=ruleset).count() + 1,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )

    def _make_singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"md-svc-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"md-svc{index:02d}",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def test_set_upserts_group_and_flat_and_audits(self):
        from .models import ManualDecision
        from .services import set_manual_decision

        set_manual_decision(
            self.version,
            manual_key="manual",
            group="G1",
            chosen=[self.singers[0].pk, self.singers[1].pk],
            created_by=self.admin,
        )
        row = ManualDecision.objects.get(
            activity=self.activity,
            ruleset_version=self.version,
            manual_key="manual",
            group="G1",
        )
        self.assertEqual(row.chosen, [str(self.singers[0].pk), str(self.singers[1].pk)])
        self.assertTrue(row.is_test_data)
        # Re-call upserts the same row instead of duplicating it.
        set_manual_decision(
            self.version,
            manual_key="manual",
            group="G1",
            chosen=[self.singers[2].pk],
            created_by=self.admin,
        )
        total_after_group_update = ManualDecision.objects.filter(
            activity=self.activity, ruleset_version=self.version
        ).count()
        self.assertEqual(total_after_group_update, 1)
        row.refresh_from_db()
        self.assertEqual(row.chosen, [str(self.singers[2].pk)])
        # A flat (non-group) source uses the "" group; a distinct row.
        set_manual_decision(
            self.version,
            manual_key="manual",
            group="",
            chosen=[self.singers[0].pk],
            created_by=self.admin,
        )
        total_after_flat_add = ManualDecision.objects.filter(
            activity=self.activity, ruleset_version=self.version
        ).count()
        self.assertEqual(total_after_flat_add, 2)
        self.assertTrue(
            AuditLog.objects.filter(
                operator=self.admin, action_type=AuditLog.ActionType.MANUAL_DECISION
            ).exists()
        )

    def test_delete_removes_and_returns_none_when_missing(self):
        from .models import ManualDecision
        from .services import delete_manual_decision, set_manual_decision

        set_manual_decision(
            self.version,
            manual_key="manual",
            group="G1",
            chosen=[str(self.singers[0].pk)],
            created_by=self.admin,
        )
        deleted = delete_manual_decision(
            self.version, manual_key="manual", group="G1", deleted_by=self.admin
        )
        self.assertIsNotNone(deleted)
        remaining = ManualDecision.objects.filter(
            activity=self.activity, ruleset_version=self.version
        )
        self.assertFalse(remaining.exists())
        self.assertIsNone(
            delete_manual_decision(
                self.version, manual_key="manual", group="G1", deleted_by=self.admin
            )
        )

    def test_set_rejects_non_frozen_version(self):
        from .services import set_manual_decision

        draft = RulesetVersion.objects.create(
            ruleset=self.ruleset,
            definition=json.dumps(self._manual_definition(), ensure_ascii=False),
            version=RulesetVersion.objects.filter(ruleset=self.ruleset).count() + 1,
            is_current=False,
            status=RulesetVersion.Status.DRAFT,
        )
        with self.assertRaises(ValidationError):
            set_manual_decision(
                draft,
                manual_key="manual",
                group="G1",
                chosen=[str(self.singers[0].pk)],
                created_by=self.admin,
            )

    def test_set_rejects_undeclared_manual_key(self):
        from .services import set_manual_decision

        with authority_write(RULESET_FREEZE):
            no_manual = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=json.dumps(
                    {"schema_version": 1, "nodes": [{"key": "roster", "type": "ROSTER"}]},
                    ensure_ascii=False,
                ),
                version=RulesetVersion.objects.filter(ruleset=self.ruleset).count() + 1,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
            )
        with self.assertRaises(ValidationError):
            set_manual_decision(
                no_manual,
                manual_key="manual",
                group="G1",
                chosen=[str(self.singers[0].pk)],
                created_by=self.admin,
            )

    def test_set_rejects_wrong_activity(self):
        from .services import set_manual_decision

        other_ruleset = ContestRuleset.objects.create(
            activity=self.other, name="他规则", is_test_data=True
        )
        other_version = self._frozen(other_ruleset, self._manual_definition())
        with self.assertRaises(ValidationError):
            set_manual_decision(
                other_version,
                manual_key="manual",
                group="G1",
                chosen=[str(self.singers[0].pk)],
                created_by=self.admin,
                activity=self.activity,
            )

    def test_set_rejects_foreign_chosen_singer(self):
        from .services import set_manual_decision

        foreign = SingerRegistration.objects.create(
            activity=self.other,
            user=User.objects.create_user(username="md-svc-f", password="pass"),
            name="外人",
            student_id="md-svc99",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )
        with self.assertRaises(ValidationError):
            set_manual_decision(
                self.version,
                manual_key="manual",
                group="G1",
                chosen=[str(foreign.pk)],
                created_by=self.admin,
            )

    def test_direct_save_is_refused_unless_authorized(self):
        from .models import ManualDecision

        md = ManualDecision(
            activity=self.activity,
            ruleset_version=self.version,
            manual_key="manual",
            group="",
            chosen=[str(self.singers[0].pk)],
            is_test_data=True,
        )
        with self.assertRaises(ValidationError):
            md.save()
        with _manual_write_ctx():
            md.save()
        saved = ManualDecision.objects.filter(activity=self.activity, ruleset_version=self.version)
        self.assertTrue(saved.exists())


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class ManualDecisionMutationConcurrencyTests(TransactionTestCase):
    """M1-R9 §三: a ManualDecision set racing a StageResult confirm is linearized by the
    Activity lock — exactly one committing order wins, so a CONFIRMED stage can never hold
    a manual pick that was mutated underneath it."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="md-race-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="并发人工选择",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PENDING,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="并发规则", is_test_data=True
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=json.dumps(self._manual_definition(), ensure_ascii=False),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "选拔"},
            )
        self.singers = [self._make_singer(i) for i in range(3)]
        from .services import recompute_activity_result, set_manual_decision

        set_manual_decision(
            self.version,
            manual_key="manual",
            group="",
            chosen=[str(self.singers[0].pk)],
            created_by=self.admin,
        )
        self.stage = recompute_activity_result(self.activity, self.admin, ruleset=self.ruleset)
        self.assertEqual(self.stage.status, StageResult.Status.READY_TO_CONFIRM)

    @staticmethod
    def _manual_definition():
        return {
            "schema_version": 1,
            "nodes": [
                {"key": "roster", "type": "ROSTER"},
                {
                    "key": "manual",
                    "type": "MANUAL_SELECT",
                    "source": "roster",
                    "groups": 1,
                    "quota": 2,
                },
            ],
        }

    def _make_singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"md-race-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"md-race{index:02d}",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def test_set_vs_confirm_race_only_one_wins(self):
        from .services import confirm_stage_result, set_manual_decision

        outcomes = []
        guard = threading.Lock()

        def do_set():
            close_old_connections()
            try:
                set_manual_decision(
                    self.version,
                    manual_key="manual",
                    group="",
                    chosen=[str(self.singers[1].pk)],
                    created_by=self.admin,
                )
                with guard:
                    outcomes.append("set-ok")
            except Exception as exc:  # noqa: BLE001 — record any failure for the assertion
                with guard:
                    outcomes.append(f"set-error:{type(exc).__name__}")
            finally:
                close_old_connections()

        def do_confirm():
            close_old_connections()
            try:
                confirm_stage_result(self.stage, confirmed_by=self.admin)
                with guard:
                    outcomes.append("confirm-ok")
            except Exception as exc:  # noqa: BLE001
                with guard:
                    outcomes.append(f"confirm-error:{type(exc).__name__}")
            finally:
                close_old_connections()

        threads = [threading.Thread(target=do_set), threading.Thread(target=do_confirm)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(timeout=30)

        self.assertEqual(len(outcomes), 2, outcomes)
        set_ok = any(o == "set-ok" for o in outcomes)
        confirm_ok = any(o == "confirm-ok" for o in outcomes)
        # The Activity FOR UPDATE lock serializes the two conflicting writes, so exactly
        # one committing order wins — a CONFIRMED stage never coexists with a mutated pick.
        self.assertFalse(set_ok and confirm_ok, outcomes)
        stage = StageResult.objects.get(pk=self.stage.pk)
        if stage.status == StageResult.Status.CONFIRMED:
            self.assertFalse(set_ok, outcomes)


class ConfirmedDependencyClosureTests(TestCase):
    """M1-R8 Commit 3: a CONFIRMED stage result pins its raw facts (round/vote/manual).

    Once a stage is 核定, the upstream round scores, vote counts, and manual picks it
    read are no longer subject to un-wind. ``unlock_round`` / ``reset_round_to_draft`` /
    ``unlock_vote_session`` / ``ManualDecision`` mutation must refuse while any
    CONFIRMED stage of the same activity/version consumed the raw fact. Only unlocking
    the stage result first (reverting to READY_TO_CONFIRM) releases finality.
    """

    def setUp(self):
        self.user = User.objects.create_user(
            username="closure-admin", password="pass", role=User.Role.STAFF
        )
        self.activity = Activity.objects.create(
            title="闭环节",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PENDING,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="闭环规则", is_test_data=True
        )
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="初赛",
        )
        self.judge = Judge.objects.create(activity=self.activity, name="评委甲")
        self.singers = [self._make_singer(i) for i in range(1, 4)]
        self.vs = VoteSession.objects.create(
            activity=self.activity,
            name="大众投票",
            passcode="0000",
            start_time=timezone.now(),
            end_time=timezone.now() + timezone.timedelta(hours=1),  # type: ignore[attr-defined]
            is_test_data=True,
        )

    def _make_singer(self, index):
        singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"closure-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"close{index:02d}",
            college="学院",
            class_name="班级",
            song_name="歌曲",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )
        ScoreRecord.objects.create(
            round=self.round,
            singer=singer,
            judge=self.judge,
            score=Decimal(80 + index),
            is_test_data=True,
        )
        return singer

    def _frozen_version(self, definition, binding):
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=json.dumps(definition, ensure_ascii=False),
                version=RulesetVersion.objects.filter(ruleset=self.ruleset).count() + 1,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding=binding,
            )
        # 核定 (R9-5) finalizes only a result on the current FROZEN authority, so make the
        # version we confirm the single current runner.
        _promote_version_to_current(version)
        return version

    def _confirm_stage(self, definition, binding, lock_round=False, lock_vote=False):
        from .services import confirm_stage_result, run_ruleset

        version = self._frozen_version(definition, binding)
        kwargs = {"stage_key": "选拔", "computed_by": self.user, "round_keys": {"r1": self.round}}
        if lock_round:
            self.round.is_locked = True
            self.round.status = ContestRound.Status.LOCKED
            self.round.save(update_fields=["is_locked", "status"])
        if lock_vote:
            self.vs.is_locked = True
            self.vs.save(update_fields=["is_locked"])
        stage = run_ruleset(version, self.activity, **kwargs, preview=False)  # type: ignore[arg-type]
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        return confirm_stage_result(stage, confirmed_by=self.user)  # type: ignore[arg-type]

    def test_confirmed_stage_blocks_unlock_and_reset_round(self):
        from .services import reset_round_to_draft, unlock_round, unlock_stage_result

        definition = {
            "schema_version": 1,
            "nodes": [
                {"key": "assess", "type": "ASSESS", "source": "entry", "round": "r1"},
                {"key": "ranked", "type": "RANK", "source": "assess", "descending": True},
            ],
        }
        stage = self._confirm_stage(
            definition, {"stage_key": "选拔", "round_keys": {"r1": self.round.pk}}, lock_round=True
        )
        self.assertEqual(stage.status, StageResult.Status.CONFIRMED)
        with self.assertRaisesMessage(ValidationError, "该原始数据已被已核定赛段结果使用"):
            unlock_round(self.round, self.user)
        with self.assertRaisesMessage(ValidationError, "该原始数据已被已核定赛段结果使用"):
            reset_round_to_draft(self.round, self.user, reason="admin unwind")
        # Unlocking the stage result releases finality so the round is editable again.
        unlock_stage_result(StageResult.objects.get(pk=stage.pk), operator=self.user)
        unlocked = unlock_round(self.round, self.user)
        self.assertFalse(unlocked.is_locked)
        self.assertEqual(unlocked.status, ContestRound.Status.SCORING)

    def test_confirm_records_rich_audit_in_same_tx(self):
        """M1-R8 Commit 4: the CONFIRM audit is recorded by the service in the same txn."""
        from common.models import AuditLog

        from .services import confirm_stage_result, run_ruleset

        definition = {
            "schema_version": 1,
            "nodes": [
                {"key": "assess", "type": "ASSESS", "source": "entry", "round": "r1"},
                {"key": "ranked", "type": "RANK", "source": "assess", "descending": True},
            ],
        }
        version = self._frozen_version(
            definition, {"stage_key": "选拔", "round_keys": {"r1": self.round.pk}}
        )
        self.round.is_locked = True
        self.round.status = ContestRound.Status.LOCKED
        self.round.save(update_fields=["is_locked", "status"])
        stage = run_ruleset(
            version,
            self.activity,
            stage_key="选拔",
            computed_by=self.user,
            round_keys={"r1": self.round},
            preview=False,
        )
        confirmed = confirm_stage_result(stage, confirmed_by=self.user)  # type: ignore[arg-type]
        audit = AuditLog.objects.get(
            action_type=AuditLog.ActionType.CONFIRM_STAGE_RESULT,
            target=f"StageResult:{confirmed.pk}",
        )
        self.assertEqual(audit.operator, self.user)
        self.assertEqual(audit.old_value, StageResult.Status.READY_TO_CONFIRM)
        payload = json.loads(audit.new_value)
        self.assertEqual(payload["status"], StageResult.Status.CONFIRMED)
        self.assertEqual(payload["stage_key"], "选拔")
        self.assertEqual(payload["result_version"], confirmed.result_version)
        self.assertEqual(payload["ruleset_version"], version.pk)
        self.assertEqual(payload["authority_hash"], version.authority_hash)
        self.assertEqual(payload["input_fingerprint"], confirmed.input_fingerprint)
        self.assertEqual(payload["confirmed_by"], self.user.pk)

    def test_confirmed_vote_sourced_stage_blocks_unlock_vote_session(self):
        from voting.services import unlock_vote_session

        from .services import ensure_vote_not_consumed_by_confirmed_stage

        self.vs.is_locked = True
        self.vs.save(update_fields=["is_locked"])
        definition = {
            "schema_version": 1,
            "nodes": [
                {"key": "agg", "type": "ASSESS", "source": "entry", "vote_source": "audience"},
                {"key": "ranked", "type": "RANK", "source": "agg", "descending": True},
            ],
        }
        version = self._frozen_version(
            definition, {"stage_key": "选拔", "vote_keys": {"audience": self.vs.pk}}
        )
        with authority_write(STAGE_RESULT_CONFIRM):
            StageResult.objects.create(
                activity=self.activity,
                ruleset_version=version,
                created_by=self.user,
                stage_key="选拔",
                status=StageResult.Status.CONFIRMED,
                input_fingerprint="x",
                result_version=1,
                is_test_data=True,
                confirmed_by=self.user,
                confirmed_at=timezone.now(),
            )
        with self.assertRaisesMessage(ValidationError, "该原始数据已被已核定赛段结果使用"):
            ensure_vote_not_consumed_by_confirmed_stage(self.vs)
        with self.assertRaisesMessage(ValidationError, "该原始数据已被已核定赛段结果使用"):
            unlock_vote_session(self.vs, self.user)

    def test_confirmed_stage_blocks_manual_decision_mutation(self):
        from .services import (
            ensure_manual_not_consumed_by_confirmed_stage,
            set_manual_decision,
        )

        definition = {
            "schema_version": 1,
            "nodes": [
                {"key": "roster", "type": "ROSTER"},
                {
                    "key": "manual",
                    "type": "MANUAL_SELECT",
                    "source": "roster",
                    "groups": 1,
                    "quota": 2,
                },
            ],
        }
        version = self._frozen_version(definition, {"stage_key": "选拔"})
        with authority_write(STAGE_RESULT_CONFIRM):
            StageResult.objects.create(
                activity=self.activity,
                ruleset_version=version,
                created_by=self.user,
                stage_key="选拔",
                status=StageResult.Status.CONFIRMED,
                input_fingerprint="x",
                result_version=1,
                is_test_data=True,
                confirmed_by=self.user,
                confirmed_at=timezone.now(),
            )
        with self.assertRaisesMessage(ValidationError, "该原始数据已被已核定赛段结果使用"):
            ensure_manual_not_consumed_by_confirmed_stage(version, "manual")
        with self.assertRaisesMessage(ValidationError, "该原始数据已被已核定赛段结果使用"):
            set_manual_decision(
                version,
                manual_key="manual",
                group="G1",
                chosen=[str(self.singers[0].pk)],
                created_by=self.user,
            )


class AutoResolveCheckpointTests(TestCase):
    """M1-INTEGRATION-2: progressive auto-resolution publishes READY_TO_CONFIRM stages."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="auto-resolve", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="院十佳",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="院十佳规则", is_test_data=True, stage_key="院十佳"
        )
        self.judge = Judge.objects.create(activity=self.activity, name="评委A")
        self.singers = [self._singer(i) for i in range(1, 16)]
        self.rounds = {}
        for idx in range(1, 5):
            self.rounds[f"r{idx}"] = ContestRound.objects.create(
                activity=self.activity,
                round_type=ContestRound.RoundType.PRELIMINARY,
                name=f"轮次{idx}",
                sequence=idx,
                roster_source=ContestRound.RosterSource.APPROVED,
            )
        from ruleset.templates import GOLDEN_SCHIDUI

        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=GOLDEN_SCHIDUI,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                binding={
                    "stage_key": "院十佳",
                    "round_keys": {k: v.pk for k, v in self.rounds.items()},
                    "audience_keys": {"audience1": "aud1set", "audience4": "aud4set"},
                },
            )

    def _singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"auto-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"90{index:03d}",
            college="学院",
            class_name="班级",
            song_name="歌",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def _round_scores(self, round_, base):
        for idx, singer in enumerate(self.singers, 1):
            ScoreRecord.objects.create(
                round=round_,
                singer=singer,
                judge=self.judge,
                score=Decimal(base) - idx,
                is_test_data=True,
            )

    def _audience(self, set_name, base):
        for singer in self.singers:
            AudienceScore.objects.create(
                activity=self.activity,
                stage_key=set_name,
                singer=singer,
                score=Decimal(base),
                is_test_data=True,
            )

    def test_stage1_then_stage2_then_final_resolve_progressively(self):
        from .services import maybe_resolve_checkpoints

        self._round_scores(self.rounds["r1"], 100)
        self._round_scores(self.rounds["r2"], 90)
        self._audience("aud1set", 50)

        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["stage1"])
        s1 = StageResult.objects.get(activity=self.activity, stage_key="stage1")
        self.assertEqual(s1.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(StageResult.objects.filter(activity=self.activity).count(), 1)

        self._round_scores(self.rounds["r3"], 100)
        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["stage2"])
        s2 = StageResult.objects.get(activity=self.activity, stage_key="stage2")
        self.assertEqual(s2.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(StageResult.objects.filter(activity=self.activity).count(), 2)

        self._round_scores(self.rounds["r4"], 100)
        self._audience("aud4set", 90)
        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["stage3"])
        s3 = StageResult.objects.get(activity=self.activity, stage_key="stage3")
        self.assertEqual(s3.status, StageResult.Status.READY_TO_CONFIRM)

    def test_idempotent_rerun_does_not_duplicate_stage_result(self):
        from .services import maybe_resolve_checkpoints

        self._round_scores(self.rounds["r1"], 100)
        self._round_scores(self.rounds["r2"], 90)
        self._audience("aud1set", 50)

        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["stage1"])
        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), [])
        self.assertEqual(
            StageResult.objects.filter(activity=self.activity, stage_key="stage1").count(), 1
        )


class AudienceCompositeFlowTests(TestCase):
    """M1-INTEGRATION-2: a staff-entered AudienceScore flows into the composite as a 0-100 value."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="aud-comp", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="院十佳",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="院十佳规则", is_test_data=True, stage_key="院十佳"
        )
        self.judge = Judge.objects.create(activity=self.activity, name="评委A")
        self.singers = [self._singer(i) for i in range(1, 11)]
        self.round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="初赛",
            sequence=1,
            roster_source=ContestRound.RosterSource.APPROVED,
        )
        definition = json.dumps(
            {
                "schema_version": 1,
                "nodes": [
                    {"key": "assess_r1", "type": "ASSESS", "source": "entry", "round": "r1"},
                    {
                        "key": "assess_a1",
                        "type": "ASSESS",
                        "source": "entry",
                        "vote_source": "audience1",
                    },
                    {
                        "key": "composite",
                        "type": "AGGREGATE",
                        "aggregate": {
                            "type": "weighted_sum",
                            "components": [
                                {"source": "assess_r1", "weight": 0.70},
                                {"source": "assess_a1", "weight": 0.30},
                            ],
                        },
                    },
                    {"key": "rank", "type": "RANK", "source": "composite", "descending": True},
                    {"key": "win", "type": "SELECT", "source": "rank", "count": 10},
                ],
            },
            ensure_ascii=False,
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=definition,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                binding={
                    "stage_key": "院十佳",
                    "round_keys": {"r1": self.round.pk},
                    "audience_keys": {"audience1": "aud1set"},
                },
            )

    def _singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"aud-comp-{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"ac{index:03d}",
            college="学院",
            class_name="班级",
            song_name="歌",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def test_audience_score_flows_into_composite_result(self):
        from .services import maybe_resolve_checkpoints

        for singer in self.singers:
            ScoreRecord.objects.create(
                round=self.round,
                singer=singer,
                judge=self.judge,
                score=Decimal("100"),
                is_test_data=True,
            )
            AudienceScore.objects.create(
                activity=self.activity,
                stage_key="aud1set",
                singer=singer,
                score=Decimal("90"),
                is_test_data=True,
            )

        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["院十佳"])
        stage = StageResult.objects.get(activity=self.activity, stage_key="院十佳")
        self.assertEqual(stage.status, StageResult.Status.READY_TO_CONFIRM)
        composite = CompositeResult.objects.get(
            stage_result=stage, node_key="composite", singer=self.singers[0]
        )
        sources = {c["source"]: c for c in composite.components}
        self.assertEqual(Decimal(sources["assess_a1"]["value"]), Decimal("90"))
        self.assertEqual(Decimal(sources["assess_a1"]["contribution"]), Decimal("27.0"))
        self.assertEqual(composite.value, Decimal("97"))


class VoteBoundaryRegressionTests(TestCase):
    """M1-INTEGRATION-2: seal that raw vote counts never masquerade as a normalized score."""

    def setUp(self):
        from datetime import timedelta

        from django.utils import timezone
        from voting.models import VoteOption

        self.activity = Activity.objects.create(
            title="投票边界",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.singers = [
            SingerRegistration.objects.create(
                activity=self.activity,
                user=User.objects.create_user(username=f"vbound-{i}", password="pass"),
                name=f"选手{i}",
                student_id=f"vb{i:03d}",
                college="学院",
                class_name="班级",
                song_name="歌",
                is_test_data=True,
            )
            for i in range(1, 3)
        ]
        self.vs = VoteSession.objects.create(
            activity=self.activity,
            name="人气奖",
            passcode="x",
            start_time=timezone.now(),
            end_time=timezone.now() + timedelta(days=1),
            is_test_data=True,
        )
        self.options = {
            singer.pk: VoteOption.objects.create(
                vote_session=self.vs, singer=singer, is_test_data=True
            )
            for singer in self.singers
        }

    def test_source_vote_scores_returns_raw_counts_not_normalized(self):
        from voting.models import VoteRecord

        from .services import _source_vote_scores

        for i in range(3):
            VoteRecord.objects.create(
                vote_session=self.vs,
                vote_option=self.options[self.singers[0].pk],
                browser_session_key=f"mb{i}",
                ip_address="127.0.0.1",
                is_test_data=True,
            )
        VoteRecord.objects.create(
            vote_session=self.vs,
            vote_option=self.options[self.singers[1].pk],
            browser_session_key="mb-solo",
            ip_address="127.0.0.1",
            is_test_data=True,
        )
        out = _source_vote_scores(self.activity, {"vote_keys": {"pop": self.vs.pk}})
        self.assertEqual(out["pop"][str(self.singers[0].pk)], Decimal("3"))
        self.assertEqual(out["pop"][str(self.singers[1].pk)], Decimal("1"))


class ShadowRehearsalTests(TestCase):
    """2025 院十佳 Shadow Rehearsal — the composed production acceptance chain.

    Drives the M1-J front-half acceptance end-to-end on the frozen ``GOLDEN_SCHIDUI``:
    15 approved singers → R1/R2 judge scores + onsite audience → auto-READY ``stage1`` →
    核定 → roster-materialised R3 (Top10) → R3 scores → auto-READY ``stage2`` → 核定 →
    roster-materialised R4 (Top5) → R4 + audience4 → auto-READY final → 核定 → result
    board (three CONFIRMED stages) + host paper handcard (Top3). The 2025 校十佳 stress
    path (unresolved fallback) stays a separate must-fail gate in ``ruleset/test_compiler``;
    this class only proves the composed happy path is automatic, not hallucinated.
    """

    JUDGE_COUNT = 5

    def setUp(self):
        self.admin = User.objects.create_user(
            username="shadow-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="院十佳",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PENDING,
            is_test_mode=True,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity, name="院十佳规则", is_test_data=True, stage_key="院十佳"
        )
        self.judges = [
            Judge.objects.create(
                activity=self.activity, name=f"评委{chr(0x41 + i)}", is_active=True
            )
            for i in range(self.JUDGE_COUNT)
        ]
        self.singers = [self._singer(i) for i in range(1, 16)]
        # R1/R2 is the full APPROVED field; R3/R4 draw their entry from a confirmed stage.
        self.r1 = self._round(1, ContestRound.RosterSource.APPROVED)
        self.r2 = self._round(2, ContestRound.RosterSource.APPROVED)
        self.r3 = self._round(3, ContestRound.RosterSource.STAGE, roster_source_stage="stage1")
        self.r4 = self._round(4, ContestRound.RosterSource.STAGE, roster_source_stage="stage2")
        from ruleset.templates import GOLDEN_SCHIDUI

        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=GOLDEN_SCHIDUI,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
                binding={
                    "stage_key": "院十佳",
                    "round_keys": {
                        "r1": self.r1.pk,
                        "r2": self.r2.pk,
                        "r3": self.r3.pk,
                        "r4": self.r4.pk,
                    },
                    "audience_keys": {"audience1": "aud1set", "audience4": "aud4set"},
                },
            )

    def _singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"shadow-s{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"95{index:03d}",
            college="学院",
            class_name="班级",
            song_name="歌",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )

    def _round(self, sequence, roster_source, roster_source_stage=None):
        return ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name=f"轮次{sequence}",
            sequence=sequence,
            roster_source=roster_source,
            roster_source_stage=roster_source_stage or "",
        )

    def _seed_scores(self, round_, base, limit=None):
        singers = self.singers if limit is None else self.singers[:limit]
        for idx, singer in enumerate(singers):
            value = Decimal(base) - idx
            for judge in self.judges:
                ScoreRecord.objects.create(
                    round=round_, singer=singer, judge=judge, score=value, is_test_data=True
                )

    def _seed_audience(self, set_name, base, limit):
        for idx, singer in enumerate(self.singers[:limit]):
            AudienceScore.objects.create(
                activity=self.activity,
                stage_key=set_name,
                singer=singer,
                score=Decimal(base),
                is_test_data=True,
            )

    def _lock(self, round_):
        round_.is_locked = True
        round_.status = ContestRound.Status.LOCKED
        round_.save(update_fields=["is_locked", "status"])

    def _entry_ids(self, round_):
        return list(
            RoundEntry.objects.filter(round=round_)
            .order_by("pk")
            .values_list("singer_id", flat=True)
        )

    def test_composed_acceptance_chain(self):
        from .services import confirm_stage_result, maybe_resolve_checkpoints

        # 1) R1 + R2 + onsite audience → stage1 auto-READY; its advancees roll into R3.
        self._seed_scores(self.r1, 100)
        self._seed_scores(self.r2, 90)
        self._seed_audience("aud1set", 80, limit=15)
        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["stage1"])
        s1 = StageResult.objects.get(activity=self.activity, stage_key="stage1")
        self.assertEqual(s1.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(self._entry_ids(self.r3), [s.pk for s in self.singers[:10]])

        # 30/60/10 composite: a real AudienceScore (80) contributes 8.0 to a 92.0 total.
        c1 = {c.singer_id: c for c in s1.composites.filter(node_key="stage1")}
        comp1 = {c["source"]: c for c in c1[self.singers[0].pk].components}
        self.assertEqual(Decimal(comp1["assess_a1"]["value"]), Decimal("80"))
        self.assertEqual(Decimal(comp1["assess_a1"]["contribution"]), Decimal("8.0"))
        self.assertEqual(c1[self.singers[0].pk].value, Decimal("92.0"))

        # 2) 核定 stage1 → R3 (Top10) is now the locked-in entry roster for round 3.
        self._lock(self.r1)
        self._lock(self.r2)
        confirmed1 = confirm_stage_result(s1, confirmed_by=self.admin)
        self.assertEqual(confirmed1.status, StageResult.Status.CONFIRMED)

        # 3) R3 scores (only the 10 advancees) → stage2 auto-READY; advancees roll into R4.
        self._seed_scores(self.r3, 100, limit=10)
        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["stage2"])
        s2 = StageResult.objects.get(activity=self.activity, stage_key="stage2")
        self.assertEqual(s2.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertEqual(self._entry_ids(self.r4), [s.pk for s in self.singers[:5]])

        # 4) 核定 stage2 → R4 (Top5) is the locked-in entry roster for round 4.
        self._lock(self.r3)
        confirmed2 = confirm_stage_result(s2, confirmed_by=self.admin)
        self.assertEqual(confirmed2.status, StageResult.Status.CONFIRMED)

        # 5) R4 scores + audience4 (only the 5 advancees) → final auto-READY.
        self._seed_scores(self.r4, 90, limit=5)
        self._seed_audience("aud4set", 85, limit=5)
        self.assertEqual(maybe_resolve_checkpoints(self.activity, self.admin), ["stage3"])
        s3 = StageResult.objects.get(activity=self.activity, stage_key="stage3")
        self.assertEqual(s3.status, StageResult.Status.READY_TO_CONFIRM)

        # 6) 核定 final → result board shows all three CONFIRMED stages + host card Top3.
        self._lock(self.r4)
        confirmed3 = confirm_stage_result(s3, confirmed_by=self.admin)
        self.assertEqual(confirmed3.status, StageResult.Status.CONFIRMED)
        board = list(
            StageResult.objects.filter(activity=self.activity, status=StageResult.Status.CONFIRMED)
            .order_by("pk")
            .values_list("stage_key", flat=True)
        )
        self.assertEqual(board, ["stage1", "stage2", "stage3"])

        # 30/50/20 final composite: real audience4 (85) contributes 17.0 to a 92.0 total.
        c3 = {c.singer_id: c for c in s3.composites.filter(node_key="final")}
        comp3 = {c["source"]: c for c in c3[self.singers[0].pk].components}
        self.assertEqual(Decimal(comp3["assess_a4"]["value"]), Decimal("85"))
        self.assertEqual(Decimal(comp3["assess_a4"]["contribution"]), Decimal("17.0"))
        self.assertEqual(c3[self.singers[0].pk].value, Decimal("92.0"))

        # Host paper handcard: the advancing top3, read in rank order.
        handcard = (
            s3.decisions.filter(outcome_code="direct")
            .order_by("rank", "pk")
            .values_list("singer_id", flat=True)
        )
        self.assertEqual(list(handcard), [s.pk for s in self.singers[:3]])
