import json
import os
import shutil
import tempfile
import threading
import time
import zipfile
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, cast
from unittest import skipUnless
from unittest.mock import patch

from accounts.models import User
from archive.models import ArchivePackage
from common.authority import RULESET_FREEZE, STAGE_RESULT_CONFIRM, authority_write
from common.models import AuditLog
from core.models import Activity
from core.services import unarchive_activity
from django import forms
from django.core.exceptions import PermissionDenied
from django.core.files.storage import Storage
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections, connection, transaction
from django.db.models import Max
from django.http import FileResponse
from django.test import RequestFactory, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from exports.models import ArticleTemplate, GeneratedDocument
from exports.services import (
    archive_activity,
    build_archive_package,
    generate_persistent_document,
)
from farewell_show.models import Program
from files.models import MaterialCheck, MaterialRequirement, SubmissionFile
from files.services import reconcile_singer_material_checks, store_submission_file
from incidents.models import IncidentRecord
from openpyxl import load_workbook
from public_portal.models import PublicPost
from singer_contest.models import (
    Award,
    ContestRound,
    Judge,
    RoundEntry,
    RoundJudge,
    ScoreRecord,
    ScoreSummary,
    SingerRegistration,
    StageDecision,
    StageResult,
)
from singer_contest.services import apply_scores, prepare_round, stage_decisions_by_blocks
from voting.models import VoteBallot, VoteOption, VoteRecord, VoteSession

from staff_panel.forms import CREATE_PHASE_CHOICES, ActivityForm
from staff_panel.views import post_edit


def login_admin(client, user):
    client.force_login(user)
    session = client.session
    session["artflow_admin_verified"] = True
    session["artflow_admin_verified_at"] = timezone.now().isoformat()
    session.save()


def _close_file_response_resources(response: Any):
    closers = response._resource_closers
    filelikes = []
    filelike = response.file_to_stream
    if filelike is not None and hasattr(filelike, "read") and hasattr(filelike, "close"):
        filelikes.append(filelike)
    for closer in closers:
        resource = getattr(closer, "__self__", None)
        if (
            resource is not None
            and hasattr(resource, "read")
            and hasattr(resource, "close")
            and all(resource is not existing for existing in filelikes)
        ):
            filelikes.append(resource)
    for filelike in filelikes:
        filelike.close()
    response.file_to_stream = None
    closers.clear()


class PathlessStorage(Storage):
    """Delegate file reads while modeling a storage backend without paths."""

    def __init__(self, delegate: Storage):
        self.delegate = delegate

    def _open(self, name: str, mode: str = "rb"):
        return self.delegate.open(name, mode)

    def _save(self, name: str, content: Any):
        return self.delegate.save(name, content)

    def exists(self, name: str) -> bool:
        return self.delegate.exists(name)

    def url(self, name: str | None) -> str:
        return self.delegate.url(name)

    def path(self, name: str) -> str:
        raise NotImplementedError("This storage does not expose local paths.")


class StaffPanelSmokeTests(TestCase):
    def setUp(self):
        self.media_root = tempfile.mkdtemp()
        self.override = override_settings(MEDIA_ROOT=self.media_root)
        self.override.enable()
        self.admin = User.objects.create_user(
            username="admin_user",
            password="pass",
            role=User.Role.ADMIN,
        )
        self.staff = User.objects.create_user(
            username="staff_user",
            password="pass",
            role=User.Role.STAFF,
        )
        self.participant = User.objects.create_user(
            username="participant",
            password="pass",
            role=User.Role.PARTICIPANT,
        )
        self.singer_activity = Activity.objects.create(
            title="Singer Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.farewell_activity = Activity.objects.create(
            title="Farewell Show",
            activity_type=Activity.Type.FAREWELL_SHOW,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media_root, ignore_errors=True)

    def test_staff_form_pages_render(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        program = Program.objects.create(
            activity=self.farewell_activity,
            user=self.participant,
            name="Dance",
            program_type=Program.ProgramType.DANCE,
            contact_name="Li Hua",
            contact_phone="13800000000",
            class_name="CS1",
        )
        self.client.force_login(self.staff)
        paths = [
            reverse("staff:post_create"),
            reverse("staff:singer_registration_detail", args=[registration.pk]),
            reverse("staff:program_detail", args=[program.pk]),
            reverse("staff:round_create"),
            reverse("staff:vote_session_create"),
            reverse("staff:incident_create"),
        ]
        for path in paths:
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)

    def test_admin_can_create_activity_from_staff_panel(self):
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_create"),
            {
                "title": "New Contest",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.REGISTRATION_OPEN,
                "is_test_mode": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Activity.objects.filter(title="New Contest").exists())

    def test_activity_create_always_defaults_to_draft(self):
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_create"),
            {
                "title": "Started Mid-Lifecycle",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.REGISTRATION_OPEN,
                "is_test_mode": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        created = Activity.objects.get(title="Started Mid-Lifecycle")
        self.assertEqual(created.phase, Activity.Phase.DRAFT)

    def test_create_form_phase_choices_exclude_live_and_terminal_phases(self):
        form = ActivityForm(phase_choices=CREATE_PHASE_CHOICES)
        offered = [
            value
            for value, _label in cast(
                list[tuple[str, str]], cast(forms.ChoiceField, form.fields["phase"]).choices
            )
        ]
        for excluded in (
            Activity.Phase.LIVE,
            Activity.Phase.RESULTS_PENDING,
            Activity.Phase.RESULTS_PUBLISHED,
            Activity.Phase.ARCHIVED,
        ):
            self.assertNotIn(excluded, offered)

    def test_activity_edit_cannot_change_lifecycle(self):
        formal_activity = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=False,
        )
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_edit", args=[formal_activity.pk]),
            {
                "title": "Renamed Contest",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.REGISTRATION_OPEN,
                "is_test_mode": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        formal_activity.refresh_from_db()
        self.assertEqual(formal_activity.title, "Renamed Contest")
        self.assertEqual(formal_activity.data_lifecycle, Activity.DataLifecycle.FORMAL)
        self.assertFalse(formal_activity.is_test_mode)

    def test_formal_activity_cannot_be_reopened_in_test_mode(self):
        formal_activity = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=False,
        )
        login_admin(self.client, self.admin)
        self.client.raise_request_exception = False

        response = self.client.post(
            reverse("staff:activity_test_toggle", args=[formal_activity.pk])
        )

        self.assertEqual(response.status_code, 403)
        formal_activity.refresh_from_db()
        self.assertEqual(formal_activity.data_lifecycle, Activity.DataLifecycle.FORMAL)
        self.assertFalse(formal_activity.is_test_mode)

    def test_export_center_hides_test_controls_for_formal_activity(self):
        formal_activity = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=False,
        )
        login_admin(self.client, self.admin)

        response = self.client.get(reverse("staff:export_center"))

        self.assertNotContains(
            response,
            reverse("staff:activity_test_toggle", args=[formal_activity.pk]),
        )
        self.assertNotContains(
            response,
            reverse("staff:activity_clear_test_data", args=[formal_activity.pk]),
        )

    def test_staff_cannot_create_activity(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:activity_create"))
        self.assertEqual(response.status_code, 403)

    def test_public_homepage_uses_published_posts(self):
        PublicPost.objects.create(
            title="Registration Open",
            post_type=PublicPost.PostType.ANNOUNCEMENT,
            status=PublicPost.Status.PUBLISHED,
            published_at=timezone.now(),
        )
        response = self.client.get(reverse("public_portal:home"))
        self.assertContains(response, "Registration Open")

    def test_singer_registration_uploads_material_and_checks_completeness(self):
        self.client.force_login(self.participant)
        response = self.client.post(
            reverse("singer_contest:apply"),
            {
                "activity_id": self.singer_activity.pk,
                "name": "Li Hua",
                "student_id": "20260001",
                "college": "Info",
                "class_name": "CS1",
                "phone": "13800000000",
                "song_name": "Song",
                "accompaniment": SimpleUploadedFile(
                    "song.mp3", b"audio", content_type="audio/mpeg"
                ),
            },
        )
        self.assertEqual(response.status_code, 302)
        registration = SingerRegistration.objects.get(user=self.participant)
        self.assertTrue(
            registration.files.filter(file_purpose=SubmissionFile.Purpose.ACCOMPANIMENT).exists()
        )
        self.assertTrue(
            MaterialCheck.objects.filter(
                singer_registration=registration,
                item_name="伴奏文件",
                status=MaterialCheck.Status.UPLOADED,
            ).exists()
        )

    def test_farewell_program_uploads_material_and_checks_completeness(self):
        self.client.force_login(self.participant)
        response = self.client.post(
            reverse("farewell_show:apply"),
            {
                "activity_id": self.farewell_activity.pk,
                "name": "Dance",
                "program_type": Program.ProgramType.DANCE,
                "contact_name": "Li Hua",
                "contact_phone": "13800000000",
                "class_name": "CS1",
                "accompaniment": SimpleUploadedFile(
                    "dance.mp3", b"audio", content_type="audio/mpeg"
                ),
            },
        )
        self.assertEqual(response.status_code, 302)
        program = Program.objects.get(user=self.participant)
        self.assertTrue(
            program.files.filter(file_purpose=SubmissionFile.Purpose.ACCOMPANIMENT).exists()
        )
        self.assertTrue(
            MaterialCheck.objects.filter(
                program=program,
                item_name="伴奏文件",
                status=MaterialCheck.Status.UPLOADED,
            ).exists()
        )

    def test_scoring_and_voting_basic_flow(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            scoring_mode=ContestRound.ScoringMode.AVERAGE,
            advance_count=1,
        )
        prepare_round(round_, self.staff)
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:round_scores_api", args=[round_.pk]),
            data=json.dumps(
                {
                    "base_version": 0,
                    "cells": [
                        {
                            "singer_id": registration.pk,
                            "judge_id": judge.pk,
                            "score": "91",
                        }
                    ],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ScoreSummary.objects.get(round=round_, singer=registration).rank, 1)

        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
        )
        option = VoteOption.objects.create(vote_session=vote_session, singer=registration)
        visitor = self.client_class()
        self.assertEqual(
            visitor.post(
                reverse("voting:vote_entry", args=[vote_session.pk]), {"passcode": "1234"}
            ).status_code,
            302,
        )
        self.assertEqual(
            visitor.post(
                reverse("voting:vote_cast", args=[vote_session.pk]),
                {"selected_option": [str(option.pk)]},
            ).status_code,
            302,
        )
        self.assertEqual(VoteRecord.objects.filter(vote_session=vote_session).count(), 1)

        qr_page = self.client.get(reverse("staff:qr_generate", args=[self.singer_activity.pk]))
        self.assertEqual(qr_page.status_code, 200)
        qr_image = self.client.get(
            reverse("staff:qr_image", args=[self.singer_activity.pk, "registration"])
        )
        self.assertEqual(qr_image.status_code, 200)
        self.assertEqual(qr_image["Content-Type"], "image/png")

    def test_round_score_entry_requires_prepared_round(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260007",
            college="Info",
            class_name="CS1",
            phone="13800000006",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.client.force_login(self.staff)

        entry_response = self.client.get(reverse("staff:round_score_entry", args=[round_.pk]))

        response = self.client.post(
            reverse("staff:round_score_entry", args=[round_.pk]),
            {f"score_{registration.pk}_{judge.pk}": "91"},
        )

        self.assertContains(entry_response, "请先准备比赛轮次")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "请先准备比赛轮次")
        self.assertFalse(ScoreRecord.objects.filter(round=round_).exists())

    def test_staff_can_prepare_round(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Ready Singer",
            student_id="20260015",
            college="Info",
            class_name="CS1",
            phone="13800000014",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Ready Judge")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.client.force_login(self.staff)

        get_response = self.client.get(reverse("staff:round_prepare", args=[round_.pk]))
        response = self.client.post(reverse("staff:round_prepare", args=[round_.pk]))

        self.assertEqual(get_response.status_code, 405)
        self.assertRedirects(response, reverse("staff:round_list"))
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.PREPARED)
        self.assertFalse(round_.is_locked)
        self.assertEqual(
            list(round_.entries.values_list("singer_id", flat=True)), [registration.pk]
        )
        self.assertEqual(list(round_.round_judges.values_list("judge_id", flat=True)), [judge.pk])
        self.assertTrue(
            AuditLog.objects.filter(
                operator=self.staff,
                action_type=AuditLog.ActionType.UPDATE_STATUS,
                target=f"ContestRound:{round_.pk}",
            ).exists()
        )

    def test_formal_score_update_stays_formal(self):
        formal_activity = Activity.objects.create(
            title="Formal Score Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        registration = SingerRegistration.objects.create(
            activity=formal_activity,
            user=self.participant,
            name="Formal Scored Singer",
            student_id="20260019",
            college="Info",
            class_name="CS1",
            phone="13800000018",
            song_name="Formal Score Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=False,
        )
        judge = Judge.objects.create(activity=formal_activity, name="Formal Score Judge")
        round_ = ContestRound.objects.create(
            activity=formal_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        ScoreRecord.objects.create(
            round=round_,
            singer=registration,
            judge=judge,
            score=91,
            is_test_data=True,
        )

        apply_scores(round_, {(registration.pk, judge.pk): "91"}, self.staff)

        self.assertFalse(
            ScoreRecord.objects.get(round=round_, singer=registration, judge=judge).is_test_data
        )
        self.assertFalse(ScoreSummary.objects.get(round=round_, singer=registration).is_test_data)

    def test_unprepared_round_cannot_accept_scores(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Draft Singer",
            student_id="20260016",
            college="Info",
            class_name="CS1",
            phone="13800000015",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Draft Judge")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("staff:round_score_entry", args=[round_.pk]),
            {f"score_{registration.pk}_{judge.pk}": "91"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "请先准备比赛轮次")
        self.assertFalse(ScoreRecord.objects.filter(round=round_).exists())

    def test_round_lock_sets_locked_status(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Lockable Singer",
            student_id="20260017",
            college="Info",
            class_name="CS1",
            phone="13800000016",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Lockable Judge")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        ScoreRecord.objects.create(round=round_, singer=registration, judge=judge, score=91)
        self.client.force_login(self.staff)

        response = self.client.post(reverse("staff:round_lock", args=[round_.pk]))

        self.assertRedirects(response, reverse("staff:round_ranking", args=[round_.pk]))
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.LOCKED)
        self.assertTrue(round_.is_locked)
        self.assertTrue(
            AuditLog.objects.filter(
                operator=self.staff,
                action_type=AuditLog.ActionType.RELOCK_RESULT,
                target=f"ContestRound:{round_.pk}",
            ).exists()
        )

    def test_round_unlock_preserves_snapshots(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Unlockable Singer",
            student_id="20260018",
            college="Info",
            class_name="CS1",
            phone="13800000017",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Unlockable Judge")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        ScoreRecord.objects.create(round=round_, singer=registration, judge=judge, score=91)
        login_admin(self.client, self.admin)
        self.client.post(reverse("staff:round_lock", args=[round_.pk]))
        snapshot_counts = (round_.entries.count(), round_.round_judges.count())

        response = self.client.post(
            reverse("staff:round_unlock", args=[round_.pk]), {"note": "Correction needed"}
        )

        self.assertRedirects(response, reverse("staff:round_ranking", args=[round_.pk]))
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.SCORING)
        self.assertFalse(round_.is_locked)
        self.assertEqual((round_.entries.count(), round_.round_judges.count()), snapshot_counts)
        self.assertTrue(
            AuditLog.objects.filter(
                operator=self.admin,
                action_type=AuditLog.ActionType.UNLOCK_RESULT,
                target=f"ContestRound:{round_.pk}",
                note="Correction needed",
            ).exists()
        )

    def test_round_unlock_rejects_downstream_active(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Downstream Singer",
            student_id="20260020",
            college="Info",
            class_name="CS1",
            phone="13800000019",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Downstream Judge")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        ScoreRecord.objects.create(round=round_, singer=registration, judge=judge, score=91)
        login_admin(self.client, self.admin)
        self.client.post(reverse("staff:round_lock", args=[round_.pk]))

        ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.SEMI_FINAL,
            status=ContestRound.Status.PREPARED,
        )

        response = self.client.post(
            reverse("staff:round_unlock", args=[round_.pk]), {"note": "Correction needed"}
        )

        self.assertEqual(response.status_code, 403)
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.LOCKED)
        self.assertTrue(round_.is_locked)

    def test_round_reset_prepared_round_to_draft(self):
        SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Reset Singer",
            student_id="20260019",
            college="Info",
            class_name="CS1",
            phone="13800000017",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        Judge.objects.create(activity=self.singer_activity, name="Reset Judge")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        login_admin(self.client, self.admin)

        response = self.client.post(
            reverse("staff:round_reset", args=[round_.pk]), {"reason": "Admin unwind"}
        )

        self.assertRedirects(response, reverse("staff:round_list"))
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.DRAFT)
        self.assertFalse(round_.entries.exists())
        self.assertFalse(round_.round_judges.exists())
        self.assertTrue(
            AuditLog.objects.filter(
                operator=self.admin,
                target=f"ContestRound:{round_.pk}",
                note__contains="Admin unwind",
            ).exists()
        )

    def test_round_reset_requires_reason(self):
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        login_admin(self.client, self.admin)

        response = self.client.post(reverse("staff:round_reset", args=[round_.pk]), {"reason": ""})

        self.assertEqual(response.status_code, 403)

    def test_score_entry_and_template_use_prepared_round_snapshots(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Prepared Singer",
            student_id="20260008",
            college="Info",
            class_name="CS1",
            phone="13800000007",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Prepared Judge")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        late_user = User.objects.create_user(username="participant-late", password="pass")
        SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=late_user,
            name="Late Singer",
            student_id="20260009",
            college="Info",
            class_name="CS1",
            phone="13800000008",
            song_name="Late Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        Judge.objects.create(activity=self.singer_activity, name="Late Judge")
        self.client.force_login(self.staff)

        entry_response = self.client.get(reverse("staff:round_score_entry", args=[round_.pk]))
        template_response = self.client.get(reverse("staff:excel_score_template", args=[round_.pk]))

        self.assertContains(entry_response, registration.name)
        self.assertContains(entry_response, judge.name)
        self.assertNotContains(entry_response, "Late Singer")
        self.assertNotContains(entry_response, "Late Judge")
        workbook = load_workbook(BytesIO(template_response.content))
        worksheet = workbook.active
        assert worksheet is not None
        self.assertEqual(
            list(worksheet.values),
            [
                ("选手ID", "姓名", f"J{judge.pk} {judge.name}"),
                (registration.pk, registration.name, None),
            ],
        )
        self.assertIn("ArtFlowMeta", workbook.sheetnames)
        self.assertEqual(workbook["ArtFlowMeta"].sheet_state, "hidden")

    def test_locked_activity_blocks_staff_score_entry(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        self.singer_activity.is_locked = True
        self.singer_activity.save(update_fields=["is_locked"])
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:round_scores_api", args=[round_.pk]),
            data=json.dumps(
                {
                    "base_version": 0,
                    "cells": [
                        {
                            "singer_id": registration.pk,
                            "judge_id": judge.pk,
                            "score": "91",
                        }
                    ],
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(ScoreRecord.objects.exists())

    def test_locked_activity_blocks_award_creation(self):
        self.singer_activity.is_locked = True
        self.singer_activity.save()
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:award_create"),
            {
                "activity_id": self.singer_activity.pk,
                "singer_id": registration.pk,
                "name": "Top Singer",
            },
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(Award.objects.exists())

    def test_activity_unlock_does_not_unlock_locked_round(self):
        contest_round = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            status=ContestRound.Status.LOCKED,
            is_locked=True,
        )
        login_admin(self.client, self.admin)

        self.client.post(reverse("staff:activity_lock", args=[self.singer_activity.pk]))
        response = self.client.post(
            reverse("staff:activity_unlock", args=[self.singer_activity.pk]),
            {"note": "Activity correction"},
        )

        self.assertEqual(response.status_code, 302)
        contest_round.refresh_from_db()
        self.assertEqual(contest_round.status, ContestRound.Status.LOCKED)
        self.assertTrue(contest_round.is_locked)
        self.client.force_login(self.staff)
        self.assertEqual(
            self.client.post(
                reverse("staff:round_scores_api", args=[contest_round.pk]),
                data=json.dumps({"base_version": 0, "cells": []}),
                content_type="application/json",
            ).status_code,
            400,
        )

    def test_activity_unlock_does_not_unlock_locked_vote_session(self):
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Locked popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_locked=True,
            is_open=False,
        )
        login_admin(self.client, self.admin)

        self.client.post(reverse("staff:activity_lock", args=[self.singer_activity.pk]))
        response = self.client.post(
            reverse("staff:activity_unlock", args=[self.singer_activity.pk]),
            {"note": "Activity correction"},
        )

        self.assertEqual(response.status_code, 302)
        vote_session.refresh_from_db()
        self.assertTrue(vote_session.is_locked)
        self.assertFalse(vote_session.is_open)
        self.client.force_login(self.staff)
        self.assertEqual(
            self.client.post(
                reverse("staff:vote_session_open", args=[vote_session.pk])
            ).status_code,
            403,
        )

    def test_activity_lock_does_not_destroy_child_open_state(self):
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Open popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
            is_locked=False,
        )
        login_admin(self.client, self.admin)

        response = self.client.post(reverse("staff:activity_lock", args=[self.singer_activity.pk]))

        self.assertEqual(response.status_code, 302)
        vote_session.refresh_from_db()
        self.assertTrue(vote_session.is_open)
        self.assertFalse(vote_session.is_locked)
        self.client.force_login(self.staff)
        self.assertEqual(
            self.client.post(
                reverse("staff:vote_session_close", args=[vote_session.pk])
            ).status_code,
            403,
        )

    def test_admin_can_unlock_vote_session_and_audit_is_recorded(self):
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_locked=True,
        )
        self.client.force_login(self.staff)
        self.assertEqual(
            self.client.post(
                reverse("staff:vote_session_unlock", args=[vote_session.pk])
            ).status_code,
            403,
        )

        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:vote_session_unlock", args=[vote_session.pk]), {"note": "fix typo"}
        )
        self.assertEqual(response.status_code, 302)
        vote_session.refresh_from_db()
        self.assertFalse(vote_session.is_locked)
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UNLOCK_RESULT,
                target__contains="VoteSession",
                note="fix typo",
            ).exists()
        )

    def test_admin_can_unlock_vote_session_without_note(self):
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity no note",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_locked=True,
        )
        login_admin(self.client, self.admin)

        response = self.client.post(reverse("staff:vote_session_unlock", args=[vote_session.pk]))

        self.assertEqual(response.status_code, 302)
        vote_session.refresh_from_db()
        self.assertFalse(vote_session.is_locked)
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UNLOCK_RESULT,
                target=f"VoteSession:{vote_session.pk}",
            ).exists()
        )

    def test_vote_session_open_and_close_update_state(self):
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Toggleable",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_locked=False,
            is_open=False,
        )
        login_admin(self.client, self.admin)

        self.client.post(reverse("staff:vote_session_open", args=[vote_session.pk]))
        vote_session.refresh_from_db()
        self.assertTrue(vote_session.is_open)

        self.client.post(reverse("staff:vote_session_close", args=[vote_session.pk]))
        vote_session.refresh_from_db()
        self.assertFalse(vote_session.is_open)
        self.assertFalse(vote_session.is_locked)

    def test_admin_can_unlock_activity_without_note(self):
        login_admin(self.client, self.admin)
        self.client.post(reverse("staff:activity_lock", args=[self.singer_activity.pk]))

        response = self.client.post(
            reverse("staff:activity_unlock", args=[self.singer_activity.pk])
        )

        self.assertEqual(response.status_code, 302)
        self.singer_activity.refresh_from_db()
        self.assertFalse(self.singer_activity.is_locked)

    def test_round_unlock_rejects_while_activity_locked(self):
        ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            status=ContestRound.Status.LOCKED,
            is_locked=True,
        )
        login_admin(self.client, self.admin)
        self.client.post(reverse("staff:activity_lock", args=[self.singer_activity.pk]))

        response = self.client.post(
            reverse(
                "staff:round_unlock",
                args=[ContestRound.objects.get(activity=self.singer_activity).pk],
            ),
        )

        self.assertEqual(response.status_code, 403)

    def test_test_cleanup_cannot_delete_formal_registration(self):
        self.singer_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        formal_registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Formal",
            student_id="20260002",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Formal Song",
            is_test_data=False,
        )
        test_user = User.objects.create_user(username="participant-test", password="pass")
        test_registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=test_user,
            name="Test",
            student_id="20260003",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Test Song",
            is_test_data=True,
        )
        VoteSession.objects.create(
            activity=self.singer_activity,
            name="Formal Vote",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_test_data=False,
        )
        VoteSession.objects.create(
            activity=self.singer_activity,
            name="Test Vote",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_test_data=True,
        )
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_clear_test_data", args=[self.singer_activity.pk])
        )
        self.assertEqual(response.status_code, 302)
        self.assertTrue(SingerRegistration.objects.filter(pk=formal_registration.pk).exists())
        self.assertFalse(SingerRegistration.objects.filter(pk=test_registration.pk).exists())
        self.assertTrue(VoteSession.objects.filter(name="Formal Vote").exists())
        self.assertFalse(VoteSession.objects.filter(name="Test Vote").exists())

    def test_cleanup_rejects_formal_vote_dependents_under_test_session(self):
        self.singer_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Test Session Singer",
            student_id="20260021",
            college="Info",
            class_name="CS1",
            phone="13800000020",
            song_name="Test Session Song",
            is_test_data=True,
        )
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Test Vote",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_test_data=True,
        )
        option = VoteOption.objects.create(
            vote_session=vote_session,
            singer=registration,
            is_test_data=False,
        )
        ballot = VoteBallot.objects.create(
            vote_session=vote_session,
            browser_session_key="formal-ballot",
            ip_address="127.0.0.1",
            is_test_data=False,
        )
        record = VoteRecord.objects.create(
            ballot=ballot,
            vote_session=vote_session,
            vote_option=option,
            browser_session_key="formal-ballot",
            ip_address="127.0.0.1",
            is_test_data=False,
        )
        login_admin(self.client, self.admin)
        self.client.raise_request_exception = False

        response = self.client.post(
            reverse("staff:activity_clear_test_data", args=[self.singer_activity.pk])
        )

        self.assertEqual(response.status_code, 403)
        self.assertTrue(VoteSession.objects.filter(pk=vote_session.pk).exists())
        self.assertFalse(VoteOption.objects.get(pk=option.pk).is_test_data)
        self.assertFalse(VoteBallot.objects.get(pk=ballot.pk).is_test_data)
        self.assertFalse(VoteRecord.objects.get(pk=record.pk).is_test_data)

    def test_cleanup_rejects_formal_vote_graph_referencing_test_singer(self):
        self.singer_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        test_registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Test Singer With Formal Vote",
            student_id="20260023",
            college="Info",
            class_name="CS1",
            phone="13800000022",
            song_name="Test Singer Song",
            is_test_data=True,
        )
        formal_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Formal Vote For Test Singer",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_test_data=False,
        )
        formal_option = VoteOption.objects.create(
            vote_session=formal_session,
            singer=test_registration,
            is_test_data=False,
        )
        formal_ballot = VoteBallot.objects.create(
            vote_session=formal_session,
            browser_session_key="formal-singer-ballot",
            ip_address="127.0.0.1",
            is_test_data=False,
        )
        formal_record = VoteRecord.objects.create(
            ballot=formal_ballot,
            vote_session=formal_session,
            vote_option=formal_option,
            browser_session_key="formal-singer-ballot",
            ip_address="127.0.0.1",
            is_test_data=False,
        )
        login_admin(self.client, self.admin)
        self.client.raise_request_exception = False

        response = self.client.post(
            reverse("staff:activity_clear_test_data", args=[self.singer_activity.pk])
        )

        self.assertEqual(response.status_code, 403)
        self.assertTrue(SingerRegistration.objects.filter(pk=test_registration.pk).exists())
        self.assertFalse(VoteSession.objects.get(pk=formal_session.pk).is_test_data)
        self.assertFalse(VoteOption.objects.get(pk=formal_option.pk).is_test_data)
        self.assertFalse(VoteBallot.objects.get(pk=formal_ballot.pk).is_test_data)
        self.assertFalse(VoteRecord.objects.get(pk=formal_record.pk).is_test_data)

    def test_cleanup_rejects_formal_file_under_test_owner(self):
        self.singer_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Test File Singer",
            student_id="20260022",
            college="Info",
            class_name="CS1",
            phone="13800000021",
            song_name="Test File Song",
            is_test_data=True,
        )
        submission = SubmissionFile.objects.create(
            singer_registration=registration,
            file=SimpleUploadedFile("formal.mp3", b"audio", content_type="audio/mpeg"),
            original_name="formal.mp3",
            file_size=5,
            file_purpose=SubmissionFile.Purpose.ACCOMPANIMENT,
            is_test_data=False,
            uploaded_by=self.participant,
        )
        login_admin(self.client, self.admin)
        self.client.raise_request_exception = False

        response = self.client.post(
            reverse("staff:activity_clear_test_data", args=[self.singer_activity.pk])
        )

        self.assertEqual(response.status_code, 403)
        self.assertTrue(SingerRegistration.objects.filter(pk=registration.pk).exists())
        self.assertFalse(SubmissionFile.objects.get(pk=submission.pk).is_test_data)

    def test_formal_file_creation_stays_formal(self):
        formal_activity = Activity.objects.create(
            title="Formal File Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=False,
        )
        registration = SingerRegistration.objects.create(
            activity=formal_activity,
            user=self.participant,
            name="Formal File Singer",
            student_id="20260020",
            college="Info",
            class_name="CS1",
            phone="13800000019",
            song_name="Formal File Song",
            is_test_data=False,
        )

        submission = store_submission_file(
            owner=registration,
            uploaded_file=SimpleUploadedFile("formal.mp3", b"audio", content_type="audio/mpeg"),
            purpose=SubmissionFile.Purpose.ACCOMPANIMENT,
            uploaded_by=self.participant,
        )

        self.assertFalse(submission.is_test_data)

    def test_leaving_test_mode_rejects_residual_test_data(self):
        self.singer_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        test_registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Residual Test",
            student_id="20260007",
            college="Info",
            class_name="CS1",
            phone="13800000006",
            song_name="Test Song",
            is_test_data=True,
        )
        login_admin(self.client, self.admin)

        response = self.client.post(
            reverse("staff:activity_test_toggle", args=[self.singer_activity.pk])
        )

        self.assertEqual(response.status_code, 403)
        self.singer_activity.refresh_from_db()
        self.assertTrue(self.singer_activity.is_test_mode)
        self.assertTrue(SingerRegistration.objects.filter(pk=test_registration.pk).exists())

    def test_clear_test_data_removes_test_file_object(self):
        self.singer_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        test_registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="File Test",
            student_id="20260012",
            college="Info",
            class_name="CS1",
            phone="13800000011",
            song_name="Test Song",
            is_test_data=True,
        )
        submission = store_submission_file(
            owner=test_registration,
            uploaded_file=SimpleUploadedFile("test.mp3", b"audio", content_type="audio/mpeg"),
            purpose=SubmissionFile.Purpose.ACCOMPANIMENT,
            uploaded_by=self.participant,
        )
        stored_name = submission.file.name
        login_admin(self.client, self.admin)

        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(
                reverse("staff:activity_clear_test_data", args=[self.singer_activity.pk])
            )

        self.assertEqual(response.status_code, 302)
        self.assertFalse(submission.file.storage.exists(stored_name))

    def test_file_response_cleanup_closes_file_without_closing_response(self):
        response = FileResponse(BytesIO(b"audio"))
        filelike = response.file_to_stream
        assert filelike is not None

        with patch("django.core.signals.request_finished.send") as request_finished:
            _close_file_response_resources(response)

        self.assertTrue(filelike.closed)
        self.assertIsNone(response.file_to_stream)
        self.assertEqual(getattr(response, "_resource_closers"), [])
        self.assertFalse(response.closed)
        request_finished.assert_not_called()
        self.assertEqual(Activity.objects.count(), 2)

    def test_submission_file_media_url_is_owner_or_staff_only(self):
        other_participant = User.objects.create_user(
            username="other",
            password="pass",
            role=User.Role.PARTICIPANT,
        )
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        uploaded = SubmissionFile.objects.create(
            singer_registration=registration,
            file=SimpleUploadedFile("song.mp3", b"audio"),
            original_name="song.mp3",
            file_size=5,
            uploaded_by=self.participant,
        )
        self.client.force_login(other_participant)
        self.assertEqual(self.client.get(uploaded.file.url).status_code, 403)

        self.client.force_login(self.participant)
        owner_response = self.client.get(uploaded.file.url)
        try:
            self.assertEqual(owner_response.status_code, 200)
        finally:
            _close_file_response_resources(owner_response)

        self.client.force_login(self.staff)
        staff_response = self.client.get(uploaded.file.url)
        try:
            self.assertEqual(staff_response.status_code, 200)
        finally:
            _close_file_response_resources(staff_response)

    def test_packages_include_contest_and_farewell_operational_indexes(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            advance_count=1,
        )
        ScoreRecord.objects.create(round=round_, singer=registration, judge=judge, score=91)
        ScoreSummary.objects.create(
            round=round_, singer=registration, average_score=91, rank=1, is_advanced=True
        )
        Award.objects.create(activity=self.singer_activity, singer=registration, name="Top Singer")
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
        )
        VoteOption.objects.create(vote_session=vote_session, singer=registration)
        SubmissionFile.objects.create(
            singer_registration=registration,
            file=SimpleUploadedFile("song.mp3", b"audio"),
            original_name="song.mp3",
            file_size=5,
            uploaded_by=self.participant,
        )
        PublicPost.objects.create(
            title="Result",
            post_type=PublicPost.PostType.RESULT_PUBLICATION,
            status=PublicPost.Status.PUBLISHED,
            related_activity=self.singer_activity,
            published_at=timezone.now(),
        )
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:archive_package_create", args=[self.singer_activity.pk])
        )
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.content)) as zf:
            names = set(zf.namelist())
        self.assertIn("score_results.xlsx", names)
        self.assertIn("vote_results.xlsx", names)
        self.assertIn("award_list.xlsx", names)
        self.assertIn("attachment_index.xlsx", names)
        self.assertIn("public_content_index.xlsx", names)

        Program.objects.create(
            activity=self.farewell_activity,
            user=self.participant,
            name="Dance",
            program_type=Program.ProgramType.DANCE,
            contact_name="Li Hua",
            contact_phone="13800000000",
            class_name="CS1",
            sort_order=1,
        )
        response = self.client.get(
            reverse("staff:execution_package", args=[self.farewell_activity.pk])
        )
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.content)) as zf:
            names = set(zf.namelist())
        self.assertIn("program_list.xlsx", names)
        self.assertIn("contact_list.xlsx", names)
        self.assertIn("material_checklist.xlsx", names)

    def test_archive_package_excludes_test_data_from_formal_indexes(self):
        formal = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Formal Singer",
            student_id="20260008",
            college="Info",
            class_name="CS1",
            phone="13800000007",
            song_name="Formal Song",
            is_test_data=False,
        )
        test_user = User.objects.create_user(username="participant-test", password="pass")
        SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=test_user,
            name="Test Singer",
            student_id="20260009",
            college="Info",
            class_name="CS1",
            phone="13800000008",
            song_name="Test Song",
            is_test_data=True,
        )
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("staff:archive_package_create", args=[self.singer_activity.pk])
        )

        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            workbook = load_workbook(BytesIO(archive.read("registration_list.xlsx")))
        worksheet = workbook.active
        assert worksheet is not None
        values = [cell.value for row in worksheet.iter_rows() for cell in row]
        self.assertIn(formal.name, values)
        self.assertNotIn("Test Singer", values)

    def test_formal_archive_excludes_test_generated_document(self):
        formal_activity = Activity.objects.create(
            title="Formal Singer Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=False,
        )
        formal_document = GeneratedDocument.objects.create(
            activity=formal_activity,
            title="Formal announcement",
            file=SimpleUploadedFile("formal.docx", b"formal document"),
            created_by=self.staff,
            is_test_data=False,
        )
        test_document = GeneratedDocument.objects.create(
            activity=formal_activity,
            title="Test announcement",
            file=SimpleUploadedFile("test.docx", b"test document"),
            created_by=self.staff,
            is_test_data=True,
        )
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("staff:archive_package_create", args=[formal_activity.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(GeneratedDocument.objects.filter(pk=formal_document.pk).exists())
        self.assertTrue(GeneratedDocument.objects.filter(pk=test_document.pk).exists())
        self.assertTrue(formal_document.file.storage.exists(formal_document.file.name or ""))
        self.assertTrue(test_document.file.storage.exists(test_document.file.name or ""))
        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            names = set(archive.namelist())
        self.assertIn(f"推文_{formal_document.pk}.docx", names)
        self.assertNotIn(f"推文_{test_document.pk}.docx", names)

    def test_formal_archive_reads_generated_document_without_storage_path(self):
        document = GeneratedDocument.objects.create(
            activity=self.singer_activity,
            title="Storage-backed announcement",
            file=SimpleUploadedFile("storage.docx", b"storage document"),
            created_by=self.staff,
            is_test_data=False,
        )
        file_field = GeneratedDocument._meta.get_field("file")
        pathless_storage = PathlessStorage(file_field.storage)
        self.client.force_login(self.staff)

        with patch.object(file_field, "storage", pathless_storage):
            response = self.client.post(
                reverse("staff:archive_package_create", args=[self.singer_activity.pk])
            )

        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            self.assertEqual(
                archive.read(f"推文_{document.pk}.docx"),
                b"storage document",
            )

    def _make_finalized_round_activity(self):
        activity = Activity.objects.create(
            title="Archive Ready Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REHEARSAL,
            is_test_mode=False,
        )
        singer = SingerRegistration.objects.create(
            activity=activity,
            user=self.participant,
            name="Final Singer",
            student_id="20260088",
            college="Info",
            class_name="CS1",
            phone="13800000001",
            song_name="Final Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=False,
        )
        judge = Judge.objects.create(activity=activity, name="Judge F", is_active=True)
        contest_round = ContestRound.objects.create(
            activity=activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            advance_count=1,
        )
        prepare_round(contest_round, self.staff)
        activity.phase = Activity.Phase.RESULTS_PUBLISHED
        activity.save(update_fields=["phase"])
        locked_round = ContestRound.objects.get(pk=contest_round.pk)
        ScoreRecord.objects.create(round=locked_round, singer=singer, judge=judge, score=95)
        locked_round.status = ContestRound.Status.LOCKED
        locked_round.is_locked = True
        locked_round.save(update_fields=["status", "is_locked"])
        return activity

    def test_execution_package_omits_archive_only_sheets(self):
        activity = Activity.objects.create(
            title="Exec Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REHEARSAL,
            is_test_mode=False,
        )
        SingerRegistration.objects.create(
            activity=activity,
            user=self.participant,
            name="Exec Singer",
            student_id="20260077",
            college="Info",
            class_name="CS1",
            phone="13800000002",
            song_name="Exec Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=False,
        )
        Judge.objects.create(activity=activity, name="Judge X", is_active=True)
        contest_round = ContestRound.objects.create(
            activity=activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            advance_count=1,
        )
        prepare_round(contest_round, self.staff)
        vote_session = VoteSession.objects.create(
            activity=activity,
            name="Pop",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
            is_test_data=False,
        )
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:execution_package", args=[activity.pk]))
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.content)) as zf:
            names = set(zf.namelist())
        self.assertIn("registration_list.xlsx", names)
        self.assertIn("contact_list.xlsx", names)
        self.assertIn("program_list.xlsx", names)
        self.assertIn("material_checklist.xlsx", names)
        self.assertIn("missing_materials.xlsx", names)
        self.assertIn("host_script.xlsx", names)
        self.assertIn("incident_empty_form.xlsx", names)
        self.assertIn(f"score_template_{contest_round.pk}.xlsx", names)
        self.assertIn(f"vote_qr_{vote_session.pk}.png", names)
        for omitted in [
            "score_results.xlsx",
            "vote_results.xlsx",
            "award_list.xlsx",
            "attachment_index.xlsx",
            "public_content_index.xlsx",
        ]:
            self.assertNotIn(omitted, names)

    def test_archive_activity_rejects_test_activity(self):
        activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PUBLISHED,
            is_test_mode=True,
        )
        with self.assertRaises(PermissionDenied):
            archive_activity(activity, self.admin)

    def test_archive_activity_rejects_residual_test_data(self):
        activity = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PUBLISHED,
            is_test_mode=False,
        )
        SingerRegistration.objects.create(
            activity=activity,
            user=self.participant,
            name="Test Singer",
            student_id="20260066",
            college="Info",
            class_name="CS1",
            phone="13800000003",
            song_name="Test Song",
            is_test_data=True,
        )
        with self.assertRaises(PermissionDenied):
            archive_activity(activity, self.admin)

    def test_archive_activity_rejects_unlocked_round(self):
        activity = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REHEARSAL,
            is_test_mode=False,
        )
        SingerRegistration.objects.create(
            activity=activity,
            user=self.participant,
            name="Singer",
            student_id="20260065",
            college="Info",
            class_name="CS1",
            phone="13800000004",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=False,
        )
        Judge.objects.create(activity=activity, name="Judge", is_active=True)
        contest_round = ContestRound.objects.create(
            activity=activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            advance_count=1,
        )
        prepare_round(contest_round, self.staff)
        activity.phase = Activity.Phase.RESULTS_PUBLISHED
        activity.save(update_fields=["phase"])
        with self.assertRaises(PermissionDenied):
            archive_activity(activity, self.admin)

    def test_archive_activity_rejects_missing_scores(self):
        activity = self._make_finalized_round_activity()
        ScoreRecord.objects.all().delete()
        with self.assertRaises(PermissionDenied):
            archive_activity(activity, self.admin)

    def test_archive_activity_rejects_unlocked_vote(self):
        activity = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PUBLISHED,
            is_test_mode=False,
        )
        VoteSession.objects.create(
            activity=activity,
            name="Pop",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=False,
            is_locked=False,
            is_test_data=False,
        )
        with self.assertRaises(PermissionDenied):
            archive_activity(activity, self.admin)

    def test_archive_activity_admin_finalizes_and_locks(self):
        activity = self._make_finalized_round_activity()
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_archive", args=[activity.pk]), {"note": "event over"}
        )
        self.assertEqual(response.status_code, 302)
        activity.refresh_from_db()
        self.assertEqual(activity.phase, Activity.Phase.ARCHIVED)
        self.assertTrue(activity.is_locked)
        self.assertTrue(ArchivePackage.objects.filter(activity=activity).exists())
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.ARCHIVE_ACTIVITY,
                target=f"Activity:{activity.pk}",
            ).exists()
        )

    def test_archive_activity_staff_forbidden(self):
        activity = self._make_finalized_round_activity()
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:activity_archive", args=[activity.pk]), {"note": "event over"}
        )
        self.assertEqual(response.status_code, 403)
        activity.refresh_from_db()
        self.assertNotEqual(activity.phase, Activity.Phase.ARCHIVED)
        self.assertFalse(ArchivePackage.objects.filter(activity=activity).exists())

    def test_archive_activity_rejects_duplicate_archive(self):
        activity = self._make_finalized_round_activity()
        archive_activity(activity, self.admin)
        self.assertEqual(ArchivePackage.objects.filter(activity=activity).count(), 1)
        with self.assertRaises(PermissionDenied):
            archive_activity(activity, self.admin)
        self.assertEqual(ArchivePackage.objects.filter(activity=activity).count(), 1)

    def _make_archived_activity(self):
        return Activity.objects.create(
            title="Archived Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.ARCHIVED,
            is_test_mode=False,
            is_locked=True,
        )

    def test_activity_unlock_rejects_archived_activity(self):
        activity = self._make_archived_activity()
        login_admin(self.client, self.admin)
        response = self.client.post(reverse("staff:activity_unlock", args=[activity.pk]))
        self.assertEqual(response.status_code, 403)
        activity.refresh_from_db()
        self.assertTrue(activity.is_locked)

    def test_activity_edit_rejects_archived_activity(self):
        activity = self._make_archived_activity()
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_edit", args=[activity.pk]),
            {
                "title": "Edited Archived",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.ARCHIVED,
                "subtitle": "should not persist",
            },
        )
        self.assertEqual(response.status_code, 403)
        activity.refresh_from_db()
        self.assertEqual(activity.title, "Archived Contest")

    def test_material_requirement_create_rejects_archived_activity(self):
        activity = self._make_archived_activity()
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:activity_material_requirements", args=[activity.pk]),
            {"applies_to": "singer", "item_name": "伴奏"},
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(MaterialRequirement.objects.filter(activity=activity).exists())

    def test_material_requirement_delete_rejects_archived_activity(self):
        activity = self._make_archived_activity()
        requirement = MaterialRequirement.objects.create(
            activity=activity,
            applies_to=MaterialRequirement.AppliesTo.SINGER,
            item_name="伴奏",
        )
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse(
                "staff:activity_material_requirement_delete", args=[activity.pk, requirement.pk]
            )
        )
        self.assertEqual(response.status_code, 403)
        self.assertTrue(MaterialRequirement.objects.filter(pk=requirement.pk).exists())

    def test_activity_unarchive_staff_forbidden(self):
        activity = self._make_archived_activity()
        self.client.force_login(self.staff)
        response = self.client.post(reverse("staff:activity_unarchive", args=[activity.pk]))
        self.assertEqual(response.status_code, 403)
        activity.refresh_from_db()
        self.assertEqual(activity.phase, Activity.Phase.ARCHIVED)

    def test_activity_unarchive_admin_restores_and_audits(self):
        activity = self._make_archived_activity()
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_unarchive", args=[activity.pk]), {"note": "reopen"}
        )
        self.assertEqual(response.status_code, 302)
        activity.refresh_from_db()
        self.assertEqual(activity.phase, Activity.Phase.RESULTS_PUBLISHED)
        self.assertFalse(activity.is_locked)
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UNARCHIVE_ACTIVITY,
                target=f"Activity:{activity.pk}",
                old_value=Activity.Phase.ARCHIVED,
                new_value=Activity.Phase.RESULTS_PUBLISHED,
            ).exists()
        )

    def test_unarchive_demotes_archive_package_and_rearchive_bumps_version(self):
        activity = self._make_finalized_round_activity()
        archive_activity(activity, self.admin)
        first_package = ArchivePackage.objects.get(activity=activity, is_current=True)
        self.assertEqual(first_package.version, 1)

        unarchive_activity(activity, actor=self.admin)
        first_package.refresh_from_db()
        self.assertFalse(first_package.is_current)
        self.assertEqual(activity.phase, Activity.Phase.RESULTS_PUBLISHED)

        archive_activity(activity, self.admin)
        second_package = ArchivePackage.objects.get(activity=activity, is_current=True)
        self.assertEqual(second_package.version, 2)
        self.assertTrue(second_package.is_current)

    def test_failed_archive_leaves_no_orphan_row_or_file(self):
        activity = self._make_finalized_round_activity()
        ScoreRecord.objects.all().delete()
        with self.assertRaises(PermissionDenied):
            archive_activity(activity, self.admin)
        self.assertFalse(ArchivePackage.objects.filter(activity=activity).exists())
        archives_dir = os.path.join(self.media_root, "archives")
        orphan_files = (
            [name for name in os.listdir(archives_dir) if name.startswith(f"archive_{activity.pk}")]
            if os.path.isdir(archives_dir)
            else []
        )
        self.assertEqual(orphan_files, [])

    def test_archive_preview_does_not_change_state(self):
        activity = self._make_finalized_round_activity()
        self.client.force_login(self.staff)
        response = self.client.post(reverse("staff:archive_package_create", args=[activity.pk]))
        self.assertEqual(response.status_code, 200)
        with zipfile.ZipFile(BytesIO(response.content)) as zf:
            self.assertIn("score_results.xlsx", zf.namelist())
        activity.refresh_from_db()
        self.assertNotEqual(activity.phase, Activity.Phase.ARCHIVED)
        self.assertFalse(activity.is_locked)
        self.assertFalse(ArchivePackage.objects.filter(activity=activity).exists())
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.EXPORT,
                target=f"archive_preview: {activity.title}",
            ).exists()
        )

    def test_formal_word_generation_creates_formal_document(self):
        formal_activity = Activity.objects.create(
            title="Formal Singer Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            is_test_mode=False,
        )
        template = ArticleTemplate.objects.create(
            name="Formal notice",
            template_type=ArticleTemplate.TemplateType.PRELIMINARY_NOTICE,
            body="{title}",
        )
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("staff:word_generate", args=[template.pk, formal_activity.pk])
        )

        self.assertEqual(response.status_code, 200)
        document = GeneratedDocument.objects.get(template=template, activity=formal_activity)
        self.assertFalse(document.is_test_data)
        self.assertTrue(document.file.storage.exists(document.file.name or ""))

    def test_test_word_generation_creates_test_document(self):
        self.singer_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        template = ArticleTemplate.objects.create(
            name="Test notice",
            template_type=ArticleTemplate.TemplateType.VOTE_GUIDE,
            body="{title}",
        )
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("staff:word_generate", args=[template.pk, self.singer_activity.pk])
        )

        self.assertEqual(response.status_code, 200)
        document = GeneratedDocument.objects.get(template=template, activity=self.singer_activity)
        self.assertTrue(document.is_test_data)
        self.assertTrue(document.file.storage.exists(document.file.name or ""))

    def test_locking_vote_session_generates_popularity_award(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
        )
        option = VoteOption.objects.create(vote_session=vote_session, singer=registration)
        VoteRecord.objects.create(
            vote_session=vote_session,
            vote_option=option,
            browser_session_key="visitor-1",
            ip_address="127.0.0.1",
        )
        self.client.force_login(self.staff)
        response = self.client.post(reverse("staff:vote_session_lock", args=[vote_session.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(
            Award.objects.filter(
                activity=self.singer_activity, singer=registration, name="最佳人气奖"
            ).exists()
        )

    def test_relocking_vote_session_reconciles_changed_popularity_winner(self):
        first = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="First Winner",
            student_id="20260010",
            college="Info",
            class_name="CS1",
            phone="13800000009",
            song_name="Song 1",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        second_user = User.objects.create_user(username="participant-two", password="pass")
        second = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=second_user,
            name="Second Winner",
            student_id="20260011",
            college="Info",
            class_name="CS1",
            phone="13800000010",
            song_name="Song 2",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity Recount",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
        )
        first_option = VoteOption.objects.create(vote_session=vote_session, singer=first)
        second_option = VoteOption.objects.create(vote_session=vote_session, singer=second)
        VoteRecord.objects.create(
            vote_session=vote_session,
            vote_option=first_option,
            browser_session_key="winner-1",
            ip_address="127.0.0.1",
        )
        login_admin(self.client, self.admin)
        self.client.post(reverse("staff:vote_session_lock", args=[vote_session.pk]))
        self.client.post(
            reverse("staff:vote_session_unlock", args=[vote_session.pk]),
            {"note": "recount"},
        )
        VoteRecord.objects.create(
            vote_session=vote_session,
            vote_option=second_option,
            browser_session_key="winner-2",
            ip_address="127.0.0.2",
        )
        VoteRecord.objects.create(
            vote_session=vote_session,
            vote_option=second_option,
            browser_session_key="winner-3",
            ip_address="127.0.0.3",
        )
        self.client.post(reverse("staff:vote_session_lock", args=[vote_session.pk]))

        self.assertFalse(
            Award.objects.filter(
                activity=self.singer_activity, singer=first, name="最佳人气奖"
            ).exists()
        )
        self.assertEqual(
            Award.objects.filter(
                activity=self.singer_activity, singer=second, name="最佳人气奖"
            ).count(),
            1,
        )

    def _make_popularity_tie_session(self):
        second_user = User.objects.create_user(username="participant-second", password="pass")
        reg1 = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Tie One",
            student_id="20260101",
            college="Info",
            class_name="CS1",
            phone="13800000001",
            song_name="Song 1",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        reg2 = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=second_user,
            name="Tie Two",
            student_id="20260102",
            college="Info",
            class_name="CS1",
            phone="13800000002",
            song_name="Song 2",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Tie Popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
        )
        option1 = VoteOption.objects.create(vote_session=vote_session, singer=reg1)
        option2 = VoteOption.objects.create(vote_session=vote_session, singer=reg2)
        VoteRecord.objects.create(
            vote_session=vote_session,
            vote_option=option1,
            browser_session_key="tie-1",
            ip_address="127.0.0.1",
        )
        VoteRecord.objects.create(
            vote_session=vote_session,
            vote_option=option2,
            browser_session_key="tie-2",
            ip_address="127.0.0.2",
        )
        return vote_session

    def test_lock_vote_session_with_popularity_tie_does_not_award(self):
        vote_session = self._make_popularity_tie_session()
        self.client.force_login(self.staff)
        response = self.client.post(reverse("staff:vote_session_lock", args=[vote_session.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertFalse(
            Award.objects.filter(activity=self.singer_activity, name="最佳人气奖").exists()
        )

    def test_vote_session_detail_shows_popularity_tie_warning(self):
        vote_session = self._make_popularity_tie_session()
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:vote_session_detail", args=[vote_session.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["popularity_tie"])

    def test_round_lock_rejects_incomplete_scores(self):
        singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Incomplete Singer",
            student_id="20260002",
            college="Info",
            class_name="CS1",
            phone="13800000001",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.client.force_login(self.staff)

        response = self.client.post(reverse("staff:round_lock", args=[round_.pk]))

        self.assertEqual(response.status_code, 403)
        round_.refresh_from_db()
        self.assertFalse(round_.is_locked)
        self.assertFalse(ScoreRecord.objects.filter(singer=singer, judge=judge).exists())

    def test_round_lock_and_unlock_update_round_status(self):
        singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Scored Singer",
            student_id="20260010",
            college="Info",
            class_name="CS1",
            phone="13800000009",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        ScoreRecord.objects.create(round=round_, singer=singer, judge=judge, score=91)
        login_admin(self.client, self.admin)

        lock_response = self.client.post(reverse("staff:round_lock", args=[round_.pk]))
        round_.refresh_from_db()
        self.assertEqual(lock_response.status_code, 302)
        self.assertEqual(round_.status, ContestRound.Status.LOCKED)
        unlock_response = self.client.post(
            reverse("staff:round_unlock", args=[round_.pk]), {"note": "correction"}
        )
        round_.refresh_from_db()

        self.assertEqual(round_.status, ContestRound.Status.SCORING)
        self.assertEqual(unlock_response.status_code, 302)
        self.assertFalse(round_.is_locked)

    def test_round_lock_rejects_round_already_locked(self):
        singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Locked Singer",
            student_id="20260011",
            college="Info",
            class_name="CS1",
            phone="13800000010",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        ScoreRecord.objects.create(round=round_, singer=singer, judge=judge, score=91)
        round_.status = ContestRound.Status.LOCKED
        round_.is_locked = True
        round_.save(update_fields=["status", "is_locked"])
        self.client.force_login(self.staff)

        response = self.client.post(reverse("staff:round_lock", args=[round_.pk]))

        self.assertEqual(response.status_code, 403)
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.LOCKED)
        self.assertTrue(round_.is_locked)

    def test_round_ranking_excludes_singers_outside_prepared_snapshot(self):
        singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Prepared Singer",
            student_id="20260012",
            college="Info",
            class_name="CS1",
            phone="13800000011",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        _judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        late_user = User.objects.create_user(username="participant-late", password="pass")
        late_singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=late_user,
            name="Late Singer",
            student_id="20260013",
            college="Info",
            class_name="CS1",
            phone="13800000012",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        other_activity = Activity.objects.create(
            title="Other Contest", activity_type=Activity.Type.SINGER_CONTEST
        )
        foreign_singer = SingerRegistration.objects.create(
            activity=other_activity,
            user=self.participant,
            name="Foreign Singer",
            student_id="20260014",
            college="Info",
            class_name="CS1",
            phone="13800000013",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        ScoreSummary.objects.create(round=round_, singer=singer, average_score=90, rank=1)
        ScoreSummary.objects.create(round=round_, singer=late_singer, average_score=99, rank=1)
        ScoreSummary.objects.create(round=round_, singer=foreign_singer, average_score=100, rank=1)
        self.client.force_login(self.staff)

        response = self.client.get(reverse("staff:round_ranking", args=[round_.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            list(response.context["summaries"].values_list("singer_id", flat=True)),
            [singer.pk],
        )
        self.assertContains(response, singer.name)
        self.assertNotContains(response, late_singer.name)
        self.assertNotContains(response, foreign_singer.name)

    def test_round_lock_get_is_rejected_without_mutating_state(self):
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        self.client.force_login(self.staff)

        response = self.client.get(reverse("staff:round_lock", args=[round_.pk]))

        self.assertEqual(response.status_code, 405)
        round_.refresh_from_db()
        self.assertFalse(round_.is_locked)

    def _make_boundary_tie_round(self):
        singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Tie Singer",
            student_id="20260021",
            college="Info",
            class_name="CS1",
            phone="13800000020",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        extra_user = User.objects.create_user(username="participant-tie", password="pass")
        extra = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=extra_user,
            name="Tie Extra",
            student_id="20260022",
            college="Info",
            class_name="CS1",
            phone="13800000021",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        Judge.objects.create(activity=self.singer_activity, name="Judge A", is_active=True)
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            advance_count=1,
        )
        prepare_round(round_, self.staff)
        judge_ids = list(RoundJudge.objects.filter(round=round_).values_list("judge_id", flat=True))
        score_values = {}
        for judge_id in judge_ids:
            score_values[(singer.pk, judge_id)] = 91
            score_values[(extra.pk, judge_id)] = 91
        apply_scores(round_, score_values, self.staff)
        round_.refresh_from_db()
        self.assertEqual(round_.advancement_status, ContestRound.AdvancementStatus.NEEDS_REVIEW)
        return round_, singer, extra

    def test_round_lock_rejects_boundary_tie_until_finalized(self):
        round_, singer, _extra = self._make_boundary_tie_round()
        self.client.force_login(self.staff)

        lock_response = self.client.post(reverse("staff:round_lock", args=[round_.pk]))
        self.assertEqual(lock_response.status_code, 403)
        round_.refresh_from_db()
        self.assertFalse(round_.is_locked)
        self.assertNotEqual(round_.status, ContestRound.Status.LOCKED)

        finalize_response = self.client.post(
            reverse("staff:round_finalize_advancement", args=[round_.pk]),
            data={"selected_singer_ids": [str(singer.pk)]},
        )
        self.assertEqual(finalize_response.status_code, 302)
        round_.refresh_from_db()
        self.assertEqual(round_.advancement_status, ContestRound.AdvancementStatus.FINALIZED)

        relock_response = self.client.post(reverse("staff:round_lock", args=[round_.pk]))
        self.assertEqual(relock_response.status_code, 302)
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.LOCKED)
        self.assertTrue(round_.is_locked)
        self.assertTrue(
            AuditLog.objects.filter(
                operator=self.staff,
                action_type=AuditLog.ActionType.FINALIZE_ADVANCEMENT,
                target=f"ContestRound:{round_.pk}",
            ).exists()
        )

    def test_award_creation_rejects_singer_from_another_activity(self):
        foreign_singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Foreign Singer",
            student_id="20260005",
            college="Info",
            class_name="CS1",
            phone="13800000004",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        other_activity = Activity.objects.create(
            title="Other Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        login_admin(self.client, self.admin)

        response = self.client.post(
            reverse("staff:award_create"),
            {
                "activity_id": other_activity.pk,
                "singer_id": foreign_singer.pk,
                "name": "Invalid Award",
            },
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(Award.objects.filter(name="Invalid Award").exists())

    def test_vote_session_creation_rejects_singer_from_another_activity(self):
        foreign_singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Foreign Vote Singer",
            student_id="20260006",
            college="Info",
            class_name="CS1",
            phone="13800000005",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        other_activity = Activity.objects.create(
            title="Other Vote Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        login_admin(self.client, self.admin)

        response = self.client.post(
            reverse("staff:vote_session_create"),
            {
                "activity_id": other_activity.pk,
                "name": "Invalid Vote",
                "passcode": "1234",
                "start_time": "2026-08-27T10:00",
                "end_time": "2026-08-27T11:00",
                "singers": [str(foreign_singer.pk)],
            },
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(VoteSession.objects.filter(name="Invalid Vote").exists())

    def test_incident_creation_rejects_singer_from_another_activity(self):
        foreign_singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Foreign Incident Singer",
            student_id="20260013",
            college="Info",
            class_name="CS1",
            phone="13800000012",
            song_name="Song",
        )
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("staff:incident_create"),
            {
                "activity_id": self.farewell_activity.pk,
                "occurred_at": timezone.now().isoformat(),
                "event_type": IncidentRecord.EventType.OTHER,
                "singer_id": foreign_singer.pk,
                "resolution": "invalid",
            },
        )

        self.assertEqual(response.status_code, 403)
        self.assertFalse(IncidentRecord.objects.filter(resolution="invalid").exists())

    def test_manual_score_submission_rolls_back_when_one_score_is_invalid(self):
        first = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="First Singer",
            student_id="20260003",
            college="Info",
            class_name="CS1",
            phone="13800000002",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        second_user = User.objects.create_user(username="participant-two", password="pass")
        second = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=second_user,
            name="Second Singer",
            student_id="20260004",
            college="Info",
            class_name="CS1",
            phone="13800000003",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge B")
        round_ = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(round_, self.staff)
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("staff:round_score_entry", args=[round_.pk]),
            {
                f"score_{first.pk}_{judge.pk}": "90",
                f"score_{second.pk}_{judge.pk}": "not-a-score",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "分数")
        self.assertFalse(ScoreRecord.objects.filter(round=round_).exists())

    def test_vote_cast_allows_distinct_browsers_sharing_an_ip(self):
        registration = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
        )
        option = VoteOption.objects.create(vote_session=vote_session, singer=registration)
        first = self.client_class()
        second = self.client_class()
        first.post(
            reverse("voting:vote_entry", args=[vote_session.pk]),
            {"passcode": "1234"},
            REMOTE_ADDR="127.0.0.1",
        )
        self.assertEqual(
            first.post(
                reverse("voting:vote_cast", args=[vote_session.pk]),
                {"selected_option": [str(option.pk)]},
                REMOTE_ADDR="127.0.0.1",
            ).status_code,
            302,
        )
        second.post(
            reverse("voting:vote_entry", args=[vote_session.pk]),
            {"passcode": "1234"},
            REMOTE_ADDR="127.0.0.1",
        )
        response = second.post(
            reverse("voting:vote_cast", args=[vote_session.pk]),
            {"selected_option": [str(option.pk)]},
            REMOTE_ADDR="127.0.0.1",
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(VoteRecord.objects.filter(vote_session=vote_session).count(), 2)

    def test_activity_clone_preserves_configuration_not_runtime_data(self):
        MaterialRequirement.objects.create(
            activity=self.singer_activity,
            applies_to=MaterialRequirement.AppliesTo.SINGER,
            item_name="Lyrics",
            file_purpose=SubmissionFile.Purpose.LYRICS_SCRIPT,
        )
        judge = Judge.objects.create(activity=self.singer_activity, name="Judge A")
        Judge.objects.create(activity=self.singer_activity, name="Judge B")
        Judge.objects.create(activity=self.singer_activity, name="Judge C")
        contest_round = ContestRound.objects.create(
            activity=self.singer_activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            scoring_mode=ContestRound.ScoringMode.DROP_HIGH_LOW,
            advance_count=2,
            name="Preliminary",
        )
        singer = SingerRegistration.objects.create(
            activity=self.singer_activity,
            user=self.participant,
            name="Runtime Data",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        SubmissionFile.objects.create(
            singer_registration=singer,
            file=SimpleUploadedFile("runtime.txt", b"runtime data"),
            original_name="runtime.txt",
            file_size=12,
            file_purpose=SubmissionFile.Purpose.OTHER,
            uploaded_by=self.participant,
        )
        prepare_round(contest_round, self.staff)
        ScoreRecord.objects.create(
            round=contest_round,
            singer=singer,
            judge=judge,
            score=91,
        )
        ScoreSummary.objects.create(
            round=contest_round,
            singer=singer,
            average_score=91,
            rank=1,
        )
        vote_session = VoteSession.objects.create(
            activity=self.singer_activity,
            name="Popularity",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
        )
        option = VoteOption.objects.create(vote_session=vote_session, singer=singer)
        ballot = VoteBallot.objects.create(
            vote_session=vote_session,
            browser_session_key="runtime-browser",
            ip_address="127.0.0.1",
        )
        VoteRecord.objects.create(
            ballot=ballot,
            vote_session=vote_session,
            vote_option=option,
            browser_session_key="runtime-browser",
            ip_address="127.0.0.1",
        )
        GeneratedDocument.objects.create(
            activity=self.singer_activity,
            title="Runtime document",
            file=SimpleUploadedFile("runtime.docx", b"runtime document"),
            created_by=self.staff,
        )
        login_admin(self.client, self.admin)
        response = self.client.post(reverse("staff:activity_clone", args=[self.singer_activity.pk]))
        self.assertEqual(response.status_code, 302)
        clone = Activity.objects.exclude(pk=self.singer_activity.pk).get(
            activity_type=Activity.Type.SINGER_CONTEST
        )
        self.assertTrue(
            MaterialRequirement.objects.filter(activity=clone, item_name="Lyrics").exists()
        )
        self.assertTrue(Judge.objects.filter(activity=clone, name="Judge A").exists())
        self.assertTrue(
            ContestRound.objects.filter(
                activity=clone, round_type=ContestRound.RoundType.PRELIMINARY
            ).exists()
        )
        self.assertFalse(VoteSession.objects.filter(activity=clone, name="Popularity").exists())
        self.assertTrue(clone.is_test_mode)
        self.assertEqual(clone.data_lifecycle, Activity.DataLifecycle.TEST)
        self.assertEqual(clone.phase, Activity.Phase.DRAFT)
        self.assertFalse(SingerRegistration.objects.filter(activity=clone).exists())
        self.assertFalse(
            SubmissionFile.objects.filter(singer_registration__activity=clone).exists()
        )
        self.assertFalse(RoundEntry.objects.filter(round__activity=clone).exists())
        self.assertFalse(RoundJudge.objects.filter(round__activity=clone).exists())
        self.assertFalse(ScoreRecord.objects.filter(round__activity=clone).exists())
        self.assertFalse(ScoreSummary.objects.filter(round__activity=clone).exists())
        self.assertFalse(VoteOption.objects.filter(vote_session__activity=clone).exists())
        self.assertFalse(VoteBallot.objects.filter(vote_session__activity=clone).exists())
        self.assertFalse(VoteRecord.objects.filter(vote_session__activity=clone).exists())
        self.assertFalse(GeneratedDocument.objects.filter(activity=clone).exists())


class RuntimeLifecycleMatrixTests(TestCase):
    """TEST/FORMAL runtime data must never mix when binding a child to an activity.

    The candidate pools scope each singer row to its own activity lifecycle, and
    POST validation rejects binding a singer whose test marker differs from the
    target activity. This is the lifecycle matrix the reviewer asked for.
    """

    def setUp(self):
        self.staff = User.objects.create_user(
            username="matrix-staff", password="pass", role=User.Role.STAFF
        )
        self.admin = User.objects.create_user(
            username="matrix-admin", password="pass", role=User.Role.ADMIN
        )
        self.test_activity = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.formal_activity = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        # Each activity carries one correctly-marked and one wrongly-marked singer.
        self.test_singer = self._singer(self.test_activity, "Test Match", True, "20260101")
        self.test_mismatch = self._singer(self.test_activity, "Test Mismatch", False, "20260102")
        self.formal_singer = self._singer(self.formal_activity, "Formal Match", False, "20260103")
        self.formal_mismatch = self._singer(
            self.formal_activity, "Formal Mismatch", True, "20260104"
        )

    def _singer(self, activity, name, is_test_data, student_id):
        user = User.objects.create_user(username=f"matrix-{student_id}", password="pass")
        return SingerRegistration.objects.create(
            activity=activity,
            user=user,
            name=name,
            student_id=student_id,
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=is_test_data,
        )

    def test_candidate_pool_binds_each_singer_to_its_own_activity(self):
        self.client.force_login(self.staff)
        for path in (
            reverse("staff:award_create"),
            reverse("staff:vote_session_create"),
            reverse("staff:incident_create"),
        ):
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, self.test_singer.name)
                self.assertContains(response, self.formal_singer.name)
                self.assertNotContains(response, self.test_mismatch.name)
                self.assertNotContains(response, self.formal_mismatch.name)

    def test_award_create_rejects_wrong_lifecycle_singer(self):
        activity = Activity.objects.create(
            title="Award Test",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.DRAFT,
            is_test_mode=True,
        )
        matching = self._singer(activity, "Award Match", True, "20260201")
        mismatched = self._singer(activity, "Award Mismatch", False, "20260202")
        login_admin(self.client, self.admin)

        rejected = self.client.post(
            reverse("staff:award_create"),
            {"activity_id": activity.pk, "singer_id": mismatched.pk, "name": "Bad Award"},
        )
        self.assertEqual(rejected.status_code, 403)
        self.assertFalse(Award.objects.filter(name="Bad Award").exists())

        accepted = self.client.post(
            reverse("staff:award_create"),
            {"activity_id": activity.pk, "singer_id": matching.pk, "name": "Good Award"},
        )
        self.assertEqual(accepted.status_code, 302)
        award = Award.objects.get(name="Good Award")
        self.assertTrue(award.is_test_data)

    def test_vote_session_create_rejects_wrong_lifecycle_singer(self):
        activity = Activity.objects.create(
            title="Vote Test",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        matching = self._singer(activity, "Vote Match", True, "20260301")
        mismatched = self._singer(activity, "Vote Mismatch", False, "20260302")
        login_admin(self.client, self.admin)
        base = {
            "activity_id": activity.pk,
            "passcode": "1234",
            "start_time": "2026-08-27T10:00",
            "end_time": "2026-08-27T11:00",
        }

        rejected = self.client.post(
            reverse("staff:vote_session_create"),
            {**base, "name": "Bad Vote", "singers": [str(mismatched.pk)]},
        )
        self.assertEqual(rejected.status_code, 403)
        self.assertFalse(VoteSession.objects.filter(name="Bad Vote").exists())

        accepted = self.client.post(
            reverse("staff:vote_session_create"),
            {**base, "name": "Good Vote", "singers": [str(matching.pk)]},
        )
        self.assertEqual(accepted.status_code, 302)
        session = VoteSession.objects.get(name="Good Vote")
        self.assertTrue(session.is_test_data)
        self.assertEqual(VoteOption.objects.filter(vote_session=session).count(), 1)

    def test_incident_create_rejects_wrong_lifecycle_singer(self):
        activity = Activity.objects.create(
            title="Incident Test",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        matching = self._singer(activity, "Incident Match", True, "20260401")
        mismatched = self._singer(activity, "Incident Mismatch", False, "20260402")
        self.client.force_login(self.staff)
        base = {
            "activity_id": activity.pk,
            "occurred_at": timezone.now().isoformat(),
            "event_type": IncidentRecord.EventType.OTHER,
        }

        rejected = self.client.post(
            reverse("staff:incident_create"),
            {**base, "singer_id": mismatched.pk, "resolution": "bad"},
        )
        self.assertEqual(rejected.status_code, 403)
        self.assertFalse(IncidentRecord.objects.filter(resolution="bad").exists())

        accepted = self.client.post(
            reverse("staff:incident_create"),
            {**base, "singer_id": matching.pk, "resolution": "good"},
        )
        self.assertEqual(accepted.status_code, 302)
        self.assertTrue(IncidentRecord.objects.get(resolution="good").is_test)

    def test_prepare_round_uses_approved_singers_scoped_to_activity(self):
        activity = Activity.objects.create(
            title="Round Test",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        test_match = self._singer(activity, "Round Test Match", True, "20260501")
        self._singer(activity, "Round Test Mismatch", False, "20260502")
        Judge.objects.create(activity=activity, name="Judge", is_active=True)
        round_ = ContestRound.objects.create(
            activity=activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            advance_count=1,
        )

        prepare_round(round_, self.staff)

        entries = set(RoundEntry.objects.filter(round=round_).values_list("singer_id", flat=True))
        self.assertEqual(entries, {test_match.pk})


class AdminAuthBoundaryTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="boundary_admin",
            password="pass",
            role=User.Role.ADMIN,
        )
        self.staff = User.objects.create_user(
            username="boundary_staff",
            password="pass",
            role=User.Role.STAFF,
        )
        self.participant = User.objects.create_user(
            username="boundary_participant",
            password="pass",
            role=User.Role.PARTICIPANT,
        )

    def _force_login_with_marker(self, user):
        self.client.force_login(user)
        session = self.client.session
        session["artflow_admin_verified"] = True
        session["artflow_admin_verified_at"] = timezone.now().isoformat()
        session.save()

    def test_unauthenticated_staff_redirects_to_artflow_login(self):
        response = self.client.get(reverse("staff:dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response["Location"].startswith(reverse("accounts:login")))

    def test_unauthenticated_staff_does_not_redirect_to_django_admin(self):
        response = self.client.get(reverse("staff:dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("/admin/login/", response["Location"])

    def test_admin_without_verification_is_rejected_on_admin_endpoint(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("staff:activity_create"))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response["Location"].startswith(reverse("accounts:admin_login")))

    def test_admin_with_verification_can_access_admin_endpoint(self):
        self._force_login_with_marker(self.admin)
        response = self.client.post(
            reverse("staff:activity_create"),
            {
                "title": "Verified Admin Activity",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.REGISTRATION_OPEN,
                "is_test_mode": "on",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], reverse("staff:activity_list"))
        self.assertTrue(Activity.objects.filter(title="Verified Admin Activity").exists())

    def test_staff_can_access_staff_endpoint(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:dashboard"))
        self.assertEqual(response.status_code, 200)

    def test_staff_cannot_access_admin_endpoint(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:activity_create"))
        self.assertEqual(response.status_code, 403)

    def test_participant_cannot_access_staff_endpoint(self):
        self.client.force_login(self.participant)
        response = self.client.get(reverse("staff:dashboard"))
        self.assertEqual(response.status_code, 403)


class ActivityPhaseEditTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(
            username="phase-admin",
            password="pass",
            role=User.Role.ADMIN,
        )

    def test_admin_edit_advances_phase_through_service(self):
        activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.DRAFT,
        )
        login_admin(self.client, self.admin)
        response = self.client.post(
            reverse("staff:activity_edit", args=[activity.pk]),
            {
                "title": "Contest",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.REGISTRATION_OPEN,
            },
        )
        self.assertEqual(response.status_code, 302)
        activity.refresh_from_db()
        self.assertEqual(activity.phase, Activity.Phase.REGISTRATION_OPEN)
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.PHASE_TRANSITION,
                operator=self.admin,
                old_value=Activity.Phase.DRAFT,
                new_value=Activity.Phase.REGISTRATION_OPEN,
            ).exists()
        )

    def test_admin_edit_rejects_backward_phase_transition(self):
        activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.LIVE,
        )
        login_admin(self.client, self.admin)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:activity_edit", args=[activity.pk]),
            {
                "title": "Contest",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.DRAFT,
            },
        )
        self.assertEqual(response.status_code, 403)
        activity.refresh_from_db()
        self.assertEqual(activity.phase, Activity.Phase.LIVE)

    def test_admin_edit_cannot_directly_archive(self):
        activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PUBLISHED,
        )
        login_admin(self.client, self.admin)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:activity_edit", args=[activity.pk]),
            {
                "title": "Contest",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.ARCHIVED,
            },
        )
        self.assertEqual(response.status_code, 403)
        activity.refresh_from_db()
        self.assertEqual(activity.phase, Activity.Phase.RESULTS_PUBLISHED)
        self.assertFalse(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.PHASE_TRANSITION,
                new_value=Activity.Phase.ARCHIVED,
            ).exists()
        )

    def test_admin_create_cannot_create_archived(self):
        # ARCHIVED is not a creatable phase; the form rejects it before the
        # view can run, so the create must not succeed or persist a row.
        login_admin(self.client, self.admin)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:activity_create"),
            {
                "title": "Fake Archived",
                "activity_type": Activity.Type.SINGER_CONTEST,
                "phase": Activity.Phase.ARCHIVED,
                "is_test_mode": "on",
            },
        )
        self.assertNotEqual(response.status_code, 302)
        self.assertFalse(Activity.objects.filter(title="Fake Archived").exists())


class ActivityPhaseViewEnforcementTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(
            username="enforce-staff",
            password="pass",
            role=User.Role.STAFF,
        )
        self.participant = User.objects.create_user(
            username="enforce-participant",
            password="pass",
            role=User.Role.PARTICIPANT,
        )

    def _approved_registration(self, activity, student_id="20260001", is_test_data=True):
        return SingerRegistration.objects.create(
            activity=activity,
            user=self.participant,
            name="Singer",
            student_id=student_id,
            college="Music",
            class_name="Class A",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=is_test_data,
        )

    def test_draft_activity_cannot_enter_scores(self):
        activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        self._approved_registration(activity)
        Judge.objects.create(activity=activity, name="Judge A")
        round_ = ContestRound.objects.create(
            activity=activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            scoring_mode=ContestRound.ScoringMode.AVERAGE,
        )
        with self.assertRaises(PermissionDenied):
            prepare_round(round_, self.staff)
        round_.refresh_from_db()
        self.assertEqual(round_.status, ContestRound.Status.DRAFT)
        self.assertEqual(round_.entries.count(), 0)
        self.assertEqual(round_.round_judges.count(), 0)

    def test_archived_activity_cannot_review_registration(self):
        activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.ARCHIVED,
        )
        registration = self._approved_registration(activity)
        self.client.force_login(self.staff)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:singer_registration_detail", args=[registration.pk]),
            {
                "pre_status": SingerRegistration.PreStatus.APPROVED,
                "live_status": SingerRegistration.LiveStatus.NOT_CHECKED_IN,
            },
        )
        self.assertEqual(response.status_code, 403)

    def test_results_published_activity_cannot_change_registration_status(self):
        activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.RESULTS_PUBLISHED,
        )
        registration = self._approved_registration(activity)
        self.client.force_login(self.staff)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:singer_registration_detail", args=[registration.pk]),
            {
                "pre_status": SingerRegistration.PreStatus.APPROVED,
                "live_status": SingerRegistration.LiveStatus.SCORED,
            },
        )
        self.assertEqual(response.status_code, 403)


class MaterialReviewAndRequirementTests(TestCase):
    def setUp(self):
        self.media_root = tempfile.mkdtemp()
        self.override = override_settings(MEDIA_ROOT=self.media_root)
        self.override.enable()
        self.staff = User.objects.create_user(
            username="staff_user",
            password="pass",
            role=User.Role.STAFF,
        )
        self.participant = User.objects.create_user(
            username="participant",
            password="pass",
            role=User.Role.PARTICIPANT,
        )
        self.activity = Activity.objects.create(
            title="Singer Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REVIEWING,
            is_test_mode=False,
        )
        self.registration = SingerRegistration.objects.create(
            activity=self.activity,
            user=self.participant,
            name="Li Hua",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
        )
        reconcile_singer_material_checks(self.registration)
        self.client.raise_request_exception = False

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media_root, ignore_errors=True)

    def test_staff_approve_material_check(self):
        self.client.force_login(self.staff)
        check = self.registration.material_checks.get(item_name="伴奏文件")
        response = self.client.post(
            reverse("staff:material_check_review"),
            {"check_id": check.pk, "result": "approve", "note": "清晰"},
        )
        self.assertEqual(response.status_code, 302)
        check.refresh_from_db()
        self.assertEqual(check.status, MaterialCheck.Status.APPROVED)
        self.assertEqual(check.review_note, "清晰")
        self.assertEqual(check.reviewed_by, self.staff)

    def test_staff_request_supplement_material_check(self):
        self.client.force_login(self.staff)
        check = self.registration.material_checks.get(item_name="伴奏文件")
        response = self.client.post(
            reverse("staff:material_check_review"),
            {"check_id": check.pk, "result": "supplement", "note": "请上传无现场人声版本"},
        )
        self.assertEqual(response.status_code, 302)
        check.refresh_from_db()
        self.assertEqual(check.status, MaterialCheck.Status.NEEDS_SUPPLEMENT)
        self.assertEqual(check.review_note, "请上传无现场人声版本")

    def test_staff_material_review_denied_when_activity_locked(self):
        self.activity.is_locked = True
        self.activity.save()
        self.client.force_login(self.staff)
        check = self.registration.material_checks.get(item_name="伴奏文件")
        response = self.client.post(
            reverse("staff:material_check_review"),
            {"check_id": check.pk, "result": "approve"},
        )
        self.assertEqual(response.status_code, 403)
        check.refresh_from_db()
        self.assertNotEqual(check.status, MaterialCheck.Status.APPROVED)

    def test_add_and_delete_material_requirement(self):
        self.client.force_login(self.staff)
        url = reverse("staff:activity_material_requirements", args=[self.activity.pk])
        response = self.client.post(
            url,
            {
                "applies_to": MaterialRequirement.AppliesTo.SINGER,
                "item_name": "歌词/台词",
                "file_purpose": SubmissionFile.Purpose.LYRICS_SCRIPT,
                "sort_order": "5",
            },
        )
        self.assertEqual(response.status_code, 302)
        requirement = MaterialRequirement.objects.get(activity=self.activity, item_name="歌词/台词")
        self.assertEqual(requirement.file_purpose, SubmissionFile.Purpose.LYRICS_SCRIPT)

        delete_url = reverse(
            "staff:activity_material_requirement_delete",
            args=[self.activity.pk, requirement.pk],
        )
        response = self.client.post(delete_url)
        self.assertEqual(response.status_code, 302)
        self.assertFalse(MaterialRequirement.objects.filter(pk=requirement.pk).exists())


class UserRoleAdministrationTests(TestCase):
    def setUp(self):
        self.actor = User.objects.create_user(
            username="role-admin",
            password="pass",
            role=User.Role.ADMIN,
        )
        self.participant = User.objects.create_user(
            username="role-target",
            password="pass",
            role=User.Role.PARTICIPANT,
        )

    def test_participant_gets_403_on_user_list(self):
        self.client.force_login(self.participant)
        self.assertEqual(self.client.get(reverse("staff:user_list")).status_code, 403)

    def test_staff_gets_403_on_user_list(self):
        staff = User.objects.create_user(
            username="role-staff",
            password="pass",
            role=User.Role.STAFF,
        )
        self.client.force_login(staff)
        self.assertEqual(self.client.get(reverse("staff:user_list")).status_code, 403)

    def test_verified_admin_can_promote_participant_to_staff(self):
        login_admin(self.client, self.actor)
        response = self.client.post(
            reverse("staff:user_role_update", args=[self.participant.pk]),
            {"new_role": User.Role.STAFF},
        )
        self.assertEqual(response.status_code, 302)
        self.participant.refresh_from_db()
        self.assertEqual(self.participant.role, User.Role.STAFF)
        self.assertTrue(self.participant.is_staff)

    def test_demoted_staff_loses_staff_endpoint_on_next_request(self):
        staff = User.objects.create_user(
            username="role-lose",
            password="pass",
            role=User.Role.STAFF,
        )
        login_admin(self.client, self.actor)
        response = self.client.post(
            reverse("staff:user_role_update", args=[staff.pk]),
            {"new_role": User.Role.PARTICIPANT},
        )
        self.assertEqual(response.status_code, 302)
        staff.refresh_from_db()
        self.assertFalse(staff.is_staff)
        self.client.force_login(staff)
        self.assertEqual(self.client.get(reverse("staff:dashboard")).status_code, 403)

    def test_role_change_writes_full_audit_old_new(self):
        login_admin(self.client, self.actor)
        self.client.post(
            reverse("staff:user_role_update", args=[self.participant.pk]),
            {"new_role": User.Role.STAFF},
        )
        audit = AuditLog.objects.get(
            action_type=AuditLog.ActionType.UPDATE_PERMISSION,
            target=f"User:{self.participant.pk}",
        )
        self.assertEqual(audit.operator, self.actor)
        self.assertEqual(audit.old_value, User.Role.PARTICIPANT)
        self.assertEqual(audit.new_value, User.Role.STAFF)

    def test_last_active_admin_cannot_be_demoted(self):
        login_admin(self.client, self.actor)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:user_role_update", args=[self.actor.pk]),
            {"new_role": User.Role.PARTICIPANT},
        )
        self.assertEqual(response.status_code, 302)
        self.actor.refresh_from_db()
        self.assertEqual(self.actor.role, User.Role.ADMIN)
        self.assertTrue(self.actor.is_active)
        self.assertFalse(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UPDATE_PERMISSION,
                target=f"User:{self.actor.pk}",
            ).exists()
        )

    def test_last_active_admin_cannot_be_deactivated(self):
        login_admin(self.client, self.actor)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:user_set_active", args=[self.actor.pk]),
            {"active": "0"},
        )
        self.assertEqual(response.status_code, 302)
        self.actor.refresh_from_db()
        self.assertTrue(self.actor.is_active)

    def test_get_user_role_endpoint_does_not_change_role(self):
        login_admin(self.client, self.actor)
        self.assertEqual(
            self.client.get(
                reverse("staff:user_role_update", args=[self.participant.pk])
            ).status_code,
            405,
        )
        self.participant.refresh_from_db()
        self.assertEqual(self.participant.role, User.Role.PARTICIPANT)

    def test_invalid_role_does_not_change_database(self):
        login_admin(self.client, self.actor)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:user_role_update", args=[self.participant.pk]),
            {"new_role": "superstar"},
        )
        self.assertEqual(response.status_code, 302)
        self.participant.refresh_from_db()
        self.assertEqual(self.participant.role, User.Role.PARTICIPANT)
        self.assertFalse(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.UPDATE_PERMISSION,
                target=f"User:{self.participant.pk}",
            ).exists()
        )


class ExportPrivacyTests(TestCase):
    """Exports that carry personal data default to activity scope, not cross-activity."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username="staff_export", password="pass", role=User.Role.STAFF
        )
        self.admin = User.objects.create_user(
            username="admin_export", password="pass", role=User.Role.ADMIN
        )
        self.participant = User.objects.create_user(username="participant_export", password="pass")
        self.contest = Activity.objects.create(
            title="Ten Best Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.other = Activity.objects.create(
            title="Other Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )

    def _registration(self, activity, name, student_id):
        return SingerRegistration.objects.create(
            activity=activity,
            user=self.participant,
            name=name,
            student_id=student_id,
            college="College",
            class_name="Class",
            phone="13800000000",
            wechat="wx",
            song_name=f"{name} Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )

    def _program(self, activity, name, contact):
        return Program.objects.create(
            activity=activity,
            user=self.participant,
            name=name,
            program_type=Program.ProgramType.SONG,
            contact_name=contact,
            contact_phone="13900000000",
            class_name="Class",
        )

    def _first_column_from(self, response):
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        assert ws is not None
        return [ws.cell(row=idx, column=1).value for idx in range(2, ws.max_row + 1)]

    def test_staff_export_registrations_requires_activity_scope(self):
        self._registration(self.contest, "Contest Singer", "S1")
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:export_registrations"))
        self.assertEqual(response.status_code, 403)

    def test_staff_export_registrations_scopes_to_selected_activity(self):
        self._registration(self.contest, "Contest Singer", "S1")
        self._registration(self.other, "Other Singer", "S2")
        self.client.force_login(self.staff)
        response = self.client.get(
            reverse("staff:export_registrations"), {"activity_id": self.other.pk}
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._first_column_from(response), ["Other Singer"])
        log = AuditLog.objects.filter(
            action_type=AuditLog.ActionType.EXPORT, operator=self.staff
        ).get()
        self.assertIn('"activity_id": %d' % self.other.pk, log.note)
        self.assertIn('"row_count": 1', log.note)

    def test_admin_export_registrations_can_export_all_activities(self):
        self._registration(self.contest, "Contest Singer", "S1")
        self._registration(self.other, "Other Singer", "S2")
        login_admin(self.client, self.admin)
        response = self.client.get(reverse("staff:export_registrations"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("Contest Singer", self._first_column_from(response))
        self.assertIn("Other Singer", self._first_column_from(response))

    def test_staff_export_programs_scopes_to_selected_activity(self):
        self._program(self.contest, "Contest Program", "C1")
        self._program(self.other, "Other Program", "C2")
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:export_programs"), {"activity_id": self.other.pk})
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        ws = wb.active
        assert ws is not None
        names = [ws.cell(row=idx, column=2).value for idx in range(2, ws.max_row + 1)]
        self.assertEqual(names, ["Other Program"])

    def test_staff_incident_export_scopes_to_selected_activity(self):
        IncidentRecord.objects.create(
            activity=self.contest,
            occurred_at=timezone.now(),
            event_type=IncidentRecord.EventType.OTHER,
        )
        IncidentRecord.objects.create(
            activity=self.other,
            occurred_at=timezone.now(),
            event_type=IncidentRecord.EventType.OTHER,
        )
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:incident_export"), {"activity_id": self.other.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self._first_column_from(response), [self.other.title])


class SensitiveExportAuditSweepTests(TestCase):
    """§17 P1 — every sensitive export must emit a typed EXPORT audit.

    Sweeps each reachable data-export endpoint and asserts every one records an
    EXPORT AuditLog whose note names the exact export_type. This turns the
    manual "who exported what" audit sweep into a regression invariant: a future
    export that forgets audit_export fails this test.
    """

    def setUp(self):
        self.staff = User.objects.create_user(
            username="sweep-staff", password="pass", role=User.Role.STAFF
        )
        self.participant = User.objects.create_user(username="sweep-participant", password="pass")
        self.activity = Activity.objects.create(
            title="Sweep Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.registration = SingerRegistration.objects.create(
            activity=self.activity,
            user=self.participant,
            name="Sweep Singer",
            student_id="20260001",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            wechat="wx",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        self.program = Program.objects.create(
            activity=self.activity,
            user=self.participant,
            name="Sweep Program",
            program_type=Program.ProgramType.SONG,
            contact_name="Contact",
            contact_phone="13900000000",
            class_name="Class",
        )
        self.judge = Judge.objects.create(activity=self.activity, name="Sweep Judge")
        self.round_ = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
        )
        prepare_round(self.round_, self.staff)
        self.vote_session = VoteSession.objects.create(
            activity=self.activity,
            name="Sweep Vote",
            passcode="1234",
            start_time=timezone.now() - timedelta(minutes=1),
            end_time=timezone.now() + timedelta(minutes=10),
            is_open=True,
        )
        VoteOption.objects.create(vote_session=self.vote_session, singer=self.registration)
        IncidentRecord.objects.create(
            activity=self.activity,
            occurred_at=timezone.now(),
            event_type=IncidentRecord.EventType.OTHER,
        )
        self.client.force_login(self.staff)

    def _assert_typed_audit(self, expected_export_type: str):
        log = (
            AuditLog.objects.filter(action_type=AuditLog.ActionType.EXPORT, operator=self.staff)
            .order_by("-pk")
            .first()
        )
        self.assertIsNotNone(log)
        self.assertIn(expected_export_type, log.note)  # type: ignore[union-attr]
        self.assertIn(f'"activity_id": {self.activity.pk}', log.note)  # type: ignore[union-attr]

    def test_registrations_export_records_typed_audit(self):
        self.client.get(reverse("staff:export_registrations"), {"activity_id": self.activity.pk})
        self._assert_typed_audit("registration_list")

    def test_programs_export_records_typed_audit(self):
        self.client.get(reverse("staff:export_programs"), {"activity_id": self.activity.pk})
        self._assert_typed_audit("program_list")

    def test_material_checklist_export_records_typed_audit(self):
        self.client.get(reverse("staff:excel_material_checklist", args=[self.activity.pk]))
        self._assert_typed_audit("material_checklist")

    def test_score_template_export_records_typed_audit(self):
        self.client.get(reverse("staff:excel_score_template", args=[self.round_.pk]))
        self._assert_typed_audit("score_template")

    def test_vote_result_export_records_typed_audit(self):
        self.client.get(reverse("staff:vote_session_export", args=[self.vote_session.pk]))
        self._assert_typed_audit("vote_result")

    def test_incident_export_records_typed_audit(self):
        self.client.get(reverse("staff:incident_export"), {"activity_id": self.activity.pk})
        self._assert_typed_audit("incident_list")


class WordGenerateArchiveAuthorityTests(TestCase):
    """M0-Y: persistent Word assets stop at the archive boundary.

    A GeneratedDocument is a persistent business asset of the Activity. Once the
    activity is ARCHIVED it must not grow; only a preview/download remains, which
    never creates a row. Source data is re-read under the Activity lock, so a
    generation racing a phase change cannot freeze stale data as the latest doc.
    """

    def setUp(self):
        self.staff = User.objects.create_user(
            username="word-staff", password="pass", role=User.Role.STAFF
        )
        self.template = ArticleTemplate.objects.create(
            name="Notice",
            template_type=ArticleTemplate.TemplateType.PRELIMINARY_NOTICE,
            body="{title}",
        )
        self.activity = Activity.objects.create(
            title="Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REHEARSAL,
            is_test_mode=False,
        )
        self.media_root = tempfile.mkdtemp()
        self.override = override_settings(MEDIA_ROOT=self.media_root)
        self.override.enable()

    def tearDown(self):
        self.override.disable()
        shutil.rmtree(self.media_root, ignore_errors=True)
        super().tearDown()

    def test_persistent_generation_creates_document(self):
        doc_obj, content = generate_persistent_document(self.activity, self.template, self.staff)
        self.assertEqual(doc_obj.activity, self.activity)
        self.assertFalse(doc_obj.is_test_data)
        self.assertTrue(doc_obj.file.storage.exists(doc_obj.file.name or ""))
        self.assertTrue(content)
        self.assertTrue(
            AuditLog.objects.filter(
                action_type=AuditLog.ActionType.EXPORT,
                target=f"GeneratedDocument:{doc_obj.pk}",
            ).exists()
        )

    def test_persistent_generation_rejects_archived(self):
        self.activity.phase = Activity.Phase.ARCHIVED
        self.activity.is_locked = True
        self.activity.save(update_fields=["phase", "is_locked"])
        with self.assertRaises(PermissionDenied):
            generate_persistent_document(self.activity, self.template, self.staff)
        self.assertFalse(GeneratedDocument.objects.filter(activity=self.activity).exists())

    def test_persistent_generation_rejects_locked(self):
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        with self.assertRaises(PermissionDenied):
            generate_persistent_document(self.activity, self.template, self.staff)
        self.assertFalse(GeneratedDocument.objects.filter(activity=self.activity).exists())

    def test_archived_word_generate_is_preview_download_only(self):
        # Case C: an archived activity may produce a downloadable preview but must
        # not grow a persistent GeneratedDocument.
        self.activity.phase = Activity.Phase.ARCHIVED
        self.activity.is_locked = True
        self.activity.save(update_fields=["phase", "is_locked"])
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:word_generate", args=[self.template.pk, self.activity.pk])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertFalse(GeneratedDocument.objects.filter(activity=self.activity).exists())

    def test_locked_word_generate_rejects(self):
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        self.client.force_login(self.staff)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:word_generate", args=[self.template.pk, self.activity.pk])
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(GeneratedDocument.objects.filter(activity=self.activity).exists())

    def test_persisted_document_is_seen_by_archive_package(self):
        # Case A: the document persists before the archive, so the authoritative
        # archive package (built from DB state) must include it.
        doc_obj, _ = generate_persistent_document(self.activity, self.template, self.staff)
        names = [a.name for a in build_archive_package(self.activity)]
        self.assertIn(f"推文_{doc_obj.pk}.docx", names)


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class WordGenerateArchiveConcurrencyTests(TransactionTestCase):
    """M0-Y Case B/D: a generation racing an archive commit must not persist."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username="arch-word-staff", password="pass", role=User.Role.STAFF
        )
        self.template = ArticleTemplate.objects.create(
            name="Notice",
            template_type=ArticleTemplate.TemplateType.PRELIMINARY_NOTICE,
            body="{title}",
        )
        self.activity = Activity.objects.create(
            title="Concurrent Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REHEARSAL,
            is_test_mode=False,
        )

    def test_generation_never_lands_after_archive_commits(self):
        hold_held = threading.Event()
        release_lock = threading.Event()
        holder_error: dict[str, object] = {}

        def commit_archive():
            try:
                with transaction.atomic():
                    activity = Activity.objects.select_for_update().get(pk=self.activity.pk)
                    activity.phase = Activity.Phase.ARCHIVED
                    activity.is_locked = True
                    activity.save(update_fields=["phase", "is_locked"])
                    hold_held.set()
                    release_lock.wait(timeout=10)
            except Exception as error:  # pragma: no cover - diagnostic only
                holder_error["error"] = error
            finally:
                close_old_connections()

        holder = threading.Thread(target=commit_archive)
        holder.start()
        self.assertTrue(hold_held.wait(timeout=10))

        gen_result: dict[str, object] = {}

        def try_generate():
            close_old_connections()
            try:
                generate_persistent_document(self.activity, self.template, self.staff)
                gen_result["done"] = True
            except PermissionDenied as error:
                gen_result["rejected"] = str(error)
            except Exception as error:  # pragma: no cover - diagnostic only
                gen_result["error"] = repr(error)
            finally:
                close_old_connections()

        worker = threading.Thread(target=try_generate)
        worker.start()
        time.sleep(1)  # let the generation block on the Activity row lock
        release_lock.set()
        holder.join(timeout=10)
        worker.join(timeout=10)

        self.assertFalse(holder_error, holder_error)
        self.activity.refresh_from_db()
        self.assertEqual(self.activity.phase, Activity.Phase.ARCHIVED)
        self.assertNotIn("done", gen_result, gen_result)
        self.assertIn("rejected", gen_result, gen_result)
        self.assertFalse(GeneratedDocument.objects.filter(activity=self.activity).exists())


class PublicPostMoveLockTests(TestCase):
    """Moving a post between activities must respect the lock of both activities."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username="staff_post", password="pass", role=User.Role.STAFF
        )
        self.activity = Activity.objects.create(
            title="A",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )

    def test_staff_post_edit_cannot_move_from_locked_activity(self):
        unlocked = Activity.objects.create(
            title="B",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.activity.is_locked = True
        self.activity.save()
        post = PublicPost.objects.create(
            title="Post", related_activity=self.activity, created_by=self.staff
        )
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:post_edit", args=[post.pk]),
            {
                "title": "Post",
                "subtitle": "",
                "content": "",
                "post_type": PublicPost.PostType.NORMAL_ARTICLE,
                "status": PublicPost.Status.DRAFT,
                "sort_order": "0",
                "related_activity_id": str(unlocked.pk),
            },
        )
        self.assertEqual(response.status_code, 403)
        post.refresh_from_db()
        self.assertEqual(post.related_activity_id, self.activity.pk)

    def _edit_payload(self, *, related_activity_id, title="Post", status=None):
        payload = {
            "title": title,
            "subtitle": "",
            "content": "",
            "post_type": PublicPost.PostType.NORMAL_ARTICLE,
            "status": status or PublicPost.Status.DRAFT,
            "sort_order": "0",
            "related_activity_id": str(related_activity_id),
        }
        return payload

    def test_staff_post_edit_same_parent_edit_succeeds(self):
        post = PublicPost.objects.create(
            title="Post", related_activity=self.activity, created_by=self.staff
        )
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:post_edit", args=[post.pk]),
            self._edit_payload(related_activity_id=self.activity.pk, title="Updated"),
        )
        self.assertEqual(response.status_code, 302)
        post.refresh_from_db()
        self.assertEqual(post.title, "Updated")
        self.assertEqual(post.related_activity_id, self.activity.pk)

    def test_staff_post_edit_reparent_succeeds(self):
        target = Activity.objects.create(
            title="B",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REHEARSAL,
            is_test_mode=False,
        )
        post = PublicPost.objects.create(
            title="Post", related_activity=self.activity, created_by=self.staff
        )
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:post_edit", args=[post.pk]),
            self._edit_payload(related_activity_id=target.pk),
        )
        self.assertEqual(response.status_code, 302)
        post.refresh_from_db()
        self.assertEqual(post.related_activity_id, target.pk)

    def test_staff_post_edit_cannot_move_to_locked_activity(self):
        target = Activity.objects.create(
            title="B",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        target.is_locked = True
        target.save()
        post = PublicPost.objects.create(
            title="Post", related_activity=self.activity, created_by=self.staff
        )
        self.client.force_login(self.staff)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:post_edit", args=[post.pk]),
            self._edit_payload(related_activity_id=target.pk),
        )
        self.assertEqual(response.status_code, 403)
        post.refresh_from_db()
        self.assertEqual(post.related_activity_id, self.activity.pk)


class PublicPostOptimisticConcurrencyTests(TestCase):
    """§17 P1 — a stale post-edit form must never silently overwrite a co-editor.

    The existing select_for_update only guards reparent races; a field-level
    stale form (A saves, then B's untouched form submits) would overwrite A's
    content. The version snapshot makes that an explicit "已过期" rejection.
    """

    def setUp(self):
        self.staff = User.objects.create_user(
            username="optimistic-staff", password="pass", role=User.Role.STAFF
        )
        self.activity = Activity.objects.create(
            title="A",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.post = PublicPost.objects.create(
            title="Post", related_activity=self.activity, created_by=self.staff
        )

    def _payload(self, *, title="Post", base_version=None, status=None):
        payload = {
            "title": title,
            "subtitle": "",
            "content": "",
            "post_type": PublicPost.PostType.NORMAL_ARTICLE,
            "status": status or PublicPost.Status.DRAFT,
            "sort_order": "0",
            "related_activity_id": str(self.activity.pk),
        }
        if base_version is not None:
            payload["base_version"] = str(base_version)
        return payload

    def test_edit_with_current_version_bumps_and_saves(self):
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:post_edit", args=[self.post.pk]),
            self._payload(title="Current", base_version=self.post.version),
        )
        self.assertEqual(response.status_code, 302)
        self.post.refresh_from_db()
        self.assertEqual(self.post.title, "Current")
        self.assertEqual(self.post.version, 1)

    def test_edit_with_stale_version_is_rejected_without_overwrite(self):
        self.client.force_login(self.staff)
        # First save moves the post to version 1.
        first = self.client.post(
            reverse("staff:post_edit", args=[self.post.pk]),
            self._payload(title="Co-author's save", base_version=0),
        )
        self.assertEqual(first.status_code, 302)
        # A second staff still holds the stale base_version=0 form.
        stale = self.client.post(
            reverse("staff:post_edit", args=[self.post.pk]),
            self._payload(title="Stale overwrite attempt", base_version=0),
        )
        self.assertEqual(stale.status_code, 200)
        self.assertContains(stale, "已被其他人更新")
        self.post.refresh_from_db()
        self.assertEqual(self.post.title, "Co-author's save")
        self.assertEqual(self.post.version, 1)

    def test_edit_without_base_version_is_accepted_for_legacy_clients(self):
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:post_edit", args=[self.post.pk]),
            self._payload(title="Legacy save"),
        )
        self.assertEqual(response.status_code, 302)
        self.post.refresh_from_db()
        self.assertEqual(self.post.title, "Legacy save")
        self.assertEqual(self.post.version, 1)

    def test_create_initializes_version_to_zero(self):
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:post_create"),
            self._payload(title="Created"),
        )
        self.assertEqual(response.status_code, 302)
        created = PublicPost.objects.get(title="Created")
        self.assertEqual(created.version, 0)


@skipUnless(connection.vendor == "postgresql", "requires PostgreSQL row locks")
class PublicPostReparentConcurrencyTests(TransactionTestCase):
    """M0-X: a stale reparent must never bypass the current parent's authority."""

    def setUp(self):
        self.staff = User.objects.create_user(
            username="reparent-staff", password="pass", role=User.Role.STAFF
        )
        self.a = Activity.objects.create(
            title="A",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.b = Activity.objects.create(
            title="B",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.c = Activity.objects.create(
            title="C",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.post = PublicPost.objects.create(
            title="Post", related_activity=self.a, created_by=self.staff
        )

    def _reparent_request(self, target_id):
        return RequestFactory().post(
            "/x",
            {
                "title": "Post",
                "subtitle": "",
                "content": "",
                "post_type": PublicPost.PostType.NORMAL_ARTICLE,
                "status": PublicPost.Status.DRAFT,
                "sort_order": "0",
                "related_activity_id": str(target_id),
            },
        )

    def _run_stale_reparent(self):
        """T2: attempt a reparent off a stale hint (A) while the post lock is held.

        The holder thread commits A->B under the post row lock and only releases
        after T2 is already blocked on that same row, so T2's pre-lock hint is
        guaranteed to be the stale A.
        """
        lock_held = threading.Event()
        release_lock = threading.Event()
        holder_error: dict[str, object] = {}

        def commit_reparent_to_b():
            try:
                with transaction.atomic():
                    locked_post = PublicPost.objects.select_for_update().get(pk=self.post.pk)
                    locked_post.related_activity_id = self.b.pk
                    locked_post.save(update_fields=["related_activity_id"])
                    lock_held.set()
                    release_lock.wait(timeout=10)
            except Exception as error:  # pragma: no cover - diagnostic only
                holder_error["error"] = error
            finally:
                close_old_connections()

        holder = threading.Thread(target=commit_reparent_to_b)
        holder.start()
        self.assertTrue(lock_held.wait(timeout=10))

        t2_result: dict[str, object] = {}

        def try_stale_reparent():
            close_old_connections()
            try:
                request = self._reparent_request(self.c.pk)
                request.user = self.staff
                post_edit(request, self.post.pk)
                t2_result["done"] = True
            except PermissionDenied as error:
                t2_result["rejected"] = str(error)
            except Exception as error:  # pragma: no cover - diagnostic only
                t2_result["error"] = repr(error)
            finally:
                close_old_connections()

        t2 = threading.Thread(target=try_stale_reparent)
        t2.start()
        time.sleep(1)  # let T2 read the stale hint (A) and block on the post row
        release_lock.set()
        holder.join(timeout=10)
        t2.join(timeout=10)
        return holder_error, t2_result

    def test_stale_reparent_rejected_and_final_parent_holds(self):
        # T1 moves A->B; T2, based on a stale A->C attempt, must be rejected.
        holder_error, t2_result = self._run_stale_reparent()
        self.assertFalse(holder_error, holder_error)
        self.post.refresh_from_db()
        self.assertEqual(self.post.related_activity_id, self.b.pk)
        self.assertNotIn("done", t2_result, t2_result)
        self.assertIn("rejected", t2_result, t2_result)

    def test_stale_reparent_cannot_bypass_locked_new_parent(self):
        # After T1 A->B, B is locked. A stale T2 must not bypass B's lock to move
        # to C: it never locks B, so only the stale-parent check can refuse it.
        self.b.is_locked = True
        self.b.save()
        holder_error, t2_result = self._run_stale_reparent()
        self.assertFalse(holder_error, holder_error)
        self.post.refresh_from_db()
        self.b.refresh_from_db()
        self.assertEqual(self.post.related_activity_id, self.b.pk)
        self.assertTrue(self.b.is_locked)
        self.assertNotIn("done", t2_result, t2_result)
        self.assertIn("rejected", t2_result, t2_result)


class PublicPortalPublicationTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(
            username="pub-staff", password="pass", role=User.Role.STAFF
        )
        self.formal = Activity.objects.create(
            title="Formal Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.testing = Activity.objects.create(
            title="Test Contest",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )

    def _post_payload(self, *, status, related_activity_id):
        return {
            "title": "A Post",
            "subtitle": "",
            "content": "",
            "post_type": PublicPost.PostType.NORMAL_ARTICLE,
            "status": status,
            "sort_order": "0",
            "related_activity_id": str(related_activity_id),
        }

    def test_post_create_rejects_published_for_test_activity(self):
        self.client.force_login(self.staff)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:post_create"),
            self._post_payload(
                status=PublicPost.Status.PUBLISHED, related_activity_id=self.testing.pk
            ),
        )
        self.assertEqual(response.status_code, 403)
        self.assertFalse(PublicPost.objects.filter(title="A Post").exists())

    def test_post_create_allows_draft_for_test_activity(self):
        self.client.force_login(self.staff)
        response = self.client.post(
            reverse("staff:post_create"),
            self._post_payload(status=PublicPost.Status.DRAFT, related_activity_id=self.testing.pk),
        )
        self.assertEqual(response.status_code, 302)
        post = PublicPost.objects.get(title="A Post")
        self.assertEqual(post.status, PublicPost.Status.DRAFT)
        self.assertEqual(post.related_activity_id, self.testing.pk)

    def test_post_edit_rejects_published_when_moving_to_test_activity(self):
        post = PublicPost.objects.create(
            title="A Post", related_activity=self.formal, created_by=self.staff
        )
        self.client.force_login(self.staff)
        self.client.raise_request_exception = False
        response = self.client.post(
            reverse("staff:post_edit", args=[post.pk]),
            self._post_payload(
                status=PublicPost.Status.PUBLISHED, related_activity_id=self.testing.pk
            ),
        )
        self.assertEqual(response.status_code, 403)
        post.refresh_from_db()
        self.assertEqual(post.related_activity_id, self.formal.pk)

    def test_post_preview_staff_accessible_public_forbidden(self):
        draft = PublicPost.objects.create(
            title="Preview Draft",
            status=PublicPost.Status.DRAFT,
            related_activity=self.testing,
            created_by=self.staff,
        )
        self.client.force_login(self.staff)
        response = self.client.get(reverse("staff:post_preview", args=[draft.pk]))
        self.assertEqual(response.status_code, 200)
        # A staff preview must not expose a non-published post on the public route.
        self.client.force_login(User.objects.create_user(username="post-pub-anon", password="pass"))
        self.client.raise_request_exception = False
        public_response = self.client.get(reverse("public_portal:post_detail", args=[draft.pk]))
        self.assertEqual(public_response.status_code, 404)


class RoundScoresApiTests(TestCase):
    """M1-H rapid-entry API: grid, sparse cell save, stale-edit conflict, auto re-resolve."""

    def setUp(self):
        self.admin = User.objects.create_user(
            username="rapid-api-admin", password="pass", role=User.Role.ADMIN
        )
        self.activity = Activity.objects.create(
            title="Rapid Entry Activity",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.round = ContestRound.objects.create(
            activity=self.activity, round_type=ContestRound.RoundType.PRELIMINARY
        )
        self.singer = SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username="rapid-api-s", password="pass"),
            name="快速选手",
            student_id="2026rapid01",
            college="Info",
            class_name="CS1",
            phone="13800000000",
            song_name="Song",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        self.judge = Judge.objects.create(activity=self.activity, name="Judge A")
        RoundEntry.objects.create(round=self.round, singer=self.singer)
        RoundJudge.objects.create(round=self.round, judge=self.judge)
        self.round.status = ContestRound.Status.PREPARED
        self.round.save(update_fields=["status"])
        self.client.force_login(self.admin)

    def _post(self, payload):
        return self.client.post(
            reverse("staff:round_scores_api", args=[self.round.pk]),
            data=json.dumps(payload),
            content_type="application/json",
        )

    def test_get_returns_grid_and_version(self):
        response = self.client.get(reverse("staff:round_scores_api", args=[self.round.pk]))
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["version"], 0)
        self.assertFalse(data["matrix_complete"])
        self.assertEqual(len(data["grid"]), 1)
        self.assertEqual(data["grid"][0]["singer_id"], self.singer.pk)

    def test_post_applies_cell_and_bumps_version(self):
        response = self._post(
            {
                "base_version": 0,
                "cells": [{"singer_id": self.singer.pk, "judge_id": self.judge.pk, "score": "91"}],
            }
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["version"], 1)
        self.assertTrue(data["matrix_complete"])
        self.assertEqual(ScoreRecord.objects.get(round=self.round).score, Decimal("91"))

    def test_post_stale_base_version_is_conflict(self):
        self._post(
            {
                "base_version": 0,
                "cells": [{"singer_id": self.singer.pk, "judge_id": self.judge.pk, "score": "91"}],
            }
        )
        response = self._post(
            {
                "base_version": 0,
                "cells": [{"singer_id": self.singer.pk, "judge_id": self.judge.pk, "score": "92"}],
            }
        )
        self.assertEqual(response.status_code, 409)
        self.assertTrue(response.json()["conflict"])
        # No silent overwrite: the stored value is unchanged.
        self.assertEqual(ScoreRecord.objects.get(round=self.round).score, Decimal("91"))

    def test_post_out_of_range_score_is_400(self):
        response = self._post(
            {
                "base_version": 0,
                "cells": [{"singer_id": self.singer.pk, "judge_id": self.judge.pk, "score": "150"}],
            }
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(ScoreRecord.objects.filter(round=self.round).exists())

    def test_post_matrix_complete_triggers_resolve(self):
        from ruleset.models import ContestRuleset, RulesetVersion
        from singer_contest.models import StageResult

        ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="Rapid Ruleset",
            is_test_data=False,
            stage_key="快速赛段",
            round_keys={"r1": self.round.pk},
        )
        with authority_write(RULESET_FREEZE):
            RulesetVersion.objects.create(
                ruleset=ruleset,
                definition=json.dumps(
                    {
                        "schema_version": 1,
                        "nodes": [
                            {
                                "key": "assess_r1",
                                "type": "ASSESS",
                                "source": "entry",
                                "round": "r1",
                            },
                            {
                                "key": "rank1",
                                "type": "RANK",
                                "source": "assess_r1",
                                "descending": True,
                            },
                            {"key": "top1", "type": "SELECT", "source": "rank1", "count": 1},
                        ],
                    }
                ),
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        response = self._post(
            {
                "base_version": 0,
                "cells": [{"singer_id": self.singer.pk, "judge_id": self.judge.pk, "score": "97"}],
            }
        )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["matrix_complete"])
        self.assertEqual(data["resolved_status"], "ready_to_confirm")
        self.assertTrue(
            StageResult.objects.filter(stage_key="快速赛段", status="ready_to_confirm").exists()
        )


class ResultBoardTests(TestCase):
    """M1-H result board: latest-per-stage banners and block-grouped handcard."""

    _DEFINITION = json.dumps(
        {
            "schema_version": 1,
            "nodes": [
                {"key": "assess_r1", "type": "ASSESS", "source": "entry", "round": "r1"},
                {"key": "rank1", "type": "RANK", "source": "assess_r1", "descending": True},
                {"key": "top1", "type": "SELECT", "source": "rank1", "count": 1},
            ],
        }
    )

    def setUp(self):
        from ruleset.models import ContestRuleset, RulesetVersion

        self.staff = User.objects.create_user(
            username="result-board-staff", password="pass", role=User.Role.STAFF
        )
        self.activity = Activity.objects.create(
            title="Board Activity",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="Board Ruleset",
            is_test_data=False,
            stage_key="院十佳",
            announcement_blocks=[
                {"label": "直接晋级第三轮", "outcome_codes": ["direct"]},
                {"label": "进入复活赛", "outcome_codes": ["repechage"]},
                {"label": "本轮淘汰", "outcome_codes": ["eliminated"]},
            ],
        )
        with authority_write(RULESET_FREEZE):
            self.version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self._DEFINITION,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        self.client.force_login(self.staff)

    def _singer(self, index):
        return SingerRegistration.objects.create(
            activity=self.activity,
            user=User.objects.create_user(username=f"rb-{index}", password="pass"),
            name=f"选手{index}",
            student_id=f"7{index:05d}",
            college="Info",
            class_name="CS1",
            phone=f"139{index:08d}",
            song_name=f"Song {index}",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )

    def _stage(self, *, status, ruleset_hash, reasons=None, stage_key="院十佳"):
        # R9-2 added a unique (activity, stage_key, result_version) publication number, so
        # a helper that mints several rows for one stage_key must increment the version.
        last = (
            StageResult.objects.filter(activity=self.activity, stage_key=stage_key).aggregate(
                m=Max("result_version")
            )["m"]
            or 0
        )
        with authority_write(STAGE_RESULT_CONFIRM):
            return StageResult.objects.create(
                activity=self.activity,
                ruleset_version=self.version,
                created_by=self.staff,
                stage_key=stage_key,
                status=status,
                reasons=reasons or [],
                ruleset_hash=ruleset_hash,
                result_version=last + 1,
                # Identity is (ruleset_version, input_fingerprint): every distinct result
                # needs its own fingerprint so one version can publish multiple rows.
                input_fingerprint=f"fp-{ruleset_hash}",
                is_test_data=False,
                # M1-R9-Final: a confirmed row must carry its confirming trail (DB CHECK).
                confirmed_at=timezone.now() if status == StageResult.Status.CONFIRMED else None,
                confirmed_by=self.staff if status == StageResult.Status.CONFIRMED else None,
            )

    def _confirm(self, stage):
        # Decisions are created while the result is non-confirmed; this flips a READY
        # stage to CONFIRMED (with its trail) under the confirm authority, mirroring
        # confirm_stage_result() without its finality checks (these fixtures use mock
        # ruleset hashes).
        with authority_write(STAGE_RESULT_CONFIRM):
            stage.status = StageResult.Status.CONFIRMED
            stage.confirmed_at = timezone.now()
            stage.confirmed_by = self.staff
            stage.save(update_fields=["status", "confirmed_at", "confirmed_by"])

    def test_board_lists_latest_stage_per_stage_key(self):
        self._stage(
            status=StageResult.Status.REVIEW,
            ruleset_hash="hash-1",
            reasons=["缺少第三轮"],
        )
        ready = self._stage(status=StageResult.Status.READY_TO_CONFIRM, ruleset_hash="hash-2")
        StageDecision.objects.create(
            stage_result=ready,
            singer=self._singer(1),
            outcome_code="direct",
            rank=1,
            score=Decimal("91.00"),
            is_test_data=False,
        )
        self._confirm(ready)
        response = self.client.get(reverse("staff:activity_result_board", args=[self.activity.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "院十佳")
        self.assertContains(response, "可抄手卡")
        self.assertNotContains(response, "待人工核定")
        self.assertNotContains(response, "缺少第三轮")

    def test_board_shows_status_and_reasons_for_hold(self):
        self._stage(
            status=StageResult.Status.HOLD,
            ruleset_hash="hash-h",
            reasons=["缺少第五轮评分"],
        )
        response = self.client.get(reverse("staff:activity_result_board", args=[self.activity.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "待齐数据")
        self.assertContains(response, "缺少第五轮评分")

    def test_detail_groups_by_announcement_blocks_in_order(self):
        ready = self._stage(status=StageResult.Status.READY_TO_CONFIRM, ruleset_hash="hash-d")
        direct = self._singer(2)
        repechage = self._singer(3)
        eliminated = self._singer(4)
        StageDecision.objects.create(
            stage_result=ready,
            singer=direct,
            outcome_code="direct",
            rank=1,
            score=Decimal("92.00"),
            is_test_data=False,
        )
        StageDecision.objects.create(
            stage_result=ready,
            singer=repechage,
            outcome_code="repechage",
            rank=2,
            score=Decimal("85.00"),
            is_test_data=False,
        )
        StageDecision.objects.create(
            stage_result=ready,
            singer=eliminated,
            outcome_code="eliminated",
            rank=3,
            score=Decimal("70.00"),
            is_test_data=False,
        )
        self._confirm(ready)
        response = self.client.get(reverse("staff:stage_result_detail", args=[ready.pk]))
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertContains(response, "直接晋级第三轮")
        self.assertContains(response, "进入复活赛")
        self.assertContains(response, "本轮淘汰")
        self.assertContains(response, direct.name)
        self.assertContains(response, repechage.name)
        self.assertContains(response, eliminated.name)
        self.assertLess(body.index("直接晋级第三轮"), body.index("进入复活赛"))
        self.assertLess(body.index("进入复活赛"), body.index("本轮淘汰"))

    def test_detail_fallback_groups_by_outcome_when_unset(self):
        from ruleset.models import ContestRuleset, RulesetVersion

        # UNIQUE(activity): a second ContestRuleset needs its own activity.
        activity_b = Activity.objects.create(
            title="Fallback Activity",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=False,
        )
        ruleset_b = ContestRuleset.objects.create(
            activity=activity_b,
            name="No Blocks",
            is_test_data=False,
            stage_key="无分组",
        )
        with authority_write(RULESET_FREEZE):
            version_b = RulesetVersion.objects.create(
                ruleset=ruleset_b,
                definition=self._DEFINITION,
                is_current=True,
                status=RulesetVersion.Status.FROZEN,
            )
        with authority_write(STAGE_RESULT_CONFIRM):
            stage = StageResult.objects.create(
                activity=activity_b,
                ruleset_version=version_b,
                created_by=self.staff,
                stage_key="无分组",
                status=StageResult.Status.READY_TO_CONFIRM,
                reasons=[],
                ruleset_hash="hash-fb",
                is_test_data=False,
            )
        singer = SingerRegistration.objects.create(
            activity=activity_b,
            user=User.objects.create_user(username="fb-singer", password="pass"),
            name="选手5",
            student_id="70005",
            college="Info",
            class_name="CS1",
            song_name="Song 5",
            pre_status=SingerRegistration.PreStatus.APPROVED,
        )
        StageDecision.objects.create(
            stage_result=stage,
            singer=singer,
            outcome_code="direct",
            rank=1,
            score=Decimal("88.00"),
            is_test_data=False,
        )
        self._confirm(stage)
        response = self.client.get(reverse("staff:stage_result_detail", args=[stage.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "直接晋级")

    def test_stage_result_confirm_flips_to_confirmed(self):
        """§36-37: the 核定 POST locks a READY_TO_CONFIRM result into its handcard state.

        The stage is built through the real resolver (M1-R6 finality): the frozen version
        records its binding so confirm can recompute the current input fingerprint, and
        the consumed round is locked so finality checks hold.
        """
        from ruleset.models import RulesetVersion
        from singer_contest.models import ContestRound, Judge, ScoreRecord
        from singer_contest.services import run_ruleset

        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        contest_round = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="初赛",
        )
        judge = Judge.objects.create(activity=self.activity, name="评委甲")
        singer = self._singer(7)
        ScoreRecord.objects.create(
            round=contest_round, singer=singer, judge=judge, score=Decimal("91.00")
        )
        with authority_write(RULESET_FREEZE):
            version = RulesetVersion.objects.create(
                ruleset=self.ruleset,
                definition=self._DEFINITION,
                version=2,
                is_current=False,
                status=RulesetVersion.Status.FROZEN,
                binding={"stage_key": "院十佳", "round_keys": {"r1": contest_round.pk}},
            )
        contest_round.is_locked = True
        contest_round.status = ContestRound.Status.LOCKED
        contest_round.save(update_fields=["is_locked", "status"])
        # R9-4/R9-5: 核定 finalizes only a result grounded on the *current* FROZEN
        # authority, so demote the setUp current version, make this frozen v2 the single
        # current runner, then persist a formal StageResult (not an isolated preview).
        RulesetVersion._base_manager.filter(ruleset_id=version.ruleset_id).exclude(
            pk=version.pk
        ).update(is_current=False)
        RulesetVersion._base_manager.filter(pk=version.pk).update(is_current=True)
        version.refresh_from_db()
        ready = run_ruleset(
            version,
            self.activity,
            stage_key="院十佳",
            computed_by=self.staff,
            round_keys={"r1": contest_round},
        )
        self.assertEqual(ready.status, StageResult.Status.READY_TO_CONFIRM)
        response = self.client.post(reverse("staff:stage_result_confirm", args=[ready.pk]))  # type: ignore[union-attr]
        self.assertRedirects(response, reverse("staff:stage_result_detail", args=[ready.pk]))  # type: ignore[union-attr]
        ready.refresh_from_db()  # type: ignore[union-attr]
        self.assertEqual(ready.status, StageResult.Status.CONFIRMED)
        self.assertEqual(ready.confirmed_by, self.staff)  # type: ignore[union-attr]
        self.assertIsNotNone(ready.confirmed_at)  # type: ignore[union-attr]
        # M1-R8 Commit 4: the CONFIRM audit is recorded once by the service (rich payload),
        # not duplicated by the view with a stale READY_TO_CONFIRM status.
        from common.models import AuditLog

        audits = AuditLog.objects.filter(
            action_type=AuditLog.ActionType.CONFIRM_STAGE_RESULT,
            target=f"StageResult:{ready.pk}",  # type: ignore[union-attr]
        )
        self.assertEqual(audits.count(), 1)
        payload = json.loads(audits.get().new_value)
        self.assertEqual(payload["status"], StageResult.Status.CONFIRMED)
        self.assertEqual(payload["stage_key"], "院十佳")
        self.assertEqual(payload["result_version"], ready.result_version)
        self.assertEqual(payload["ruleset_version"], ready.ruleset_version_id)  # type: ignore[union-attr]
        self.assertEqual(payload["input_fingerprint"], ready.input_fingerprint)

    def test_stage_result_confirm_rejects_unresolved(self):
        """§36-37: 核定 a HOLD stage is refused and leaves it unchanged."""
        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        hold = self._stage(
            status=StageResult.Status.HOLD,
            ruleset_hash="hash-chold",
            reasons=["缺评分"],
        )
        response = self.client.post(reverse("staff:stage_result_confirm", args=[hold.pk]))
        self.assertRedirects(response, reverse("staff:stage_result_detail", args=[hold.pk]))
        hold.refresh_from_db()
        self.assertEqual(hold.status, StageResult.Status.HOLD)
        self.assertIsNone(hold.confirmed_by)

    def test_stage_result_unlock_requires_admin_and_reverts(self):
        """§38: only an admin may unlock a confirmed stage; it reverts to READY."""
        self.activity.phase = Activity.Phase.RESULTS_PENDING
        self.activity.save(update_fields=["phase"])
        confirmed = self._stage(status=StageResult.Status.CONFIRMED, ruleset_hash="hash-unlock")
        # A non-admin staff member is denied and the lock is preserved.
        self.client.post(reverse("staff:stage_result_unlock", args=[confirmed.pk]), {"note": "x"})
        confirmed.refresh_from_db()
        self.assertEqual(confirmed.status, StageResult.Status.CONFIRMED)
        # An admin can unlock (with the step-up admin verification), reverting the marker.
        admin = User.objects.create_user(
            username="result-unlock-admin", password="pass", role=User.Role.ADMIN
        )
        login_admin(self.client, admin)
        response = self.client.post(
            reverse("staff:stage_result_unlock", args=[confirmed.pk]), {"note": "核对录错了"}
        )
        self.assertRedirects(response, reverse("staff:stage_result_detail", args=[confirmed.pk]))
        confirmed.refresh_from_db()
        self.assertEqual(confirmed.status, StageResult.Status.READY_TO_CONFIRM)
        self.assertIsNone(confirmed.confirmed_by)
        self.assertIsNone(confirmed.confirmed_at)

    def test_board_shows_ready_to_confirm_banner(self):
        """§36-37: a machine-computed (not yet 核定) stage shows 待核定, not 可抄手卡."""
        self._stage(status=StageResult.Status.READY_TO_CONFIRM, ruleset_hash="hash-rtc")
        response = self.client.get(reverse("staff:activity_result_board", args=[self.activity.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "待核定")
        self.assertNotContains(response, "可抄手卡")

    def test_stage_decisions_by_blocks_helper(self):
        ready = self._stage(status=StageResult.Status.READY_TO_CONFIRM, ruleset_hash="hash-helper")
        singer = self._singer(6)
        StageDecision.objects.create(
            stage_result=ready,
            singer=singer,
            outcome_code="direct",
            rank=1,
            score=Decimal("90.00"),
            is_test_data=False,
        )
        self._confirm(ready)
        blocks = stage_decisions_by_blocks(ready)
        self.assertEqual(blocks[0]["label"], "直接晋级第三轮")
        self.assertEqual(blocks[0]["decisions"][0].singer, singer)


class RulesetTemplateLibraryTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(
            username="lib-staff",
            password="pass",
            role=User.Role.STAFF,
        )
        self.client.force_login(self.staff)
        from ruleset.templates import seed_ruleset_templates

        seed_ruleset_templates(self.staff)

    def test_template_list_renders_library(self):
        response = self.client.get(reverse("staff:ruleset_template_list"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "赛制模板库")
        self.assertContains(response, "院十佳")
        self.assertContains(response, "校十佳屏峰")

    def test_template_list_requires_staff(self):
        self.client.logout()
        response = self.client.get(reverse("staff:ruleset_template_list"))
        self.assertEqual(response.status_code, 302)

    def test_template_detail_renders_definition_nodes(self):
        from ruleset.models import RulesetTemplate

        template = RulesetTemplate.objects.get(name="院十佳")
        response = self.client.get(reverse("staff:ruleset_template_detail", args=[template.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, template.name)
        self.assertContains(response, "assess_r1")
        self.assertContains(response, "AGGREGATE")
        self.assertContains(response, "节点")

    def _clone_activity(self, title="院十佳2026"):
        return Activity.objects.create(
            title=title,
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )

    def test_clone_from_template_creates_ruleset_and_version(self):
        from ruleset.models import ContestRuleset, RulesetTemplate, RulesetVersion

        activity = self._clone_activity()
        template = RulesetTemplate.objects.get(name="院十佳")
        response = self.client.post(
            reverse("staff:ruleset_clone_from_template", args=[template.pk]),
            {"activity": activity.pk, "name": "院十佳2026克隆"},
        )
        ruleset = ContestRuleset.objects.get(activity=activity)
        self.assertEqual(ruleset.name, "院十佳2026克隆")
        self.assertEqual(ruleset.source_template, template)
        version = RulesetVersion.objects.get(ruleset=ruleset)
        self.assertEqual(version.status, "draft")
        self.assertEqual(version.definition, template.definition)
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[version.pk]))

    def test_clone_into_activity_with_existing_ruleset_reuses_it(self):
        from ruleset.models import ContestRuleset, RulesetTemplate, RulesetVersion

        activity = self._clone_activity("院十佳2027")
        existing = ContestRuleset.objects.create(
            activity=activity,
            name="已有赛制",
            is_test_data=True,
            created_by=self.staff,
        )
        template = RulesetTemplate.objects.get(name="院十佳")
        response = self.client.post(
            reverse("staff:ruleset_clone_from_template", args=[template.pk]),
            {"activity": activity.pk, "name": "院十佳2027复制"},
        )
        self.assertEqual(ContestRuleset.objects.filter(activity=activity).count(), 1)
        ruleset = ContestRuleset.objects.get(activity=activity)
        self.assertEqual(ruleset.pk, existing.pk)
        self.assertEqual(ruleset.name, "院十佳2027复制")
        self.assertEqual(ruleset.source_template, template)
        version = RulesetVersion.objects.get(ruleset=ruleset)
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[version.pk]))

    def test_clone_twice_into_same_activity_avoids_one_current_collision(self):
        from ruleset.models import ContestRuleset, RulesetTemplate, RulesetVersion

        activity = self._clone_activity("院十佳2028")
        template = RulesetTemplate.objects.get(name="院十佳")
        self.client.post(
            reverse("staff:ruleset_clone_from_template", args=[template.pk]),
            {"activity": activity.pk, "name": "赛制A"},
        )
        response = self.client.post(
            reverse("staff:ruleset_clone_from_template", args=[template.pk]),
            {"activity": activity.pk, "name": "赛制B"},
        )
        ruleset = ContestRuleset.objects.get(activity=activity)
        versions = list(RulesetVersion.objects.filter(ruleset=ruleset).order_by("version"))
        self.assertEqual(len(versions), 2)
        self.assertEqual([v.version for v in versions], [1, 2])
        # A cloned DRAFT is an editing target, never the official authority (M1-R8),
        # so repeated clones never claim is_current until a successor is frozen.
        self.assertEqual(RulesetVersion.objects.filter(ruleset=ruleset, is_current=True).count(), 0)
        latest = versions[-1]
        latest.refresh_from_db()
        self.assertFalse(latest.is_current)
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[latest.pk]))

    def test_clone_last_year_picks_golden_template_by_name(self):
        from ruleset.models import ContestRuleset, RulesetVersion

        activity = self._clone_activity("校十佳2026")
        response = self.client.post(
            reverse("staff:ruleset_clone_last_year"),
            {"activity": activity.pk, "name": "校十佳2026"},
        )
        ruleset = ContestRuleset.objects.get(activity=activity)
        self.assertEqual(ruleset.source_template.name, "院十佳")  # type: ignore[union-attr]
        version = RulesetVersion.objects.get(ruleset=ruleset)
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[version.pk]))

    def test_clone_requires_staff_and_post(self):
        activity = self._clone_activity()
        from ruleset.models import RulesetTemplate

        template = RulesetTemplate.objects.get(name="院十佳")
        self.client.logout()
        response = self.client.post(
            reverse("staff:ruleset_clone_from_template", args=[template.pk]),
            {"activity": activity.pk, "name": "x"},
        )
        self.assertEqual(response.status_code, 302)


def _acceptance_definition_json():
    """院十佳2026 walkthrough: 16人 R1 40% R2 60% → Top12 → fin Top6."""
    import json

    from ruleset.schema import ENTRY_KEY

    return json.dumps(
        {
            "schema_version": 1,
            "nodes": [
                {"key": "assess_r1", "type": "ASSESS", "source": ENTRY_KEY, "round": "r1"},
                {"key": "assess_r2", "type": "ASSESS", "source": ENTRY_KEY, "round": "r2"},
                {
                    "key": "stage1",
                    "type": "AGGREGATE",
                    "within": ENTRY_KEY,
                    "aggregate": {
                        "type": "weighted_sum",
                        "components": [
                            {"source": "assess_r1", "weight": 0.4},
                            {"source": "assess_r2", "weight": 0.6},
                        ],
                    },
                },
                {"key": "rank1", "type": "RANK", "source": "stage1", "descending": True},
                {"key": "top12", "type": "SELECT", "source": "rank1", "count": 12},
                {"key": "assess_r3", "type": "ASSESS", "source": "top12", "round": "r3"},
                {
                    "key": "fin",
                    "type": "AGGREGATE",
                    "within": "top12",
                    "aggregate": {
                        "type": "weighted_sum",
                        "components": [
                            {"source": "stage1", "weight": 0.5},
                            {"source": "assess_r3", "weight": 0.5},
                        ],
                    },
                },
                {"key": "rank2", "type": "RANK", "source": "fin", "descending": True},
                {"key": "top6", "type": "SELECT", "source": "rank2", "count": 6},
            ],
        },
        ensure_ascii=False,
    )


def _seed_bound_acceptance(activity, ruleset):
    """Lay down the DB facts the acceptance walkthrough needs to bound-compile.

    Always-bound freeze (P0-2) derives entry_size / round scales from the DB, so the
    walkthrough must have enough approved singers and each referenced round bound with a
    rubric whose highest criterion max_score is >= 90 (→ hundred scale). No judges are
    needed because the walkthrough uses mean mode.
    """
    from singer_contest.models import (
        ContestRound,
        RubricCriterion,
        ScoringRubric,
        SingerRegistration,
    )

    for i in range(12):
        SingerRegistration.objects.create(
            activity=activity,
            user=User.objects.create_user(username=f"ab-entry-{i}", password="pass"),
            name=f"Entry{i}",
            student_id=f"{i:04d}",
            pre_status=SingerRegistration.PreStatus.APPROVED,
            is_test_data=True,
        )
    bound = {}
    for key in ("r1", "r2", "r3"):
        rubric = ScoringRubric.objects.create(
            activity=activity, name=f"{key}评分", is_test_data=True
        )
        RubricCriterion.objects.create(rubric=rubric, name="总分", max_score=100, is_test_data=True)
        round_ = ContestRound.objects.create(
            activity=activity, round_type=ContestRound.RoundType.PRELIMINARY, name=key
        )
        round_.rubric = rubric
        round_.save(update_fields=["rubric"])
        bound[key] = round_.pk
    ruleset.round_keys = bound
    ruleset.save(update_fields=["round_keys"])


def _invalid_definition_json():
    """Parse-valid but compiles with an ERROR (AGGREGATE weights sum to 0.9)."""
    import json

    from ruleset.schema import ENTRY_KEY

    return json.dumps(
        {
            "schema_version": 1,
            "nodes": [
                {"key": "assess_r1", "type": "ASSESS", "source": ENTRY_KEY, "round": "r1"},
                {"key": "assess_r2", "type": "ASSESS", "source": ENTRY_KEY, "round": "r2"},
                {
                    "key": "comp",
                    "type": "AGGREGATE",
                    "within": ENTRY_KEY,
                    "aggregate": {
                        "type": "weighted_sum",
                        "components": [
                            {"source": "assess_r1", "weight": 0.5},
                            {"source": "assess_r2", "weight": 0.4},
                        ],
                    },
                },
                {"key": "rank", "type": "RANK", "source": "comp", "descending": True},
                {"key": "top", "type": "SELECT", "source": "rank", "count": 5},
            ],
        },
        ensure_ascii=False,
    )


class RulesetEditorTests(TestCase):
    def setUp(self):
        from ruleset.models import ContestRuleset, RulesetVersion

        self.admin = User.objects.create_user(
            username="editor-admin",
            password="pass",
            role=User.Role.ADMIN,
        )
        self.activity = Activity.objects.create(
            title="院十佳2026",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        self.client.force_login(self.admin)
        self.ruleset = ContestRuleset.objects.create(
            activity=self.activity,
            name="院十佳2026规则",
            is_test_data=True,
            created_by=self.admin,
        )
        self.version = RulesetVersion.objects.create(
            ruleset=self.ruleset,
            definition=_acceptance_definition_json(),
            created_by=self.admin,
        )

    def _nodes(self):
        import json

        self.version.refresh_from_db()
        return json.loads(self.version.definition)["nodes"]

    def test_contest_ruleset_create_reuses_single_activity_authority(self):
        from ruleset.models import ContestRuleset

        # §16: one activity owns exactly one ContestRuleset; versions hold the history.
        # Re-creating a ruleset for an activity that already has one must reuse that single
        # authority (editing its DRAFT) instead of building a competing one.
        response = self.client.post(
            reverse("staff:contest_ruleset_create"),
            {"activity": self.activity.pk, "name": "新赛制2026"},
        )
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[self.version.pk]))
        self.assertEqual(ContestRuleset.objects.filter(activity=self.activity).count(), 1)
        self.assertTrue(ContestRuleset.objects.filter(name="院十佳2026规则").exists())
        self.assertFalse(ContestRuleset.objects.filter(name="新赛制2026").exists())

    def test_contest_ruleset_create_requires_unlocked_activity(self):
        """R0: ruleset creation must re-validate the activity lock (DRAFT/FORMAL)."""
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        response = self.client.post(
            reverse("staff:contest_ruleset_create"),
            {"activity": self.activity.pk, "name": "锁定活动禁止建赛制"},
        )
        self.assertEqual(response.status_code, 403)

    def test_rule_edit_requires_unlocked_activity(self):
        """R0: ruleset edit must re-validate the activity lock before mutating."""
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        response = self.client.post(
            reverse("staff:ruleset_edit", args=[self.version.pk]),
            {"action": "add", "new_type": "ASSESS"},
        )
        self.assertEqual(response.status_code, 403)

    def test_editor_validate_clean_on_acceptance_definition(self):
        response = self.client.post(reverse("staff:ruleset_validate", args=[self.version.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "校验通过")

    def test_editor_preview_shows_candidate_pool_sizes(self):
        response = self.client.post(reverse("staff:ruleset_preview", args=[self.version.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "READY")
        self.assertContains(response, "候选池规模")

    def test_editor_add_node_roundtrips(self):
        before = len(self._nodes())
        response = self.client.post(
            reverse("staff:ruleset_edit", args=[self.version.pk]),
            {"action": "add", "new_type": "ASSESS"},
        )
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[self.version.pk]))
        nodes = self._nodes()
        self.assertEqual(len(nodes), before + 1)
        self.assertEqual(nodes[-1]["type"], "ASSESS")

    def test_editor_move_node_reorders(self):
        response = self.client.post(
            reverse("staff:ruleset_edit", args=[self.version.pk]),
            {"action": "move_up", "key": "assess_r2"},
        )
        self.assertEqual(response.status_code, 302)
        keys = [n["key"] for n in self._nodes()]
        self.assertLess(keys.index("assess_r2"), keys.index("assess_r1"))

    def test_editor_delete_node_removes(self):
        before = len(self._nodes())
        response = self.client.post(
            reverse("staff:ruleset_edit", args=[self.version.pk]),
            {"action": "delete", "key": "top6"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(len(self._nodes()), before - 1)

    def test_editor_refuses_frozen_version_edit(self):
        self.version.status = "frozen"
        with authority_write(RULESET_FREEZE):
            self.version.save(update_fields=["status"])
        response = self.client.get(reverse("staff:ruleset_edit", args=[self.version.pk]))
        self.assertEqual(response.status_code, 403)

    def test_editor_hides_unsupported_runtime_node_types(self):
        # §28 capability matrix: BRANCH/AWARD are compile-rejected
        # (NODE_UNSUPPORTED_RUNTIME), so the editor must not offer them in the
        # add-node dropdown.
        response = self.client.get(reverse("staff:ruleset_edit", args=[self.version.pk]))
        self.assertEqual(response.status_code, 200)
        node_types = list(response.context["node_types"])
        self.assertNotIn("BRANCH", node_types)
        self.assertNotIn("AWARD", node_types)

    def test_editor_add_action_rejects_unsupported_runtime_node(self):
        before = len(self._nodes())
        response = self.client.post(
            reverse("staff:ruleset_edit", args=[self.version.pk]),
            {"action": "add", "new_type": "BRANCH"},
        )
        # ValueError in _edit_nodes -> messages.error + redirect, no mutation.
        self.assertEqual(response.status_code, 302)
        self.assertEqual(len(self._nodes()), before)

    def test_freeze_refuses_invalid_and_accepts_valid(self):
        from ruleset.models import RulesetVersion

        invalid = RulesetVersion.objects.create(
            ruleset=self.ruleset,
            definition=_invalid_definition_json(),
            version=99,
            is_current=False,
            created_by=self.admin,
        )
        response = self.client.post(reverse("staff:ruleset_freeze", args=[invalid.pk]))
        self.assertEqual(response.status_code, 302)
        invalid.refresh_from_db()
        self.assertEqual(invalid.status, "draft")

        # Always-bound freeze (P0-2) only passes once the activity carries the facts the
        # walkthrough references: enough approved singers + each round bound with a scale.
        _seed_bound_acceptance(self.activity, self.ruleset)
        response = self.client.post(reverse("staff:ruleset_freeze", args=[self.version.pk]))
        self.assertEqual(response.status_code, 302)
        self.version.refresh_from_db()
        self.assertEqual(self.version.status, "frozen")
        self.assertTrue(self.version.is_current)
        self.assertTrue(self.version.execution_plan)

    def _bind_round(self):
        from singer_contest.models import ContestRound

        return ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="r1",
        )

    def _bind_vote(self):
        from voting.models import VoteSession

        return VoteSession.objects.create(
            activity=self.activity,
            name="大众投票",
            passcode="0000",
            start_time=timezone.now(),
            end_time=timezone.now() + timedelta(hours=1),
            is_test_data=True,
        )

    def test_ruleset_bind_persists_json_fields(self):
        round_ = self._bind_round()
        vote = self._bind_vote()
        response = self.client.post(
            reverse("staff:ruleset_bind", args=[self.version.pk]),
            {
                "stage_key": "院十佳",
                "round_keys": json.dumps({"r1": round_.pk}),
                "vote_keys": json.dumps({"audience": vote.pk}),
                "group_keys": json.dumps({"initial_group": round_.pk}),
                "announcement_blocks": json.dumps([{"label": "晋级", "outcome_codes": ["direct"]}]),
                "announcement_blocks_by_checkpoint": json.dumps(
                    {"stage2": [{"label": "赛段专属", "outcome_codes": ["advanced"]}]}
                ),
            },
        )
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[self.version.pk]))
        self.ruleset.refresh_from_db()
        self.assertEqual(self.ruleset.stage_key, "院十佳")
        self.assertEqual(self.ruleset.round_keys, {"r1": round_.pk})
        self.assertEqual(self.ruleset.vote_keys, {"audience": vote.pk})
        self.assertEqual(self.ruleset.group_keys, {"initial_group": round_.pk})
        self.assertEqual(
            self.ruleset.announcement_blocks,
            [{"label": "晋级", "outcome_codes": ["direct"]}],
        )
        self.assertEqual(
            self.ruleset.announcement_blocks_by_checkpoint,
            {"stage2": [{"label": "赛段专属", "outcome_codes": ["advanced"]}]},
        )

    def test_ruleset_editor_bind_form_prerenders_json(self):
        from singer_contest.models import ContestRound

        round_ = ContestRound.objects.create(
            activity=self.activity,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="r1",
        )
        self.ruleset.stage_key = "院十佳"
        self.ruleset.round_keys = {"r1": round_.pk}
        self.ruleset.group_keys = {"by": round_.pk}
        self.ruleset.save()
        response = self.client.get(reverse("staff:ruleset_edit", args=[self.version.pk]))
        self.assertEqual(response.status_code, 200)
        # The bind form must pre-render maps as JSON (autoescaped to &quot;) so a later
        # POST of the same page round-trips through json.loads rather than a dict repr.
        self.assertContains(response, "{&quot;r1&quot;: %d}" % round_.pk)

    def test_ruleset_bind_bad_json_errors_without_mutation(self):
        response = self.client.post(
            reverse("staff:ruleset_bind", args=[self.version.pk]),
            {
                "stage_key": "院十佳",
                "round_keys": "{not-json",
                "vote_keys": "{}",
                "group_keys": "{}",
                "announcement_blocks": "[]",
            },
        )
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[self.version.pk]))
        self.ruleset.refresh_from_db()
        self.assertEqual(self.ruleset.round_keys, {})
        self.assertEqual(self.ruleset.stage_key, "")

    def test_ruleset_bind_rejects_foreign_round(self):
        from singer_contest.models import ContestRound

        other = Activity.objects.create(
            title="其他活动",
            activity_type=Activity.Type.SINGER_CONTEST,
            phase=Activity.Phase.REGISTRATION_OPEN,
            is_test_mode=True,
        )
        foreign_round = ContestRound.objects.create(
            activity=other,
            round_type=ContestRound.RoundType.PRELIMINARY,
            name="f",
        )
        response = self.client.post(
            reverse("staff:ruleset_bind", args=[self.version.pk]),
            {
                "stage_key": "院十佳",
                "round_keys": json.dumps({"r1": foreign_round.pk}),
                "vote_keys": "{}",
                "group_keys": "{}",
                "announcement_blocks": "[]",
            },
        )
        self.assertRedirects(response, reverse("staff:ruleset_edit", args=[self.version.pk]))
        self.ruleset.refresh_from_db()
        self.assertEqual(self.ruleset.round_keys, {})

    def test_ruleset_bind_requires_unlocked_activity(self):
        self.activity.is_locked = True
        self.activity.save(update_fields=["is_locked"])
        response = self.client.post(
            reverse("staff:ruleset_bind", args=[self.version.pk]),
            {"stage_key": "院十佳"},
        )
        self.assertEqual(response.status_code, 403)
