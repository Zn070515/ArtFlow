import io
import json
from decimal import Decimal, InvalidOperation
from typing import cast
from urllib.parse import urlsplit, urlunsplit
from uuid import uuid4

from accounts.decorators import admin_required, staff_required
from accounts.models import User
from accounts.services import (
    admin_verification_is_valid,
    change_user_role,
    require_current_admin,
    set_user_active,
)
from common.audit import audit_export, log_action
from common.authority import ACTIVITY_STATE, authority_write
from common.business_rules import (
    ensure_activity_unlocked,
    ensure_lifecycle_consistent,
    ensure_round_unlocked,
    ensure_same_activity,
)
from common.lifecycle import (
    runtime_approved_singers,
    runtime_is_test,
    scope_lifecycle,
    scope_runtime,
)
from common.models import AuditLog
from common.test_data import (
    clear_activity_test_data,
    get_test_data_counts,
    leave_test_mode,
    lock_activity_for_runtime_data,
)
from core.models import Activity
from core.policies import ActivityAction, ensure_activity_action_allowed
from core.services import (
    lock_activity_for_action,
    phase_choices_for_activity,
    transition_activity_phase,
    unarchive_activity,
)
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.views import redirect_to_login
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.db.models import Count, Prefetch, QuerySet
from django.http import Http404, HttpResponse, JsonResponse
from django.http.request import QueryDict
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from exports.models import ArticleTemplate
from exports.pdf import render_stage_result_pdf
from exports.services import (
    INCIDENT_LIST_HEADERS,
    archive_activity,
    build_archive_package,
    build_execution_package,
    build_package_zip,
    build_score_template_workbook,
    generate_persistent_document,
    incident_list_rows,
    render_document_bytes,
)
from farewell_show.models import Program
from files.models import MaterialCheck, MaterialRequirement, StaffNote, SubmissionFile
from files.services import (
    questionnaire_material_authority_active,
    reconcile_activity_material_checks,
    reconcile_program_material_checks,
    reconcile_singer_material_checks,
    review_material_check,
    store_submission_file,
)
from incidents.models import IncidentRecord
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from public_portal.models import PublicPost, ResultRelease
from public_portal.services import (
    qr_print_host_warning,
    release_result_post,
    revoke_result_release,
)
from questionnaire.projection import (
    generic_song_label,
    prime_questionnaire_answers,
    questionnaire_song_for_round,
)
from realtime.events import schedule_activity_event
from ruleset import editor as ruleset_editor
from ruleset.business_editor import infer_business_stages
from ruleset.compiler import ValidationReport, compile_definition
from ruleset.live_format import (
    PREPARATION_TASK_ROW_KEYS,
    _bound_first_round,
    apply_live_2026_counts,
    bind_live_2026_group_stage,
    clone_first_round_for_successor,
    create_live_2026_group_stage,
    first_round_configuration_signature,
    freeze_blocker,
    is_live_2026_template,
    live_2026_preparation_tasks,
    live_2026_readiness,
    live_challenge_counts_for_activity,
    live_challenge_counts_from_binding,
    provision_live_2026_ruleset,
    save_live_2026_preparation,
)
from ruleset.models import ContestRuleset, RulesetTemplate, RulesetVersion
from ruleset.planning import (
    PlanningCompileResult,
    planning_projection_from_definition,
)
from ruleset.schema import ENTRY_KEY, NODE_TYPE_SPEC, OutputType, parse_definition
from ruleset.services import (
    LIVE_CHALLENGE_FORMAT_KEY,
    RulesetInvalidError,
    _binding_signature,
    _snapshot_binding,
    apply_planning_to_ruleset,
    business_round_configuration_signature,
    create_ruleset_version,
    current_frozen_round_ids,
    editable_ruleset_version,
    freeze_ruleset_version,
    newer_ruleset_version,
    prepare_planning_compile_for_version,
    prepare_planning_preflight,
    read_planning_draft,
    save_business_ruleset,
    save_planning_draft,
    supersede_ruleset_version,
    update_ruleset_binding,
    update_ruleset_definition_section,
    validate_definition_completeness,
    version_is_editable,
)
from ruleset.templates import GOLDEN_SCHIDUI_BUILTIN_KEY, LIVE_2026_PODIUM_KEYS
from singer_contest.challenge_services import (
    advance_challenge_performance,
    challenge_session_revoke_blocker,
    confirm_challenge_match_result,
    confirm_challenge_pairing,
    confirm_challenge_session,
    create_challenge_session,
    create_final_round_from_challenge,
    lock_challenge_match_scores,
    lock_challenge_performance,
    prepare_challenge_match,
    record_challenge_target,
    revoke_challenge_session,
    start_challenge_turn,
    waive_challenge_turn,
)
from singer_contest.group_chorus import (
    cancel_group_stage,
    confirm_group_stage,
    correct_group_stage,
    create_group_stage,
    freeze_group_stage,
    group_material_readiness,
    group_stage_roster_singers,
    record_group_stage,
    set_group_material_status,
)
from singer_contest.judge_authority import (
    advance_performance,
    advance_performance_group,
    hold_judge_panel,
    hold_performance,
    lock_performance_score,
    prepare_judge_panel,
    release_judge_seat,
    resume_judge_panel,
    resume_performance,
    set_judge_seat_display_label,
    submit_paper_score,
    submit_staff_proxy_score,
)
from singer_contest.judge_entry import judge_entry_credential
from singer_contest.materials import live_material_summary
from singer_contest.models import (
    AudienceScore,
    ChallengeMatch,
    ChallengeSession,
    ChallengeTurn,
    ContestRound,
    Group,
    GroupMembership,
    GroupStage,
    Judge,
    JudgeSeat,
    JudgeSession,
    Performance,
    PerformanceGroup,
    PerformanceRunState,
    PerformanceScoreState,
    RoundEntry,
    RoundPanelSnapshot,
    RubricCriterion,
    ScoreRecord,
    ScoreSummary,
    ScoringRubric,
    SingerRegistration,
    StageResult,
)
from singer_contest.services import (
    IdempotencyConflictError,
    ResultClosureCode,
    RoundSetupReadiness,
    StaleScoreVersionError,
    _current_frozen_version,
    _current_resolve_status,
    _eligible_singers,
    _live_group_stage_for_round,
    _version_binding,
    apply_scores,
    apply_scores_if_version,
    authoritative_panel_judges,
    build_result_closure,
    confirm_stage_result,
    create_scoring_rubric,
    current_round_roster,
    ensure_audience_not_consumed_by_confirmed_stage,
    finalize_advancement,
    latest_stage_result_queryset,
    live_first_round_group_snapshot,
    lock_round,
    maybe_resolve_checkpoints,
    missing_score_cells,
    official_stage_award_queryset,
    parse_score_workbook,
    prepare_round,
    reset_round_to_draft,
    result_closure_as_dict,
    round_execution_authority_label,
    round_execution_authority_message,
    round_roster_stage_choices,
    round_setup_readiness,
    set_manual_decision,
    set_official_contestant_number,
    set_round_groups,
    set_round_running_order,
    stage_decisions_by_blocks,
    unlock_round,
    unlock_stage_result,
    validate_roster_source_stage,
)
from voting.models import VoteOption, VoteScoringRule, VoteSession
from voting.policies import formal_singer_contest_requires_ticket
from voting.services import (
    close_vote_session,
    lock_vote_session,
    open_vote_session,
    unlock_vote_session,
    valid_record_queryset,
    valid_vote_count,
    vote_session_configuration_facts,
)

from staff_panel.forms import (
    CREATE_PHASE_CHOICES,
    ActivityForm,
    ContestRoundForm,
    GroupStageAssignmentForm,
    GroupStageCorrectionForm,
    GroupStageCreateForm,
    IncidentForm,
    JudgeBoundScoreForm,
    JudgePanelAttendanceForm,
    JudgePerformanceActionForm,
    ManualDecisionForm,
    OfficialContestantNumberForm,
    ProgramReviewForm,
    PublicPostForm,
    RapidScoreCommandForm,
    ResultReleaseForm,
    ResultReleaseRevokeForm,
    RoundGroupsForm,
    RoundRunningOrderForm,
    ScoringRubricProvisionForm,
    SingerReviewForm,
    VoteSessionForm,
)


def _with_generic_song_labels(registrations):
    """Materialize registrations with an explicit legacy/multi-round song label."""
    rows = prime_questionnaire_answers(registrations)
    for registration in rows:
        registration.song_label = generic_song_label(registration)
    return rows


def _choices(enum_class):
    return enum_class.choices


def _incident_rounds(activity_id):
    """The rounds an incident on this activity may be filed against (GOAL §16)."""
    return list(ContestRound.objects.filter(activity_id=activity_id).order_by("pk"))


def _form_error(form):
    for field_errors in form.errors.values():
        if field_errors:
            return str(field_errors[0])
    return "提交的数据无效。"


def domain_error_messages(error) -> str:
    """Return the stable user-facing text for an expected domain failure."""
    error_messages = getattr(error, "messages", None)
    if error_messages:
        return "；".join(str(message) for message in error_messages)
    return str(error)


def safe_domain_error_messages(error) -> str:
    """Return domain messages without exposing arbitrary exception text to JSON clients."""
    error_messages = getattr(error, "messages", None)
    if error_messages:
        return "；".join(str(message) for message in error_messages)
    return "请求无法完成，请检查当前状态后重试。"


_AUTO_RESOLVE_WARNING = "该赛段无法自动核定，请人工核定。"


def _require_admin(user):
    return require_current_admin(user)


def _ensure_activity_mutable(activity):
    if activity.phase == Activity.Phase.ARCHIVED:
        raise PermissionDenied("活动已归档，为只读状态。")
    ensure_activity_unlocked(activity)


def _editable_singer_activities():
    """Return activities that can accept new singer-contest configuration."""
    return (
        Activity.objects.filter(
            activity_type=Activity.Type.SINGER_CONTEST,
            is_locked=False,
        )
        .exclude(phase=Activity.Phase.ARCHIVED)
        .order_by("-created_at")
    )


def _ensure_publication_allowed(related_activity, status):
    if (
        status == PublicPost.Status.PUBLISHED
        and related_activity is not None
        and related_activity.data_lifecycle == Activity.DataLifecycle.TEST
    ):
        raise PermissionDenied("测试活动的公开内容不能直接发布，请保存为草稿或隐藏。")


def _active_worksheet(workbook):
    worksheet = workbook.active
    if not isinstance(worksheet, Worksheet):
        raise ValueError("The active workbook sheet must be a worksheet.")
    return worksheet


@staff_required
def dashboard(request):
    return render(request, "staff_panel/dashboard.html")


@staff_required
def activity_list(request):
    activities = Activity.objects.all()
    return render(request, "staff_panel/activity_list.html", {"activities": activities})


@staff_required
def live_challenge_console(request, pk):
    """One plain-language control surface for the chronological challenge stage."""
    activity = get_object_or_404(Activity, pk=pk)
    if activity.activity_type != Activity.Type.SINGER_CONTEST:
        raise PermissionDenied("只有歌手比赛活动支持现场挑战。")
    if request.method == "POST":
        action = (request.POST.get("action") or "").strip()
        try:
            command_id = request.POST.get("command_id")
            if not command_id:
                raise ValidationError("现场挑战命令标识无效，请刷新页面后重试。")
            expected_revision = int(request.POST.get("expected_revision") or "0")
            session_id = int(request.POST.get("session_id") or "0")
            if action == "create":
                source_round = get_object_or_404(
                    ContestRound,
                    pk=request.POST.get("source_round_id"),
                    activity=activity,
                )
                version = _current_frozen_version(activity)
                if version is None:
                    raise ValidationError("当前活动还没有已确认的赛制版本。")
                create_challenge_session(
                    source_round,
                    ruleset_version=version,
                    operator=request.user,
                    command_id=command_id,
                )
            elif action == "revoke":
                revoke_challenge_session(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                    reason=request.POST.get("reason") or "",
                )
            elif action == "start":
                start_challenge_turn(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action == "target":
                record_challenge_target(
                    session_id,
                    seat_position=request.POST.get("seat_position"),
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action == "pair":
                confirm_challenge_pairing(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action == "waive":
                waive_challenge_turn(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                    reason=request.POST.get("reason") or "",
                )
            elif action == "prepare":
                prepare_challenge_match(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action in {"open_challenger", "open_defender"}:
                advance_challenge_performance(
                    session_id,
                    role=action.removeprefix("open_"),
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action in {"lock_challenger", "lock_defender"}:
                lock_challenge_performance(
                    session_id,
                    role=action.removeprefix("lock_"),
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                    expected_context_version=int(
                        request.POST.get("expected_context_version") or "-1"
                    ),
                )
            elif action == "lock_scores":
                lock_challenge_match_scores(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action == "confirm_match":
                confirm_challenge_match_result(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action == "confirm_session":
                confirm_challenge_session(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            elif action == "prepare_final_round":
                create_final_round_from_challenge(
                    session_id,
                    operator=request.user,
                    command_id=command_id,
                    expected_revision=expected_revision,
                )
            else:
                raise ValidationError("现场挑战操作无效。")
        except (PermissionDenied, ValidationError, ValueError) as error:
            messages.error(request, domain_error_messages(error))
        else:
            messages.success(request, "现场挑战已更新。")
        return redirect("staff:live_challenge_console", pk=activity.pk)

    # A revoked session is history, not the console's current workflow. Hiding it lets the
    # operator immediately establish a fresh session from the same first-round result.
    live_session = (
        ChallengeSession.objects.filter(activity=activity)
        .exclude(state=ChallengeSession.State.REVOKED)
        .select_related("source_round", "ruleset_version")
        .order_by("-created_at", "-pk")
        .first()
    )
    # The console's wording follows the format's actual numbers, which the operators may have
    # changed from the 15/10/5 default.
    live_counts = live_challenge_counts_for_activity(activity)
    current_turn = None
    current_match = None
    last_confirmed_match = None
    legal_seats = []
    command_ids = {
        action: uuid4().hex
        for action in (
            "create",
            "revoke",
            "start",
            "waive",
            "pair",
            "prepare",
            "open_challenger",
            "open_defender",
            "lock_challenger",
            "lock_defender",
            "lock_scores",
            "confirm_match",
            "confirm_session",
            "prepare_final_round",
        )
    }
    if live_session is not None:
        current_turn = (
            live_session.turns.select_related("challenger", "target_seat")
            .filter(sequence=live_session.current_turn_number)
            .first()
        )
        current_match = (
            current_turn.match
            if current_turn is not None and hasattr(current_turn, "match")
            else None
        )
        last_confirmed_match = (
            live_session.matches.select_related("challenger", "defender", "winner")
            .filter(status=ChallengeMatch.Status.CONFIRMED)
            .order_by("-turn__sequence", "-pk")
            .first()
        )
        if current_turn is not None:
            legal_seats = [
                {"seat": seat, "command_id": uuid4().hex}
                for seat in live_session.seats.select_related("current_holder")
                .filter(is_challenged=False)
                .exclude(current_holder_id=current_turn.challenger_id)
                .order_by("position")
            ]
    challenge_performance_controls = {}
    if current_match is not None and current_match.score_round_id:
        for role, singer_id in (
            ("challenger", current_match.challenger_id),
            ("defender", current_match.defender_id),
        ):
            performance = (
                Performance.objects.filter(
                    round_id=current_match.score_round_id,
                    singer_id=singer_id,
                )
                .select_related("score_state")
                .first()
            )
            if performance is not None:
                challenge_performance_controls[role] = {
                    "performance": performance,
                    "score_state": PerformanceScoreState.objects.filter(
                        performance=performance
                    ).first(),
                    "lock_command_id": command_ids[f"lock_{role}"],
                }
    source_rounds = ContestRound.objects.filter(
        activity=activity,
        status=ContestRound.Status.LOCKED,
        advancement_status=ContestRound.AdvancementStatus.FINALIZED,
    ).order_by("sequence", "pk")
    revoke_blocker = (
        challenge_session_revoke_blocker(live_session) if live_session is not None else None
    )
    return render(
        request,
        "staff_panel/live_challenge_console.html",
        {
            "activity": activity,
            "session": live_session,
            "can_revoke_session": request.user.is_admin and revoke_blocker is None,
            "revoke_blocker": revoke_blocker,
            "current_turn": current_turn,
            "current_match": current_match,
            "last_confirmed_match": last_confirmed_match,
            "legal_seats": legal_seats,
            "source_rounds": source_rounds,
            "state_labels": dict(ChallengeSession.State.choices),
            "turn_state_labels": dict(ChallengeTurn.State.choices),
            "match_state_labels": dict(ChallengeMatch.Status.choices),
            "command_ids": command_ids,
            "challenge_performance_controls": challenge_performance_controls,
            "live_counts": live_counts,
        },
    )


@staff_required
def live_material_progress(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    if activity.activity_type != Activity.Type.SINGER_CONTEST:
        raise PermissionDenied("只有歌手比赛活动支持现场材料进度。")
    return render(
        request,
        "staff_panel/live_material_progress.html",
        {"activity": activity, "summary": live_material_summary(activity)},
    )


@staff_required
def activity_workspace(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    rounds = ContestRound.objects.filter(activity=activity).annotate(
        entry_count=Count("entries", distinct=True),
        configured_judge_count=Count("round_judges", distinct=True),
    )
    vote_sessions = VoteSession.objects.filter(activity=activity)
    rubrics = ScoringRubric.objects.filter(activity=activity)
    group_stages = GroupStage.objects.filter(activity=activity).prefetch_related("groups")
    can_configure_singer = (
        activity.activity_type == Activity.Type.SINGER_CONTEST
        and not activity.is_locked
        and activity.phase != Activity.Phase.ARCHIVED
    )
    # The workspace should answer "do we already have a ruleset, and where do I continue?"
    # rather than pointing at the create page and leaving staff to work it out.
    versions = RulesetVersion.objects.filter(ruleset__activity=activity)
    activity_ruleset = ContestRuleset.objects.filter(activity=activity).first()
    # A frozen authority is not a draft: it has no edit page at all, so offering it as
    # "继续准备" would land the operator on a 403. The workspace shows it as the official
    # version it now is, with the two things that are actually available — the pre-launch
    # check, and issuing a successor draft to change it.
    #
    # The draft must also be the *newest* version, not merely the newest draft: after v2 is
    # frozen, an older v1 draft is still a DRAFT, and letting it win here would label the
    # workspace with a version the activity has already moved past.
    editing_version = (
        editable_ruleset_version(activity_ruleset) if activity_ruleset is not None else None
    ) or versions.filter(is_current=True).order_by("-version").first()
    ruleset_summary = None
    if editing_version is not None:
        is_frozen = editing_version.status == RulesetVersion.Status.FROZEN
        is_live_2026 = (
            str((editing_version.ruleset.live_challenge or {}).get("format_key") or "")
            == LIVE_CHALLENGE_FORMAT_KEY
        )
        checklist = _live_2026_checklist(editing_version) if is_live_2026 and not is_frozen else []
        ruleset_summary = {
            "version": editing_version,
            "is_live_2026": is_live_2026,
            "is_frozen": is_frozen,
            "outstanding": [row for row in checklist if not row["done"] and not row["pending"]],
            "done_count": sum(1 for row in checklist if row["done"]),
            "total_count": len(checklist),
            "continue_url": (
                "" if is_frozen else _ruleset_edit_url(editing_version, is_live_2026=is_live_2026)
            ),
        }
    return render(
        request,
        "staff_panel/activity_workspace.html",
        {
            "activity": activity,
            "ruleset_summary": ruleset_summary,
            "rounds": rounds,
            "vote_sessions": vote_sessions,
            "rubrics": rubrics,
            "group_stages": group_stages,
            "registration_count": SingerRegistration.objects.filter(activity=activity).count(),
            "programs": Program.objects.filter(activity=activity),
            "incident_count": IncidentRecord.objects.filter(activity=activity).count(),
            "can_configure_singer": can_configure_singer,
        },
    )


@staff_required
def group_stage_create(request, activity_id):
    activity = get_object_or_404(Activity, pk=activity_id)
    if activity.activity_type != Activity.Type.SINGER_CONTEST:
        raise PermissionDenied("只有歌手比赛支持分组合唱赛段。")
    rounds = ContestRound.objects.filter(activity=activity).order_by("pk")
    form = GroupStageCreateForm(request.POST or None, rounds=rounds)
    if request.method == "POST" and form.is_valid():
        try:
            stage = create_group_stage(
                activity,
                stage_key=form.cleaned_data["stage_key"],
                name=form.cleaned_data["name"],
                operator=request.user,
                roster_source=form.cleaned_data["roster_source"],
                roster_source_round=form.cleaned_data["roster_source_round"],
            )
        except (PermissionDenied, ValidationError) as error:
            form.add_error(None, domain_error_messages(error))
        else:
            messages.success(request, "分组合唱赛段已创建，请录入外部分组结果。")
            return redirect("staff:group_stage_detail", pk=stage.pk)
    return render(
        request,
        "staff_panel/group_stage_form.html",
        {"activity": activity, "form": form},
    )


@staff_required
def group_stage_detail(request, pk):
    stage = get_object_or_404(
        GroupStage.objects.select_related("activity").prefetch_related(
            "groups__memberships__singer"
        ),
        pk=pk,
    )
    try:
        singers = group_stage_roster_singers(stage)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"当前分组合唱名单暂不可用：{domain_error_messages(error)}")
        return redirect("staff:activity_workspace", pk=stage.activity_id)
    initial_groups = [
        {
            "name": group.name,
            "singer_ids": list(
                group.memberships.filter(is_current=True).values_list("singer_id", flat=True)
            ),
        }
        for group in stage.groups.filter(is_active=True).order_by("group_order", "pk")
    ]
    max_group_slots = max(len(singers), len(initial_groups), 1)
    assignment_form = GroupStageAssignmentForm(
        singers=singers, initial_groups=initial_groups, max_group_slots=max_group_slots
    )
    correction_form = GroupStageCorrectionForm(
        singers=singers, initial_groups=initial_groups, max_group_slots=max_group_slots
    )
    groups = list(
        stage.groups.filter(is_active=True).prefetch_related(
            "memberships__singer", "material_checks", "files"
        )
    )
    for group in groups:
        group.current_members = [
            membership for membership in group.memberships.all() if membership.is_current
        ]
        readiness = group_material_readiness(group)
        group.readiness = readiness.value
        group.readiness_label = {
            "waiting_for_confirmation": "等待确认分组",
            "missing_member": "缺少有效成员",
            "missing_material": "缺少已审核的伴奏材料",
            "ready": "READY",
        }[readiness.value]
        current_files = {file.file_purpose for file in group.files.all() if file.is_current}
        has_performance_video = "performance_video" in current_files
        group.missing_materials = [
            check.item_name
            for check in group.material_checks.all()
            if check.required
            and check.status != MaterialCheck.Status.APPROVED
            and not (
                has_performance_video
                and not check.question_key
                and check.file_purpose == "accompaniment"
            )
        ]
        if not any(
            file_purpose in current_files for file_purpose in ("accompaniment", "performance_video")
        ):
            group.missing_materials.insert(0, "合唱伴奏（音频或视频）")
    return render(
        request,
        "staff_panel/group_stage_detail.html",
        {
            "stage": stage,
            "groups": groups,
            "assignment_form": assignment_form,
            "correction_form": correction_form,
            "material_status_choices": GroupStage.MaterialStatus.choices,
            "can_record": stage.status == GroupStage.Status.DRAFT and not stage.activity.is_locked,
            "can_correct": stage.status in (GroupStage.Status.CONFIRMED, GroupStage.Status.FROZEN)
            and not stage.activity.is_locked,
        },
    )


@staff_required
def group_stage_history(request, pk):
    stage = get_object_or_404(
        GroupStage.objects.select_related("activity"),
        pk=pk,
    )
    history = (
        GroupMembership.objects.filter(group__stage=stage, is_current=False)
        .select_related("group", "singer", "added_by", "removed_by")
        .order_by("-left_at", "-pk")
    )
    audits = (
        AuditLog.objects.filter(
            target=f"GroupStage:{stage.pk}",
            action_type__in={
                AuditLog.ActionType.CREATE_GROUP_STAGE,
                AuditLog.ActionType.UPDATE_GROUP_STAGE,
                AuditLog.ActionType.CONFIRM_GROUP_STAGE,
                AuditLog.ActionType.FREEZE_GROUP_STAGE,
                AuditLog.ActionType.CORRECT_GROUP_STAGE,
            },
        )
        .select_related("operator")
        .order_by("-created_at", "-pk")
    )
    return render(
        request,
        "staff_panel/group_stage_history.html",
        {"stage": stage, "history": history, "audits": audits},
    )


@staff_required
def group_material_history(request, pk):
    group = get_object_or_404(
        Group.objects.select_related("stage__activity"),
        pk=pk,
    )
    files = (
        SubmissionFile.objects.filter(group=group)
        .select_related("uploaded_by", "source_ruleset_version")
        .order_by("-uploaded_at", "-version", "-pk")
    )
    return render(
        request,
        "staff_panel/group_material_history.html",
        {"group": group, "files": files},
    )


@staff_required
@require_POST
def group_stage_record(request, pk):
    stage = get_object_or_404(GroupStage, pk=pk)
    singers = SingerRegistration.objects.filter(
        activity=stage.activity, pre_status=SingerRegistration.PreStatus.APPROVED
    ).order_by("pk")
    form = GroupStageAssignmentForm(
        request.POST, singers=singers, max_group_slots=max(singers.count(), 1)
    )
    if form.is_valid():
        try:
            record_group_stage(stage, form.cleaned_data["groups"], request.user)
        except (PermissionDenied, ValidationError) as error:
            messages.error(request, domain_error_messages(error))
        else:
            messages.success(request, "外部分组结果已保存。")
    else:
        messages.error(request, _form_error(form))
    return redirect("staff:group_stage_detail", pk=stage.pk)


@staff_required
@require_POST
def group_stage_confirm(request, pk):
    stage = get_object_or_404(GroupStage, pk=pk)
    try:
        confirm_group_stage(stage, request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, domain_error_messages(error))
    else:
        messages.success(request, "分组合唱分组已确认。")
    return redirect("staff:group_stage_detail", pk=stage.pk)


@staff_required
@require_POST
def group_stage_cancel(request, pk):
    stage = get_object_or_404(GroupStage, pk=pk)
    try:
        cancel_group_stage(stage, request.user, reason=request.POST.get("reason", ""))
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, domain_error_messages(error))
    else:
        messages.success(request, "分组合唱赛段已取消。")
    return redirect("staff:group_stage_detail", pk=stage.pk)


@staff_required
@require_POST
def group_stage_material_status(request, pk):
    stage = get_object_or_404(GroupStage, pk=pk)
    try:
        set_group_material_status(
            stage,
            request.POST.get("material_status", ""),
            request.user,
            note=request.POST.get("note", ""),
        )
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, domain_error_messages(error))
    else:
        messages.success(
            request, f"分组合唱材料窗口已设为：{stage.get_material_status_display()}。"
        )
    return redirect("staff:group_stage_detail", pk=stage.pk)


@staff_required
@require_POST
def group_stage_freeze(request, pk):
    stage = get_object_or_404(GroupStage, pk=pk)
    try:
        freeze_group_stage(stage, request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, domain_error_messages(error))
    else:
        messages.success(request, "分组合唱分组已冻结。")
    return redirect("staff:group_stage_detail", pk=stage.pk)


@staff_required
@require_POST
def group_stage_correct(request, pk):
    stage = get_object_or_404(GroupStage, pk=pk)
    singers = SingerRegistration.objects.filter(
        activity=stage.activity, pre_status=SingerRegistration.PreStatus.APPROVED
    ).order_by("pk")
    form = GroupStageCorrectionForm(
        request.POST, singers=singers, max_group_slots=max(singers.count(), 1)
    )
    if form.is_valid():
        try:
            correct_group_stage(
                stage,
                form.cleaned_data["groups"],
                request.user,
                reason=form.cleaned_data["reason"],
            )
        except (PermissionDenied, ValidationError) as error:
            messages.error(request, domain_error_messages(error))
        else:
            messages.success(request, "分组合唱成员修正已保存并记录审计。")
    else:
        messages.error(request, _form_error(form))
    return redirect("staff:group_stage_detail", pk=stage.pk)


@staff_required
@require_POST
@transaction.atomic
def activity_judge_entry_toggle(request, pk):
    """Open or close the judge entry for the night.

    GOAL §12.3's stable shared QR removed the per-seat grant, which left the judge entry
    with no credential at all: `/e/<code>/judge/` is public and the public activity page
    linked to it, so anyone reading the code could take a real seat and submit a real
    score. The capability in the QR is the fix; this switch is the second layer, and it
    covers the live night as well as a rehearsal — the door is shut until staff open it,
    and shutting it again after the panel is seated means a photographed QR claims
    nothing. Sessions already claimed are untouched.
    """
    # The row lock is the same Activity-first discipline every other mutation uses, and it
    # is what makes this read-modify-write safe: without it two staff clicking at the same
    # moment both read the old value and the second write silently loses the first. The
    # claim path takes the same lock, so a claim and a toggle can no longer interleave.
    activity = get_object_or_404(Activity.objects.select_for_update(), pk=pk)
    if activity.activity_type != Activity.Type.SINGER_CONTEST:
        raise PermissionDenied("只有歌手比赛支持评委入口。")
    if activity.phase == Activity.Phase.ARCHIVED or activity.is_locked:
        raise PermissionDenied("已归档或已锁定的活动不能改变评委入口。")
    activity.judge_entry_open = not activity.judge_entry_open
    with authority_write(ACTIVITY_STATE):
        activity.save(update_fields=["judge_entry_open"])
    log_action(
        request,
        AuditLog.ActionType.OTHER,
        f"Activity:{activity.pk}",
        new_value=f"judge_entry_open={activity.judge_entry_open}",
    )
    messages.success(
        request,
        "评委入口已开放。" if activity.judge_entry_open else "评委入口已关闭。",
    )
    return redirect("staff:activity_workspace", pk=activity.pk)


@staff_required
@require_POST
@transaction.atomic
def activity_judge_entry_rotate(request, pk):
    """Reissue the judge QR: the printed one stops working, the entry stays open.

    Closing the entry is the wrong tool for a leaked code — it also turns away the teacher
    who is walking up right now. Bumping the version invalidates every capability signed
    for the old one, and the new QR is printed from the QR centre as usual. Existing
    sessions keep scoring: they are seats, not codes.
    """
    # Locked for the same reason the toggle is: two staff reissuing at once would both read
    # version N and both write N+1, so the increment would be lost and a code that both of
    # them believed they had killed would still resolve.
    activity = get_object_or_404(Activity.objects.select_for_update(), pk=pk)
    if activity.activity_type != Activity.Type.SINGER_CONTEST:
        raise PermissionDenied("只有歌手比赛有评委入口二维码。")
    if activity.phase == Activity.Phase.ARCHIVED or activity.is_locked:
        raise PermissionDenied("已归档或已锁定的活动不能重发评委二维码。")
    activity.judge_entry_version += 1
    with authority_write(ACTIVITY_STATE):
        activity.save(update_fields=["judge_entry_version"])
    log_action(
        request,
        AuditLog.ActionType.OTHER,
        f"Activity:{activity.pk}",
        new_value=f"judge_entry_version={activity.judge_entry_version}",
    )
    messages.success(request, "评委二维码已重发；此前发出的那张立即失效。")
    return redirect("staff:activity_workspace", pk=activity.pk)


@admin_required
def activity_create(request):
    _require_admin(request.user)
    if request.method == "POST":
        form = ActivityForm(request.POST, include_lifecycle=True, include_phase=False)
        if not form.is_valid():
            return render(
                request,
                "staff_panel/activity_form.html",
                {
                    "error": _form_error(form),
                    "activity_types": _choices(Activity.Type),
                    "phases": CREATE_PHASE_CHOICES,
                    "data": request.POST,
                    "form_data_present": True,
                },
            )
        # A new activity always starts as a draft.  The create form deliberately does
        # not expose a lifecycle selector: entering a live phase is a later, explicit
        # transition from the activity editor.
        data = dict(form.cleaned_data)
        data["phase"] = Activity.Phase.DRAFT
        activity = Activity.objects.create(**data)
        if request.FILES.get("cover_image"):
            activity.cover_image = request.FILES["cover_image"]
            activity.save(update_fields=["cover_image"])
        AuditLog.objects.create(
            operator=request.user,
            action_type=AuditLog.ActionType.OTHER,
            target=f"创建活动: {activity.title}",
        )
        return redirect("staff:activity_list")
    return render(
        request,
        "staff_panel/activity_form.html",
        {
            "activity": None,
            "activity_types": _choices(Activity.Type),
            "phases": CREATE_PHASE_CHOICES,
            "data": {},
            "form_data_present": False,
        },
    )


@admin_required
@transaction.atomic
def activity_edit(request, pk):
    _require_admin(request.user)
    activity = get_object_or_404(Activity, pk=pk)
    if activity.phase == Activity.Phase.ARCHIVED and request.method == "GET":
        messages.info(request, "已归档活动为只读状态，请从导出中心执行解归档后再编辑。")
        return redirect("staff:export_center")
    if request.method == "POST":
        activity = lock_activity_for_action(activity)
        _ensure_activity_mutable(activity)
        form = ActivityForm(request.POST)
        if not form.is_valid():
            return render(
                request,
                "staff_panel/activity_form.html",
                {
                    "error": _form_error(form),
                    "activity": activity,
                    "activity_types": _choices(Activity.Type),
                    "phases": phase_choices_for_activity(activity),
                    "data": request.POST,
                    "form_data_present": True,
                },
            )
        data = dict(form.cleaned_data)
        if data["activity_type"] != activity.activity_type:
            form.add_error("activity_type", "活动类型创建后不可修改。")
            return render(
                request,
                "staff_panel/activity_form.html",
                {
                    "error": _form_error(form),
                    "activity": activity,
                    "activity_types": _choices(Activity.Type),
                    "phases": phase_choices_for_activity(activity),
                    "data": request.POST,
                    "form_data_present": True,
                },
                status=400,
            )
        target_phase = data.pop("phase")
        if target_phase != activity.phase:
            activity = transition_activity_phase(activity, target_phase, actor=request.user)
        for key, value in data.items():
            setattr(activity, key, value)
        if request.FILES.get("cover_image"):
            activity.cover_image = request.FILES["cover_image"]
        try:
            activity.save()
        except ValidationError as error:
            transaction.set_rollback(True)
            form.add_error(None, domain_error_messages(error))
            return render(
                request,
                "staff_panel/activity_form.html",
                {
                    "error": _form_error(form),
                    "activity": activity,
                    "activity_types": _choices(Activity.Type),
                    "phases": phase_choices_for_activity(activity),
                    "data": request.POST,
                    "form_data_present": True,
                },
                status=400,
            )
        AuditLog.objects.create(
            operator=request.user,
            action_type=AuditLog.ActionType.OTHER,
            target=f"更新活动: {activity.title}",
        )
        return redirect("staff:activity_list")
    return render(
        request,
        "staff_panel/activity_form.html",
        {
            "activity": activity,
            "activity_types": _choices(Activity.Type),
            "phases": phase_choices_for_activity(activity),
            "data": {},
            "form_data_present": False,
        },
    )


_ACTIVITY_COLLABORATIVE_FIELDS = {
    "title": 200,
    "subtitle": 400,
    "description": 10000,
}


def _activity_field_etag(value: str) -> str:
    import hashlib

    return "sha256:" + hashlib.sha256(value.encode("utf-8")).hexdigest()


@admin_required
@require_http_methods(["PATCH"])
@transaction.atomic
def activity_field_patch(request, pk, field):
    """Apply one Activity text field with value-based optimistic concurrency."""
    limit = _ACTIVITY_COLLABORATIVE_FIELDS.get(field)
    if limit is None:
        return JsonResponse({"detail": "该活动字段不支持协作保存。"}, status=404)
    try:
        payload = json.loads(request.body or "{}")
    except json.JSONDecodeError:
        return JsonResponse(
            {"detail": "请求体不是有效 JSON。", "reason_code": "INVALID_REQUEST"}, status=400
        )
    if (
        not isinstance(payload, dict)
        or set(payload) != {"base", "value"}
        or not isinstance(payload["base"], str)
        or not isinstance(payload["value"], str)
    ):
        return JsonResponse(
            {"detail": "需要 base 和 value 文本。", "reason_code": "INVALID_REQUEST"}, status=400
        )

    activity = Activity.objects.select_for_update().filter(pk=pk).first()
    if activity is None:
        return JsonResponse({"detail": "活动不存在。"}, status=404)
    _ensure_activity_mutable(activity)
    current_value = getattr(activity, field)
    if current_value != payload["base"]:
        return JsonResponse(
            {
                "detail": "该字段已被其他工作人员更新。",
                "reason_code": "FIELD_STALE",
                "field": field,
                "value": current_value,
                "etag": _activity_field_etag(current_value),
            },
            status=409,
        )
    value = payload["value"].strip() if field != "description" else payload["value"]
    if len(value) > limit:
        return JsonResponse(
            {"detail": f"{field} 长度不能超过 {limit} 个字符。", "reason_code": "INVALID_REQUEST"},
            status=400,
        )
    if value == current_value:
        return JsonResponse({"field": field, "value": value, "etag": _activity_field_etag(value)})

    setattr(activity, field, value)
    activity.save(update_fields=[field, "updated_at"])
    log_action(
        request,
        AuditLog.ActionType.OTHER,
        f"Activity:{activity.pk}:{field}",
        old_value=current_value,
        new_value=value,
    )
    schedule_activity_event(
        activity.pk,
        event="activity.field_changed",
        resource=f"activity:{activity.pk}:field:{field}",
        revision=_activity_field_etag(value),
        actor={"id": request.user.pk, "display": request.user.get_username()},
    )
    return JsonResponse({"field": field, "value": value, "etag": _activity_field_etag(value)})


@staff_required
def post_list(request):
    posts = PublicPost.objects.select_related("related_activity").prefetch_related(
        "result_releases__stage_result",
        Prefetch(
            "related_activity__stage_results",
            queryset=latest_stage_result_queryset().filter(
                status=StageResult.Status.CONFIRMED,
            ),
            to_attr="release_candidates",
        ),
    )
    return render(request, "staff_panel/post_list.html", {"posts": posts})


def _post_status_choices(request):
    choices = _choices(PublicPost.Status)
    if request.user.is_admin:
        return choices
    return [choice for choice in choices if choice[0] != PublicPost.Status.PUBLISHED]


@staff_required
def post_create(request):
    if request.method == "POST":
        form = PublicPostForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(
                request,
                "staff_panel/post_form.html",
                {
                    "error": _form_error(form),
                    "form_data": request.POST,
                    "form_data_present": True,
                    "post_types": _choices(PublicPost.PostType),
                    "statuses": _post_status_choices(request),
                    "activities": Activity.objects.all(),
                },
            )
        data = form.cleaned_data
        with transaction.atomic():
            related_activity = None
            if data["related_activity_id"]:
                related_activity = lock_activity_for_action(
                    get_object_or_404(Activity, pk=data["related_activity_id"])
                )
            _ensure_publication_allowed(related_activity, data["status"])
            publication_actor = request.user
            if data["status"] == PublicPost.Status.PUBLISHED:
                if not admin_verification_is_valid(request.session):
                    return redirect_to_login(
                        request.get_full_path(), reverse("accounts:admin_login")
                    )
                publication_actor = require_current_admin(request.user)
            post = PublicPost(
                title=data["title"],
                subtitle=data["subtitle"],
                content=data["content"],
                post_type=data["post_type"],
                status=data["status"],
                sort_order=data["sort_order"],
                is_pinned=data["is_pinned"],
                related_activity_id=data["related_activity_id"],
                created_by=publication_actor,
                updated_by=publication_actor,
            )
            if data["cover_image"]:
                post.cover_image = data["cover_image"]
            if data["status"] == PublicPost.Status.PUBLISHED:
                post.published_at = timezone.now()
            post.save()
            log_action(
                request,
                AuditLog.ActionType.PUBLISH_POST,
                f"PublicPost:{post.pk}",
                new_value=f"status={post.status}",
            )
        return redirect("staff:post_list")
    return render(
        request,
        "staff_panel/post_form.html",
        {
            "post_types": _choices(PublicPost.PostType),
            "statuses": _post_status_choices(request),
            "activities": Activity.objects.all(),
            "form_data_present": False,
        },
    )


@staff_required
def post_preview(request, pk):
    post = get_object_or_404(PublicPost, pk=pk)
    media_items = post.media_items.filter(is_published=True)
    return render(
        request,
        "public_portal/post_detail.html",
        {"post": post, "media_items": media_items, "preview": True},
    )


@staff_required
@transaction.atomic
def post_edit(request, pk):
    post = get_object_or_404(PublicPost, pk=pk)
    if request.method == "POST":
        form = PublicPostForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(
                request,
                "staff_panel/post_form.html",
                {
                    "error": _form_error(form),
                    "post": post,
                    "form_data": request.POST,
                    "form_data_present": True,
                    "post_types": _choices(PublicPost.PostType),
                    "statuses": _post_status_choices(request),
                    "activities": Activity.objects.all(),
                },
            )
        data = form.cleaned_data
        old_hint_activity_id = post.related_activity_id
        new_activity_id = data["related_activity_id"]
        old_value = f"status={post.status}; title={post.title}"
        # A reparent touches up to two activities; lock both in stable ascending
        # PK order so two concurrent reparents cannot deadlock on opposite order.
        activity_ids = sorted({a for a in (old_hint_activity_id, new_activity_id) if a is not None})
        locked_by_pk = {
            a.pk: a
            for a in Activity.objects.select_for_update().filter(pk__in=activity_ids).order_by("pk")
        }
        if len(locked_by_pk) != len(activity_ids):
            raise Http404("关联的活动不存在。")
        # Lock the post too; its current parent is the authority, not the
        # pre-lock hint. If a reparent committed between our hint read and this
        # lock the row no longer matches — reject rather than chase a newly
        # revealed parent, which would break the stable Activity→PublicPost order.
        locked_post = PublicPost.objects.select_for_update().get(pk=post.pk)
        if locked_post.related_activity_id != old_hint_activity_id:
            raise PermissionDenied("页面内容已发生并发修改，请刷新后重新编辑。")
        if ResultRelease.objects.filter(
            post=locked_post, status=ResultRelease.Status.ACTIVE
        ).exists():
            raise PermissionDenied("结果已公示，编辑或移动前必须先撤销当前结果发布。")
        # Optimistic concurrency: if the client's base_version is behind the
        # current one, another staff member already saved changes. Refuse to
        # silently overwrite them; render the stale form back with an explicit
        # outdated notice so the staff can refresh and re-apply their intent.
        base_version = data.get("base_version")
        if base_version is not None and locked_post.version != base_version:
            return render(
                request,
                "staff_panel/post_form.html",
                {
                    "error": "该内容已被其他人更新，请刷新后重新编辑。",
                    "post": locked_post,
                    "form_data": request.POST,
                    "form_data_present": True,
                    "post_types": _choices(PublicPost.PostType),
                    "statuses": _post_status_choices(request),
                    "activities": Activity.objects.all(),
                },
            )
        old_locked = locked_by_pk.get(old_hint_activity_id) if old_hint_activity_id else None
        new_locked = locked_by_pk.get(new_activity_id) if new_activity_id else None
        old_status = locked_post.status
        # A move from a locked activity to an unlocked one must still be blocked;
        # checking only the new activity would let staff bypass the old lock.
        for locked_activity in (old_locked, new_locked):
            if locked_activity is not None:
                ensure_activity_unlocked(locked_activity)
        _ensure_publication_allowed(new_locked, data["status"])
        publication_actor = request.user
        if (
            data["status"] == PublicPost.Status.PUBLISHED
            or old_status == PublicPost.Status.PUBLISHED
        ):
            if not admin_verification_is_valid(request.session):
                return redirect_to_login(request.get_full_path(), reverse("accounts:admin_login"))
            publication_actor = require_current_admin(request.user)
        locked_post.title = data["title"]
        locked_post.subtitle = data["subtitle"]
        locked_post.content = data["content"]
        locked_post.post_type = data["post_type"]
        locked_post.status = data["status"]
        locked_post.sort_order = data["sort_order"]
        locked_post.is_pinned = data["is_pinned"]
        locked_post.related_activity_id = data["related_activity_id"]
        locked_post.updated_by = publication_actor
        locked_post.version += 1
        if data["cover_image"]:
            locked_post.cover_image = data["cover_image"]
        if data["status"] == PublicPost.Status.PUBLISHED and not locked_post.published_at:
            locked_post.published_at = timezone.now()
        elif (
            data["status"] != PublicPost.Status.PUBLISHED
            and old_status == PublicPost.Status.PUBLISHED
        ):
            locked_post.published_at = None
        locked_post.save()
        log_action(
            request,
            AuditLog.ActionType.PUBLISH_POST,
            f"PublicPost:{locked_post.pk}",
            old_value=old_value,
            new_value=f"status={locked_post.status}; title={locked_post.title}",
        )
        return redirect("staff:post_list")

    return render(
        request,
        "staff_panel/post_form.html",
        {
            "post": post,
            "post_types": _choices(PublicPost.PostType),
            "statuses": _post_status_choices(request),
            "activities": Activity.objects.all(),
            "form_data_present": False,
        },
    )


def _result_release_form_context(post, form):
    stage_results: QuerySet[StageResult] = StageResult.objects.none()
    if post.related_activity_id:
        stage_results = latest_stage_result_queryset().filter(
            activity_id=post.related_activity_id,
            status=StageResult.Status.CONFIRMED,
        )
    return {
        "post": post,
        "form": form,
        "stage_results": stage_results,
    }


@staff_required
@require_POST
def result_release(request, post_id):
    post = get_object_or_404(PublicPost, pk=post_id)
    if not admin_verification_is_valid(request.session):
        return redirect_to_login(request.get_full_path(), reverse("accounts:admin_login"))
    current_admin = require_current_admin(request.user)
    form = ResultReleaseForm(request.POST)
    if not form.is_valid():
        return render(
            request,
            "staff_panel/result_release_form.html",
            _result_release_form_context(post, form),
            status=400,
        )
    try:
        stage_result = get_object_or_404(
            StageResult,
            pk=form.cleaned_data["stage_result_id"],
            activity_id=post.related_activity_id,
        )
        release_result_post(
            post,
            stage_result,
            current_admin,
            note=form.cleaned_data["note"],
        )
    except (PermissionDenied, ValidationError, IntegrityError) as error:
        form.add_error(None, domain_error_messages(error))
        return render(
            request,
            "staff_panel/result_release_form.html",
            _result_release_form_context(post, form),
            status=400,
        )
    messages.success(request, "结果已通过独立发布权威公示。")
    return redirect("staff:post_list")


@staff_required
@require_POST
def result_release_revoke(request, post_id):
    post = get_object_or_404(PublicPost, pk=post_id)
    if not admin_verification_is_valid(request.session):
        return redirect_to_login(request.get_full_path(), reverse("accounts:admin_login"))
    current_admin = require_current_admin(request.user)
    form = ResultReleaseRevokeForm(request.POST)
    if not form.is_valid():
        return render(
            request,
            "staff_panel/result_release_form.html",
            _result_release_form_context(post, form),
            status=400,
        )
    try:
        revoke_result_release(post, current_admin, note=form.cleaned_data["note"])
    except (PermissionDenied, ValidationError, IntegrityError) as error:
        form.add_error(None, domain_error_messages(error))
        return render(
            request,
            "staff_panel/result_release_form.html",
            _result_release_form_context(post, form),
            status=400,
        )
    messages.success(request, "结果发布已撤销，历史记录已保留。")
    return redirect("staff:post_list")


# --- Singer registration management ---


@staff_required
def singer_registration_list(request):
    registrations = SingerRegistration.objects.select_related("activity", "user")
    activity_id = request.GET.get("activity_id")
    if activity_id:
        registrations = registrations.filter(activity_id=activity_id)
    registrations = _with_generic_song_labels(registrations)
    return render(
        request,
        "staff_panel/singer_registration_list.html",
        {
            "registrations": registrations,
            "activities": Activity.objects.all(),
            "selected_activity_id": activity_id,
        },
    )


@staff_required
def singer_registration_detail(request, pk):
    reg = get_object_or_404(SingerRegistration.objects.select_related("activity", "user"), pk=pk)
    errors = []
    if request.method == "POST":
        if "set_official_number" in request.POST:
            number_form = OfficialContestantNumberForm(request.POST)
            if number_form.is_valid():
                try:
                    set_official_contestant_number(
                        reg.pk,
                        number_form.cleaned_data["official_number"],
                        operator=request.user,
                    )
                except (PermissionDenied, ValidationError, IntegrityError) as error:
                    errors = [domain_error_messages(error)]
                else:
                    messages.success(request, "官方选手编号已保存。")
                    return redirect("staff:singer_registration_detail", pk=reg.pk)
            else:
                errors = [_form_error(number_form)]
        is_upload = "upload_file" in request.POST
        action = ActivityAction.UPLOAD_MATERIAL if is_upload else ActivityAction.REVIEW_REGISTRATION
        if "set_official_number" not in request.POST and not errors:
            ensure_activity_action_allowed(reg.activity, action)
            with transaction.atomic():
                activity = lock_activity_for_action(reg.activity, action)
                locked_reg = (
                    SingerRegistration.objects.select_for_update()
                    .select_related("activity", "user")
                    .get(pk=reg.pk)
                )
                if locked_reg.activity_id != activity.pk:
                    raise PermissionDenied("报名信息不属于当前活动。")
                if is_upload:
                    f = request.FILES.get("file")
                    if f:
                        try:
                            submission_file = store_submission_file(
                                owner=locked_reg,
                                uploaded_file=f,
                                purpose=request.POST.get(
                                    "file_purpose", SubmissionFile.Purpose.OTHER
                                ),
                                uploaded_by=request.user,
                            )
                        except ValidationError as error:
                            errors.extend(error.messages)
                        else:
                            reconcile_singer_material_checks(locked_reg)
                            log_action(
                                request,
                                AuditLog.ActionType.UPLOAD_FILE,
                                f"SubmissionFile:{submission_file.pk}",
                                new_value=submission_file.original_name,
                            )
                else:
                    form = SingerReviewForm(request.POST)
                    if not form.is_valid():
                        errors = [_form_error(form)]
                    else:
                        old_value = f"pre={locked_reg.pre_status}; live={locked_reg.live_status}"
                        locked_reg.pre_status = form.cleaned_data["pre_status"]
                        locked_reg.live_status = form.cleaned_data["live_status"]
                        note = form.cleaned_data.get("staff_note", "").strip()
                        if note:
                            StaffNote.objects.create(
                                singer_registration=locked_reg,
                                content=note,
                                created_by=request.user,
                            )
                        locked_reg.save()
                        log_action(
                            request,
                            AuditLog.ActionType.UPDATE_STATUS,
                            f"SingerRegistration:{locked_reg.pk}",
                            old_value=old_value,
                            new_value=f"pre={locked_reg.pre_status}; live={locked_reg.live_status}",
                        )
                        return redirect("staff:singer_registration_detail", pk=reg.pk)
    notes = reg.staff_notes.select_related("created_by")
    files = reg.files.all()
    checks = reg.material_checks.all()
    reg = _with_generic_song_labels([reg])[0]
    return render(
        request,
        "staff_panel/singer_registration_detail.html",
        {
            "reg": reg,
            "pre_statuses": _choices(SingerRegistration.PreStatus),
            "live_statuses": _choices(SingerRegistration.LiveStatus),
            "notes": notes,
            "files": files,
            "checks": checks,
            "check_statuses": _choices(MaterialCheck.Status),
            "file_purposes": _choices(SubmissionFile.Purpose),
            "questionnaire_material_authority": questionnaire_material_authority_active(
                reg.activity
            ),
            "errors": errors,
            "official_number_form": OfficialContestantNumberForm(
                initial={"official_number": reg.official_number}
            ),
        },
    )


# --- Program management ---


@staff_required
def program_list(request):
    programs = Program.objects.select_related("activity", "user")
    activity_id = request.GET.get("activity_id")
    if activity_id:
        programs = programs.filter(activity_id=activity_id)
    return render(
        request,
        "staff_panel/program_list.html",
        {
            "programs": programs,
            "activities": Activity.objects.all(),
            "selected_activity_id": activity_id,
        },
    )


@staff_required
def program_detail(request, pk):
    prog = get_object_or_404(Program.objects.select_related("activity", "user"), pk=pk)
    errors = []
    if request.method == "POST":
        is_upload = "upload_file" in request.POST
        action = ActivityAction.UPLOAD_MATERIAL if is_upload else ActivityAction.REVIEW_REGISTRATION
        ensure_activity_action_allowed(prog.activity, action)
        with transaction.atomic():
            activity = lock_activity_for_action(prog.activity, action)
            locked_prog = (
                Program.objects.select_for_update()
                .select_related("activity", "user")
                .get(pk=prog.pk)
            )
            if locked_prog.activity_id != activity.pk:
                raise PermissionDenied("节目信息不属于当前活动。")
            if is_upload:
                f = request.FILES.get("file")
                if f:
                    try:
                        submission_file = store_submission_file(
                            owner=locked_prog,
                            uploaded_file=f,
                            purpose=request.POST.get("file_purpose", SubmissionFile.Purpose.OTHER),
                            uploaded_by=request.user,
                        )
                    except ValidationError as error:
                        errors.extend(error.messages)
                    else:
                        reconcile_program_material_checks(locked_prog)
                        log_action(
                            request,
                            AuditLog.ActionType.UPLOAD_FILE,
                            f"SubmissionFile:{submission_file.pk}",
                            new_value=submission_file.original_name,
                        )
            else:
                form = ProgramReviewForm(request.POST)
                if not form.is_valid():
                    errors = [_form_error(form)]
                else:
                    old_value = f"status={locked_prog.status}; sort_order={locked_prog.sort_order}"
                    locked_prog.status = form.cleaned_data["status"]
                    locked_prog.sort_order = form.cleaned_data["sort_order"]
                    note = form.cleaned_data.get("staff_note", "").strip()
                    if note:
                        StaffNote.objects.create(
                            program=locked_prog,
                            content=note,
                            created_by=request.user,
                        )
                    locked_prog.save()
                    log_action(
                        request,
                        AuditLog.ActionType.REVIEW_MATERIAL,
                        f"Program:{locked_prog.pk}",
                        old_value=old_value,
                        new_value=(
                            f"status={locked_prog.status}; sort_order={locked_prog.sort_order}"
                        ),
                    )
                    return redirect("staff:program_detail", pk=prog.pk)
    notes = prog.staff_notes.select_related("created_by")
    files = prog.files.all()
    checks = prog.material_checks.all()
    return render(
        request,
        "staff_panel/program_detail.html",
        {
            "prog": prog,
            "statuses": _choices(Program.Status),
            "program_types": _choices(Program.ProgramType),
            "notes": notes,
            "files": files,
            "checks": checks,
            "check_statuses": _choices(MaterialCheck.Status),
            "file_purposes": _choices(SubmissionFile.Purpose),
            "errors": errors,
        },
    )


@staff_required
@require_POST
def material_check_review(request):
    check = get_object_or_404(MaterialCheck, pk=request.POST.get("check_id"))
    owner = check.singer_registration or check.group or check.program
    if owner is None:
        messages.error(request, "材料检查项未关联有效报名，无法审核。")
        return redirect("staff:export_center")
    result = request.POST.get("result")
    if result == "approve":
        status = MaterialCheck.Status.APPROVED
    elif result == "supplement":
        status = MaterialCheck.Status.NEEDS_SUPPLEMENT
    else:
        messages.error(request, "无效的审核操作。")
        return _material_review_redirect(owner)
    try:
        review_material_check(
            check,
            status=status,
            note=request.POST.get("note", ""),
            actor=request.user,
        )
    except ValidationError as error:
        messages.error(request, domain_error_messages(error))
        return _material_review_redirect(owner)
    messages.success(request, f"「{check.item_name}」已审核。")
    return _material_review_redirect(owner)


def _material_review_redirect(owner):
    if getattr(owner, "student_id", None):
        return redirect("staff:singer_registration_detail", pk=owner.pk)
    if getattr(owner, "stage_id", None):
        return redirect("staff:group_stage_detail", pk=owner.stage_id)
    return redirect("staff:program_detail", pk=owner.pk)


@staff_required
def activity_material_requirements(request, activity_id):
    activity = get_object_or_404(Activity, pk=activity_id)
    if request.method == "POST":
        with transaction.atomic():
            locked_activity = lock_activity_for_action(activity)
            _ensure_activity_mutable(locked_activity)
            applies_to = request.POST.get("applies_to")
            item_name = request.POST.get("item_name", "").strip()
            if not item_name:
                messages.error(request, "检查项名称不能为空。")
            elif applies_to not in MaterialRequirement.AppliesTo.values:
                messages.error(request, "无效的适用范围。")
            elif (
                applies_to == MaterialRequirement.AppliesTo.SINGER
                and questionnaire_material_authority_active(locked_activity)
            ):
                messages.error(
                    request, "当前活动的选手材料由已确认问卷管理，不能再配置旧式材料项。"
                )
            else:
                # Validate before the ORM call. The form posts these as text and the
                # model fields are typed, so an unparseable sort_order or an over-long
                # name used to reach the database and surface as a 500 rather than as a
                # message on the page the operator is already looking at.
                sort_order_raw = (request.POST.get("sort_order") or "0").strip()
                file_purpose = (request.POST.get("file_purpose") or "").strip()
                if not sort_order_raw.isdigit():
                    messages.error(request, "排序必须是数字。")
                elif len(item_name) > 100:
                    messages.error(request, "检查项名称不能超过 100 个字符。")
                elif file_purpose and file_purpose not in SubmissionFile.Purpose.values:
                    messages.error(request, "无效的文件用途。")
                else:
                    MaterialRequirement.objects.update_or_create(
                        activity=locked_activity,
                        applies_to=applies_to,
                        item_name=item_name,
                        defaults={
                            "file_purpose": file_purpose,
                            "is_required": True,
                            "sort_order": int(sort_order_raw),
                        },
                    )
                    reconcile_activity_material_checks(locked_activity, applies_to)
                    messages.success(request, "材料检查项已保存。")
        return redirect("staff:activity_material_requirements", activity_id=activity.pk)
    requirements = MaterialRequirement.objects.filter(activity=activity)
    questionnaire_authority = questionnaire_material_authority_active(activity)
    questionnaire_version = _current_frozen_version(activity)
    return render(
        request,
        "staff_panel/material_requirements.html",
        {
            "activity": activity,
            "requirements": requirements,
            "applies_to_choices": _choices(MaterialRequirement.AppliesTo),
            "file_purposes": _choices(SubmissionFile.Purpose),
            "questionnaire_material_authority": questionnaire_authority,
            "questionnaire_version": questionnaire_version,
        },
    )


@staff_required
@require_POST
def activity_material_requirement_delete(request, activity_id, pk):
    activity = get_object_or_404(Activity, pk=activity_id)
    with transaction.atomic():
        locked_activity = lock_activity_for_action(activity)
        _ensure_activity_mutable(locked_activity)
        requirement = MaterialRequirement.objects.filter(pk=pk, activity=locked_activity).first()
        if requirement is None:
            messages.error(request, "材料检查项不存在，未做任何修改。")
        else:
            applies_to = requirement.applies_to
            if (
                applies_to == MaterialRequirement.AppliesTo.SINGER
                and questionnaire_material_authority_active(locked_activity)
            ):
                messages.error(request, "当前活动的选手材料由已确认问卷管理，不能删除旧式材料项。")
            else:
                # Report success only where a row was actually removed. The unconditional
                # message meant a refused delete showed the failure and the success banner
                # at the same time, so the page contradicted itself about what happened.
                requirement.delete()
                reconcile_activity_material_checks(locked_activity, applies_to)
                messages.success(request, "材料检查项已删除。")
    return redirect("staff:activity_material_requirements", activity_id=activity.pk)


# --- Excel export ---


def _export_activity(request):
    """Resolve the ?activity_id= scope, or None for a cross-activity export.

    A staff user must pick an activity; only an admin may export across all
    activities (enforced by the caller). Returns the Activity or None.
    """
    activity_id = request.GET.get("activity_id")
    if not activity_id:
        return None
    return get_object_or_404(Activity, pk=activity_id)


@staff_required
def export_registrations(request):
    activity = _export_activity(request)
    if activity is None and not request.user.is_admin:
        raise PermissionDenied("全量导出仅管理员可用；请先用活动筛选导出单个活动。")
    wb = Workbook()
    ws = _active_worksheet(wb)
    ws.title = "报名名单"
    ws.append(
        [
            "官方编号",
            "姓名",
            "学号",
            "学院",
            "班级",
            "手机",
            "微信",
            "曲目",
            "原创",
            "赛前状态",
            "赛中状态",
            "活动",
            "提交时间",
        ]
    )
    if activity is not None:
        registration_rows = scope_runtime(
            SingerRegistration.objects.select_related("activity").filter(activity=activity),
            activity,
        )
    else:
        registration_rows = SingerRegistration.objects.select_related("activity").filter(
            is_test_data=False
        )
    rows = prime_questionnaire_answers(registration_rows)
    row_count = 0
    for r in rows:
        ws.append(
            [
                r.official_number,
                r.name,
                r.student_id,
                r.college,
                r.class_name,
                r.phone,
                r.wechat,
                generic_song_label(r),
                "是" if r.is_original else "否",
                r.get_pre_status_display(),
                r.get_live_status_display(),
                r.activity.title,
                r.created_at.strftime("%Y-%m-%d %H:%M"),
            ]
        )
        row_count += 1
    audit_export(request, activity, "registration_list", row_count=row_count)
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=registration_list.xlsx"
    wb.save(response)
    return response


@staff_required
def export_programs(request):
    activity = _export_activity(request)
    if activity is None and not request.user.is_admin:
        raise PermissionDenied("全量导出仅管理员可用；请先用活动筛选导出单个活动。")
    wb = Workbook()
    ws = _active_worksheet(wb)
    ws.title = "节目单"
    ws.append(
        [
            "顺序",
            "节目名称",
            "类型",
            "负责人",
            "联系方式",
            "所属",
            "演员",
            "时长",
            "麦克风需求",
            "道具需求",
            "状态",
            "活动",
            "提交时间",
        ]
    )
    if activity is not None:
        rows = scope_runtime(
            Program.objects.select_related("activity").filter(activity=activity), activity
        )
    else:
        rows = Program.objects.select_related("activity").filter(is_test_data=False)
    row_count = 0
    for p in rows:
        ws.append(
            [
                p.sort_order,
                p.name,
                p.get_program_type_display(),
                p.contact_name,
                p.contact_phone,
                p.class_name,
                p.performers,
                p.estimated_duration,
                p.mic_requirements,
                p.prop_requirements,
                p.get_status_display(),
                p.activity.title,
                p.created_at.strftime("%Y-%m-%d %H:%M"),
            ]
        )
        row_count += 1
    audit_export(request, activity, "program_list", row_count=row_count)
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=program_list.xlsx"
    wb.save(response)
    return response


# --- Singer contest scoring ---


@staff_required
def round_list(request):
    rounds = list(
        ContestRound.objects.select_related("activity").annotate(
            entry_count=Count("entries", distinct=True),
            assigned_judge_count=Count("round_judges", distinct=True),
        )
    )
    activity_id = request.GET.get("activity_id") or ""
    if activity_id.isdigit():
        rounds = [
            contest_round
            for contest_round in rounds
            if str(contest_round.activity_id) == activity_id
        ]
    readiness_labels = {
        RoundSetupReadiness.WAITING_FOR_UPSTREAM_CONFIRMATION: "等待上游赛段核定",
        RoundSetupReadiness.MISSING_UPSTREAM: "缺少上游赛段结果",
        RoundSetupReadiness.UNBOUND_TO_FROZEN_RULESET: "未绑定当前冻结赛制",
        RoundSetupReadiness.MATERIALIZATION_REQUIRED: "名单尚未物化",
        RoundSetupReadiness.STALE: "晋级名单已变化，需重新配置",
    }
    for contest_round in rounds:
        readiness = round_setup_readiness(contest_round)
        setattr(contest_round, "setup_readiness", readiness.value)
        setattr(
            contest_round,
            "setup_readiness_label",
            readiness_labels.get(readiness, ""),
        )
        # The readiness above is about this round's own roster, so a round the current
        # official format does not bind reads as 就绪 here and then refuses to start. The
        # button is the promise; it is not offered where the click would be refused.
        setattr(
            contest_round,
            "execution_authority_label",
            round_execution_authority_label(contest_round),
        )
        setattr(
            contest_round,
            "execution_authority_message",
            round_execution_authority_message(contest_round),
        )
    return render(
        request,
        "staff_panel/round_list.html",
        {"rounds": rounds, "selected_activity_id": activity_id},
    )


@staff_required
def round_create(request):
    activities = _editable_singer_activities()
    if request.method == "POST":
        activity = get_object_or_404(activities, pk=request.POST.get("activity_id"))
        stage_choices = round_roster_stage_choices(activity)
        form = ContestRoundForm(
            request.POST,
            rubrics=ScoringRubric.objects.filter(activity=activity),
            stage_choices=stage_choices,
        )
        if not form.is_valid():
            return render(
                request,
                "staff_panel/round_form.html",
                {
                    "form": form,
                    "error": _form_error(form),
                    "activities": activities,
                    "round_types": _choices(ContestRound.RoundType),
                    "scoring_modes": _choices(ContestRound.ScoringMode),
                    "order_policies": _choices(ContestRound.OrderPolicy),
                    "tie_order_policies": _choices(ContestRound.TieOrderPolicy),
                    "roster_sources": [("", "自动判定"), *_choices(ContestRound.RosterSource)],
                    "rubrics": ScoringRubric.objects.select_related("activity").filter(
                        activity=activity
                    ),
                    "selected_activity_id": activity.pk,
                },
            )
        if form.cleaned_data["roster_source"] == ContestRound.RosterSource.STAGE:
            try:
                validate_roster_source_stage(activity, form.cleaned_data["roster_source_stage"])
            except ValidationError as error:
                form.add_error("roster_source_stage", domain_error_messages(error))
                return render(
                    request,
                    "staff_panel/round_form.html",
                    {
                        "form": form,
                        "error": _form_error(form),
                        "activities": activities,
                        "round_types": _choices(ContestRound.RoundType),
                        "scoring_modes": _choices(ContestRound.ScoringMode),
                        "order_policies": _choices(ContestRound.OrderPolicy),
                        "tie_order_policies": _choices(ContestRound.TieOrderPolicy),
                        "roster_sources": [("", "自动判定"), *_choices(ContestRound.RosterSource)],
                        "rubrics": ScoringRubric.objects.select_related("activity").filter(
                            activity=activity
                        ),
                        "selected_activity_id": activity.pk,
                    },
                )
        with transaction.atomic():
            locked_activity = lock_activity_for_action(activity)
            sequence = form.cleaned_data["sequence"]
            if (
                sequence is not None
                and ContestRound.objects.filter(
                    activity=locked_activity, sequence=sequence
                ).exists()
            ):
                form.add_error("sequence", "轮次序号已存在，请填写其他序号。")
                return render(
                    request,
                    "staff_panel/round_form.html",
                    {
                        "form": form,
                        "error": _form_error(form),
                        "activities": activities,
                        "round_types": _choices(ContestRound.RoundType),
                        "scoring_modes": _choices(ContestRound.ScoringMode),
                        "order_policies": _choices(ContestRound.OrderPolicy),
                        "tie_order_policies": _choices(ContestRound.TieOrderPolicy),
                        "roster_sources": [("", "自动判定"), *_choices(ContestRound.RosterSource)],
                        "rubrics": ScoringRubric.objects.select_related("activity").filter(
                            activity=activity
                        ),
                        "selected_activity_id": activity.pk,
                    },
                )
            try:
                with transaction.atomic():
                    create_kwargs = {
                        "activity": locked_activity,
                        "round_type": form.cleaned_data["round_type"],
                        "scoring_mode": form.cleaned_data["scoring_mode"],
                        "score_completion_mode": form.cleaned_data["score_completion_mode"],
                        "judge_count": form.cleaned_data["judge_count"],
                        "minimum_judge_count": form.cleaned_data["minimum_judge_count"]
                        or form.cleaned_data["judge_count"],
                        "name": form.cleaned_data["name"],
                        "advance_count": form.cleaned_data["advance_count"],
                        "order_policy": form.cleaned_data["order_policy"],
                        "tie_order_policy": form.cleaned_data["tie_order_policy"],
                        "roster_source": form.cleaned_data["roster_source"],
                        "roster_source_stage": form.cleaned_data["roster_source_stage"],
                        "rubric": form.cleaned_data["rubric"],
                    }
                    if sequence is not None:
                        create_kwargs["sequence"] = sequence
                    contest_round = ContestRound.objects.create(**create_kwargs)
            except IntegrityError:
                form.add_error("sequence", "轮次序号已存在，请填写其他序号。")
                return render(
                    request,
                    "staff_panel/round_form.html",
                    {
                        "form": form,
                        "error": _form_error(form),
                        "activities": activities,
                        "round_types": _choices(ContestRound.RoundType),
                        "scoring_modes": _choices(ContestRound.ScoringMode),
                        "order_policies": _choices(ContestRound.OrderPolicy),
                        "tie_order_policies": _choices(ContestRound.TieOrderPolicy),
                        "roster_sources": [("", "自动判定"), *_choices(ContestRound.RosterSource)],
                        "rubrics": ScoringRubric.objects.select_related("activity").filter(
                            activity=activity
                        ),
                        "selected_activity_id": activity.pk,
                    },
                )
            log_action(
                request,
                AuditLog.ActionType.OTHER,
                f"ContestRound:{contest_round.pk}",
                new_value=contest_round.name,
            )
        return redirect("staff:round_list")

    requested_activity_id = request.GET.get("activity_id") or None
    selected_activity = (
        activities.filter(pk=requested_activity_id).first()
        if requested_activity_id is not None
        else activities.first()
    )
    if selected_activity is None:
        selected_activity = activities.first()
    selected_activity_id = str(selected_activity.pk) if selected_activity else ""
    rubrics = (
        ScoringRubric.objects.select_related("activity").filter(activity=selected_activity)
        if selected_activity is not None
        else ScoringRubric.objects.none()
    )
    form = ContestRoundForm(
        rubrics=rubrics,
        stage_choices=round_roster_stage_choices(selected_activity)
        if selected_activity is not None
        else (),
    )
    return render(
        request,
        "staff_panel/round_form.html",
        {
            "form": form,
            "activities": activities,
            "round_types": _choices(ContestRound.RoundType),
            "scoring_modes": _choices(ContestRound.ScoringMode),
            "order_policies": _choices(ContestRound.OrderPolicy),
            "tie_order_policies": _choices(ContestRound.TieOrderPolicy),
            "roster_sources": [("", "自动判定"), *_choices(ContestRound.RosterSource)],
            "rubrics": rubrics,
            "selected_activity_id": selected_activity_id,
        },
    )


@staff_required
def round_edit(request, pk):
    contest_round = get_object_or_404(ContestRound.objects.select_related("activity"), pk=pk)
    activities = _editable_singer_activities()
    activity = contest_round.activity
    if contest_round.status != ContestRound.Status.DRAFT or activity.is_locked:
        raise PermissionDenied("只有未准备且未锁定活动中的轮次可以修改配置。")
    rubrics = ScoringRubric.objects.filter(activity=activity)
    initial = {
        "name": contest_round.name,
        "round_type": contest_round.round_type,
        "scoring_mode": contest_round.scoring_mode,
        "score_completion_mode": contest_round.score_completion_mode,
        "judge_count": contest_round.judge_count,
        "minimum_judge_count": contest_round.minimum_judge_count,
        "advance_count": contest_round.advance_count,
        "sequence": contest_round.sequence,
        "order_policy": contest_round.order_policy,
        "tie_order_policy": contest_round.tie_order_policy,
        "roster_source": contest_round.roster_source,
        "roster_source_stage": contest_round.roster_source_stage,
        "rubric": contest_round.rubric_id,
    }
    if request.method == "POST":
        form = ContestRoundForm(
            request.POST,
            rubrics=rubrics,
            stage_choices=round_roster_stage_choices(activity),
        )
        if form.is_valid():
            if form.cleaned_data["roster_source"] == ContestRound.RosterSource.STAGE:
                try:
                    validate_roster_source_stage(activity, form.cleaned_data["roster_source_stage"])
                except ValidationError as error:
                    form.add_error("roster_source_stage", domain_error_messages(error))
            sequence = form.cleaned_data.get("sequence")
            if (
                not form.errors
                and sequence is not None
                and ContestRound.objects.filter(activity=activity, sequence=sequence)
                .exclude(pk=contest_round.pk)
                .exists()
            ):
                form.add_error("sequence", "轮次序号已存在，请填写其他序号。")
            if form.is_valid():
                with transaction.atomic():
                    locked_round = ContestRound.objects.select_for_update().get(pk=contest_round.pk)
                    if locked_round.status != ContestRound.Status.DRAFT:
                        raise PermissionDenied("轮次已准备，不能修改配置。")
                    for field in (
                        "name",
                        "round_type",
                        "scoring_mode",
                        "score_completion_mode",
                        "advance_count",
                        "sequence",
                        "order_policy",
                        "tie_order_policy",
                        "roster_source",
                        "roster_source_stage",
                        "rubric",
                    ):
                        setattr(locked_round, field, form.cleaned_data[field])
                    locked_round.judge_count = form.cleaned_data["judge_count"]
                    locked_round.minimum_judge_count = (
                        form.cleaned_data["minimum_judge_count"] or form.cleaned_data["judge_count"]
                    )
                    locked_round.save(
                        update_fields=[
                            "name",
                            "round_type",
                            "scoring_mode",
                            "score_completion_mode",
                            "judge_count",
                            "minimum_judge_count",
                            "advance_count",
                            "sequence",
                            "order_policy",
                            "tie_order_policy",
                            "roster_source",
                            "roster_source_stage",
                            "rubric",
                        ]
                    )
                log_action(
                    request,
                    AuditLog.ActionType.OTHER,
                    f"ContestRound:{contest_round.pk}",
                    new_value="configuration_updated",
                )
                return redirect("staff:round_list")
    else:
        form = ContestRoundForm(
            rubrics=rubrics,
            stage_choices=round_roster_stage_choices(activity),
            initial=initial,
        )
    return render(
        request,
        "staff_panel/round_form.html",
        {
            "form": form,
            "activities": activities,
            "round_types": _choices(ContestRound.RoundType),
            "scoring_modes": _choices(ContestRound.ScoringMode),
            "order_policies": _choices(ContestRound.OrderPolicy),
            "tie_order_policies": _choices(ContestRound.TieOrderPolicy),
            "roster_sources": [("", "自动判定"), *_choices(ContestRound.RosterSource)],
            "rubrics": rubrics,
            "selected_activity_id": str(activity.pk),
            "editing_round": contest_round,
        },
    )


@staff_required
@require_POST
def round_prepare(request, pk):
    contest_round = get_object_or_404(ContestRound.objects.select_related("activity"), pk=pk)
    try:
        ensure_activity_unlocked(contest_round.activity)
        ensure_activity_action_allowed(contest_round.activity, ActivityAction.SCORE)
        prepare_round(contest_round, request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"比赛轮次准备失败：{domain_error_messages(error)}")
    return redirect("staff:round_list")


@staff_required
def round_running_order(request, pk):
    contest_round = get_object_or_404(ContestRound.objects.select_related("activity"), pk=pk)
    try:
        singers = current_round_roster(contest_round)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"暂时无法配置人工顺序：{domain_error_messages(error)}")
        return redirect("staff:round_list")
    singer_by_id = {str(singer.pk): singer for singer in singers}
    entries = list(contest_round.entries.select_related("singer").order_by("running_order", "pk"))
    initial_ids = [str(entry.singer_id) for entry in entries] or [
        str(singer.pk) for singer in singers
    ]
    display_ids = initial_ids
    if request.method == "POST":
        form = RoundRunningOrderForm(request.POST, singers=singers)
        display_ids = request.POST.getlist("singer_ids")
        if form.is_valid():
            ordered_ids = form.cleaned_data["singer_ids"]
            move_action = request.POST.get("move_action")
            if move_action is not None:
                try:
                    direction, raw_index = move_action.split(":", 1)
                    index = int(raw_index)
                except (TypeError, ValueError):
                    direction, index = "", -1
                target = index + (-1 if direction == "up" else 1)
                if direction not in {"up", "down"}:
                    target = -1
                if 0 <= index < len(ordered_ids) and 0 <= target < len(ordered_ids):
                    ordered_ids[index], ordered_ids[target] = (
                        ordered_ids[target],
                        ordered_ids[index],
                    )
                form = RoundRunningOrderForm(initial={"singer_ids": ordered_ids}, singers=singers)
                initial_ids = ordered_ids
                display_ids = ordered_ids
            else:
                try:
                    set_round_running_order(contest_round, ordered_ids, request.user)
                except (PermissionDenied, ValidationError) as error:
                    form.add_error(None, domain_error_messages(error))
                else:
                    messages.success(request, "人工出场顺序已保存。")
                    return redirect("staff:round_list")
    else:
        form = RoundRunningOrderForm(initial={"singer_ids": initial_ids}, singers=singers)
    if form.is_bound:
        ordered_ids = cast(QueryDict, form.data).getlist("singer_ids")
    else:
        ordered_ids = form.initial.get("singer_ids", display_ids)
    if hasattr(ordered_ids, "__iter__") and not isinstance(ordered_ids, str):
        ordered_ids = [str(value) for value in ordered_ids]
    order_rows = [
        {"singer": singer_by_id[singer_id], "singer_id": singer_id}
        for singer_id in ordered_ids
        if singer_id in singer_by_id
    ]
    return render(
        request,
        "staff_panel/round_running_order.html",
        {"round": contest_round, "form": form, "entries": entries, "order_rows": order_rows},
    )


@staff_required
def round_groups(request, pk):
    contest_round = get_object_or_404(ContestRound.objects.select_related("activity"), pk=pk)
    try:
        singers = current_round_roster(contest_round)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"暂时无法配置分组：{domain_error_messages(error)}")
        return redirect("staff:round_list")
    groups = list(
        PerformanceGroup.objects.filter(round=contest_round)
        .prefetch_related("performances__singer")
        .order_by("sequence", "pk")
    )
    frozen_group_snapshot = False
    group_snapshot_error = ""
    try:
        live_snapshot = live_first_round_group_snapshot(contest_round)
    except (PermissionDenied, ValidationError) as error:
        live_snapshot = None
        group_snapshot_error = domain_error_messages(error)
    initial_groups = [
        {"name": group.name, "singer_ids": [p.singer_id for p in group.performances.all()]}
        for group in groups
    ]
    if not groups and live_snapshot is not None and request.method != "POST":
        initial_groups = [
            {"name": name, "singer_ids": list(singer_ids)} for name, singer_ids in live_snapshot
        ]
        frozen_group_snapshot = True
    if request.method == "POST":
        form = RoundGroupsForm(request.POST, singers=singers)
        if form.is_valid():
            try:
                set_round_groups(contest_round, form.cleaned_data["groups"], request.user)
            except (PermissionDenied, ValidationError) as error:
                form.add_error(None, domain_error_messages(error))
            else:
                messages.success(request, "分组与演出归属已保存。")
                return redirect("staff:round_list")
    else:
        form = RoundGroupsForm(singers=singers, initial_groups=initial_groups)
    group_slots = [
        {
            "number": index,
            "name": form[f"group_name_{index}"],
            "members": form[f"group_singer_ids_{index}"],
        }
        for index in range(1, RoundGroupsForm.MAX_GROUP_SLOTS + 1)
    ]
    return render(
        request,
        "staff_panel/round_groups.html",
        {
            "round": contest_round,
            "form": form,
            "singers": singers,
            "group_slots": group_slots,
            "frozen_group_snapshot": frozen_group_snapshot,
            "group_snapshot_error": group_snapshot_error,
        },
    )


@staff_required
def round_score_entry(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    if contest_round.status == ContestRound.Status.DRAFT:
        return render(
            request,
            "staff_panel/round_score_entry.html",
            {"round": contest_round, "preparation_required": True},
        )

    return render(
        request,
        "staff_panel/round_score_entry.html",
        {
            "round": contest_round,
            "grid_payload": _round_grid_payload(contest_round),
            "is_locked": contest_round.is_locked
            or contest_round.status == ContestRound.Status.LOCKED,
        },
    )


def _round_grid_payload(contest_round: ContestRound) -> dict:
    singers = prime_questionnaire_answers(_eligible_singers(contest_round))
    judges = list(authoritative_panel_judges(contest_round))
    scores = {
        (s.singer_id, s.judge_id): str(s.score)
        for s in ScoreRecord.objects.filter(round=contest_round)
    }
    grid = [
        {
            "singer_id": singer.pk,
            "singer_name": singer.name,
            "song": questionnaire_song_for_round(singer, contest_round, default="未填写"),
            "cells": [
                {
                    "judge_id": judge.pk,
                    "judge_name": judge.name,
                    "score": scores.get((singer.pk, judge.pk), ""),
                }
                for judge in judges
            ],
        }
        for singer in singers
    ]
    # Completeness from the rows this payload already materialised. Calling
    # missing_score_cells() here would repeat the eligible-singer, authoritative-
    # judge and ScoreRecord queries a second time for the same answer.
    expected_cells = {(singer.pk, judge.pk) for singer in singers for judge in judges}
    matrix_complete = expected_cells.issubset(scores)
    return {
        "version": contest_round.score_version,
        "matrix_complete": matrix_complete,
        "grid": grid,
        "judges": [{"id": judge.pk, "name": judge.name} for judge in judges],
    }


@staff_required
def round_scores_api(request, pk):
    """Backstage rapid-entry endpoint (M1-H).

    GET returns the current grid + ``score_version`` (the client's base version).
    POST accepts sparse ``cells`` and ``base_version``; a version mismatch is a 409
    (stale edit) so two staff never silently overwrite each other. When the save
    completes the score matrix the server re-resolves the activity's bound ruleset
    (best-effort) and reports the resulting stage status.
    """
    contest_round = get_object_or_404(ContestRound, pk=pk)

    if request.method == "GET":
        return JsonResponse({**_round_grid_payload(contest_round)})

    try:
        payload = json.loads(request.body or "{}")
    except json.JSONDecodeError:
        return JsonResponse({"detail": "请求体不是有效 JSON。"}, status=400)
    command_form = RapidScoreCommandForm(
        {"command_id": payload.get("command_id"), "base_version": payload.get("base_version")}
    )
    if not command_form.is_valid():
        return JsonResponse(
            {"detail": command_form.errors.get_json_data(), "reason_code": "INVALID_REQUEST"},
            status=400,
        )
    base_version = command_form.cleaned_data["base_version"]
    command_id = command_form.cleaned_data["command_id"]
    cells = payload.get("cells", [])
    if not isinstance(cells, list):
        return JsonResponse(
            {"detail": "cells 必须是列表。", "reason_code": "INVALID_REQUEST"}, status=400
        )

    score_values = {}
    base_values = {}
    has_cell_base = False
    has_legacy_cell = False
    for cell in cells:
        try:
            singer_id = int(cell["singer_id"])
            judge_id = int(cell["judge_id"])
        except (KeyError, TypeError, ValueError):
            return JsonResponse({"detail": "单元格缺少 singer_id/judge_id。"}, status=400)
        pair = (singer_id, judge_id)
        score_values[pair] = str(cell.get("score", "")).strip()
        if "base" in cell:
            if not isinstance(cell["base"], str):
                return JsonResponse(
                    {"detail": "单元格 base 必须是文本。", "reason_code": "INVALID_REQUEST"},
                    status=400,
                )
            has_cell_base = True
            base_values[pair] = cell["base"]
        else:
            has_legacy_cell = True

    if has_cell_base and has_legacy_cell:
        return JsonResponse(
            {
                "detail": "评分请求不能混用逐格 base 和旧版单元格格式。",
                "reason_code": "INVALID_REQUEST",
            },
            status=400,
        )

    # ``apply_scores_if_version`` owns the Activity→Round transaction and its locks;
    # the view holds no row lock itself (M0 canonical order, §13.2 stale-guard).
    try:
        result = apply_scores_if_version(
            contest_round.pk,
            base_version,
            score_values,
            request.user,
            command_id=command_id,
            base_values=base_values if has_cell_base else None,
        )
    except IdempotencyConflictError:
        return JsonResponse(
            {
                "detail": "同一 command_id 已用于不同的评分请求。",
                "reason_code": IdempotencyConflictError.reason_code,
                "conflict": True,
            },
            status=409,
        )
    except StaleScoreVersionError:
        contest_round.refresh_from_db()
        return JsonResponse(
            {
                **_round_grid_payload(contest_round),
                "conflict": True,
                "reason_code": "STALE_SCORE_VERSION",
            },
            status=409,
        )
    except ValidationError as error:
        return JsonResponse(
            {"detail": error.messages, "reason_code": "INVALID_REQUEST"}, status=400
        )
    except PermissionDenied:
        return JsonResponse({"detail": "权限不足。"}, status=403)

    resolved_status = None
    resolve_warning = None
    if result["matrix_complete"]:
        try:
            maybe_resolve_checkpoints(contest_round.activity, request.user)
        except (ValidationError, PermissionDenied):
            # Surface a ruleset/resolve problem to the operator instead of silently
            # dropping it (M1-INTEGRATION-CLOSE Item 7): the scores are saved but the
            # stage that could not auto-resolve must be visible, not invisible.
            resolve_warning = _AUTO_RESOLVE_WARNING
        # Report the stage's true current status (not just whether this save newly resolved
        # it): a no-op re-save of an already-READY stage must not read as "未计算".
        resolved_status = _current_resolve_status(contest_round.activity)
    response = {**result, "resolved_status": resolved_status}
    if resolve_warning:
        response["resolve_warning"] = resolve_warning
    return JsonResponse(response)


def _audience_pool(activity, version, audience_key):
    """The roster an audience score-group scores, derived from the frozen definition.

    ``audience_key`` is a binding key matching an ASSESS node's ``vote_source``; that
    node's ``source`` declares the pool: :data:`ENTRY_KEY` → the check-in pool, or the
    name of a SELECT node → the checkpoint that outputs it, bridged to the round that
    consumes that stage (its :class:`RoundEntry`). Invalid definitions and missing stage
    rosters return an empty queryset so an operator cannot accidentally score every singer.
    """
    try:
        parsed = parse_definition(version.definition)
    except (ValidationError, TypeError, ValueError):
        return SingerRegistration.objects.none()
    source = None
    for node in parsed["nodes"]:
        if node.get("vote_source") == audience_key:
            source = node.get("source")
            break
    if not source:
        return SingerRegistration.objects.none()
    if source == ENTRY_KEY:
        return runtime_approved_singers(activity).order_by("pk")
    checkpoint_key = None
    for cp in parsed["checkpoints"]:
        if cp.get("output") == source:
            checkpoint_key = cp.get("key")
            break
    if checkpoint_key:
        advance_round = (
            ContestRound.objects.filter(
                activity=activity,
                roster_source=ContestRound.RosterSource.STAGE,
                roster_source_stage=checkpoint_key,
                status__in=[
                    ContestRound.Status.PREPARED,
                    ContestRound.Status.SCORING,
                    ContestRound.Status.LOCKED,
                ],
            )
            .order_by("-sequence")
            .first()
        )
        if advance_round is not None:
            singer_ids = list(
                RoundEntry.objects.filter(round=advance_round).values_list("singer_id", flat=True)
            )
            return list(
                SingerRegistration.objects.filter(pk__in=singer_ids, activity=activity).order_by(
                    "pk"
                )
            )
    return SingerRegistration.objects.none()


def _audience_sets(activity, version):
    """Build the audience grid, scoped per stage to its real roster.

    Each binding ``audience_key → set_name`` exposes the roster the stage actually scores
    (its ``within`` pool), derived from the frozen definition via :func:`_audience_pool`.
    This prevents a staff member from entering audience scores for singers who are no
    longer in that stage (e.g. audience4 over all 15).
    """
    if version is None:
        return []
    binding = _version_binding(version)
    audience_keys = binding.get("audience_keys") or {}
    stored: dict[tuple[int, str], str] = {}
    provenance: dict[str, dict[str, str]] = {}
    for row in AudienceScore.objects.filter(
        activity=activity, is_test_data=runtime_is_test(activity)
    ):
        stored[(row.singer_id, row.stage_key)] = str(row.score)
        provenance.setdefault(
            row.stage_key,
            {
                "source": row.source,
                "source_label": row.get_source_display(),
                "note": row.source_note,
            },
        )
    sets = []
    for set_key, set_name in audience_keys.items():
        singers = prime_questionnaire_answers(_audience_pool(activity, version, set_key))
        for singer in singers:
            singer.song_label = generic_song_label(singer)
        sets.append(
            {
                "set_key": set_key,
                "stage_key": set_name,
                "provenance": provenance.get(set_name, {"source": "", "note": ""}),
                "rows": [
                    {
                        "singer_id": singer.pk,
                        "singer_name": singer.name,
                        "song": singer.song_label,
                        "score": stored.get((singer.pk, set_name), ""),
                    }
                    for singer in singers
                ],
            }
        )
    return sets


@staff_required
def audience_score_entry(request, activity_id):
    """Minimal staff page for entering backstage audience scores."""
    activity = get_object_or_404(Activity, pk=activity_id)
    try:
        version = _current_frozen_version(activity)
    except ValidationError as error:
        # Render the empty state with the reason. Redirecting back into this same view
        # re-raises the same error on the redirected GET, so the page used to loop until
        # the browser gave up with ERR_TOO_MANY_REDIRECTS and the grid stayed unreachable.
        messages.error(request, domain_error_messages(error))
        version = None
    sets = _audience_sets(activity, version)
    recorded = next(
        (s["provenance"] for s in sets if s["provenance"]["source"] or s["provenance"]["note"]),
        {"source": "", "note": ""},
    )
    return render(
        request,
        "staff_panel/audience_score_entry.html",
        {
            "activity": activity,
            "sets": sets,
            "api_url": reverse("staff:audience_scores_api", kwargs={"activity_id": activity_id}),
            "source_choices": AudienceScore.Source.choices,
            "default_source": recorded["source"],
            "default_source_note": recorded["note"],
        },
    )


@staff_required
def audience_scores_api(request, activity_id):
    """Backstage audience-score entry (M1-INTEGRATION-2).

    GET returns the audience grid: one group per binding ``audience_keys`` entry, each
    listing the current roster for each binding plus its stored score (scale ``hundred``).
    POST accepts sparse ``cells`` as ``[{singer_id, set_key, score}]``, maps each ``set_key``
    to its ``stage_key``, upserts :class:`~singer_contest.models.AudienceScore` rows, then
    auto-resolves any now-satisfiable checkpoint.
    """
    activity = get_object_or_404(Activity, pk=activity_id)
    try:
        version = _current_frozen_version(activity)
    except ValidationError as error:
        return JsonResponse({"detail": error.messages}, status=400)
    binding = _version_binding(version) if version is not None else {}
    audience_keys = binding.get("audience_keys") or {}
    if not audience_keys:
        return JsonResponse({"detail": "该活动未绑定观众分。"}, status=400)

    if request.method == "GET":
        return JsonResponse({"sets": _audience_sets(activity, version)})

    try:
        payload = json.loads(request.body or "{}")
    except json.JSONDecodeError:
        return JsonResponse({"detail": "请求体不是有效 JSON。"}, status=400)
    cells = payload.get("cells", [])
    if not cells:
        return JsonResponse({"detail": "缺少 cells。"}, status=400)
    # §9.8: a degraded / external audience source must name where its numbers came from.
    # Both keys stay optional so an existing caller that only sends cells keeps working.
    source = str(payload.get("source", "") or "").strip()
    source_note = str(payload.get("note", "") or "").strip()
    valid_sources = {choice for choice, _label in AudienceScore.Source.choices}
    if source not in valid_sources:
        return JsonResponse({"detail": [f"未知观众分来源 {source}。"]}, status=400)
    if len(source_note) > 200:
        return JsonResponse({"detail": ["来源备注不能超过 200 字。"]}, status=400)

    pool_ids_by_set = {
        set_key: {singer.pk for singer in _audience_pool(activity, version, set_key)}
        for set_key in audience_keys
    }
    rows: list[tuple[int, str, Decimal]] = []
    stage_keys: set[str] = set()
    errors: list[str] = []
    for cell in cells:
        try:
            singer_id = int(cell["singer_id"])
            score = Decimal(str(cell.get("score", "")).strip())
        except (KeyError, TypeError, ValueError, InvalidOperation):
            errors.append("单元格缺少 singer_id 或 score 不是有效数字。")
            continue
        set_key = str(cell.get("set_key", "")).strip()
        stage_key = audience_keys.get(set_key)
        if not stage_key:
            errors.append(f"未知观众分组 {set_key}。")
            continue
        if singer_id not in pool_ids_by_set[set_key]:
            errors.append(f"选手 {singer_id} 不属于观众分组 {set_key}。")
            continue
        if not 0 <= score <= 100:
            errors.append(f"选手 {singer_id} 的观众分必须在 0-100 之间。")
            continue
        # AudienceScore.score is Decimal(max_digits=8, decimal_places=2); reject more than
        # two decimal places so SQLite (stores verbatim) and Postgres (rounds to 0.01) agree.
        if score != score.quantize(Decimal("0.01")):
            errors.append(f"选手 {singer_id} 的观众分最多保留两位小数。")
            continue
        rows.append((singer_id, stage_key, score))
        stage_keys.add(stage_key)
    if errors:
        return JsonResponse({"detail": errors}, status=400)

    # Hold the Activity FOR UPDATE lock across the guard + write so the check and the rows
    # are serialized against a concurrent confirm_stage_result (which takes the same lock);
    # otherwise a confirm could commit a stage consuming this audience set between our guard
    # and our insert. The auto-resolve runs in its own transaction afterwards.
    try:
        with transaction.atomic():
            lock_activity_for_action(activity)
            for stage_key in stage_keys:
                ensure_audience_not_consumed_by_confirmed_stage(activity, stage_key)
            test_flag = runtime_is_test(activity)
            # §9.8: a manually entered fallback score is a formal fact, so the *first* row
            # for a slot has to say where its number came from. Correcting a number later
            # may still omit the source — the rule protects against inventing provenance,
            # not against erasing a record the caller never touched.
            existing_slots = {
                (stage_key, singer_id)
                for stage_key, singer_id in AudienceScore.objects.filter(
                    activity=activity,
                    stage_key__in=stage_keys,
                    singer_id__in=[row[0] for row in rows],
                ).values_list("stage_key", "singer_id")
            }
            if not source and any(
                (stage_key, singer_id) not in existing_slots
                for singer_id, stage_key, _score in rows
            ):
                raise ValidationError("首次录入观众分必须注明来源（外部表单 / 纸质统计 / 其他）。")
            for singer_id, stage_key, score in rows:
                defaults: dict[str, object] = {
                    "score": score,
                    "entered_by": request.user,
                    "is_test_data": test_flag,
                }
                # Only restamp provenance when a source was supplied, so re-saving numbers
                # without touching the source control does not erase what was recorded.
                if source:
                    defaults["source"] = source
                    defaults["source_note"] = source_note
                AudienceScore.objects.update_or_create(
                    activity=activity,
                    stage_key=stage_key,
                    singer_id=singer_id,
                    defaults=defaults,
                )
            # §14: a manually verified fallback score is a formal fact, so entering it must
            # be traceable. The audit carries the set names and row counts only — never the
            # score values, which live in the raw fact rows themselves.
            AuditLog.objects.create(
                operator=request.user,
                action_type=AuditLog.ActionType.ENTER_SCORE,
                target=f"AudienceScore:{activity.pk}",
                old_value="",
                new_value=(
                    f"sets={sorted(stage_keys)}; rows={len(rows)}; "
                    f"source={source or AudienceScore.Source.UNRECORDED}"
                ),
                note="人工/外部核验观众分录入",
            )
    except ValidationError as error:
        return JsonResponse({"detail": error.messages}, status=400)
    except PermissionDenied:
        return JsonResponse({"detail": "权限不足。"}, status=403)

    # Resolve after the save transaction commits so a ruleset misconfiguration (missing bound
    # rounds, duplicate auto rulesets) reports an error without rolling back entered scores.
    # Surface it as resolve_warning (M1-INTEGRATION-CLOSE Item 7) so the operator sees that
    # a stage could not auto-resolve, instead of it disappearing silently.
    resolve_warning = None
    try:
        maybe_resolve_checkpoints(activity, request.user)
    except (ValidationError, PermissionDenied):
        resolve_warning = _AUTO_RESOLVE_WARNING

    response = {
        "saved": len(rows),
        "resolved_status": _current_resolve_status(activity),
    }
    if resolve_warning:
        response["resolve_warning"] = resolve_warning
    return JsonResponse(response)


@staff_required
def round_ranking(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    summaries = ScoreSummary.objects.filter(
        round=contest_round, singer__round_entries__round=contest_round
    ).select_related("singer")
    summary_rows = list(summaries)
    singers = prime_questionnaire_answers(summary.singer for summary in summary_rows)
    for singer in singers:
        singer.song_label = questionnaire_song_for_round(singer, contest_round, default="未填写")
    missing_cells = missing_score_cells(contest_round)
    return render(
        request,
        "staff_panel/round_ranking.html",
        {
            "round": contest_round,
            "summaries": summaries,
            "preparation_required": contest_round.status == ContestRound.Status.DRAFT,
            "has_missing": bool(missing_cells),
            "missing_cells": missing_cells,
        },
    )


@staff_required
def activity_result_board(request, activity_id):
    """List the latest StageResult per stage_key with HOLD/REVIEW/READY banners (M1-H)."""
    activity = get_object_or_404(Activity, pk=activity_id)
    stages = list(latest_stage_result_queryset(activity).order_by("stage_key", "-pk"))
    return render(
        request,
        "staff_panel/activity_result_board.html",
        {"activity": activity, "stages": stages},
    )


_RESULT_CLOSURE_LABELS = (
    {"code": ResultClosureCode.NO_CURRENT_FROZEN_RULESET.value, "label": "没有当前冻结赛制"},
    {"code": ResultClosureCode.RULESET_BINDING_INVALID.value, "label": "赛制绑定无效"},
    {"code": ResultClosureCode.RAW_FACTS_INCOMPLETE.value, "label": "原始输入未齐"},
    {"code": ResultClosureCode.RAW_FACTS_UNLOCKED.value, "label": "原始输入未锁定"},
    {"code": ResultClosureCode.RULE_REVIEW_REQUIRED.value, "label": "规则要求人工复核"},
    {
        "code": ResultClosureCode.UPSTREAM_CONFIRMATION_PENDING.value,
        "label": "上游赛段尚未核定",
    },
    {"code": ResultClosureCode.STALE_CANDIDATE.value, "label": "候选结果已过期"},
    {"code": ResultClosureCode.STAGE_CONFIRMATION_PENDING.value, "label": "赛段等待核定"},
    {
        "code": ResultClosureCode.ACTIVITY_OPERATIONALLY_LOCKED.value,
        "label": "活动已被操作锁定",
    },
    {
        "code": ResultClosureCode.ACTIVITY_PHASE_NOT_READY.value,
        "label": "活动尚未进入结果整理/公示阶段",
    },
    {
        "code": ResultClosureCode.SCHOOL_EXTERNAL_EVIDENCE_PENDING.value,
        "label": "学校外部证据待补齐",
    },
)


@staff_required
@require_GET
def activity_result_closure(request, activity_id):
    """Render the read-only event-day result closure checklist."""
    activity = get_object_or_404(Activity, pk=activity_id)
    closure = build_result_closure(activity)
    return render(
        request,
        "staff_panel/activity_result_closure.html",
        {
            "activity": activity,
            "closure": closure,
            "closure_payload": result_closure_as_dict(closure),
            "blocking_labels": _RESULT_CLOSURE_LABELS,
        },
    )


@staff_required
def stage_result_detail(request, pk):
    """Render a single stage result's decisions grouped into handcard blocks (M1-H)."""
    stage = get_object_or_404(
        latest_stage_result_queryset(),
        pk=pk,
    )
    blocks = stage_decisions_by_blocks(stage)
    decisions = [decision for block in blocks for decision in block["decisions"]]
    prime_questionnaire_answers(decision.singer for decision in decisions)
    for decision in decisions:
        decision.song_label = generic_song_label(decision.singer)
    return render(
        request,
        "staff_panel/stage_result_detail.html",
        {
            "stage": stage,
            "activity": stage.activity,
            "blocks": blocks,
            # Staff may perform the operational confirmation; only unlocking remains
            # an administrator authority because it reopens an already-final result.
            "can_confirm_stage_result": request.user.is_staff_or_admin,
            "can_unlock_stage_result": request.user.is_admin,
        },
    )


@staff_required
@require_GET
def stage_result_pdf(request, pk):
    """Download the stage result handcard as a PDF (GOAL §2.4).

    Renders the same blocks the staff page shows, so the printed handcard and the page
    can never disagree about what the result is.
    """
    stage = get_object_or_404(latest_stage_result_queryset(), pk=pk)
    blocks = stage_decisions_by_blocks(stage)
    decisions = [decision for block in blocks for decision in block["decisions"]]
    prime_questionnaire_answers(decision.singer for decision in decisions)
    for decision in decisions:
        decision.song_label = generic_song_label(decision.singer)
    response = HttpResponse(
        render_stage_result_pdf(stage=stage, blocks=blocks), content_type="application/pdf"
    )
    response["Content-Disposition"] = f'inline; filename="stage_result_{stage.pk}.pdf"'
    response["Cache-Control"] = "no-store"
    return response


@staff_required
@require_POST
def stage_result_confirm(request, pk):
    """核定并锁定 a stage result into its final handcard state (M1-H §36-37)."""
    stage = get_object_or_404(StageResult, pk=pk)
    try:
        ensure_activity_action_allowed(stage.activity, ActivityAction.PUBLISH_RESULT)
        confirm_stage_result(stage, confirmed_by=request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"赛段核定失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "已核定并锁定该赛段结果，可抄主持手卡。")
    return redirect("staff:stage_result_detail", pk=pk)


@admin_required
@require_POST
def stage_result_unlock(request, pk):
    """Admin-only: unlock a confirmed stage result so its raw facts can be corrected."""
    _require_admin(request.user)
    stage = get_object_or_404(StageResult, pk=pk)
    try:
        unlock_stage_result(stage, operator=request.user, note=request.POST.get("note", "").strip())
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"赛段解锁失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "已解锁该赛段结果，可修正原始数据后重新核定。")
    return redirect("staff:stage_result_detail", pk=pk)


@staff_required
@require_POST
def round_finalize_advancement(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        ensure_activity_action_allowed(contest_round.activity, ActivityAction.SCORE)
        finalize_advancement(
            contest_round, request.POST.getlist("selected_singer_ids"), request.user
        )
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"晋级名单核定失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "已核定晋级名单。")
    return redirect("staff:round_ranking", pk=pk)


@staff_required
@require_POST
def round_lock(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        lock_round(contest_round, request.user)
    except ValidationError as error:
        messages.error(request, f"轮次锁定失败：{domain_error_messages(error)}")
        return redirect("staff:round_ranking", pk=pk)
    except PermissionDenied as error:
        messages.error(request, f"轮次锁定失败：{domain_error_messages(error)}")
        return redirect("staff:round_ranking", pk=pk)
    try:
        maybe_resolve_checkpoints(contest_round.activity, request.user)
    except ValidationError as error:
        messages.warning(
            request,
            f"轮次已锁定，但赛段未能自动解析：{domain_error_messages(error)}",
        )
    except PermissionDenied as error:
        messages.warning(
            request,
            f"轮次已锁定，但赛段未能自动解析：{domain_error_messages(error)}",
        )
    return redirect("staff:round_ranking", pk=pk)


@admin_required
@require_POST
def round_unlock(request, pk):
    _require_admin(request.user)
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        unlock_round(contest_round, request.user, note=request.POST.get("note", "").strip())
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"轮次解锁失败：{domain_error_messages(error)}")
    return redirect("staff:round_ranking", pk=pk)


@admin_required
@require_POST
def round_reset(request, pk):
    _require_admin(request.user)
    contest_round = get_object_or_404(ContestRound, pk=pk)
    reason = request.POST.get("reason", "")
    if not reason.strip():
        messages.error(request, "轮次重置失败：重置轮次必须填写原因。")
        return redirect("staff:round_list")
    try:
        reset_round_to_draft(contest_round, request.user, reason=reason)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"轮次重置失败：{domain_error_messages(error)}")
    return redirect("staff:round_list")


@staff_required
def judge_list(request):
    judges = Judge.objects.select_related("activity")
    return render(request, "staff_panel/judge_list.html", {"judges": judges})


@staff_required
def judge_create(request):
    activities = _editable_singer_activities()
    if request.method == "POST":
        activity = get_object_or_404(activities, pk=request.POST["activity_id"])
        with transaction.atomic():
            locked_activity = lock_activity_for_action(activity)
            judge = Judge.objects.create(
                activity=locked_activity,
                name=request.POST["name"],
            )
            log_action(
                request, AuditLog.ActionType.OTHER, f"Judge:{judge.pk}", new_value=judge.name
            )
        return redirect("staff:judge_list")

    return render(request, "staff_panel/judge_form.html", {"activities": activities})


def _live_first_round_material_status(contest_round):
    """Non-blocking R1 group material readiness for the live control page."""
    try:
        stage = _live_group_stage_for_round(contest_round)
    except ValidationError:
        # A read-only readiness banner must not surface a broken binding as a 500.
        return None
    if stage is None:
        return None
    rows = [
        row
        for row in live_material_summary(contest_round.activity)["first_round"]["groups"]
        if row["stage_id"] == stage.pk
    ]
    if not rows:
        return None
    ready_count = sum(row["status"] == "ready" for row in rows)
    return {
        "stage_name": rows[0]["stage_name"],
        "ready_count": ready_count,
        "total_count": len(rows),
        "missing_count": len(rows) - ready_count,
    }


def _judge_control_context(
    request,
    contest_round,
    *,
    error="",
    policy_state_override="",
):
    round_judges = list(contest_round.round_judges.select_related("judge").order_by("pk"))
    panel_snapshot = (
        RoundPanelSnapshot.objects.filter(round=contest_round)
        .exclude(state=RoundPanelSnapshot.State.SUPERSEDED)
        .order_by("-version")
        .first()
    )
    panel_members = []
    panel_member_rows = []
    seat_rows = []
    if panel_snapshot is not None:
        panel_members = list(
            panel_snapshot.members.select_related("judge").filter(is_active=True).order_by("pk")
        )
        seats = {
            seat.panel_member_id: seat
            for seat in JudgeSeat.objects.filter(panel_member__in=panel_members).order_by("pk")
        }
        # Which seats currently hold a live session. GOAL §12.6 lets staff release a seat
        # when a device is lost or swapped, and that action only means anything for a seat
        # that is actually occupied, so the control page needs to say which those are.
        connected_seat_ids = set(
            JudgeSession.objects.filter(
                seat__in=[seat for seat in seats.values() if seat is not None],
                state=JudgeSession.State.ACTIVE,
                expires_at__gt=timezone.now(),
            ).values_list("seat_id", flat=True)
        )
        for member in panel_members:
            seat = seats.get(member.pk)
            seat_label = (
                seat.display_label.strip()
                if seat is not None and seat.display_label.strip()
                else member.seat_key.replace("seat-", "J")
            )
            panel_member_rows.append({"member": member, "seat_label": seat_label})
            seat_rows.append(
                {
                    "member": member,
                    "seat": seat,
                    "seat_label": seat_label,
                    "connected": seat is not None and seat.pk in connected_seat_ids,
                }
            )

    run_state = (
        PerformanceRunState.objects.filter(round=contest_round)
        .select_related("current_performance__singer")
        .first()
    )
    performances = list(
        Performance._base_manager.filter(round=contest_round)
        .select_related("singer")
        .order_by("sequence", "pk")
    )
    performance_score_states = {
        state.performance_id: state
        for state in PerformanceScoreState._base_manager.filter(
            performance_id__in=[performance.pk for performance in performances]
        )
    }
    performance_rows = [
        {"performance": performance, "score_state": performance_score_states.get(performance.pk)}
        for performance in performances
    ]
    performance_group_rows = [
        {
            "group": group,
            "performances": [
                {
                    "performance": performance,
                    "score_state": performance_score_states.get(performance.pk),
                }
                for performance in performances
                if performance.group_id == group.pk
            ],
        }
        for group in PerformanceGroup._base_manager.filter(round=contest_round).order_by(
            "sequence", "pk"
        )
    ]
    expected_judge_count = (
        contest_round.judge_count if contest_round.judge_count is not None else len(round_judges)
    )
    actual_judge_count = len(panel_members)
    minimum_judge_count = (
        panel_snapshot.minimum_judge_count
        if panel_snapshot is not None
        else (
            expected_judge_count
            if contest_round.judge_count is not None
            else contest_round.minimum_judge_count or expected_judge_count
        )
    )
    if panel_snapshot is None:
        policy_state = "PREPARE_REQUIRED"
    elif panel_snapshot.state == RoundPanelSnapshot.State.HOLD:
        policy_state = "HOLD"
    elif run_state is not None and run_state.state == PerformanceRunState.State.HOLD:
        policy_state = "HOLD"
    else:
        policy_state = "ACTIVE"
    if policy_state_override:
        policy_state = policy_state_override
    live_material_status = _live_first_round_material_status(contest_round)
    policy_state_labels = {
        "PREPARE_REQUIRED": "待准备评委组",
        "HOLD": "已暂停",
        "ACTIVE": "现场进行中",
        "INSUFFICIENT_JUDGES": "评委组人数不足，暂不可开始",
    }

    attendance_form = JudgePanelAttendanceForm(
        initial={"attending_judge_ids": [member.judge_id for member in panel_members]},
        judge_choices=[
            (round_judge.judge_id, round_judge.judge.name) for round_judge in round_judges
        ],
    )
    return {
        "round": contest_round,
        "round_judges": round_judges,
        "panel_snapshot": panel_snapshot,
        "panel_members": panel_members,
        "panel_member_rows": panel_member_rows,
        "seat_rows": seat_rows,
        "run_state": run_state,
        "performances": performances,
        "performance_rows": performance_rows,
        "performance_group_rows": performance_group_rows,
        "expected_judge_count": expected_judge_count,
        "actual_judge_count": actual_judge_count,
        "minimum_judge_count": minimum_judge_count,
        "policy_state": policy_state,
        "policy_state_label": policy_state_labels.get(policy_state, "现场状态待确认"),
        "live_material_status": live_material_status,
        "attendance_form": attendance_form,
        "performance_action_form": JudgePerformanceActionForm(),
        "score_form": JudgeBoundScoreForm(
            rubric=contest_round.rubric,
            initial={
                "performance_id": run_state.current_performance_id if run_state else "",
                "context_version": run_state.context_version if run_state else 0,
            },
        ),
        "score_actions": (
            {
                "kind": "proxy",
                "title": "工作人员代录",
                "reference_label": "故障记录或来源编号",
                "command_id": f"staff-proxy-{uuid4().hex}",
            },
            {
                "kind": "paper",
                "title": "纸面评分补录",
                "reference_label": "纸面评分编号",
                "command_id": f"staff-paper-{uuid4().hex}",
            },
        ),
        "error": error,
    }


def _render_judge_control(request, contest_round, **kwargs):
    return render(
        request,
        "staff_panel/judge_control.html",
        _judge_control_context(request, contest_round, **kwargs),
    )


@staff_required
def judge_control(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    return _render_judge_control(request, contest_round)


@staff_required
@require_POST
@transaction.atomic
def judge_prepare(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    round_judges = list(contest_round.round_judges.select_related("judge").order_by("pk"))
    form = JudgePanelAttendanceForm(
        request.POST,
        judge_choices=[
            (round_judge.judge_id, round_judge.judge.name) for round_judge in round_judges
        ],
    )
    if not form.is_valid():
        return _render_judge_control(request, contest_round, error=_form_error(form))
    try:
        prepare_judge_panel(
            contest_round.pk,
            operator=request.user,
            attending_judge_ids=(
                None
                if contest_round.judge_count is not None
                and not request.POST.getlist("attending_judge_ids")
                else form.cleaned_data["attending_judge_ids"]
            ),
        )
    except (PermissionDenied, ValidationError) as error:
        error_message = domain_error_messages(error)
        return _render_judge_control(
            request,
            contest_round,
            error=error_message,
            policy_state_override=(
                "INSUFFICIENT_JUDGES" if "INSUFFICIENT_JUDGES" in str(error) else ""
            ),
        )
    messages.success(request, "评委组已准备，评委席位和评分上下文已冻结。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_panel_hold(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        hold_judge_panel(
            contest_round.pk,
            operator=request.user,
            reason=request.POST.get("reason", ""),
        )
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"暂停评委组失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "评委组已暂停，当前评委会话已保留；恢复后可继续评分。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_panel_resume(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        resume_judge_panel(contest_round.pk, operator=request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"恢复评委组失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "评委组已恢复。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_seat_label(request, pk, seat_id):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        seat = get_object_or_404(
            JudgeSeat,
            pk=seat_id,
            panel_member__panel_snapshot__round=contest_round,
        )
        set_judge_seat_display_label(
            seat.pk,
            operator=request.user,
            display_label=request.POST.get("display_label", ""),
        )
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"保存评委席位备注失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "评委席位备注已保存。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_seat_release(request, pk, seat_id):
    """Free a seat's live session so a replacement device can take over (GOAL §12.4)."""
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        seat = get_object_or_404(
            JudgeSeat,
            pk=seat_id,
            panel_member__panel_snapshot__round=contest_round,
        )
        released = release_judge_seat(
            seat.pk,
            operator=request.user,
            reason=request.POST.get("reason", ""),
        )
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"释放评委席位失败：{domain_error_messages(error)}")
    else:
        if released is None:
            messages.info(request, "该席位当前没有已连接的评委终端。")
        else:
            messages.success(request, "评委席位已释放，该席位下一台设备可直接扫码接入。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_performance_advance(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    form = JudgePerformanceActionForm(request.POST)
    if form.is_valid():
        try:
            advance_performance(
                contest_round.pk,
                form.cleaned_data["performance_id"],
                operator=request.user,
            )
        except (PermissionDenied, ValidationError) as error:
            messages.error(request, f"推进表演失败：{domain_error_messages(error)}")
        else:
            messages.success(request, "当前表演已切换，旧评分上下文自动失效。")
    else:
        messages.error(request, f"推进表演失败：{_form_error(form)}")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_performance_group_advance(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        group_id = int(request.POST.get("group_id", ""))
    except ValueError:
        messages.error(request, "开放表演组失败：表演组无效。")
        return redirect("staff:judge_control", pk=pk)
    try:
        advance_performance_group(
            group_id=group_id,
            round_id=contest_round.pk,
            operator=request.user,
        )
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"开放表演组失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "本组独立评分已全部开放，评委可分别提交每位成员的分数。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_performance_score_lock(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    form = JudgePerformanceActionForm(request.POST)
    if form.is_valid() and form.cleaned_data["context_version"] is not None:
        try:
            lock_performance_score(
                contest_round.pk,
                form.cleaned_data["performance_id"],
                operator=request.user,
                expected_context_version=form.cleaned_data["context_version"],
            )
        except (PermissionDenied, ValidationError) as error:
            messages.error(request, f"锁定表演评分失败：{domain_error_messages(error)}")
        else:
            messages.success(request, "该表演评分已锁定，之后的评分将被拒绝。")
    else:
        messages.error(request, "锁定表演评分失败：现场版本已失效，请刷新后重试。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_performance_hold(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    form = JudgePerformanceActionForm(request.POST)
    if form.is_valid():
        try:
            hold_performance(
                contest_round.pk,
                operator=request.user,
                reason=form.cleaned_data["reason"],
                expected_performance_id=form.cleaned_data["performance_id"],
            )
        except (PermissionDenied, ValidationError) as error:
            messages.error(request, f"暂停表演失败：{domain_error_messages(error)}")
        else:
            messages.success(request, "当前表演已暂停。")
    else:
        messages.error(request, f"暂停表演失败：{_form_error(form)}")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_performance_resume(request, pk):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    try:
        resume_performance(contest_round.pk, operator=request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"恢复表演失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "当前表演已恢复。")
    return redirect("staff:judge_control", pk=pk)


def _submit_staff_judge_score(request, pk, *, paper):
    contest_round = get_object_or_404(ContestRound, pk=pk)
    form = JudgeBoundScoreForm(request.POST, rubric=contest_round.rubric)
    if not form.is_valid():
        messages.error(request, f"评分提交失败：{_form_error(form)}")
        return redirect("staff:judge_control", pk=pk)
    values = form.cleaned_data
    if contest_round.rubric is None:
        payload = {"score": values["score"], "notes": values["notes"]}
    else:
        payload = {
            "criteria": [
                {
                    "criterion_id": criterion.pk,
                    "value": values[f"criterion_{criterion.pk}"],
                }
                for criterion in contest_round.rubric.criteria.all().order_by("sequence", "pk")
            ],
            "notes": values["notes"],
        }
    try:
        if paper:
            submit_paper_score(
                contest_round.pk,
                values["performance_id"],
                values["seat_id"],
                operator=request.user,
                command_id=values["command_id"],
                expected_context_version=values["context_version"],
                score_payload=payload,
                paper_reference=values["source_reference"],
                reason=values["reason"],
            )
        else:
            submit_staff_proxy_score(
                contest_round.pk,
                values["performance_id"],
                values["seat_id"],
                operator=request.user,
                command_id=values["command_id"],
                expected_context_version=values["context_version"],
                score_payload=payload,
                source_reference=values["source_reference"],
                reason=values["reason"],
            )
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"评分提交失败：{domain_error_messages(error)}")
    else:
        messages.success(request, "工作人员评分已记录，并绑定到当前评委席位和表演上下文。")
    return redirect("staff:judge_control", pk=pk)


@staff_required
@require_POST
def judge_score_proxy(request, pk):
    return _submit_staff_judge_score(request, pk, paper=False)


@staff_required
@require_POST
def judge_score_paper(request, pk):
    return _submit_staff_judge_score(request, pk, paper=True)


@staff_required
def rubric_create(request):
    activities = _editable_singer_activities()
    form = ScoringRubricProvisionForm(request.POST or None)
    requested_activity_id = (
        request.POST.get("activity_id") or request.GET.get("activity_id") or None
    )
    selected_activity = (
        activities.filter(pk=requested_activity_id).first()
        if requested_activity_id is not None
        else activities.first()
    )
    if selected_activity is None:
        selected_activity = activities.first()
    selected_activity_id = str(selected_activity.pk) if selected_activity else ""
    if request.method == "POST" and form.is_valid():
        activity = get_object_or_404(activities, pk=request.POST.get("activity_id"))
        try:
            create_scoring_rubric(
                activity,
                name=form.cleaned_data["name"],
                description=form.cleaned_data["description"],
                criteria=form.cleaned_data["criteria"],
                operator=request.user,
            )
        except (PermissionDenied, ValidationError) as error:
            form.add_error(None, domain_error_messages(error))
        else:
            messages.success(request, "评分标准及评分项已创建。")
            return redirect(f"{reverse('staff:round_create')}?activity_id={activity.pk}")
    criterion_slots = [
        {
            "number": index,
            "name": form[f"criterion_name_{index}"],
            "max_score": form[f"criterion_max_score_{index}"],
            "description": form[f"criterion_description_{index}"],
        }
        for index in range(1, ScoringRubricProvisionForm.MAX_CRITERION_SLOTS + 1)
    ]
    return render(
        request,
        "staff_panel/rubric_form.html",
        {
            "form": form,
            "activities": activities,
            "criterion_slots": criterion_slots,
            "selected_activity_id": selected_activity_id,
        },
    )


@staff_required
def award_list(request):
    awards = official_stage_award_queryset().select_related("activity")
    return render(request, "staff_panel/award_list.html", {"awards": awards})


# --- Vote session management ---


@staff_required
def vote_session_list(request):
    sessions = VoteSession.objects.select_related("activity")
    activity_id = request.GET.get("activity_id") or ""
    if activity_id.isdigit():
        sessions = sessions.filter(activity_id=activity_id)
    return render(
        request,
        "staff_panel/vote_session_list.html",
        {"sessions": sessions, "selected_activity_id": activity_id},
    )


@staff_required
@transaction.atomic
def vote_session_create(request):
    activities = _editable_singer_activities()
    if request.method == "POST":
        activity = get_object_or_404(activities, pk=request.POST["activity_id"])
        activity = lock_activity_for_runtime_data(activity)
        ensure_activity_unlocked(activity)
        ensure_activity_action_allowed(activity, ActivityAction.MANAGE_VOTE)
        form = VoteSessionForm(request.POST, activity=activity)
        if not form.is_valid():
            return render(
                request,
                "staff_panel/vote_session_form.html",
                {
                    "error": _form_error(form),
                    "form": form,
                    "activities": activities,
                    "singers": _with_generic_song_labels(
                        scope_lifecycle(
                            SingerRegistration.objects.select_related("activity").filter(
                                activity=activity,
                                pre_status=SingerRegistration.PreStatus.APPROVED,
                            )
                        )
                    ),
                    "selected_activity_id": activity.pk,
                    "submitted_singer_ids": request.POST.getlist("singers"),
                    "selection_types": _choices(VoteSession.SelectionType),
                    "purposes": _choices(VoteSession.Purpose),
                    "formal_ticket_voting": formal_singer_contest_requires_ticket(activity),
                },
            )
        singer_ids = request.POST.getlist("singers")
        if len(singer_ids) != len(set(singer_ids)):
            raise PermissionDenied("Vote options cannot contain duplicate singers.")
        selected_singers = list(
            SingerRegistration.objects.filter(
                pk__in=singer_ids,
                activity=activity,
                pre_status=SingerRegistration.PreStatus.APPROVED,
            )
        )
        if len(selected_singers) != len(singer_ids):
            raise PermissionDenied("Every vote option must be an approved singer.")
        for singer in selected_singers:
            ensure_same_activity(activity, singer, label="Vote option singer")
            ensure_lifecycle_consistent(activity, singer, label="投票候选选手")
        vote_session = VoteSession.objects.create(
            activity=activity,
            name=form.cleaned_data["name"],
            passcode=(
                ""
                if formal_singer_contest_requires_ticket(activity)
                else form.cleaned_data["passcode"]
            ),
            start_time=form.cleaned_data["start_time"],
            end_time=form.cleaned_data["end_time"],
            selection_type=form.cleaned_data["selection_type"],
            max_selections=form.cleaned_data["max_selections"],
            purpose=form.cleaned_data["purpose"],
            requires_ticket=(
                True
                if formal_singer_contest_requires_ticket(activity)
                else form.cleaned_data["requires_ticket"]
            ),
            is_test_data=activity.is_test_mode,
        )
        for i, sid in enumerate(singer_ids):
            VoteOption.objects.create(
                vote_session=vote_session,
                singer_id=sid,
                sort_order=i,
                is_test_data=activity.is_test_mode,
            )
        log_action(
            request,
            AuditLog.ActionType.VOTE_MANAGE,
            f"VoteSession:{vote_session.pk}",
            new_value=vote_session.name,
        )
        return redirect("staff:vote_session_list")

    requested_activity_id = request.GET.get("activity_id") or None
    selected_activity = (
        activities.filter(pk=requested_activity_id).first()
        if requested_activity_id is not None
        else activities.first()
    )
    if selected_activity is None:
        selected_activity = activities.first()
    selected_activity_id = str(selected_activity.pk) if selected_activity else ""
    singer_queryset = SingerRegistration.objects.select_related("activity").filter(
        pre_status=SingerRegistration.PreStatus.APPROVED
    )
    if selected_activity_id.isdigit():
        singer_queryset = singer_queryset.filter(activity_id=selected_activity_id)
    else:
        singer_queryset = singer_queryset.none()
    singers = _with_generic_song_labels(scope_lifecycle(singer_queryset))
    form = VoteSessionForm(activity=selected_activity)
    return render(
        request,
        "staff_panel/vote_session_form.html",
        {
            "form": form,
            "activities": activities,
            "singers": singers,
            "selected_activity_id": selected_activity_id,
            "selection_types": _choices(VoteSession.SelectionType),
            "purposes": _choices(VoteSession.Purpose),
            "submitted_singer_ids": [],
            "formal_ticket_voting": bool(
                selected_activity and formal_singer_contest_requires_ticket(selected_activity)
            ),
        },
    )


@staff_required
def vote_session_detail(request, pk):
    vote_session = get_object_or_404(VoteSession.objects.select_related("activity"), pk=pk)
    # §9.1 / §9.3: the screen has to reconcile against the *valid* ballots the conversion
    # will publish, so a revoked or voided ticket's ballot is not counted here either.
    options = list(
        vote_session.options.select_related("singer").annotate(vote_count=valid_vote_count())
    )
    _with_generic_song_labels(option.singer for option in options)
    total_votes = valid_record_queryset(
        vote_session, test_flag=runtime_is_test(vote_session.activity)
    ).count()
    _, top = _popularity_top_tie(vote_session)
    # §9: a score-component session must be reconcilable by hand — valid ballot count,
    # per-candidate votes and support rate next to the conversion mode and lock state.
    conversion = None
    conversion_rows = []
    if vote_session.purpose == VoteSession.Purpose.SCORE_COMPONENT:
        conversion = vote_session_configuration_facts(
            vote_session, test_flag=runtime_is_test(vote_session.activity)
        )
        rates = conversion["support_rate"]
        counted = conversion["candidate_votes"]
        conversion_rows = [
            {
                "singer": option.singer,
                "votes": counted.get(option.singer_id, 0),
                "rate": rates.get(str(option.singer_id)),
            }
            for option in options
        ]
    return render(
        request,
        "staff_panel/vote_session_detail.html",
        {
            "vote_session": vote_session,
            "options": options,
            "total_votes": total_votes,
            "popularity_tie": len(top) > 1,
            "conversion": conversion,
            "conversion_rows": conversion_rows,
        },
    )


@staff_required
@require_POST
def vote_session_open(request, pk):
    vote_session = get_object_or_404(VoteSession.objects.select_related("activity"), pk=pk)
    try:
        ensure_activity_action_allowed(vote_session.activity, ActivityAction.MANAGE_VOTE)
        open_vote_session(vote_session, request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"投票开放失败：{domain_error_messages(error)}")
    return redirect("staff:vote_session_detail", pk=pk)


@staff_required
@require_POST
def vote_session_close(request, pk):
    vote_session = get_object_or_404(VoteSession.objects.select_related("activity"), pk=pk)
    try:
        ensure_activity_action_allowed(vote_session.activity, ActivityAction.MANAGE_VOTE)
        close_vote_session(vote_session, request.user)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"投票关闭失败：{domain_error_messages(error)}")
    return redirect("staff:vote_session_detail", pk=pk)


@staff_required
@require_POST
@transaction.atomic
def vote_session_lock(request, pk):
    vote_session = get_object_or_404(VoteSession.objects.select_related("activity"), pk=pk)
    try:
        ensure_activity_action_allowed(vote_session.activity, ActivityAction.MANAGE_VOTE)
        lock_vote_session(vote_session, request.user)
    except (ValidationError, PermissionDenied) as error:
        messages.error(request, f"投票锁定失败：{domain_error_messages(error)}")
        return redirect("staff:vote_session_detail", pk=pk)
    try:
        maybe_resolve_checkpoints(vote_session.activity, request.user)
    except (ValidationError, PermissionDenied) as error:
        messages.warning(
            request,
            "投票已锁定，但自动重算未完成：" + domain_error_messages(error),
        )
    return redirect("staff:vote_session_detail", pk=pk)


@admin_required
@require_POST
def vote_session_unlock(request, pk):
    _require_admin(request.user)
    vote_session = get_object_or_404(VoteSession, pk=pk)
    note = request.POST.get("note", "").strip()
    try:
        unlock_vote_session(vote_session, request.user, note=note)
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, f"投票解锁失败：{domain_error_messages(error)}")
    return redirect("staff:vote_session_detail", pk=pk)


@staff_required
def vote_session_export(request, pk):
    vote_session = get_object_or_404(VoteSession, pk=pk)
    wb = Workbook()
    ws = _active_worksheet(wb)
    ws.title = "投票结果"
    ws.append(["选手", "票数"])
    # §9.3: the exported number has to be the number the result is computed from, or the
    # file a staff member hands to the judges disagrees with the result on screen.
    option_list = list(
        vote_session.options.select_related("singer").annotate(vote_count=valid_vote_count())
    )
    for opt in option_list:
        ws.append([opt.singer.name, getattr(opt, "vote_count", 0)])
    audit_export(request, vote_session.activity, "vote_result", row_count=len(option_list))
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f"attachment; filename=vote_result_{vote_session.pk}.xlsx"
    wb.save(response)
    return response


def _popularity_top_tie(vote_session):
    """Return (top_vote_count, tied_top_options).

    An empty tuple result means there are no winning ballots. The tie list is
    deliberately order-independent so a database sort can never decide who wins.
    """
    leaderboard = list(
        vote_session.options.select_related("singer")
        .annotate(vote_count=valid_vote_count())
        .filter(vote_count__gt=0)
        .order_by("-vote_count", "sort_order", "pk")
    )
    if not leaderboard:
        return None, []
    top_count = getattr(leaderboard[0], "vote_count", 0)
    return top_count, [
        option for option in leaderboard if getattr(option, "vote_count", 0) == top_count
    ]


# --- QR code center ---


@staff_required
def qr_center(request):
    activities = Activity.objects.filter(activity_type=Activity.Type.SINGER_CONTEST)
    return render(
        request,
        "staff_panel/qr_center.html",
        {"activities": activities, "qr_host_warning": qr_print_host_warning(request)},
    )


@staff_required
def qr_generate(request, pk):
    activity = get_object_or_404(
        Activity,
        pk=pk,
        activity_type=Activity.Type.SINGER_CONTEST,
    )
    return render(
        request,
        "staff_panel/qr_detail.html",
        {"activity": activity, "qr_host_warning": qr_print_host_warning(request)},
    )


@staff_required
def qr_image(request, pk, kind):
    activity = get_object_or_404(
        Activity,
        pk=pk,
        activity_type=Activity.Type.SINGER_CONTEST,
    )
    stable_paths = {
        "activity": "public_portal:activity_entry",
        "apply": "public_portal:activity_apply",
        "live": "public_portal:activity_live",
        "judge": "public_portal:activity_judge",
        # Keep already printed registration sheets usable while moving them to
        # the stable activity-scoped entry. Vote/result QR kinds are retired.
        "registration": "public_portal:activity_apply",
    }
    route_name = stable_paths.get(kind)
    if route_name is not None:
        path = reverse(route_name, kwargs={"public_code": activity.public_code})
    else:
        return HttpResponse("Unknown QR code kind", status=404)
    if kind == "judge":
        # The judge QR is the only one that carries a capability: the shared entry is the
        # one place where the URL *is* the credential. It goes in the fragment, which the
        # browser never sends, so the secret stays out of access logs and out of `Referer`.
        path = f"{path}#{judge_entry_credential(activity)}"

    import qrcode

    absolute_url = request.build_absolute_uri(path)
    if settings.APP_ENV == "production":
        parts = urlsplit(absolute_url)
        absolute_url = urlunsplit(("https", parts.netloc, parts.path, parts.query, parts.fragment))
    image = qrcode.make(absolute_url)
    buffer = io.BytesIO()
    image.save(buffer)
    response = HttpResponse(buffer.getvalue(), content_type="image/png")
    response["Content-Disposition"] = f'inline; filename="{kind}_{activity.pk}.png"'
    response["Cache-Control"] = "no-store"
    response["Pragma"] = "no-cache"
    return response


# --- Export center ---


@staff_required
def export_center(request):
    activities = list(Activity.objects.all())
    activity_rows = []
    for activity in activities:
        counts = get_test_data_counts(activity) if activity.is_test_mode else {}
        activity_rows.append(
            {
                "activity": activity,
                "test_data_counts": counts,
                "test_data_total": sum(counts.values()),
            }
        )
    templates = ArticleTemplate.objects.all()
    return render(
        request,
        "staff_panel/export_center.html",
        {
            "activities": activities,
            "activity_rows": activity_rows,
            "templates": templates,
        },
    )


@staff_required
def excel_material_checklist(request, activity_id):
    activity = get_object_or_404(Activity, pk=activity_id)
    wb = Workbook()
    ws = _active_worksheet(wb)
    ws.title = "Material Checklist"
    ws.append(["Owner Type", "Owner", "Item", "Status"])
    row_count = 0
    for c in MaterialCheck.objects.filter(
        singer_registration__activity=activity,
        singer_registration__is_test_data=runtime_is_test(activity),
    ).select_related("singer_registration"):
        ws.append(
            [
                "singer",
                c.singer_registration.name if c.singer_registration else "",
                c.item_name,
                c.get_status_display(),
            ]
        )
        row_count += 1
    for c in MaterialCheck.objects.filter(
        program__activity=activity,
        program__is_test_data=runtime_is_test(activity),
    ).select_related("program"):
        ws.append(
            [
                "program",
                c.program.name if c.program else "",
                c.item_name,
                c.get_status_display(),
            ]
        )
        row_count += 1
    audit_export(request, activity, "material_checklist", row_count=row_count)
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=material_checklist.xlsx"
    wb.save(response)
    return response


@staff_required
def excel_score_template(request, round_id):
    contest_round = get_object_or_404(ContestRound, pk=round_id)
    singers = list(_eligible_singers(contest_round))
    wb = build_score_template_workbook(contest_round)
    audit_export(request, contest_round.activity, "score_template", row_count=len(singers))
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=score_template.xlsx"
    wb.save(response)
    return response


@staff_required
def excel_import_scores(request, round_id):
    contest_round = get_object_or_404(ContestRound, pk=round_id)
    errors = []
    if request.method == "POST" and request.FILES.get("file"):
        ensure_round_unlocked(contest_round)
        try:
            scores, errors = parse_score_workbook(request.FILES["file"], contest_round)
            if not errors:
                apply_scores(contest_round, scores, request.user, note="excel import")
                return redirect("staff:round_score_entry", pk=round_id)
        except ValidationError as error:
            errors.extend(error.messages)
        except Exception:
            errors.append("文件解析失败，请上传有效的 xlsx 评分表。")
    elif request.method == "POST":
        errors.append("请选择要导入的 xlsx 评分表。")
    return render(
        request,
        "staff_panel/excel_import.html",
        {
            "round": contest_round,
            "errors": errors,
        },
    )


# --- Word generation ---


@staff_required
@require_POST
def word_generate(request, template_id, activity_id):
    template = get_object_or_404(ArticleTemplate, pk=template_id)
    activity = get_object_or_404(Activity, pk=activity_id)
    if activity.phase == Activity.Phase.ARCHIVED:
        # Read-only preview/download for an archived activity. Archived data is
        # immutable, so this never creates a persistent asset and the official
        # archive package stays the authoritative record of it.
        content = render_document_bytes(template, activity)
        audit_export(request, activity, "word_preview")
        return _docx_response(content, template, activity.pk)
    doc_obj, content = generate_persistent_document(activity, template, request.user)
    return _docx_response(content, template, activity.pk)


def _docx_response(content: bytes, template: ArticleTemplate, activity_id: int) -> HttpResponse:
    response = HttpResponse(
        content,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    response["Content-Disposition"] = (
        f"attachment; filename={template.template_type}_{activity_id}.docx"
    )
    return response


# --- Execution & archive packages ---


@staff_required
def execution_package(request, activity_id):
    activity = get_object_or_404(Activity, pk=activity_id)
    artifacts = build_execution_package(activity, request)
    response = HttpResponse(build_package_zip(artifacts), content_type="application/zip")
    response["Content-Disposition"] = f"attachment; filename=execution_{activity_id}.zip"
    audit_export(request, activity, "execution_package")
    return response


@staff_required
@require_POST
def archive_package_create(request, activity_id):
    """Produce a source-of-truth-free preview ZIP; never mutates the activity."""
    activity = get_object_or_404(Activity, pk=activity_id)
    artifacts = build_archive_package(activity)
    response = HttpResponse(build_package_zip(artifacts), content_type="application/zip")
    response["Content-Disposition"] = f"attachment; filename=archive_{activity_id}.zip"
    audit_export(request, activity, "archive_preview")
    return response


@admin_required
@require_POST
def activity_archive(request, pk):
    activity = get_object_or_404(Activity, pk=pk)
    try:
        archive_activity(activity, request.user, note=request.POST.get("note", "").strip())
    except (PermissionDenied, ValidationError) as error:
        messages.error(request, domain_error_messages(error))
        return redirect("staff:export_center")
    messages.success(request, "活动已归档并锁定。")
    return redirect("staff:export_center")


# --- Incident records ---


@staff_required
def incident_list(request):
    incidents = IncidentRecord.objects.select_related("activity", "singer", "handled_by", "round")
    return render(request, "staff_panel/incident_list.html", {"incidents": incidents})


@staff_required
@transaction.atomic
def incident_create(request):
    activities = Activity.objects.all()
    form = IncidentForm(request.POST or None)
    if request.method == "POST":
        activity = get_object_or_404(Activity, pk=request.POST["activity_id"])
        # An incident is deliberately NOT gated on the activity lock. GOAL §16 makes the
        # incident record a first-class object for the failures that happen around a
        # lock — a restart, a restore, a manual override — and IncidentRecord even
        # carries an authority_state of LOCKED for exactly that case. Rejecting the write
        # when the activity is locked would make the record unavailable precisely when it
        # is needed. The lock is still taken, so this row serializes with the rest of the
        # activity's writes.
        activity = lock_activity_for_runtime_data(activity)
        if not form.is_valid():
            return render(
                request,
                "staff_panel/incident_form.html",
                {
                    "error": _form_error(form),
                    "activities": activities,
                    "singers": _with_generic_song_labels(
                        scope_lifecycle(
                            SingerRegistration.objects.select_related("activity").filter(
                                activity=activity, pre_status=SingerRegistration.PreStatus.APPROVED
                            )
                        )
                    ),
                    "event_types": _choices(IncidentRecord.EventType),
                    "authority_states": _choices(IncidentRecord.AuthorityState),
                    "rounds": _incident_rounds(activity),
                    "selected_activity_id": activity.pk,
                    "selected_round_id": request.POST.get("round_id", ""),
                    "form": form,
                    "form_data": request.POST,
                    "selected_singer_id": request.POST.get("singer_id", ""),
                },
            )
        singer = None
        if request.POST.get("singer_id"):
            singer = get_object_or_404(SingerRegistration, pk=request.POST["singer_id"])
            ensure_same_activity(activity, singer, label="Incident singer")
            ensure_lifecycle_consistent(activity, singer, label="涉事选手")
        program = None
        if request.POST.get("program_id"):
            program = get_object_or_404(Program, pk=request.POST["program_id"])
            ensure_same_activity(activity, program, label="Incident program")
            ensure_lifecycle_consistent(activity, program, label="涉事节目")
        contest_round = None
        if request.POST.get("round_id"):
            contest_round = get_object_or_404(ContestRound, pk=request.POST["round_id"])
            ensure_same_activity(activity, contest_round, label="Incident round")
        incident = IncidentRecord.objects.create(
            activity=activity,
            occurred_at=form.cleaned_data["occurred_at"],
            event_type=form.cleaned_data["event_type"],
            round=contest_round,
            authority_state=form.cleaned_data.get("authority_state", ""),
            needs_review=form.cleaned_data.get("needs_review", False),
            singer=singer,
            program=program,
            handled_by_id=request.user.pk,
            problem=form.cleaned_data.get("problem", ""),
            action_taken=form.cleaned_data.get("action_taken", ""),
            resolution=form.cleaned_data.get("resolution", ""),
            remark=form.cleaned_data.get("remark", ""),
            is_test=activity.is_test_mode,
        )
        log_action(
            request,
            AuditLog.ActionType.OTHER,
            f"IncidentRecord:{incident.pk}",
            new_value=incident.event_type,
        )
        return redirect("staff:incident_list")
    selected_activity_id = request.GET.get("activity_id") or ""
    singer_queryset = SingerRegistration.objects.select_related("activity").filter(
        pre_status=SingerRegistration.PreStatus.APPROVED
    )
    if selected_activity_id.isdigit():
        singer_queryset = singer_queryset.filter(activity_id=selected_activity_id)
    else:
        singer_queryset = singer_queryset.none()
    singers = _with_generic_song_labels(scope_lifecycle(singer_queryset))
    return render(
        request,
        "staff_panel/incident_form.html",
        {
            "activities": activities,
            "singers": singers,
            "event_types": _choices(IncidentRecord.EventType),
            "authority_states": _choices(IncidentRecord.AuthorityState),
            "rounds": _incident_rounds(int(selected_activity_id))
            if selected_activity_id.isdigit()
            else [],
            "selected_activity_id": selected_activity_id,
            "selected_round_id": "",
            "form": form,
            "form_data": {},
            "selected_singer_id": "",
        },
    )


@staff_required
def incident_export(request):
    activity = _export_activity(request)
    if activity is None and not request.user.is_admin:
        raise PermissionDenied("全量导出仅管理员可用；请先用活动筛选导出单个活动。")
    wb = Workbook()
    ws = _active_worksheet(wb)
    ws.title = "异常记录"
    # Same builder as the permanent archive sheet, so the evidence a staff member sees
    # now and the evidence that survives the event cannot describe different columns.
    ws.append(list(INCIDENT_LIST_HEADERS))
    related = ("activity", "round", "singer", "program", "handled_by")
    if activity is not None:
        incidents = IncidentRecord.objects.select_related(*related).filter(
            activity=activity, is_test=runtime_is_test(activity)
        )
    else:
        incidents = IncidentRecord.objects.select_related(*related).filter(is_test=False)
    rows = incident_list_rows(incidents)
    for row in rows:
        ws.append(list(row))
    audit_export(request, activity, "incident_list", row_count=len(rows))
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = "attachment; filename=incident_list.xlsx"
    wb.save(response)
    return response


# --- Test mode management ---


@admin_required
@require_POST
def activity_test_toggle(request, pk):
    _require_admin(request.user)
    activity = get_object_or_404(Activity, pk=pk)
    if not activity.is_test_mode:
        raise PermissionDenied(
            "Formal activities cannot re-enter test mode. Clone the activity instead."
        )
    leave_test_mode(
        activity,
        operator=request.user,
        clear=request.POST.get("clear") == "on",
        reason=request.POST.get("reason", ""),
    )
    return redirect("staff:export_center")


@admin_required
@require_POST
def activity_clear_test_data(request, pk):
    _require_admin(request.user)
    activity = get_object_or_404(Activity, pk=pk)
    if not activity.is_test_mode:
        raise PermissionDenied("Test data can only be cleared while the activity is in test mode.")
    counts = clear_activity_test_data(activity, operator=request.user)
    messages.success(request, f"测试数据已清理，共 {sum(counts.values())} 条记录。")
    return redirect("staff:export_center")


# --- Activity management ---


@admin_required
@require_POST
@transaction.atomic
def activity_lock(request, pk):
    _require_admin(request.user)
    activity = get_object_or_404(Activity.objects.select_for_update(), pk=pk)
    activity.is_locked = True
    activity.locked_at = timezone.now()
    activity.locked_by = request.user
    with authority_write(ACTIVITY_STATE):
        activity.save(update_fields=["is_locked", "locked_at", "locked_by"])
    log_action(
        request, AuditLog.ActionType.RELOCK_RESULT, f"Activity:{activity.pk}", new_value="locked"
    )
    return redirect("staff:export_center")


@admin_required
@require_POST
@transaction.atomic
def activity_unlock(request, pk):
    _require_admin(request.user)
    note = request.POST.get("note", "").strip()
    activity = get_object_or_404(Activity.objects.select_for_update(), pk=pk)
    if activity.phase == Activity.Phase.ARCHIVED:
        raise PermissionDenied("活动已归档，请使用解归档功能。")
    activity.is_locked = False
    activity.locked_at = None
    activity.locked_by = None
    with authority_write(ACTIVITY_STATE):
        activity.save(update_fields=["is_locked", "locked_at", "locked_by"])
    log_action(
        request,
        AuditLog.ActionType.UNLOCK_RESULT,
        f"Activity:{activity.pk}",
        old_value="locked",
        new_value="unlocked",
        note=note,
    )
    return redirect("staff:export_center")


@admin_required
@require_POST
def activity_unarchive(request, pk):
    _require_admin(request.user)
    activity = get_object_or_404(Activity, pk=pk)
    unarchive_activity(activity, actor=request.user, note=request.POST.get("note", "").strip())
    messages.success(request, "活动已解归档并恢复到结果公示阶段。")
    return redirect("staff:export_center")


@admin_required
@require_POST
@transaction.atomic
def activity_clone(request, pk):
    _require_admin(request.user)
    original = get_object_or_404(Activity, pk=pk)
    new_activity = Activity.objects.create(
        title=f"{original.title}（副本）",
        subtitle=original.subtitle,
        activity_type=original.activity_type,
        phase=Activity.Phase.DRAFT,
        description=original.description,
        is_test_mode=True,
    )
    for req in MaterialRequirement.objects.filter(activity=original):
        MaterialRequirement.objects.create(
            activity=new_activity,
            applies_to=req.applies_to,
            item_name=req.item_name,
            file_purpose=req.file_purpose,
            is_required=req.is_required,
            sort_order=req.sort_order,
        )
    for judge in Judge.objects.filter(activity=original):
        Judge.objects.create(
            activity=new_activity,
            name=judge.name,
            is_active=judge.is_active,
        )
    rubric_map: dict[int | None, ScoringRubric] = {}
    for rubric in ScoringRubric.objects.filter(activity=original):
        new_rubric = ScoringRubric.objects.create(
            activity=new_activity,
            name=rubric.name,
            description=rubric.description,
            sequence=rubric.sequence,
            is_test_data=runtime_is_test(new_activity),
        )
        for criterion in rubric.criteria.all():
            RubricCriterion.objects.create(
                rubric=new_rubric,
                name=criterion.name,
                max_score=criterion.max_score,
                sequence=criterion.sequence,
                description=criterion.description,
                is_test_data=runtime_is_test(new_activity),
            )
        rubric_map[rubric.pk] = new_rubric
    for contest_round in ContestRound.objects.filter(activity=original):
        ContestRound.objects.create(
            activity=new_activity,
            round_type=contest_round.round_type,
            scoring_mode=contest_round.scoring_mode,
            score_completion_mode=contest_round.score_completion_mode,
            minimum_judge_count=contest_round.minimum_judge_count,
            judge_count=contest_round.judge_count,
            name=contest_round.name,
            advance_count=contest_round.advance_count,
            sequence=contest_round.sequence,
            order_policy=contest_round.order_policy,
            tie_order_policy=contest_round.tie_order_policy,
            roster_source=contest_round.roster_source,
            roster_source_stage=contest_round.roster_source_stage,
            rubric=rubric_map.get(contest_round.rubric_id),
            is_locked=False,
        )
    # Vote sessions are runtime state (passcode, start/end window, open/locked),
    # not reusable template config. Cloning would copy a stale passcode and time
    # window into the new activity, so staff create sessions fresh per activity.
    log_action(
        request,
        AuditLog.ActionType.OTHER,
        f"Activity:{original.pk}",
        new_value=f"cloned_to={new_activity.pk}",
    )
    return redirect("staff:export_center")


# --- Audit log ---


@staff_required
def audit_log_list(request):
    logs = AuditLog.objects.select_related("operator")[:200]
    return render(request, "staff_panel/audit_log_list.html", {"logs": logs})


# --- User & role administration ---


@admin_required
def user_list(request):
    _require_admin(request.user)
    users = User.objects.order_by("-date_joined", "username")
    return render(request, "staff_panel/user_list.html", {"users": users})


@admin_required
@require_POST
def user_role_update(request, pk):
    _require_admin(request.user)
    target = get_object_or_404(User, pk=pk)
    try:
        updated = change_user_role(
            target=target, new_role=request.POST.get("new_role", ""), actor=request.user
        )
    except (ValidationError, PermissionDenied) as error:
        messages.error(request, domain_error_messages(error))
    else:
        messages.success(
            request, f"已将 {updated.username} 的角色改为 {updated.get_role_display()}。"
        )
    return redirect("staff:user_list")


@admin_required
@require_POST
def user_set_active(request, pk):
    _require_admin(request.user)
    target = get_object_or_404(User, pk=pk)
    is_active = request.POST.get("active") == "1"
    try:
        updated = set_user_active(target=target, is_active=is_active, actor=request.user)
    except (ValidationError, PermissionDenied) as error:
        messages.error(request, domain_error_messages(error))
    else:
        messages.success(
            request, f"已{'启用' if updated.is_active else '停用'}账号 {updated.username}。"
        )
    return redirect("staff:user_list")


@staff_required
def ruleset_template_list(request):
    # Show the catalog with explicit capability status. Only production templates
    # expose cloning; experimental/unsupported entries must never look production-ready.
    templates = RulesetTemplate.objects.all()
    activities = _editable_singer_activities()
    return render(
        request,
        "staff_panel/ruleset_template_list.html",
        {"templates": templates, "total": templates.count(), "activities": activities},
    )


@staff_required
def ruleset_template_detail(request, pk):
    template = get_object_or_404(RulesetTemplate, pk=pk)
    nodes = parse_definition(template.definition)["nodes"] if template.definition else []
    activities = _editable_singer_activities()
    return render(
        request,
        "staff_panel/ruleset_template_detail.html",
        {"template": template, "nodes": nodes, "activities": activities},
    )


def _starter_definition():
    """A blank ruleset starts with an explicit roster and no guessed contest stage."""
    return json.dumps(
        {
            "schema_version": 1,
            "nodes": [{"key": "roster", "type": "ROSTER"}],
        },
        ensure_ascii=False,
    )


def _unique_key(nodes, base):
    used = {n["key"] for n in nodes}
    candidate = base
    counter = 1
    while candidate in used:
        candidate = f"{base}_{counter}"
        counter += 1
    return candidate


def _business_stages_from_post(post):
    """Parse the repeatable business-stage form into the service payload."""
    try:
        stage_count = int(post.get("stage_count", "0"))
    except (TypeError, ValueError) as exc:
        raise ValidationError("比赛阶段数量无效。") from exc
    if stage_count < 0 or stage_count > 50:
        raise ValidationError("比赛阶段数量必须在 0 到 50 之间。")
    stages = []
    for index in range(stage_count):
        prefix = f"stage_{index}"
        try:
            source_count = int(post.get(f"{prefix}_judge_source_count", "1"))
        except (TypeError, ValueError) as exc:
            raise ValidationError(f"第 {index + 1} 个阶段的评委来源数量无效。") from exc
        if source_count <= 0 or source_count > 20:
            raise ValidationError(f"第 {index + 1} 个阶段的评委来源数量无效。")
        judge_sources = []
        for source_index in range(source_count):
            source_prefix = f"{prefix}_judge_{source_index}"
            judge_sources.append(
                {
                    "round_key": (post.get(f"{source_prefix}_round_key") or "").strip(),
                    "round_name": (post.get(f"{source_prefix}_round_name") or "").strip(),
                    "round_id": post.get(f"{source_prefix}_round_id") or "",
                    "weight": (post.get(f"{source_prefix}_weight") or "0").strip(),
                    "create_round": post.get(f"{source_prefix}_create_round") == "on",
                    "source": (post.get(f"{source_prefix}_source") or "").strip(),
                    "source_stage_key": (
                        post.get(f"{source_prefix}_source_stage_key") or ""
                    ).strip(),
                    "source_stage_index": post.get(f"{source_prefix}_source_stage_index") or "",
                    "scoring_mode_conflict_ack": post.get(
                        f"{source_prefix}_scoring_mode_conflict_ack"
                    )
                    == "on",
                }
            )
        try:
            judge_weight = str(
                sum(
                    (Decimal(source["weight"] or "0") for source in judge_sources),
                    Decimal("0"),
                )
            )
        except (InvalidOperation, TypeError, ValueError) as exc:
            raise ValidationError(f"第 {index + 1} 个阶段的评委权重无效。") from exc
        stages.append(
            {
                "key": (post.get(f"{prefix}_key") or f"stage{index + 1}").strip(),
                "name": (post.get(f"{prefix}_name") or "").strip(),
                "round_key": judge_sources[0]["round_key"],
                "round_name": judge_sources[0]["round_name"],
                "round_id": judge_sources[0]["round_id"],
                "advance_count": post.get(f"{prefix}_advance_count") or "",
                "judge_sources": judge_sources,
                "judge_weight": (post.get(f"{prefix}_judge_weight") or "").strip() or judge_weight,
                "previous_weight": (post.get(f"{prefix}_previous_weight") or "0").strip(),
                "audience_weight": (post.get(f"{prefix}_audience_weight") or "0").strip(),
                "vote_source": (post.get(f"{prefix}_vote_source") or "").strip(),
                "audience_source": (post.get(f"{prefix}_audience_source") or "").strip(),
                "audience_source_stage_key": (
                    post.get(f"{prefix}_audience_source_stage_key") or ""
                ).strip(),
                "audience_source_stage_index": post.get(f"{prefix}_audience_source_stage_index")
                or "",
                "create_round": any(source["create_round"] for source in judge_sources),
                "scoring_mode": post.get(f"{prefix}_scoring_mode") or "average",
                "scoring_mode_conflict_ack": any(
                    source["scoring_mode_conflict_ack"] for source in judge_sources
                ),
            }
        )
    return stages


def _last_node_of(output_types, nodes):
    for node in reversed(nodes):
        if NODE_TYPE_SPEC[node["type"]].output_type in output_types:
            return node["key"]
    return ENTRY_KEY


def _default_node(new_type, nodes):
    best_spec = NODE_TYPE_SPEC.get(new_type)
    if best_spec is None:
        raise ValueError(f"unknown node type {new_type!r}")
    node = {"key": _unique_key(nodes, new_type.lower()), "type": new_type}
    expects = (best_spec.expects or {}).get("source")
    if expects and OutputType.SCOREMAP in expects:
        source = _last_node_of({OutputType.SCOREMAP, OutputType.RANKED_ROSTER}, nodes)
    elif expects and OutputType.RANKED_ROSTER in expects:
        source = _last_node_of({OutputType.RANKED_ROSTER}, nodes)
    else:
        source = ENTRY_KEY
    for field in best_spec.required:
        if field == "source":
            node["source"] = source
        elif field == "count":
            node["count"] = 1
        elif field == "sources":
            node["sources"] = [ENTRY_KEY]
        elif field == "by":
            node["by"] = "组1"
        elif field == "aggregate":
            src = _last_node_of({OutputType.SCOREMAP}, nodes)
            node["aggregate"] = {
                "type": "weighted_sum",
                "components": [{"source": src or ENTRY_KEY, "weight": 1.0}],
            }
        elif field == "branches":
            node["branches"] = [{"when": "default", "into": ENTRY_KEY}]
        elif field == "minuend":
            node["minuend"] = ENTRY_KEY
        elif field == "subtrahend":
            node["subtrahend"] = ENTRY_KEY
        elif field == "quota":
            node["quota"] = 1
        elif field == "groups":
            node["groups"] = 1
        elif field == "pairing_policy":
            node["pairing_policy"] = "ADJACENT"
        elif field == "from":
            node["from"] = ENTRY_KEY
        elif field == "into":
            node["into"] = _last_node_of({OutputType.GROUP_MAP}, nodes)
        elif field == "award":
            node["award"] = "默认奖项"
        elif field == "decision_source":
            node["decision_source"] = "manual_recorded_result"
    if new_type == "ASSESS":
        node["round"] = f"r{len(nodes) + 1}"
    return node


def _humanize_ruleset_editor_error(error):
    """Turn schema/compiler vocabulary into an operator-facing message."""
    raw = " ".join(getattr(error, "messages", None) or [str(error)])
    import re

    forward = re.search(
        r"Node (?P<node>[^:]+): source '(?P<source>[^']+)' is not defined before this node",
        raw,
    )
    if forward:
        return (
            f"节点“{forward.group('node')}”依赖“{forward.group('source')}”，"
            "当前顺序或引用不合法，请先调整依赖关系。"
        )
    missing = re.search(
        r"Node (?P<node>[^:]+): source '(?P<source>[^']+)' is not defined",
        raw,
    )
    if missing:
        return (
            f"节点“{missing.group('node')}”仍引用不存在的“{missing.group('source')}”，"
            "请先修改引用或删除下游节点。"
        )
    if "unknown node type" in raw:
        return "这个节点类型暂不支持，请从可用节点类型中选择。"
    if "duplicate node key" in raw:
        return "节点键不能重复，请修改节点名称。"
    if raw.startswith(("当前", "该", "找不到", "这个节点")):
        return raw
    return "赛制定义未通过校验，请检查节点顺序、来源和字段填写。"


_NODE_TYPE_LABELS = {
    "ROSTER": "报名名单",
    "ASSESS": "评分",
    "RANK": "排名",
    "SELECT": "选出名额",
    "PARTITION": "分组",
    "MERGE": "合并名单",
    "SUBTRACT": "排除名单",
    "PAIR": "配对",
    "DUEL": "对决",
    "AGGREGATE": "综合成绩",
    "MANUAL_SELECT": "人工选择",
    "FILL_TO_QUOTA": "补足名额",
    "AWARD": "奖项",
}


def _try_parse_nodes(nodes):
    try:
        parse_definition({"nodes": nodes})
    except ValidationError:
        return False
    return True


def _move_node_to_nearest_valid_slot(nodes, key, delta):
    index = next((i for i, n in enumerate(nodes) if n["key"] == key), None)
    if index is None:
        return list(nodes), index
    targets = range(index - 1, -1, -1) if delta < 0 else range(index + 1, len(nodes))
    for target in targets:
        candidate = list(nodes)
        item = candidate.pop(index)
        candidate.insert(target, item)
        if _try_parse_nodes(candidate):
            return candidate, target
    return list(nodes), index


def _edit_nodes(post, nodes) -> list[dict]:
    """Return the new ``nodes`` list for one editor action.

    Only the node list: the caller patches ``definition.nodes`` as a single section via
    ``update_ruleset_definition_section``, so the root's other sections survive. Returning
    a rebuilt ``{"schema_version", "nodes"}`` document here is what used to delete
    ``checkpoints`` / ``context`` / ``questionnaire`` on every save.
    """
    action = post.get("action")
    if action == "add":
        new_type = post.get("new_type", "ASSESS")
        if new_type == "BRANCH":
            raise ValueError("这个节点类型暂不支持，请从可用节点类型中选择。")
        candidate = nodes + [_default_node(new_type, nodes)]
        if not _try_parse_nodes(candidate):
            raise ValueError(f"当前节点结构不能新增“{new_type}”，请先补齐可用的上游节点。")
        nodes = candidate
    elif action == "delete":
        key = post.get("key")
        candidate = [n for n in nodes if n["key"] != key]
        if len(candidate) == len(nodes):
            raise ValueError("找不到要删除的节点，请刷新页面后重试。")
        if not _try_parse_nodes(candidate):
            raise ValueError("该节点仍被其他节点引用，请先调整下游节点的来源。")
        nodes = candidate
    elif action == "move_up":
        original_index = next((i for i, n in enumerate(nodes) if n["key"] == post.get("key")), None)
        nodes, target = _move_node_to_nearest_valid_slot(nodes, post.get("key"), -1)
        if target == original_index:
            raise ValueError("该节点不能再向上移动：它的上游依赖必须保持在前面。")
    elif action == "move_down":
        original_index = next((i for i, n in enumerate(nodes) if n["key"] == post.get("key")), None)
        nodes, target = _move_node_to_nearest_valid_slot(nodes, post.get("key"), +1)
        if target == original_index:
            raise ValueError("该节点不能再向下移动：后续节点仍依赖它。")
    elif action == "save":
        return ruleset_editor.nodes_from_form(post)
    else:
        raise ValueError(f"unknown action {action!r}")
    return nodes


def _source_options(nodes, index, field, current=None):
    spec = NODE_TYPE_SPEC.get(nodes[index].get("type"))
    expected = (spec.expects or {}).get(field) if spec else None
    options = []
    if expected and OutputType.ROSTER in expected:
        options.append(ENTRY_KEY)
    for prior in nodes[:index]:
        prior_spec = NODE_TYPE_SPEC.get(prior.get("type"))
        if prior_spec is None or (expected and prior_spec.output_type not in expected):
            continue
        options.append(prior["key"])
    if current and current not in options:
        options.append(current)
    return options


def _display_options(options, labels=None, nodes=None):
    labels = labels or {}
    node_labels = {
        node.get("key"): (
            f"{node.get('key')}（{_NODE_TYPE_LABELS.get(node.get('type'), node.get('type'))}）"
        )
        for node in (nodes or [])
        if node.get("key")
    }
    return [
        {
            "value": option,
            "label": labels.get(option)
            or ("报名选手" if option == ENTRY_KEY else node_labels.get(option, option)),
        }
        for option in options
    ]


def _selected_binding_labels(rows):
    """Return the operator-facing label for each currently selected binding."""
    labels = {}
    for row in rows:
        value = str(row.get("value") or "")
        if not value:
            continue
        label = next(
            (label for option_value, label in row.get("options", []) if str(option_value) == value),
            "",
        )
        if label:
            labels[row["key"]] = label
    return labels


def _node_field_entries(node, index, nodes, *, binding_options=None):
    """Render data-driven edit fields for one node card (select/text/checkbox/json/aggregate)."""
    labels = ruleset_editor.field_labels()
    help_text = ruleset_editor.field_help()
    binding_options = binding_options or {}
    allowed = ruleset_editor.field_allowed()
    option_labels = ruleset_editor.field_option_labels()
    spec = NODE_TYPE_SPEC.get(node.get("type"))
    fields: list[dict] = []
    if spec is None:
        return fields
    reference_fields = {
        "source",
        "within",
        "minuend",
        "subtrahend",
        "from",
        "into",
        "by",
        "ranking_source",
        "tie_break_source",
    }
    for field in list(spec.required) + list(spec.optional):
        label = labels.get(field, field)
        base = f"node_{index}_{field}"
        value = node.get(field)
        if field == "aggregate":
            aggregate = node.get("aggregate", {})
            fields.append(
                {
                    "kind": "aggregate",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "index": index,
                    "aggregate_type": aggregate.get("type", "weighted_sum"),
                    "components": [
                        {
                            "source": c.get("source", ""),
                            "weight": c.get("weight", 1.0),
                            "options": _source_options(
                                nodes, index, "aggregate.components[].source", c.get("source")
                            ),
                            "display_options": _display_options(
                                _source_options(
                                    nodes, index, "aggregate.components[].source", c.get("source")
                                ),
                                nodes=nodes,
                            ),
                        }
                        for c in aggregate.get("components", [])
                    ],
                }
            )
        elif field == "within":
            fields.append(
                {
                    "kind": "select",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": base,
                    "value": value or "",
                    "options": _source_options(nodes, index, field, value),
                    "display_options": _display_options(
                        _source_options(nodes, index, field, value), nodes=nodes
                    ),
                }
            )
        elif field in ("branches", "conversion"):
            fields.append(
                {
                    "kind": "json",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": f"{base}_json",
                    "value": json.dumps(value, ensure_ascii=False) if value else "",
                }
            )
        elif isinstance(value, bool):
            fields.append(
                {
                    "kind": "checkbox",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": base,
                    "value": value,
                }
            )
        elif field in binding_options:
            options = list(binding_options[field])
            if value and value not in options:
                options.append(value)
            fields.append(
                {
                    "kind": "select",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": base,
                    "value": value if value is not None else "",
                    "options": options,
                    "display_options": _display_options(
                        options,
                        labels={option: f"{option}（活动绑定）" for option in options},
                    ),
                }
            )
        elif allowed.get(field):
            fields.append(
                {
                    "kind": "select",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": base,
                    "value": value if value is not None else "",
                    "options": allowed[field],
                    "display_options": _display_options(
                        allowed[field], labels=option_labels.get(field)
                    ),
                }
            )
        elif field in reference_fields and (field != "by" or field in (spec.expects or {})):
            fields.append(
                {
                    "kind": "select",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": base,
                    "value": value or "",
                    "options": _source_options(nodes, index, field, value),
                    "display_options": _display_options(
                        _source_options(nodes, index, field, value), nodes=nodes
                    ),
                }
            )
        elif field == "sources":
            fields.append(
                {
                    "kind": "multi_select",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": base,
                    "value": value or [],
                    "options": _source_options(nodes, index, field),
                    "display_options": _display_options(
                        _source_options(nodes, index, field), nodes=nodes
                    ),
                }
            )
        else:
            fields.append(
                {
                    "kind": "text",
                    "label": label,
                    "help": help_text.get(field, ""),
                    "name": base,
                    "value": value if value is not None else "",
                }
            )
    return fields


def _build_card(node, index, nodes, *, binding_options=None):
    sources = ruleset_editor.available_sources(nodes, index)
    _up_preview, up_target = _move_node_to_nearest_valid_slot(nodes, node["key"], -1)
    _down_preview, down_target = _move_node_to_nearest_valid_slot(nodes, node["key"], 1)
    delete_candidate = [item for item in nodes if item["key"] != node["key"]]
    delete_allowed = len(delete_candidate) != len(nodes) and _try_parse_nodes(delete_candidate)
    return {
        "node": node,
        "index": index,
        "type": node.get("type"),
        "type_label": _NODE_TYPE_LABELS.get(node.get("type"), node.get("type")),
        "sources": sources,
        "fields": _node_field_entries(node, index, nodes, binding_options=binding_options),
        "can_move_up": up_target != index,
        "can_move_down": down_target != index,
        "can_delete": delete_allowed,
        "move_up_reason": "该节点的上游依赖必须保持在前面。",
        "move_down_reason": "后续节点仍依赖该节点。",
        "delete_reason": "该节点仍被其他节点引用。",
    }


def _node_type_options(nodes):
    options = []
    for node_type in sorted(t for t in NODE_TYPE_SPEC if t != "BRANCH"):
        try:
            candidate = nodes + [_default_node(node_type, nodes)]
            if not _try_parse_nodes(candidate):
                raise ValueError
        except (ValueError, ValidationError):
            options.append(
                {
                    "value": node_type,
                    "label": _NODE_TYPE_LABELS.get(node_type, node_type),
                    "enabled": False,
                    "reason": "当前缺少可用的上游节点。",
                }
            )
        else:
            options.append(
                {
                    "value": node_type,
                    "label": _NODE_TYPE_LABELS.get(node_type, node_type),
                    "enabled": True,
                    "reason": "",
                }
            )
    return options


def _planning_result_json(result: PlanningCompileResult) -> dict[str, object]:
    return {
        "can_apply": result.safe_to_apply,
        "business_complete": result.business_complete,
        "binding_ready": result.binding_ready,
        "freeze_ready": result.freeze_ready,
        "issues": [issue.to_dict() for issue in result.issues],
        "business_stages": list(result.business_stages),
        "change_summary": list(result.change_summary),
    }


def _planning_request_payload(request):
    try:
        body = json.loads(request.body.decode("utf-8") or "{}")
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValidationError("规划数据不是合法 JSON。") from exc
    if not isinstance(body, dict):
        raise ValidationError("规划数据必须是 JSON 对象。")
    payload = body.get("payload")
    if not isinstance(payload, dict):
        raise ValidationError("规划草稿缺少有效内容。")
    raw_revision = body.get("revision")
    if isinstance(raw_revision, bool) or not isinstance(raw_revision, (int, str)):
        raise ValidationError("规划草稿版本号无效。")
    try:
        revision = int(raw_revision)
    except ValueError as exc:
        raise ValidationError("规划草稿版本号无效。") from exc
    if revision < 0:
        raise ValidationError("规划草稿版本号无效。")
    return payload, revision


def _planning_error_response(error, *, payload=None):
    return JsonResponse(
        {
            "ok": False,
            "error": safe_domain_error_messages(error),
            "payload": payload,
        },
        status=400,
        json_dumps_params={"ensure_ascii": False},
    )


@staff_required
@require_http_methods(["GET", "POST"])
def ruleset_planning(request, pk):
    version = get_object_or_404(RulesetVersion.objects.select_related("ruleset__activity"), pk=pk)
    if request.method == "POST":
        raw_payload = None
        try:
            raw_payload, revision = _planning_request_payload(request)
            draft = save_planning_draft(
                version,
                payload=raw_payload,
                operator=request.user,
                expected_revision=revision,
            )
        except (PermissionDenied, ValidationError) as exc:
            response = _planning_error_response(exc, payload=raw_payload)
            if isinstance(exc, PermissionDenied):
                response.status_code = 403
            return response
        return JsonResponse({"ok": True, "revision": draft.revision, "payload": draft.payload})

    display_payload, planning_revision = read_planning_draft(version)
    projection = planning_projection_from_definition(version.definition)
    if not display_payload.get("stages") and projection.editable:
        display_payload = projection.payload
    compile_result = None
    if display_payload.get("stages"):
        compile_result = prepare_planning_preflight(version, display_payload)
    return render(
        request,
        "staff_panel/ruleset_planning.html",
        {
            "activity": version.ruleset.activity,
            "ruleset": version.ruleset,
            "version": version,
            "planning_payload": display_payload,
            "planning_revision": planning_revision,
            "planning_projection": projection,
            "planning_result": compile_result,
            "planning_result_json": _planning_result_json(compile_result)
            if compile_result is not None
            else {
                "can_apply": False,
                "business_complete": False,
                "binding_ready": False,
                "freeze_ready": False,
                "issues": [],
                "change_summary": [],
            },
            # Editable means the *newest* draft, not merely a draft: a version a newer one has
            # replaced is history, and the page used to offer it as an editing target anyway.
            "planning_read_only": projection.read_only or not version_is_editable(version),
        },
    )


@staff_required
@require_http_methods(["GET", "POST"])
def ruleset_planning_preview(request, pk):
    version = get_object_or_404(RulesetVersion, pk=pk)
    try:
        if request.method == "POST":
            payload, _revision = _planning_request_payload(request)
            result = prepare_planning_preflight(version, payload)
        else:
            result = prepare_planning_compile_for_version(version)
    except (PermissionDenied, ValidationError) as exc:
        response = _planning_error_response(exc)
        if isinstance(exc, PermissionDenied):
            response.status_code = 403
        return response
    return JsonResponse({"ok": True, "result": _planning_result_json(result)})


@staff_required
@require_POST
def ruleset_planning_apply(request, pk):
    version = get_object_or_404(RulesetVersion, pk=pk)
    try:
        _payload, revision = _planning_request_payload(request)
        saved = apply_planning_to_ruleset(
            version,
            operator=request.user,
            expected_revision=revision,
        )
    except (PermissionDenied, ValidationError) as exc:
        response = _planning_error_response(exc)
        if isinstance(exc, PermissionDenied):
            response.status_code = 403
        return response
    return JsonResponse({"ok": True, "redirect": reverse("staff:ruleset_edit", args=[saved.pk])})


def _provision_builtin_live_format(request, *, template, activity, name):
    """Provision the 2026 built-in format, or return ``None`` for an ordinary template.

    The system built-in behaves exactly like any other ruleset creation — same page, same
    editor, everything still editable — except that the first round and the format's fixed
    binding keys arrive already filled in.
    """

    if not is_live_2026_template(template):
        return None
    try:
        version = provision_live_2026_ruleset(
            activity, template=template, name=name, operator=request.user
        )
    except ValidationError as error:
        # Provisioning re-validates this activity's whole binding surface, so an existing
        # stale reference (a deactivated judge, a deleted round) surfaces here. The binding
        # editor reports those as messages; this entry point must not answer with a 500.
        messages.error(request, f"无法建立内置赛制：{domain_error_messages(error)}")
        return redirect("staff:contest_ruleset_create")
    messages.success(
        request,
        f"已建立「{template.name}」：第一轮与赛制形态已预填，请补充两套权重、三轮评委与评分表。",
    )
    # Neither the planning workbench (its graph has no AGGREGATE node, so it renders
    # read-only) nor the advanced editor is where this work is done: the 2026 page exists
    # for exactly what provisioning leaves open, and it links on to the editor from there.
    return redirect("staff:ruleset_live_2026_config", pk=version.pk)


@staff_required
@transaction.atomic
def contest_ruleset_create(request):
    activities = _editable_singer_activities()
    if request.method == "POST":
        # Keep the candidate lookup broad enough for the authoritative lock check to
        # produce its stable 403 response.  Filtering ``is_locked=False`` here would
        # turn a valid mutation attempt against a locked activity into a misleading
        # 404 before ``lock_activity_for_action`` can re-validate it.
        activity = get_object_or_404(
            Activity,
            pk=request.POST.get("activity"),
            activity_type=Activity.Type.SINGER_CONTEST,
        )
        activity = lock_activity_for_action(activity)
        # The clone entry point has always refused an archived activity; this one now creates
        # a physical first round as well, so it has to answer the same question the same way.
        _ensure_activity_mutable(activity)
        name = (request.POST.get("name") or "").strip()
        if not name:
            messages.error(request, "请填写赛制名称。")
            return redirect("staff:contest_ruleset_create")
        template = None
        template_pk = request.POST.get("template")
        if template_pk:
            # The page only lists production templates, but a POST is not bound by what
            # the page rendered. `staff:ruleset_clone_from_template` already filters on
            # capability status; this second clone entry point has to apply the same
            # gate, or an experimental/unsupported template can be cloned into a live
            # workflow by picking its pk directly.
            template = get_object_or_404(
                RulesetTemplate,
                pk=template_pk,
                capability_status=RulesetTemplate.CapabilityStatus.PRODUCTION,
            )
        provisioned = _provision_builtin_live_format(
            request, template=template, activity=activity, name=name
        )
        if provisioned is not None:
            return provisioned
        # One activity owns one ContestRuleset (versions live on RulesetVersion). If a
        # ruleset already exists, reuse it instead of creating a competing authority that
        # would make recompute_activity_result's single-authority assumption ambiguous.
        existing = ContestRuleset.objects.filter(activity=activity).first()
        if existing is not None:
            if template is not None:
                existing.source_template = template
                existing.name = name or existing.name
                existing.is_test_data = runtime_is_test(activity)
                existing.save()
                version = create_ruleset_version(
                    existing, definition=template.definition, created_by=request.user
                )
                messages.success(request, f"已将「{template.name}」克隆到该活动，进入准备。")
            else:
                # Reuse the single authority: continue on the newest DRAFT, or supersede a
                # frozen current into a fresh editing DRAFT — never redirect into a frozen
                # version, and never revive a draft an older version has moved past.
                draft = editable_ruleset_version(existing)
                if draft is None:
                    frozen = (
                        existing.versions.filter(status=RulesetVersion.Status.FROZEN)
                        .order_by("-version")
                        .first()
                    )
                    version = (
                        supersede_ruleset_version(frozen, created_by=request.user)
                        if frozen is not None
                        else create_ruleset_version(
                            existing,
                            definition=_starter_definition(),
                            created_by=request.user,
                        )
                    )
                else:
                    version = draft
                messages.info(request, "该活动已有赛制，进入现有赛制编辑。")
            # The planning workbench cannot express the 2026 format's graph, so it renders
            # read-only there; sending the operator to the configuration page instead is the
            # difference between "you cannot edit this here" and "here is where you edit".
            if _is_live_2026_ruleset(version.ruleset):
                return redirect("staff:ruleset_live_2026_config", pk=version.pk)
            return redirect("staff:ruleset_planning", pk=version.pk)
        definition = template.definition if template else _starter_definition()
        ruleset = ContestRuleset.objects.create(
            activity=activity,
            name=name,
            source_template=template,
            is_test_data=runtime_is_test(activity),
            created_by=request.user,
        )
        version = create_ruleset_version(ruleset, definition=definition, created_by=request.user)
        messages.success(request, "赛制已创建，进入准备。")
        return redirect("staff:ruleset_planning", pk=version.pk)
    templates = RulesetTemplate.objects.filter(
        capability_status=RulesetTemplate.CapabilityStatus.PRODUCTION
    ).order_by("name")
    requested_activity_id = request.GET.get("activity") or None
    selected_activity = (
        activities.filter(pk=requested_activity_id).first()
        if requested_activity_id is not None
        else activities.first()
    )
    if selected_activity is None:
        selected_activity = activities.first()
    selected_activity_id = str(selected_activity.pk) if selected_activity else ""
    return render(
        request,
        "staff_panel/ruleset_create.html",
        {
            "activities": activities,
            "templates": templates,
            "selected_activity_id": selected_activity_id,
        },
    )


@staff_required
@transaction.atomic
def ruleset_edit(request, pk):
    version = get_object_or_404(RulesetVersion.objects.select_related("ruleset__activity"), pk=pk)
    # The generic editor is the business editor and the raw-definition editor at once, so it
    # is the other door onto the state the 2026 pages guard: a version with anything newer
    # beside it is history, whatever its status.
    stale_response = _refuse_frozen_or_superseded(request, version)
    if stale_response is not None:
        return stale_response
    if request.method == "POST":
        if request.POST.get("action") == "business_save":
            try:
                business_stages = _business_stages_from_post(request.POST)
                save_business_ruleset(
                    version,
                    stages=business_stages,
                    operator=request.user,
                    base_content_hash=request.POST.get("base_content_hash") or None,
                    base_binding=request.POST.get("base_binding") or None,
                    base_round_configuration=request.POST.get("base_round_configuration") or None,
                )
            except PermissionDenied as exc:
                messages.error(request, domain_error_messages(exc))
            except ValidationError as exc:
                messages.error(request, domain_error_messages(exc))
            else:
                messages.success(request, "赛制已保存。")
            return redirect("staff:ruleset_edit", pk=pk)
        try:
            current = ruleset_editor.definition_nodes(version.definition)
            nodes = _edit_nodes(request.POST, current)
            parse_definition({"nodes": nodes})
        except (ValueError, ValidationError) as exc:
            messages.error(request, f"保存失败：{_humanize_ruleset_editor_error(exc)}")
            return redirect("staff:ruleset_edit", pk=pk)
        try:
            update_ruleset_definition_section(
                version,
                section="nodes",
                value=nodes,
                operator=request.user,
                base_content_hash=request.POST.get("base_content_hash") or None,
            )
        except ValidationError as exc:
            messages.error(request, domain_error_messages(exc))
            return redirect("staff:ruleset_edit", pk=pk)
        messages.success(request, "赛制已保存。")
        return redirect("staff:ruleset_edit", pk=pk)
    try:
        nodes = parse_definition(version.definition)["nodes"]
    except ValidationError:
        nodes = []
    ruleset = version.ruleset
    binding_options = {
        "round": list((ruleset.round_keys or {}).keys()),
        "vote_source": list((ruleset.vote_keys or {}).keys()),
    }
    cards = [
        _build_card(node, i, nodes, binding_options=binding_options) for i, node in enumerate(nodes)
    ]
    current_round_keys = {str(node.get("round")) for node in nodes if node.get("round")}
    current_vote_keys = {str(node.get("vote_source")) for node in nodes if node.get("vote_source")}
    current_group_keys = {
        str(node.get("by")) for node in nodes if node.get("type") == "PARTITION" and node.get("by")
    }
    round_candidates = list(
        ContestRound.objects.filter(activity=ruleset.activity).order_by("sequence", "pk")
    )
    vote_candidates = list(
        VoteSession.objects.filter(activity=ruleset.activity).order_by("-created_at", "pk")
    )
    scoring_rule_candidates = list(
        VoteScoringRule.objects.filter(vote_session__activity=ruleset.activity)
        .select_related("vote_session")
        .order_by("vote_session_id")
    )
    round_binding = ruleset.round_keys or {}
    vote_binding = ruleset.vote_keys or {}
    scoring_binding = ruleset.vote_scoring_rule_keys or {}
    group_binding = ruleset.group_keys or {}
    group_stage_binding = ruleset.group_stage_keys or {}
    group_stage_candidates = list(
        GroupStage.objects.filter(
            activity=ruleset.activity,
            status=GroupStage.Status.FROZEN,
        ).order_by("stage_key", "pk")
    )
    binding_maps = {
        "rounds": [
            {
                "key": key,
                "value": str(round_binding.get(key, "")),
                "options": [
                    (
                        str(item.pk),
                        (
                            f"{item.name or item.get_round_type_display()}"
                            f"（第 {item.sequence or '-'} 轮）"
                        ),
                    )
                    for item in round_candidates
                ],
            }
            for key in sorted(current_round_keys | set(round_binding))
        ],
        "votes": [
            {
                "key": key,
                "value": str(vote_binding.get(key, "")),
                "options": [(str(item.pk), item.name) for item in vote_candidates],
            }
            for key in sorted(current_vote_keys | set(vote_binding))
        ],
        "scoring_rules": [
            {
                "key": key,
                "value": str(scoring_binding.get(key, "")),
                "options": [
                    (str(item.pk), f"{item.vote_session.name}（{item.get_mode_display()}）")
                    for item in scoring_rule_candidates
                    if item.vote_session_id == vote_binding.get(key)
                    or str(item.vote_session_id) == str(vote_binding.get(key, ""))
                ],
            }
            for key in sorted(set(scoring_binding) | current_vote_keys)
        ],
        "groups": [
            {
                "key": key,
                "value": str(group_binding.get(key, "")),
                "options": [
                    (
                        str(item.pk),
                        f"{item.name or item.get_round_type_display()}（分组来源）",
                    )
                    for item in round_candidates
                ],
            }
            for key in sorted(current_group_keys | set(group_binding))
        ],
        "group_stages": [
            {
                "key": key,
                "value": str(group_stage_binding.get(key, "")),
                "options": [
                    (str(item.pk), f"{item.name}（{item.stage_key}）")
                    for item in group_stage_candidates
                ],
            }
            for key in sorted(current_group_keys | set(group_stage_binding))
        ],
    }
    try:
        human_summary = ruleset_editor.human_ruleset_summary(
            version.definition,
            round_labels=_selected_binding_labels(binding_maps["rounds"]),
            vote_labels=_selected_binding_labels(binding_maps["votes"]),
        )
    except (ValueError, ValidationError):
        human_summary = None
    if human_summary:
        node_titles = human_summary.get("node_titles", {})
        for card in cards:
            card["display_label"] = node_titles.get(card["node"]["key"], card["type_label"])
        for bucket in ("rounds", "votes"):
            for row in binding_maps[bucket]:
                row["usage_label"] = next(
                    (
                        component["label"]
                        for stage in human_summary.get("stages", [])
                        for component in stage.get("components", [])
                        if component.get("binding_key") == row["key"]
                    ),
                    "",
                )
    business_editor = infer_business_stages(version.definition)
    business_editor["content_hash"] = version.content_hash
    business_editor["binding_signature"] = _binding_signature(ruleset)
    business_editor["existing_round_keys"] = sorted(str(key) for key in (ruleset.round_keys or {}))
    business_editor["round_configuration_signature"] = business_round_configuration_signature(
        round_candidates
    )
    if business_editor.get("editable"):
        prepared_round_options = [
            (
                str(item.pk),
                f"{item.name or item.get_round_type_display()}（{item.get_status_display()}）",
            )
            for item in round_candidates
        ]
        vote_by_id = {item.pk: item for item in vote_candidates}
        scoring_mode_labels = dict(ContestRound.ScoringMode.choices)
        vote_options = [
            (
                str(key),
                f"{vote_by_id[int(value)].name}（{vote_by_id[int(value)].get_purpose_display()}）"
                if str(value).isdigit() and int(value) in vote_by_id
                else f"{key}（暂不绑定）",
            )
            for key, value in (ruleset.vote_keys or {}).items()
        ]
        for stage in business_editor.get("stages", []):
            stage["previous_weight"] = stage.get("previous_weight", "0")
            stage["audience_weight"] = stage.get("audience_weight", "0")
            stage["vote_options"] = vote_options
            stage["scoring_modes"] = list(ContestRound.ScoringMode.choices)
            for source in stage.get("judge_sources", []):
                source["round_id"] = str(ruleset.round_keys.get(source["round_key"], ""))
                source["round_options"] = prepared_round_options
                source["create_round"] = not source["round_id"]
                bound_round = next(
                    (item for item in round_candidates if str(item.pk) == source["round_id"]),
                    None,
                )
                source["bound_round_scoring_mode"] = (
                    bound_round.scoring_mode if bound_round is not None else ""
                )
                source["bound_round_scoring_mode_label"] = (
                    scoring_mode_labels.get(bound_round.scoring_mode, bound_round.scoring_mode)
                    if bound_round is not None
                    else ""
                )
                source["scoring_mode_conflict"] = bool(
                    bound_round is not None and bound_round.scoring_mode != source["scoring_mode"]
                )
        business_editor["round_options"] = prepared_round_options
        business_editor["vote_options"] = vote_options
        business_editor["scoring_modes"] = list(ContestRound.ScoringMode.choices)
    live_challenge = dict(ruleset.live_challenge or {})
    for weight_key in (
        "first_round_weight",
        "challenge_round_weight",
        "podium_first_round_weight",
        "podium_third_round_weight",
    ):
        if weight_key in live_challenge:
            try:
                live_challenge[weight_key] = f"{Decimal(str(live_challenge[weight_key])) * 100:g}%"
            except (InvalidOperation, ValueError):
                pass
    live_judges = list(
        Judge.objects.filter(activity=ruleset.activity, is_active=True).order_by("name", "pk")
    )
    return render(
        request,
        "staff_panel/ruleset_editor.html",
        {
            "version": version,
            "ruleset": ruleset,
            "cards": cards,
            # §28 capability matrix: BRANCH remains unsupported at runtime and is
            # not offered. Other types are shown but disabled when their default
            # configuration cannot be valid with the current upstream graph.
            "node_types": sorted(t for t in NODE_TYPE_SPEC if t != "BRANCH"),
            "node_type_options": _node_type_options(nodes),
            "binding_json": {
                "round_keys": json.dumps(ruleset.round_keys or {}, ensure_ascii=False),
                "vote_keys": json.dumps(ruleset.vote_keys or {}, ensure_ascii=False),
                "vote_scoring_rule_keys": json.dumps(
                    ruleset.vote_scoring_rule_keys or {}, ensure_ascii=False
                ),
                "group_keys": json.dumps(ruleset.group_keys or {}, ensure_ascii=False),
                "group_stage_keys": json.dumps(ruleset.group_stage_keys or {}, ensure_ascii=False),
                "audience_keys": json.dumps(ruleset.audience_keys or {}, ensure_ascii=False),
                "live_challenge": json.dumps(ruleset.live_challenge or {}, ensure_ascii=False),
                "announcement_blocks": json.dumps(
                    ruleset.announcement_blocks or [], ensure_ascii=False
                ),
                "announcement_blocks_by_checkpoint": json.dumps(
                    ruleset.announcement_blocks_by_checkpoint or {},
                    ensure_ascii=False,
                ),
            },
            "live_challenge": live_challenge,
            "live_judges": live_judges,
            "binding_signature": _binding_signature(ruleset),
            "binding_maps": binding_maps,
            "human_summary": human_summary,
            "can_freeze": request.user.is_admin,
            "business_editor": business_editor,
        },
    )


def _refuse_frozen_or_superseded(request, version, *, route: str | None = None):
    """Return a response when ``version`` must not be mutated, else ``None``.

    A frozen version is read-only by definition. So is any version with a newer one beside it:
    every version of a ruleset shares one ``ContestRuleset`` binding surface, so two live
    pages both claim the same state — save 8 seats from v1 and v2's definition still says 10,
    a disagreement nobody created, reported later by the freeze contract.

    The newer version decides what the operator should see. A newer *draft* is where they
    continue, so a page view hands over to it (``route`` names the page to hand over on). A
    newer *frozen* authority has no edit page at all, so there is nothing to hand over to and
    the workspace — which shows the real state and offers a successor draft — is the answer.
    """

    if version.status == RulesetVersion.Status.FROZEN:
        raise PermissionDenied("已冻结赛制版本不可编辑。")
    newer = newer_ruleset_version(version)
    if newer is None:
        return None
    if request.method == "POST":
        raise PermissionDenied(
            f"此版本已被 v{newer.version} 替代，请在最新版本上修改。"
            if newer.status != RulesetVersion.Status.FROZEN
            else f"此版本已被 v{newer.version} 正式冻结取代，不能再编辑；"
            "如需调整，请从工作区创建后继草稿。"
        )
    if route is not None and newer.status == RulesetVersion.Status.DRAFT:
        messages.info(request, "此版本已被更新的草稿替代，已转到最新草稿。")
        return redirect(route, pk=newer.pk)
    messages.info(request, f"此版本已被 v{newer.version} 替代，已转到该活动的工作区。")
    return redirect("staff:activity_workspace", pk=version.ruleset.activity_id)


@staff_required
@require_POST
@transaction.atomic
def ruleset_bind(request, pk):
    """Save the activity-owning binding (stage/round/vote/group/announcement) to a ruleset (§32).

    Only the editable input surface (ContestRuleset) is written here; the snapshot runs at
    freeze time via ``_snapshot_binding``.
    """
    version = get_object_or_404(RulesetVersion, pk=pk)
    lock_activity_for_action(version.ruleset.activity)
    stale_response = _refuse_frozen_or_superseded(request, version, route="staff:ruleset_bind")
    if stale_response is not None:
        return stale_response

    def _parse_json(name, empty):
        raw = (request.POST.get(name) or "").strip()
        if not raw:
            return empty
        try:
            return json.loads(raw)
        except json.JSONDecodeError:
            raise ValidationError(f"{name} 不是合法 JSON。")

    def _parse_pairs(key_name, value_name, legacy_name, empty):
        keys = request.POST.getlist(key_name)
        values = request.POST.getlist(value_name)
        if not keys and not values:
            return _parse_json(legacy_name, empty)
        if len(keys) != len(values):
            raise ValidationError(f"{key_name} 与 {value_name} 数量不一致。")
        return {
            key.strip(): value
            for key, value in zip(keys, values, strict=True)
            if key.strip() and value
        }

    try:
        base_live_challenge = _parse_json("base_live_challenge", {})
        if not isinstance(base_live_challenge, dict):
            raise ValidationError("base_live_challenge 不是合法 JSON。")

        live_enabled = request.POST.get("live_format_key") == "2026_live_challenge"

        def _live_value(post_name, key):
            value = (request.POST.get(post_name) or "").strip()
            return value if value else base_live_challenge.get(key, "")

        def _judge_id_list(post_name, key):
            values = request.POST.getlist(post_name)
            if live_enabled:
                return [
                    item.strip() for value in values for item in value.split(",") if item.strip()
                ]
            if values:
                return [
                    item.strip() for value in values for item in value.split(",") if item.strip()
                ]
            if request.POST.get(post_name) is not None:
                return []
            return base_live_challenge.get(key, [])

        live_challenge_binding = {}
        if live_enabled:
            live_challenge_binding = {
                "format_key": "2026_live_challenge",
                "participant_count": _live_value("live_participant_count", "participant_count"),
                "seat_count": _live_value("live_seat_count", "seat_count"),
                # 踢馆人数 is derived (参赛人数 − 晋级数) and deliberately not read from the
                # form: a value posted under this key is refused by ``validate_binding``.
                "group_count": _live_value("live_group_count", "group_count"),
                "group_size": _live_value("live_group_size", "group_size"),
                "first_round_weight": _live_value("live_first_weight", "first_round_weight"),
                "challenge_round_weight": _live_value(
                    "live_challenge_weight", "challenge_round_weight"
                ),
                "order_policy": "official_number",
                # These are product invariants, not operator-configurable DSL keys. The
                # built-in's own pre-fill reads the same dict, so a change to the format's
                # identity cannot make a hand-saved binding disagree with a provisioned one.
                **LIVE_2026_PODIUM_KEYS,
                "podium_first_round_weight": _live_value(
                    "live_podium_first_weight", "podium_first_round_weight"
                ),
                "podium_third_round_weight": _live_value(
                    "live_podium_third_weight", "podium_third_round_weight"
                ),
                "challenge_judge_ids": _judge_id_list(
                    "live_challenge_judge_ids", "challenge_judge_ids"
                ),
                "challenge_minimum_judge_count": _live_value(
                    "live_challenge_minimum_judges", "challenge_minimum_judge_count"
                ),
                "final_judge_ids": _judge_id_list("live_final_judge_ids", "final_judge_ids"),
                "final_minimum_judge_count": _live_value(
                    "live_final_minimum_judges", "final_minimum_judge_count"
                ),
            }

        payload = {
            "stage_key": request.POST.get("stage_key"),
            "round_keys": _parse_pairs(
                "round_binding_key", "round_binding_value", "round_keys", {}
            ),
            "vote_keys": _parse_pairs("vote_binding_key", "vote_binding_value", "vote_keys", {}),
            "vote_scoring_rule_keys": _parse_pairs(
                "scoring_binding_key", "scoring_binding_value", "vote_scoring_rule_keys", {}
            ),
            "group_keys": _parse_pairs(
                "group_binding_key", "group_binding_value", "group_keys", {}
            ),
            "group_stage_keys": _parse_pairs(
                "group_stage_binding_key",
                "group_stage_binding_value",
                "group_stage_keys",
                {},
            ),
            "audience_keys": _parse_json("audience_keys", {}),
            "live_challenge": live_challenge_binding,
            "announcement_blocks": _parse_json("announcement_blocks", []),
            "announcement_blocks_by_checkpoint": _parse_json(
                "announcement_blocks_by_checkpoint", {}
            ),
        }
        # One unit, exactly as the 2026 configuration page does: if the counts cannot be
        # carried through to the round and the definition, the binding edit must not stand.
        with transaction.atomic():
            update_ruleset_binding(
                version.ruleset,
                binding=payload,
                operator=request.user,
                base_binding=request.POST.get("base_binding") or None,
            )
            apply_live_2026_counts(version.ruleset, version, operator=request.user)
    except (ValidationError, PermissionDenied) as exc:
        messages.error(request, domain_error_messages(exc))
        return redirect("staff:ruleset_edit", pk=pk)
    messages.success(request, "生产绑定已保存。")
    return redirect("staff:ruleset_edit", pk=pk)


def _live_weight_percent(raw) -> str:
    """Render a stored weight fraction as the percentage the operators think in."""

    try:
        value = Decimal(str(raw)) * 100
    except (InvalidOperation, ValueError):
        return ""
    if value == value.to_integral_value():
        return str(int(value))
    return f"{value.normalize():f}"


def _is_live_2026_ruleset(ruleset) -> bool:
    """Whether a ruleset's editable surface declares the 2026 live format."""

    return str((ruleset.live_challenge or {}).get("format_key") or "") == LIVE_CHALLENGE_FORMAT_KEY


def _ruleset_edit_url(version, *, is_live_2026: bool) -> str:
    """Where a *draft* ruleset's operator continues working."""

    route = "staff:ruleset_live_2026_config" if is_live_2026 else "staff:ruleset_edit"
    return reverse(route, args=[version.pk])


def _live_2026_checklist(version) -> list[dict]:
    """The 准备清单 rows — the freeze's own verdicts, each linking to the page that fixes it.

    The link is resolved per row rather than from one shared ``(route, args)`` pair, because
    the rows live on different pages with different parameters: ``round_list`` takes no
    argument while ``round_edit`` takes a round, and a row whose route cannot be reversed
    takes down every page that renders the checklist.
    """

    ruleset = version.ruleset
    config_url = reverse("staff:ruleset_live_2026_config", args=[version.pk])
    edit_url = reverse("staff:ruleset_edit", args=[version.pk])
    questionnaire_url = reverse("staff:ruleset_questionnaire", args=[version.pk])
    bound_round = _bound_first_round(ruleset)
    round_url = (
        "#section-people"
        if bound_round is not None
        else reverse("staff:round_create") + f"?activity_id={ruleset.activity_id}"
    )
    group_url = _live_2026_group_url(version)
    links = {
        "counts": ("#section-people", "本页填写"),
        "weights": ("#section-scoring", "本页填写"),
        "podium": ("#section-scoring", "本页填写"),
        "judges": ("#section-judges", "本页填写"),
        "binding": (config_url, "本页填写"),
        "definition": (edit_url, "去编辑器"),
        "round_seats": (round_url, "本页填写"),
        "round_judges": ("#section-judges", "本页填写"),
        "rubric": ("#section-rubric", "本页填写"),
        "groups": (group_url, "去录入分组"),
        "round_contract": (edit_url, "去编辑器"),
        "questionnaire": (questionnaire_url, "去填写"),
        "successor": (questionnaire_url, "去检查"),
        # 已核定结果要靠撤回赛段结果来释放，那不是一页配置。
        "finality": ("", ""),
    }
    rows = []
    for item in live_2026_readiness(version):
        url, action = links.get(item.key, ("", ""))
        rows.append(
            {
                "key": item.key,
                "label": item.label,
                "done": item.ready,
                # A row whose own prerequisite has not been met is neither done nor missing;
                # the user-task layer counts it as unfinished without pretending to know why.
                "pending": item.unknown,
                "detail": item.message,
                "url": url,
                "action": action,
            }
        )
    return rows


def _bound_live_2026_group_stage(version):
    """The *explicitly bound* first-round group stage, or ``None``.

    Never fall back to "the newest GroupStage in the activity": an older, cancelled or
    test grouping would send the operator to the wrong object. The workbench section owns
    creation and explicit binding when this returns ``None``.
    """

    ruleset = version.ruleset
    first_round_key = str((ruleset.live_challenge or {}).get("podium_first_round_key") or "r1")
    stage_keys = ruleset.group_stage_keys or {}
    stage_id = stage_keys.get(first_round_key)
    if stage_id in (None, "") and len(stage_keys) == 1:
        stage_id = next(iter(stage_keys.values()))
    if stage_id in (None, ""):
        return None
    return GroupStage.objects.filter(pk=stage_id, activity=ruleset.activity_id).first()


def _live_2026_group_url(version) -> str:
    """Where the group task is completed: the bound stage, or the workbench section."""

    stage = _bound_live_2026_group_stage(version)
    if stage is not None:
        return reverse("staff:group_stage_detail", args=[stage.pk])
    return "#section-groups"


def _live_2026_tasks(version) -> list[dict]:
    """The stable seven-task operator view of the preparation checklist."""

    links = {
        "counts": ("#section-people", "本页填写"),
        "challenge_weights": ("#section-scoring", "本页填写"),
        "podium_weights": ("#section-scoring", "本页填写"),
        "judges": ("#section-judges", "本页填写"),
        "rubric": ("#section-rubric", "本页填写"),
        "groups": (_live_2026_group_url(version), "去录入分组"),
        "questionnaire": (
            reverse("staff:ruleset_questionnaire", args=[version.pk]),
            "去填写",
        ),
    }
    rows = []
    for task in live_2026_preparation_tasks(version):
        url, action = links.get(task.key, ("", ""))
        rows.append(
            {
                "key": task.key,
                "label": task.label,
                "ready": task.ready,
                "pending": task.pending,
                "detail": task.detail,
                "url": url,
                "action": action,
            }
        )
    return rows


def _parse_rubric_criteria(text: str) -> list[dict]:
    """Parse the workbench's "名称,满分" lines into rubric criteria."""

    criteria: list[dict] = []
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip().replace("，", ",")
        if not line:
            continue
        name, separator, score = line.partition(",")
        if not separator or not name.strip() or not score.strip():
            raise ValidationError(
                f"评分项第 {line_number} 行格式应为「名称,满分」，例如「演唱技巧,40」。"
            )
        criteria.append({"name": name.strip(), "max_score": score.strip()})
    return criteria


@staff_required
def ruleset_live_2026_config(request, pk):
    """The 傻瓜式 2026 决赛赛制配置页: one page for the format's few real choices."""

    version = get_object_or_404(RulesetVersion.objects.select_related("ruleset__activity"), pk=pk)
    stale_response = _refuse_frozen_or_superseded(
        request, version, route="staff:ruleset_live_2026_config"
    )
    if stale_response is not None:
        return stale_response
    ruleset = version.ruleset
    if str((ruleset.live_challenge or {}).get("format_key") or "") != LIVE_CHALLENGE_FORMAT_KEY:
        messages.error(
            request,
            "当前赛制不是 2026 现场踢馆赛制，请先在绑定时选择「2026 现场挑战」。",
        )
        return redirect("staff:ruleset_edit", pk=version.pk)

    if request.method == "POST":
        _ensure_activity_mutable(ruleset.activity)
        action = (request.POST.get("action") or "").strip()
        if action == "create_group_stage":
            try:
                stage = create_live_2026_group_stage(version, operator=request.user)
            except (PermissionDenied, ValidationError) as exc:
                messages.error(request, domain_error_messages(exc))
                return redirect("staff:ruleset_live_2026_config", pk=version.pk)
            return redirect("staff:group_stage_detail", pk=stage.pk)
        if action == "bind_group_stage":
            selected_stage = GroupStage.objects.filter(
                pk=request.POST.get("group_stage_id"), activity=ruleset.activity
            ).first()
            if selected_stage is None:
                messages.error(request, "选择的合并分组不属于当前活动。")
                return redirect("staff:ruleset_live_2026_config", pk=version.pk)
            try:
                bind_live_2026_group_stage(
                    version,
                    selected_stage,
                    operator=request.user,
                    base_binding=request.POST.get("base_binding") or None,
                )
            except (PermissionDenied, ValidationError) as exc:
                messages.error(request, domain_error_messages(exc))
            else:
                messages.success(request, "第一轮分组已绑定到本版赛制。")
            return redirect("staff:ruleset_live_2026_config", pk=version.pk)

        live_challenge = {
            "participant_count": request.POST.get("live_participant_count"),
            "seat_count": request.POST.get("live_seat_count"),
            "group_count": request.POST.get("live_group_count"),
            "group_size": request.POST.get("live_group_size"),
            "first_round_weight": request.POST.get("live_first_weight"),
            "challenge_round_weight": request.POST.get("live_challenge_weight"),
            "podium_first_round_weight": request.POST.get("live_podium_first_weight"),
            "podium_third_round_weight": request.POST.get("live_podium_third_weight"),
            "challenge_judge_ids": request.POST.getlist("live_challenge_judge_ids"),
            "challenge_minimum_judge_count": request.POST.get("live_challenge_minimum_judges"),
            "final_judge_ids": request.POST.getlist("live_final_judge_ids"),
            "final_minimum_judge_count": request.POST.get("live_final_minimum_judges"),
        }
        rubric_id = (request.POST.get("live_first_rubric_id") or "").strip()
        first_rubric = None
        if rubric_id:
            first_rubric = ScoringRubric.objects.filter(
                pk=rubric_id, activity=ruleset.activity
            ).first()
            if first_rubric is None:
                messages.error(request, "选择的评分标准不属于当前活动。")
                return redirect("staff:ruleset_live_2026_config", pk=version.pk)
        try:
            new_rubric_criteria = _parse_rubric_criteria(
                request.POST.get("live_new_rubric_criteria") or ""
            )
        except ValidationError as exc:
            messages.error(request, domain_error_messages(exc))
            return redirect("staff:ruleset_live_2026_config", pk=version.pk)

        before = {task.key: task for task in live_2026_preparation_tasks(version)}
        try:
            result = save_live_2026_preparation(
                version,
                operator=request.user,
                live_challenge=live_challenge,
                base_binding=request.POST.get("base_binding") or None,
                base_first_round_configuration=request.POST.get("base_first_round_configuration")
                or None,
                first_judge_count=request.POST.get("live_first_judge_count"),
                first_minimum_judges=request.POST.get("live_first_minimum_judges"),
                first_rubric=first_rubric,
                new_rubric_name=request.POST.get("live_new_rubric_name") or "",
                new_rubric_criteria=new_rubric_criteria,
            )
        except (ValidationError, PermissionDenied) as exc:
            messages.error(request, domain_error_messages(exc))
        else:
            after = live_2026_preparation_tasks(result.version)
            newly = [task.label for task in after if task.ready and not before[task.key].ready]
            remaining = [task.label for task in after if not task.ready]
            completed = sum(1 for task in after if task.ready)
            parts = ["已保存。"]
            if newly:
                parts.append("本次完成：" + "、".join(newly) + "。")
            parts.append(f"当前进度：{completed} / {len(after)}。")
            if remaining:
                parts.append("还需完成：" + "、".join(remaining) + "。")
            messages.success(request, "".join(parts))
        return redirect("staff:ruleset_live_2026_config", pk=version.pk)

    config = dict(ruleset.live_challenge or {})
    counts = live_challenge_counts_from_binding(_snapshot_binding(ruleset))
    first_round_key = str(config.get("podium_first_round_key") or "r1")
    round_id = (ruleset.round_keys or {}).get(first_round_key)
    first_round = (
        ContestRound.objects.filter(pk=round_id, activity_id=ruleset.activity_id).first()
        if round_id not in (None, "")
        else None
    )
    checklist = _live_2026_checklist(version)
    tasks = _live_2026_tasks(version)
    completed_count = sum(1 for task in tasks if task["ready"])
    bound_group_stage = _bound_live_2026_group_stage(version)
    unbound_group_stages = (
        GroupStage.objects.filter(activity=ruleset.activity)
        .exclude(pk=bound_group_stage.pk if bound_group_stage else -1)
        .order_by("-pk")
    )
    resumable_group_stage = (
        GroupStage.objects.filter(
            activity=ruleset.activity,
            managed_by=GroupStage.ManagedBy.LIVE_2026,
            status__in=[GroupStage.Status.DRAFT, GroupStage.Status.CONFIRMED],
        )
        .order_by("-pk")
        .first()
    )
    first_round_editable = bool(
        first_round
        and first_round.status == ContestRound.Status.DRAFT
        and not first_round.is_locked
        and first_round.pk not in current_frozen_round_ids(ruleset, excluding_version=version)
    )
    return render(
        request,
        "staff_panel/ruleset_live_2026_config.html",
        {
            "version": version,
            "ruleset": ruleset,
            "activity": ruleset.activity,
            "counts": counts,
            "config": config,
            "weights": {
                "first_round_weight": _live_weight_percent(config.get("first_round_weight")),
                "challenge_round_weight": _live_weight_percent(
                    config.get("challenge_round_weight")
                ),
                "podium_first_round_weight": _live_weight_percent(
                    config.get("podium_first_round_weight")
                ),
                "podium_third_round_weight": _live_weight_percent(
                    config.get("podium_third_round_weight")
                ),
            },
            "live_judges": list(
                Judge.objects.filter(activity=ruleset.activity, is_active=True).order_by(
                    "name", "pk"
                )
            ),
            "challenge_judge_ids": [int(v) for v in config.get("challenge_judge_ids") or []],
            "final_judge_ids": [int(v) for v in config.get("final_judge_ids") or []],
            "first_round": first_round,
            # 第一轮一旦准备过就不能再改晋级人数。 Two very different situations produce that,
            # and offering the same button for both is the trap: when the round belongs to the
            # running FROZEN authority, resetting it destroys the prepared first round *and*
            # the guard still refuses the new number afterwards. Reset is only ever the way
            # out for a round nothing authoritative is using; the protected case gets the
            # clone-and-rebind action instead.
            "first_round_is_protected": bool(
                first_round
                and first_round.pk in current_frozen_round_ids(ruleset, excluding_version=version)
            ),
            "first_round_needs_reset": bool(
                first_round
                and (first_round.status != ContestRound.Status.DRAFT or first_round.is_locked)
                and first_round.pk
                not in current_frozen_round_ids(ruleset, excluding_version=version)
            ),
            "checklist": checklist,
            "system_checklist": [
                row for row in checklist if row["key"] not in PREPARATION_TASK_ROW_KEYS
            ],
            "tasks": tasks,
            "completed_count": completed_count,
            "total_count": len(tasks),
            "remaining_count": len(tasks) - completed_count,
            "rubrics": ScoringRubric.objects.filter(
                activity=ruleset.activity,
                is_test_data=runtime_is_test(ruleset.activity),
            ).order_by("sequence", "pk"),
            "first_judge_count": first_round.judge_count if first_round else None,
            "first_minimum_judges": first_round.minimum_judge_count if first_round else None,
            "first_rubric_id": first_round.rubric_id if first_round else None,
            "first_round_editable": first_round_editable,
            "base_first_round_configuration": first_round_configuration_signature(first_round),
            "bound_group_stage": bound_group_stage,
            "resumable_group_stage": resumable_group_stage,
            "ready_group_stages": [
                stage for stage in unbound_group_stages if stage.status == GroupStage.Status.FROZEN
            ],
            "unfinished_group_stages": [
                stage for stage in unbound_group_stages if stage.status != GroupStage.Status.FROZEN
            ],
            # The authoritative answer, from the walk the freeze itself performs: the
            # checklist below is more granular, but this is the line that cannot disagree
            # with the freeze button.
            "blocker": freeze_blocker(version),
            "binding_signature": _binding_signature(ruleset),
            "can_freeze": request.user.is_admin,
        },
    )


@staff_required
@require_POST
def ruleset_live_2026_clone_first_round(request, pk):
    """Give a successor draft its own first round, so its seats can move.

    The one action behind the "当前正式赛制仍在使用第一轮" refusal: copy the round's
    configuration onto a fresh draft round, rebind this version to it, and come back here to
    change the number. Nothing from the running round's runtime state comes across.
    """

    version = get_object_or_404(RulesetVersion.objects.select_related("ruleset__activity"), pk=pk)
    stale_response = _refuse_frozen_or_superseded(
        request, version, route="staff:ruleset_live_2026_config"
    )
    if stale_response is not None:
        return stale_response
    try:
        _ensure_activity_mutable(version.ruleset.activity)
        cloned = clone_first_round_for_successor(version.ruleset, version, operator=request.user)
    except (PermissionDenied, ValidationError) as exc:
        messages.error(request, domain_error_messages(exc))
    else:
        messages.success(
            request,
            f"已为本版赛制新建「{cloned.name}」并改绑第一轮，现在可以直接修改晋级人数。",
        )
    return redirect("staff:ruleset_live_2026_config", pk=version.pk)


@admin_required
def manual_decision(request, pk):
    version = get_object_or_404(RulesetVersion.objects.select_related("ruleset__activity"), pk=pk)
    form = ManualDecisionForm(
        request.POST or None, version=version, activity=version.ruleset.activity
    )
    if request.method == "POST" and form.is_valid():
        try:
            set_manual_decision(
                version,
                manual_key=form.cleaned_data["manual_key"],
                group=form.cleaned_data["group"],
                chosen=form.cleaned_data["chosen"],
                created_by=request.user,
            )
        except (PermissionDenied, ValidationError) as error:
            form.add_error(None, domain_error_messages(error))
        else:
            messages.success(request, "人工晋级选择已保存。")
            return redirect("staff:ruleset_edit", pk=pk)
    return render(
        request,
        "staff_panel/manual_decision_form.html",
        {"form": form, "version": version, "activity": version.ruleset.activity},
    )


@staff_required
def ruleset_validate(request, pk):
    version = get_object_or_404(RulesetVersion.objects.select_related("ruleset"), pk=pk)
    report, plan = compile_definition(version.definition)
    completeness_issues = validate_definition_completeness(version.definition)
    if completeness_issues:
        report = ValidationReport(report.issues + completeness_issues)
    is_live_2026 = (
        str((version.ruleset.live_challenge or {}).get("format_key") or "")
        == LIVE_CHALLENGE_FORMAT_KEY
    )
    return render(
        request,
        "staff_panel/ruleset_validate.html",
        {
            "version": version,
            "passes": report.passes(),
            "issues": [i.to_dict() for i in report.issues],
            "counts": report.counts(),
            "summary": plan.summary if plan else None,
            # For a 2026 ruleset the compile report is only the graph. The gate that decides
            # whether freeze will pass is the preflight, which runs freeze's own validators.
            "live_2026": is_live_2026,
            "preflight": _live_2026_checklist(version) if is_live_2026 else [],
            # The authoritative verdict, straight from the freeze's own walk. The checklist
            # above can only be read as "what is left"; this is "冻结会不会通过".
            "blocker": freeze_blocker(version) if is_live_2026 else None,
            "config_url": (
                reverse("staff:ruleset_live_2026_config", args=[version.pk]) if is_live_2026 else ""
            ),
        },
    )


@staff_required
def ruleset_preview(request, pk):
    version = get_object_or_404(RulesetVersion, pk=pk)
    report, plan, result = ruleset_editor.preview_definition(version.definition)
    return render(
        request,
        "staff_panel/ruleset_preview.html",
        {
            "version": version,
            "passes": report.passes(),
            "issues": [i.to_dict() for i in report.issues],
            "summary": plan.summary if plan else None,
            "status": result.status.value if result else "—",
            "node_values": result.node_values if result else None,
            "decisions": [d.to_dict() for d in result.decisions][:20] if result else None,
        },
    )


@admin_required
@require_POST
def ruleset_freeze(request, pk):
    version = get_object_or_404(RulesetVersion.objects.select_related("ruleset"), pk=pk)
    try:
        freeze_ruleset_version(version, request.user)
    except PermissionDenied as exc:
        messages.error(request, domain_error_messages(exc))
    except RulesetInvalidError as exc:
        errors = [i.message for i in exc.report.issues if i.severity.value == "error"]
        messages.error(request, "赛制排版未通过校验：" + ("；".join(errors) or "存在 ERROR"))
    except ValidationError as exc:
        messages.error(request, domain_error_messages(exc))
    else:
        messages.success(request, "赛制已冻结。")
    return _post_freeze_redirect(version)


def _post_freeze_redirect(version):
    """Where to land after a freeze attempt — never on a page that will refuse it.

    A successful freeze used to redirect into the editor, which refuses frozen versions: the
    operator watched the row freeze and then got a 403 for it, with the success message
    hidden behind the error page. The workspace is the page that can answer "and now what?"
    for a frozen version, and it is also the safe landing when the freeze failed because the
    version is not editable — an older draft, or one that was already frozen.
    """

    version.refresh_from_db(fields=["status"])
    if not version_is_editable(version):
        return redirect("staff:activity_workspace", pk=version.ruleset.activity_id)
    return redirect(_ruleset_edit_url(version, is_live_2026=_is_live_2026_ruleset(version.ruleset)))


@staff_required
@require_POST
@transaction.atomic
def ruleset_clone_from_template(request, template_pk):
    template = get_object_or_404(
        RulesetTemplate,
        pk=template_pk,
        capability_status=RulesetTemplate.CapabilityStatus.PRODUCTION,
    )
    activity = get_object_or_404(
        Activity,
        pk=request.POST.get("activity"),
        activity_type=Activity.Type.SINGER_CONTEST,
    )
    activity = lock_activity_for_action(activity)
    _ensure_activity_mutable(activity)
    name = (request.POST.get("name") or "").strip() or template.name
    provisioned = _provision_builtin_live_format(
        request, template=template, activity=activity, name=name
    )
    if provisioned is not None:
        return provisioned
    ruleset, _created = ContestRuleset.objects.get_or_create(
        activity=activity,
        defaults={
            "name": name,
            "source_template": template,
            "is_test_data": runtime_is_test(activity),
            "created_by": request.user,
        },
    )
    ruleset.name = name
    ruleset.source_template = template
    ruleset.is_test_data = runtime_is_test(activity)
    ruleset.save()
    version = create_ruleset_version(
        ruleset, definition=template.definition, created_by=request.user
    )
    messages.success(request, f"已从「{template.name}」克隆到活动，进入准备。")
    return redirect("staff:ruleset_planning", pk=version.pk)


@staff_required
@require_POST
def ruleset_clone_last_year(request):
    # Identity, not display name: the 2025 historical template is found by its stable
    # builtin_key so a future rename can never break "clone last year" (§11.2).
    template = get_object_or_404(RulesetTemplate, builtin_key=GOLDEN_SCHIDUI_BUILTIN_KEY)
    return ruleset_clone_from_template(request, template.pk)
